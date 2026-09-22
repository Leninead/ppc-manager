"""PPC Insights page: the rows the AI table shows, what the tab says about the analysis on screen, and Recalcular."""
import io
from datetime import date, datetime, timezone
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from core.ai_analysis.store import StoredAnalysis
from core.amazon_ads.advertised_asins import FROM_AD_GROUP, SEVERAL_ASINS, WITHOUT_ASIN
from core.helpers import kpi_card
from core.ppc_insights.analysis import build_analysis_input
from core.ppc_insights.asin_health import InsightsAnalysisParams
from core.ppc_insights.asin_health import ATTRIBUTED, FROM_FILE, NO_ASINS
from core.search_term.file import SearchTermFileError
from core.search_term.frame import SOURCE_FILE, console_columns
from modules.pages import ppc_insights, search_term_source
from tests.test_str_analysis_job import FakeRest

WINDOW = (date(2026, 9, 8), date(2026, 9, 14))


def _stored(analysis_id=1, *, window=WINDOW, params=None, input_digest="d-now", agent_version="v-now"):
    return StoredAnalysis(
        id=analysis_id, module="ppc_insights", subject_id="111", window_start=window[0], window_end=window[1],
        lang="es", params=params or {"target_acos": 25, "price": 15.0}, params_digest="", input_digest=input_digest,
        agent_version=agent_version, status="done", trigger="manual", requested_by="ana", job_id=None,
        source_last_success_at=None, result={"asins": [], "synthesis": {"situation": "s"}}, model="",
        duration_ms=None, created_at=None, finished_at=datetime(2026, 9, 15, 12, 0, tzinfo=timezone.utc),
        records=[{"asin": "B0CYLMJJJC"}])


def _records():
    frame = pd.DataFrame([{"Customer Search Term": "vitamin a cream", "Campaign Name": "DG", "_ad_group_id": "AG1",
                           "Spend": 30.0, "7 Day Total Sales": 0.0, "7 Day Total Orders (#)": 0, "Clicks": 12,
                           "Impressions": 900}])
    for column in console_columns(7):
        if column not in frame.columns:
            frame[column] = 0
    br = pd.DataFrame({"(Child) ASIN": ["B0CYLMJJJC"], "Sessions - Total": ["100"],
                       "Featured Offer (Buy Box) Percentage": ["90%"], "Units Ordered": ["5"]})
    camp = pd.DataFrame({"Campaign Name": ["DG - B0CYLMJJJC - SP - AUTO"], "State": ["enabled"]})
    return build_analysis_input(frame, params=InsightsAnalysisParams(25, 15.0), account_label="dg", period_label="",
                                currency_code="MXN", ad_group_asins={"AG1": frozenset({"B0CYLMJJJC"})},
                                br_df=br, camp_df=camp).records


# ── the AI table ─────────────────────────────────────────────────────────────

def test_the_ai_rows_show_the_figures_in_the_account_currency_and_speak_in_plain_names():
    records = _records()
    opinions = [{"row_id": "P01", "razon": "gasto_sin_venta alto con 12 clicks", "foco": "DESPERDICIO",
                 "confianza": "baja", "advertencia": None},
                {"row_id": "P07", "razon": "no existe", "foco": "ACOS", "confianza": "alta", "advertencia": None}]

    rows = ppc_insights.insights_ai_rows(opinions, records, ppc_insights.INSIGHTS_FIELD_NAMES["es"], "MXN")

    assert len(rows) == 1
    assert rows[0]["item"] == "B0CYLMJJJC" and rows[0]["badges"] == ["DESPERDICIO"]
    assert "gasto MX$30.00" in rows[0]["metrics"] and "sin venta MX$30.00" in rows[0]["metrics"]
    assert rows[0]["reasoning"] == "gasto sin venta de sus peores términos alto con 12 clicks"
    assert rows[0]["confidence"] == "BAJA"


def test_the_glossary_names_every_column_the_agent_reads_in_both_languages():
    columns = set(_records()[0])
    glossary = ppc_insights.INSIGHTS_FIELD_NAMES

    assert columns - {"asin"} <= set(glossary["es"])
    assert set(glossary["es"]) == set(glossary["en"])


# ── what the tab says about the analysis on screen ───────────────────────────

def test_the_difference_names_the_period_and_parameters_that_changed():
    stored = _stored(window=(date(2026, 9, 1), date(2026, 9, 7)), params={"target_acos": 30, "price": None})

    text = ppc_insights.stored_difference_text(stored, InsightsAnalysisParams(25, 15.0), *WINDOW)

    assert "es del 1 – 7 sep 2026" in text and "usa un target de 30%" in text and "no tenía precio" in text


def test_the_same_period_and_parameters_mean_new_data_arrived():
    text = ppc_insights.stored_difference_text(_stored(), InsightsAnalysisParams(25, 15.0), *WINDOW)

    assert "datos nuevos" in text


@pytest.mark.parametrize("source, share, expected", [
    (ATTRIBUTED, {FROM_AD_GROUP: 70.0, SEVERAL_ASINS: 20.0, WITHOUT_ASIN: 10.0},
     "Gasto por origen del ASIN: producto anunciado del ad group 70.0% · ad groups con varios ASINs sin ASIN en el "
     "nombre 20.0% · sin ASIN 10.0%. El 30.0% sin ASIN no entra en las cards."),
    (FROM_FILE, {FROM_FILE: 100.0}, ""),
    (NO_ASINS, {WITHOUT_ASIN: 100.0}, ""),
])
def test_the_coverage_caption_says_where_the_asins_came_from_only_when_it_matters(source, share, expected):
    assert ppc_insights.asin_coverage_caption(source, share) == expected


def test_the_no_asin_notice_names_the_cause_the_spend_actually_has():
    several_only = ppc_insights.no_asin_notice({SEVERAL_ASINS: 100.0})
    both = ppc_insights.no_asin_notice({WITHOUT_ASIN: 30.0, SEVERAL_ASINS: 70.0})

    assert "ad groups que anuncian varios ASINs (100.0%)" in several_only
    assert "no vio" not in several_only
    assert "varios ASINs (70.0%) y ad groups que el listado de productos anunciados no vio (30.0%)" in both
    assert ppc_insights.no_asin_notice({}).endswith("una sola fila (ALL).")


def test_the_spend_card_says_of_how_much_only_when_the_cards_leave_spend_out():
    assert ppc_insights.spend_kpi(750.0, 1000.0, "MXN") == ("Spend en cards", "MX$750.00", "de MX$1,000.00 (75.0%)")
    assert ppc_insights.spend_kpi(1000.0, 1000.0, "USD") == ("Spend total", "$1,000.00", None)


def test_the_spend_caption_keeps_its_amounts_out_of_latex():
    card = kpi_card("Spend en cards", "$750.00", caption="de $1,000.00 (75.0%)")

    assert "de &#36;1,000.00 (75.0%)" in card
    # Without a caption the card keeps its three blocks, as every other page draws it.
    assert kpi_card("Spend total", "$1,000.00").count("<div") == 3


def test_a_grouping_row_shows_how_many_asins_it_groups():
    records = [{"asin": "B0FAMILY01", "health_score": 40, "spend": 90.0, "acos": None, "cvr": 0.0,
                "gasto_sin_venta": 0.0, "asins_agrupados": 14}]
    opinions = [{"row_id": "P01", "razon": "r", "foco": "ACOS", "confianza": "alta", "advertencia": None}]

    rows = ppc_insights.insights_ai_rows(opinions, records, ppc_insights.INSIGHTS_FIELD_NAMES["es"], "USD")

    assert "agrupa 14 ASINs" in rows[0]["metrics"]


def test_the_inputs_signature_changes_with_the_data_the_target_or_a_file_and_only_with_them():
    source = SimpleNamespace(signature="profile-1:v3")
    upload = SimpleNamespace(getvalue=lambda: b"sqp bytes")
    base = ppc_insights._inputs_signature(source, {"sqp": upload, "br": None}, 25)

    assert ppc_insights._inputs_signature(source, {"sqp": upload, "br": None}, 25) == base
    assert ppc_insights._inputs_signature(source, {"sqp": upload, "br": None}, 30) != base
    assert ppc_insights._inputs_signature(source, {"sqp": None, "br": None}, 25) != base
    assert ppc_insights._inputs_signature(SimpleNamespace(signature="profile-1:v4"),
                                          {"sqp": upload, "br": None}, 25) != base


def test_an_unreadable_manual_file_is_refused_with_a_message():
    with pytest.raises(SearchTermFileError, match="Error al leer el STR"):
        ppc_insights._read_manual_str(b"not an excel file", "report.xlsx")


def test_the_manual_file_goes_through_the_modules_own_reader():
    source = search_term_source._source_from_module_reader(
        lambda file_bytes, file_name: pd.DataFrame({"Customer Search Term": ["a"]}), b"csv bytes", "report.csv")

    assert (source.source, source.label, source.currency_code) == (SOURCE_FILE, "report.csv", "")
    assert source.signature.endswith(":module-reader") and list(source.frame.columns) == ["Customer Search Term"]


def test_the_excel_writes_amounts_in_the_account_currency():
    asin_data = {"B0CYLMJJJC": {
        "spend": 30.0, "sales": 90.0, "orders": 3.0, "clicks": 12.0, "acos": 33.3, "cvr": 25.0, "wasted_spend": 0.0,
        "top_kws": pd.DataFrame(), "bleeders": pd.DataFrame(), "imp_share": None, "purchase_share": None,
        "sqp_gaps": 0, "sessions": None, "buybox": None, "br_units": None, "n_campaigns": None,
        "campaign_types": None, "funnel_complete": None, "health_score": 60}}

    for currency, expected in (("MXN", '"MX$"#,##0.00'), ("JPY", '"¥"#,##0')):
        book = openpyxl.load_workbook(io.BytesIO(ppc_insights._build_insights_excel(asin_data, "dg", 25, currency)
                                                 .getvalue()))
        assert book["Resumen"].cell(row=3, column=3).number_format == expected


# ── the stored analysis tab ──────────────────────────────────────────────────

class _AnalysisRest(FakeRest):
    def rpc(self, name, args, **_):
        self.rpc_calls.append((name, args))
        if name == "request_ai_analysis":
            return [{"job_id": 55, "created": True, "reason": "created"}]
        if name == "save_ai_analysis_settings":
            return True
        return super().rpc(name, args)


_TAB_SCRIPT = """
import streamlit as st
from datetime import date
from types import SimpleNamespace
from core.ai_analysis import stored_tab
from core.ppc_insights.asin_health import InsightsAnalysisParams
st.cache_data.clear()
fake = st.session_state["fake_rest"]
params = InsightsAnalysisParams(25, 15.0)
result = stored_tab.render_recalculable_analysis(
    module="ppc_insights", key_prefix="t", source=SimpleNamespace(
        profile_id="111", window_start=date(2026, 9, 8), window_end=date(2026, 9, 14)),
    input_digest="d-now", agent_version="v-now", params=params, account_params=params,
    open_rest=lambda: fake, current_username=lambda: "ana",
    render_result=lambda stored: st.markdown(f"RENDERED {stored.id}"),
    describe_difference=lambda stored: "DIFFERENT DATA")
st.session_state["tab_state"] = result.state
"""


def _tab(fake):
    app = AppTest.from_string(_TAB_SCRIPT, default_timeout=30)
    app.session_state["fake_rest"] = fake
    app.run()
    return app


def _texts(app):
    return " ".join(str(element.value) for element in app.markdown)


def _analysis_row(stored):
    return {"id": stored.id, "module": stored.module, "subject_id": stored.subject_id,
            "window_start": stored.window_start.isoformat(), "window_end": stored.window_end.isoformat(),
            "lang": "es", "params": stored.params, "input_digest": stored.input_digest,
            "agent_version": stored.agent_version, "status": "done", "trigger": "manual", "requested_by": "ana",
            "result": stored.result, "records": stored.records, "finished_at": stored.finished_at.isoformat()}


def test_an_account_without_analyses_offers_to_generate_the_first_one():
    app = _tab(_AnalysisRest())

    assert "Todavía no hay un análisis IA de esta cuenta" in _texts(app)
    assert [button.label for button in app.button] == ["Generar análisis IA"]


def test_the_latest_analysis_of_other_data_stays_on_screen_with_recalcular():
    fake = _AnalysisRest()
    fake.tables["ai_analyses"].append(_analysis_row(_stored(7, input_digest="d-older")))

    app = _tab(fake)

    assert "DIFFERENT DATA" in _texts(app) and "RENDERED 7" in _texts(app)
    assert [button.label for button in app.button] == ["Recalcular"]
    app.button[0].click().run()
    [(name, args)] = [call for call in fake.rpc_calls if call[0] == "request_ai_analysis"]
    assert args["p_input_digest"] == "d-now" and "p_agent_version" not in args


def test_the_analysis_of_exactly_these_data_by_this_version_needs_no_recalcular():
    fake = _AnalysisRest()
    fake.tables["ai_analyses"].append(_analysis_row(_stored(8)))

    app = _tab(fake)

    assert "RENDERED 8" in _texts(app)
    assert len(app.button) == 0
    assert app.session_state["tab_state"] == "current"


def test_the_same_data_read_by_an_older_version_can_be_recalculated_with_the_new_one():
    fake = _AnalysisRest()
    fake.tables["ai_analyses"].append(_analysis_row(_stored(9, agent_version="v-old")))

    app = _tab(fake)

    assert "Hay una versión más nueva del análisis IA" in _texts(app) and "RENDERED 9" in _texts(app)
    app.button[0].click().run()
    [(name, args)] = [call for call in fake.rpc_calls if call[0] == "request_ai_analysis"]
    assert args["p_agent_version"] == "v-now"


def test_while_the_new_analysis_runs_the_previous_one_stays_on_screen():
    fake = _AnalysisRest()
    fake.tables["ai_analyses"].append(_analysis_row(_stored(10, input_digest="d-older")))
    fake.tables["integration_sync_jobs"].append({
        "id": 60, "integration_slug": "amazon_ads", "job_kind": "ai_ppc_insights_analysis", "status": "running",
        "external_account_id": "111", "params": {"input_digest": "d-now"},
        "created_at": "2026-09-15T12:00:00+00:00"})

    app = _tab(fake)

    assert app.session_state["tab_state"] == "running"
    assert "RENDERED 10" in _texts(app)
    assert len(app.button) == 0
    assert any("Generando el análisis IA" in str(status.label) for status in app.status)
