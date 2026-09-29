"""Stock in the Weekly Client Report: the Pricing Dashboard snapshot per ASIN and the STOCK group of WoW Comparison."""
from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.weekly_report import stock as stock_mod
from core.weekly_report.stock import (
    REPORT_COLUMNS, AsinStock, StockSnapshot, days_before_week, latest_stock, stock_by_asin,
)
from modules.pages import weekly_client_report as wcr

DASH = "—"
FIRST_STOCK_COL = 24  # after TACoS, the sheet's last fixed column

_parse_br_wow = wcr._parse_br_wow.__wrapped__
_parse_br_daily_wow = wcr._parse_br_daily_wow.__wrapped__

# 14 days, 3-16 Aug 2026: the report's week starts on 2026-08-10.
_BR_DAILY = "".join(
    ["Date,Ordered Product Sales,Units Ordered,Sessions - Total,Order Item Session Percentage\n"]
    + [f'8/{d}/26,"MX$500.00",5,25,20.00%\n' for d in range(3, 17)]
)
_BR_CHILD = """\
(Parent) ASIN,(Child) ASIN,Title,Sessions - Total,Featured Offer (Buy Box) Percentage,Units Ordered,Unit Session Percentage,Ordered Product Sales
B0PARENT001,B0TEST0001,Producto Uno,50,100.00%,10,20.00%,"MX$1,000.00"
B0PARENT002,B0TEST0002,Producto Dos,100,100.00%,20,20.00%,"MX$2,000.00"
B0PARENT003,B0TEST0003,Producto Tres,40,100.00%,4,10.00%,"MX$400.00"
"""


def _snapshot(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(rows)


def _wow_sheet(buf):
    wb = load_workbook(buf)
    return next(ws for ws in wb.worksheets if ws.title.endswith("WoW Comparison"))


def _excel(stock, client="Dermaglos", daily=True):
    br_tw = _parse_br_wow(io.BytesIO(_BR_CHILD.encode()))
    br_daily = _parse_br_daily_wow(io.BytesIO(_BR_DAILY.encode())) if daily else None
    buf = wcr._build_weekly_excel(br_tw, {}, {}, {}, "TEST", "es", br_daily, stock=stock, stock_client=client)
    return _wow_sheet(buf)


def _row_of(ws, asin: str) -> int:
    return next(r for r in range(4, ws.max_row + 1) if ws.cell(r, 2).value == asin)


def _stock_cells(ws, row: int) -> list:
    return [ws.cell(row, FIRST_STOCK_COL + i).value for i in range(len(REPORT_COLUMNS))]


def _notes(ws) -> list[str]:
    return [str(ws.cell(r, 1).value) for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value]


# ── stock_by_asin ────────────────────────────────────────────────────────────


def test_the_skus_of_one_asin_add_up():
    stock = stock_by_asin(_snapshot([
        {"sku": "A-1", "asin": "B0TEST0001", "fba_available": 10, "awd_available": 5, "izzi_available": 0,
         "total_stock": 15},
        {"sku": "A-2", "asin": " b0test0001 ", "fba_available": 3, "awd_available": 0, "izzi_available": 2,
         "total_stock": 5},
    ]), "2026-08-12")

    assert stock.for_asin("B0TEST0001") == AsinStock(fba=13, awd=5, izzi=2, total=20)
    assert stock.snapshot_date == "2026-08-12"


def test_skus_without_an_asin_are_left_out_and_counted():
    stock = stock_by_asin(_snapshot([
        {"sku": "A-1", "asin": "B0TEST0001", "fba_available": 4, "total_stock": 4},
        {"sku": "A-2", "asin": "", "fba_available": 9, "total_stock": 9},
        {"sku": "A-3", "asin": None, "fba_available": 7, "total_stock": 7},
        {"sku": "A-4", "asin": float("nan"), "fba_available": 1, "total_stock": 1},
    ]), "2026-08-12")

    assert list(stock.by_asin) == ["B0TEST0001"]
    assert stock.skus_without_asin == 3
    assert stock.account.fba == 4


def test_a_missing_value_stays_unknown_and_a_real_zero_stays_zero():
    stock = stock_by_asin(_snapshot([
        {"sku": "A-1", "asin": "B0TEST0001", "fba_available": 0, "awd_available": None, "total_stock": 0},
        {"sku": "B-1", "asin": "B0TEST0002", "fba_available": 6, "awd_available": None, "total_stock": 6},
        {"sku": "B-2", "asin": "B0TEST0002", "fba_available": None, "awd_available": 4, "total_stock": 4},
    ]), "2026-08-12")

    assert stock.for_asin("B0TEST0001") == AsinStock(fba=0, awd=None, izzi=None, total=0)
    assert stock.for_asin("B0TEST0002") == AsinStock(fba=6, awd=4, izzi=None, total=10)


def test_an_asin_the_snapshot_does_not_have_is_unknown():
    stock = stock_by_asin(_snapshot([{"sku": "A-1", "asin": "B0TEST0001", "fba_available": 1}]), "2026-08-12")

    assert stock.for_asin("B0OTHER") == AsinStock()


def test_an_older_snapshot_without_a_stock_column_leaves_it_unknown():
    stock = stock_by_asin(_snapshot([{"sku": "A-1", "asin": "B0TEST0001", "fba_available": 8}]), "2026-08-12")

    assert stock.for_asin("B0TEST0001") == AsinStock(fba=8)


def test_the_account_adds_up_every_asin():
    stock = stock_by_asin(_snapshot([
        {"sku": "A-1", "asin": "B0TEST0001", "fba_available": 10, "total_stock": 10},
        {"sku": "B-1", "asin": "B0TEST0002", "fba_available": 5, "total_stock": 12},
    ]), "2026-08-12")

    assert stock.account == AsinStock(fba=15, awd=None, izzi=None, total=22)


def test_without_a_snapshot_there_is_no_date_and_no_stock():
    assert stock_by_asin(None, "2026-08-12") == StockSnapshot()


def test_a_snapshot_without_an_asin_column_leaves_every_sku_out():
    stock = stock_by_asin(_snapshot([{"sku": "A-1", "fba_available": 3}, {"sku": "A-2", "fba_available": 1}]),
                          "2026-08-12")

    assert stock.by_asin == {}
    assert stock.skus_without_asin == 2
    assert stock.snapshot_date == "2026-08-12"


# ── latest_stock ─────────────────────────────────────────────────────────────


def test_latest_stock_reads_the_newest_snapshot(monkeypatch):
    loaded = []
    monkeypatch.setattr(stock_mod, "_list_periods", lambda area, cliente, modulo: ["2026-08-01", "2026-08-12"])

    def load(area, cliente, modulo, period):
        loaded.append((area, cliente, modulo, period))
        return _snapshot([{"sku": "A-1", "asin": "B0TEST0001", "fba_available": 2}])

    monkeypatch.setattr(stock_mod, "_load_snapshot", load)

    stock = latest_stock("dermaglos")

    assert loaded == [("account-health", "dermaglos", "pricing-dashboard", "2026-08-12")]
    assert stock.snapshot_date == "2026-08-12"
    assert stock.for_asin("B0TEST0001").fba == 2


def test_latest_stock_without_snapshots_reads_nothing(monkeypatch):
    monkeypatch.setattr(stock_mod, "_list_periods", lambda *args: [])
    monkeypatch.setattr(stock_mod, "_load_snapshot", lambda *args: pytest.fail("no snapshot to load"))

    assert latest_stock("dermaglos") == StockSnapshot()


# ── days_before_week ─────────────────────────────────────────────────────────


@pytest.mark.parametrize("snapshot_date, week_start, expected", [
    ("2026-08-01", "2026-08-10", 9),
    ("2026-08-10", "2026-08-10", None),
    ("2026-08-14", "2026-08-10", None),
    (None, "2026-08-10", None),
    ("2026-08-01", None, None),
])
def test_days_before_week(snapshot_date, week_start, expected):
    assert days_before_week(snapshot_date, week_start) == expected


# ── the STOCK group in WoW Comparison ────────────────────────────────────────


def _stock_for_report(date="2026-08-12"):
    return stock_by_asin(_snapshot([
        {"sku": "A-1", "asin": "B0TEST0001", "fba_available": 10, "awd_available": 5, "izzi_available": 1,
         "total_stock": 16},
        {"sku": "B-1", "asin": "B0TEST0002", "fba_available": 0, "awd_available": 0, "izzi_available": 0,
         "total_stock": 0},
        {"sku": "X-1", "asin": "", "fba_available": 7, "total_stock": 7},
    ]), date)


def test_without_stock_the_sheet_ends_at_tacos():
    ws = _excel(None)

    assert ws.max_column == FIRST_STOCK_COL - 1
    assert not any("Stock" in note for note in _notes(ws))


def test_the_stock_group_is_one_fba_column_after_tacos_with_its_date():
    ws = _excel(_stock_for_report())

    assert ws.cell(2, FIRST_STOCK_COL).value == "STOCK al 2026-08-12"
    assert _stock_cells(ws, 3) == ["FBA"]
    assert ws.max_column == FIRST_STOCK_COL
    assert ws.cell(3, 23).value == "Esta semana", "TACoS stays in column 23"


def test_each_asin_shows_its_units_a_real_zero_and_a_dash_when_unknown():
    ws = _excel(_stock_for_report())

    assert _stock_cells(ws, _row_of(ws, "B0TEST0001")) == [10]
    assert _stock_cells(ws, _row_of(ws, "B0TEST0002")) == [0]
    assert _stock_cells(ws, _row_of(ws, "B0TEST0003")) == [DASH]


def test_the_account_row_adds_up_the_snapshot_fba():
    ws = _excel(_stock_for_report())

    assert ws.cell(4, 1).value == "▶ CUENTA TOTAL"
    assert _stock_cells(ws, 4) == [10]


def test_the_note_names_the_source_date_and_skus_left_out():
    ws = _excel(_stock_for_report())

    stock_note = next(note for note in _notes(ws) if note.startswith("Stock:"))
    assert "Pricing Dashboard de Dermaglos (2026-08-12)" in stock_note
    assert "SKUs sin ASIN fuera del stock: 1" in stock_note
    assert "AWD e Izzi todavía no se integran en el Pricing Dashboard: el stock es solo FBA" in stock_note
    assert "anterior a la semana" not in stock_note


def test_a_snapshot_older_than_the_week_is_flagged():
    ws = _excel(_stock_for_report(date="2026-08-01"))

    stock_note = next(note for note in _notes(ws) if note.startswith("Stock:"))
    assert "9 días anterior a la semana del reporte (desde 2026-08-10)" in stock_note


def test_a_client_without_snapshots_gets_dashes_and_says_so():
    ws = _excel(StockSnapshot())

    assert ws.cell(2, FIRST_STOCK_COL).value == "STOCK (sin snapshot)"
    assert _stock_cells(ws, 4) == [DASH]
    assert _stock_cells(ws, _row_of(ws, "B0TEST0001")) == [DASH]
    assert any("Dermaglos no tiene snapshots" in note for note in _notes(ws))


def test_without_the_daily_report_the_stock_still_shows_per_asin():
    ws = _excel(_stock_for_report(), daily=False)

    assert _stock_cells(ws, _row_of(ws, "B0TEST0001")) == [10]
    assert any(note.startswith("Stock:") for note in _notes(ws))
