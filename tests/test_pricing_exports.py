"""Tests M30 F3.5 — exports XLSX (_build_resumen_excel + _build_vista_excel).

Valida bytes parseables con openpyxl (load_workbook desde BytesIO), nombres de hoja,
headers/orden, conteo de filas, orden por clasificación, defaults y anchos.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from modules.pages.pricing_dashboard import (
    _build_resumen_excel,
    _RESUMEN_HEADERS,
)


def _rec(**over) -> dict:
    r = {
        "sku": "SKU1",
        "classification": "mantener",
        "fba_available": 10.0,
        "awd_available": 0.0,
        "izzi_available": 0.0,
        "total_stock": 10.0,
        "price": 20.0,
        "suggestedPrice": 19.0,
        "t7": 3.0,
    }
    r.update(over)
    return r


def _wb(data: bytes):
    return load_workbook(BytesIO(data))


# =====================================================================
# _build_resumen_excel
# =====================================================================
class TestBuildResumenExcel:
    def test_devuelve_bytes_parseables(self):
        data = _build_resumen_excel([_rec()], "2026-04-17")
        assert isinstance(data, bytes) and len(data) > 0
        wb = _wb(data)
        assert wb.sheetnames == ["Pricing"]

    def test_headers_orden_exacto(self):
        wb = _wb(_build_resumen_excel([_rec()], "2026-04-17"))
        ws = wb["Pricing"]
        fila1 = [ws.cell(row=1, column=j).value for j in range(1, 10)]
        assert fila1 == _RESUMEN_HEADERS
        assert len(_RESUMEN_HEADERS) == 9

    def test_filas_igual_a_records_mas_header(self):
        recs = [_rec(sku="A"), _rec(sku="B"), _rec(sku="C")]
        ws = _wb(_build_resumen_excel(recs, "2026-04-17"))["Pricing"]
        assert ws.max_row == len(recs) + 1

    def test_orden_por_clasificacion(self):
        # mezclados: debe salir bajar < subir < liquidar < mantener (por _RESUMEN_ORDER)
        recs = [
            _rec(sku="MANT", classification="mantener"),
            _rec(sku="LIQ", classification="liquidar"),
            _rec(sku="BAJ", classification="bajar"),
            _rec(sku="SUB", classification="subir"),
        ]
        ws = _wb(_build_resumen_excel(recs, "2026-04-17"))["Pricing"]
        skus = [ws.cell(row=i, column=1).value for i in range(2, 6)]
        assert skus == ["BAJ", "SUB", "LIQ", "MANT"]

    def test_status_label_mapea(self):
        ws = _wb(_build_resumen_excel([_rec(classification="bajar")], "2026-04-17"))["Pricing"]
        assert ws.cell(row=2, column=2).value == "Bajar Precio"

    def test_status_label_fallback_crudo(self):
        # clasificación desconocida -> se usa el valor crudo
        ws = _wb(_build_resumen_excel([_rec(classification="raro")], "2026-04-17"))["Pricing"]
        assert ws.cell(row=2, column=2).value == "raro"

    def test_defaults_record_sin_keys(self):
        # record vacío -> stocks 0, precios '', t7 0, sku '', status ''
        ws = _wb(_build_resumen_excel([{}], "2026-04-17"))["Pricing"]
        row = [ws.cell(row=2, column=j).value for j in range(1, 10)]
        # SKU '', Status '', FBA 0, AWD 0, IZZI 0, Total 0, Precio '', Sugerido '', T7 0
        assert row[0] in ("", None)            # SKU '' (openpyxl puede leer '' como None)
        assert row[1] in ("", None)            # Status ''
        assert row[2] == 0 and row[3] == 0 and row[4] == 0 and row[5] == 0
        assert row[6] in ("", None)            # Precio Actual ''
        assert row[7] in ("", None)            # Precio Sugerido ''
        assert row[8] == 0                     # Ventas T7

    def test_valores_reales(self):
        ws = _wb(_build_resumen_excel(
            [_rec(sku="X", fba_available=5.0, awd_available=2.0, total_stock=7.0,
                  price=30.0, suggestedPrice=28.5, t7=4.0)], "2026-04-17"))["Pricing"]
        row = [ws.cell(row=2, column=j).value for j in range(1, 10)]
        assert row[0] == "X"
        assert row[2] == 5.0 and row[3] == 2.0 and row[5] == 7.0
        assert row[6] == 30.0 and row[7] == 28.5 and row[8] == 4.0

    def test_anchos_seteados(self):
        ws = _wb(_build_resumen_excel([_rec()], "2026-04-17"))["Pricing"]
        assert ws.column_dimensions["A"].width == 26
        assert ws.column_dimensions["I"].width == 11


# =====================================================================
# _build_vista_excel
# =====================================================================
from modules.pages.pricing_dashboard import _build_vista_excel  # noqa: E402


class TestBuildVistaExcel:
    def _df(self):
        return pd.DataFrame(
            {"SKU": ["A", "B"], "Precio": [10.0, 20.0], "Margen": [15.0, float("nan")]}
        )

    def test_bytes_parseables(self):
        data = _build_vista_excel(self._df(), "Liquidar")
        assert isinstance(data, bytes) and len(data) > 0
        assert load_workbook(BytesIO(data))  # abre sin error

    def test_sheetname(self):
        wb = _wb(_build_vista_excel(self._df(), "Liquidar"))
        assert wb.sheetnames == ["Liquidar"]

    def test_sheetname_truncado_31(self):
        nombre = "X" * 40
        wb = _wb(_build_vista_excel(self._df(), nombre))
        assert wb.sheetnames == [nombre[:31]]
        assert len(wb.sheetnames[0]) == 31

    def test_header_en_orden(self):
        ws = _wb(_build_vista_excel(self._df(), "V"))["V"]
        fila1 = [ws.cell(row=1, column=j).value for j in range(1, 4)]
        assert fila1 == ["SKU", "Precio", "Margen"]

    def test_filas_igual_len_mas_header(self):
        ws = _wb(_build_vista_excel(self._df(), "V"))["V"]
        assert ws.max_row == len(self._df()) + 1  # 2 datos + 1 header

    def test_df_vacio_solo_header(self):
        df_vacio = pd.DataFrame(columns=["SKU", "Precio"])
        ws = _wb(_build_vista_excel(df_vacio, "V"))["V"]
        assert ws.max_row == 1  # solo header
        assert [ws.cell(row=1, column=j).value for j in range(1, 3)] == ["SKU", "Precio"]

    def test_nan_float_celda_vacia(self):
        # Margen fila 2 (B) es NaN -> celda None/vacía
        ws = _wb(_build_vista_excel(self._df(), "V"))["V"]
        assert ws.cell(row=3, column=3).value is None  # B.Margen = NaN
