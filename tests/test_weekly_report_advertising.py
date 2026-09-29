"""The Weekly Client Report's Advertising sheet from the account's campaign totals over the report's days."""
import pandas as pd
from openpyxl import load_workbook

from core.amazon_ads.product_provider import NewToBrand
from core.weekly_report.advertising import TOP_CAMPAIGNS, advertising_summary
from modules.pages import weekly_client_report as wcr

COLUMNS = ["product", "campaign_id", "campaign", "portfolio", "spend", "sales", "orders", "clicks", "impressions",
           "sales_clicks", "orders_clicks", "ntb_orders", "ntb_sales"]
NAN = float("nan")
ROWS = [
    ("SP", "1", "Brand exact", "Brand", 120.0, 600.0, 12, 40, 900, 600.0, 12, NAN, NAN),
    ("SP", "2", "Old auto", "", 30.0, 0.0, 0, 12, 400, 0.0, 0, NAN, NAN),
    ("SB", "3", "Video", "Brand", 10.0, 40.0, 2, 8, 1400, 20.0, 1, 1, 20.0),
    ("SD", "4", "Retargeting", "Generic", 40.0, 50.0, 4, 10, 2000, 25.0, 2, 2, 30.0),
]
SOURCE = "Fuente: Luna Kids · US · 03–16 ago 2026 · 14 de 14 días del BR con datos de ads · SP · SB · SD."


def _frame(rows=ROWS) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=COLUMNS)


def test_the_totals_add_every_product_and_leave_acos_unknown_without_sales():
    summary = advertising_summary(_frame())

    assert summary.totals == {"Impressions": 4700.0, "Clicks": 70.0, "Spend": 200.0, "Sales": 690.0, "Orders": 18.0,
                              "CTR": 1.49, "CPC": 2.86, "ACoS": 28.99}
    assert summary.campaign_count == 4
    assert advertising_summary(_frame([ROWS[1]])).totals["ACoS"] is None


def test_the_top_campaigns_come_most_spend_first_and_a_campaign_without_sales_has_no_acos():
    summary = advertising_summary(_frame())

    assert [(campaign["Campaign"], campaign["ACoS"]) for campaign in summary.campaigns] == [
        ("Brand exact", 20.0), ("Retargeting", 80.0), ("Old auto", None), ("Video", 25.0)]
    assert summary.campaigns[0] == {"Campaign": "Brand exact", "Product": "SP", "Impressions": 900, "Clicks": 40,
                                    "CTR": 4.44, "Spend": 120.0, "Sales": 600.0, "ACoS": 20.0, "Orders": 12}


def test_the_alarms_are_the_campaigns_over_60_and_the_ones_that_spent_without_selling():
    summary = advertising_summary(_frame())

    assert [campaign["Campaign"] for campaign in summary.alarms] == ["Retargeting", "Old auto"]


def test_only_the_top_campaigns_are_listed_but_every_campaign_counts():
    rows = [("SP", str(index), f"Campaign {index}", "", float(index), 10.0, 1, 1, 10, 10.0, 1, NAN, NAN)
            for index in range(1, TOP_CAMPAIGNS + 6)]

    summary = advertising_summary(_frame(rows))

    assert len(summary.campaigns) == TOP_CAMPAIGNS and summary.campaign_count == TOP_CAMPAIGNS + 5
    assert summary.campaigns[0]["Campaign"] == f"Campaign {TOP_CAMPAIGNS + 5}"


def test_the_portfolios_add_their_campaigns_and_name_the_ones_without_a_portfolio():
    summary = advertising_summary(_frame())

    assert summary.portfolios == [{"Portfolio": "Brand", "Spend": 130.0, "Sales": 640.0, "ACoS": 20.3},
                                  {"Portfolio": "Generic", "Spend": 40.0, "Sales": 50.0, "ACoS": 80.0},
                                  {"Portfolio": "(Sin Portfolio)", "Spend": 30.0, "Sales": 0.0, "ACoS": None}]


def test_new_to_brand_adds_brands_and_display_over_their_own_orders():
    summary = advertising_summary(_frame())

    assert summary.new_to_brand == NewToBrand(3, 50.0)
    assert summary.new_to_brand_share == 50.0


def test_new_to_brand_is_unknown_when_a_display_day_was_never_measured_or_nothing_credits_it():
    unmeasured = [row if row[0] != "SD" else row[:11] + (NAN, NAN) for row in ROWS]

    assert advertising_summary(_frame(unmeasured)).new_to_brand is None
    assert advertising_summary(_frame(ROWS[:2])).new_to_brand is None


def _advertising_sheet(**kwargs):
    buf = wcr._build_weekly_excel({}, {}, {}, {}, "Luna", kwargs.pop("lang", "es"), None, **kwargs)
    return load_workbook(buf)["\U0001f4e3 Advertising"]


def test_the_sheet_shows_the_account_figures_in_its_currency_with_their_source():
    sheet = _advertising_sheet(advertising=advertising_summary(_frame()), ads_currency="USD", ads_source=SOURCE)

    assert sheet["A1"].value == "Luna — Advertising Overview"
    assert sheet["A2"].value == SOURCE
    assert [cell.value for cell in sheet[4]][:8] == ["4,700", "70", "1.49%", "$2.86", "$200.00", "$690.00", "29.0%",
                                                     "18"]
    assert sheet["A5"].value == ("New-to-brand (SB y SD): 3 órdenes (50.0%) · ventas $50.00  |  Vistas de la página "
                                 "de detalle: — (no se sincronizan)")
    rows = {row[0].value: [cell.value for cell in row] for row in sheet.iter_rows(min_row=7) if row[0].value}
    assert rows["Brand exact"] == ["Brand exact", "SP", 900, 40, 4.44, 120.0, 600.0, 20.0, 12]
    assert rows["Old auto"][7] == "—"
    assert rows["  Old auto — sin ventas | Spend $30.00 | Sales $0.00"][0] is not None
    assert rows["(Sin Portfolio)"][:4] == ["(Sin Portfolio)", 30.0, 0.0, "—"]
    brand_row = next(row for row in sheet.iter_rows(min_row=7) if row[0].value == "Brand exact")
    assert brand_row[5].number_format == '"$"#,##0.00'


def test_new_to_brand_that_is_not_known_reads_as_a_dash():
    sheet = _advertising_sheet(advertising=advertising_summary(_frame(ROWS[:2])), ads_currency="MXN",
                               ads_source=SOURCE, lang="en")

    assert sheet["A5"].value == "New-to-brand (SB and SD): —  |  Detail page views: — (not synced)"
    assert sheet["E4"].value == "MX$150.00"


def test_without_the_account_the_sheet_says_why():
    sheet = _advertising_sheet(ads_note="no se eligió la cuenta de Amazon Ads del BR")

    assert sheet["A1"].value == "⚠️ Sin datos de Amazon Ads: no se eligió la cuenta de Amazon Ads del BR."
