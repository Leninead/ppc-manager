"""PPC Audit Pro over an in-memory PostgREST: one account choice feeds the search terms and the whole structure.

No network: `_open_rest` is replaced before every script run, uploads are faked, and the AI tab runs disabled except
where a test fakes the provider.
"""
import io
import time
from datetime import date

import openpyxl
import pytest
from streamlit.proto.WidgetStates_pb2 import WidgetState
from streamlit.testing.v1 import AppTest
from streamlit.testing.v1.element_tree import ButtonGroup

import ai.client as ai_client
import ai.runtime as ai_runtime
import modules.pages.search_term_source as search_term_source
from core.chat.screen_selection import FROM_AMAZON_ADS, FROM_HAND_UPLOAD, HAND_UPLOAD_NOTE, OLDER_DATA_NOTE
from core.date_labels import date_range_label
from core.ppc_audit.checks import run_audit
from core.ppc_audit.synced_reads import NOT_LISTED_REASON
from modules.pages import audit_source
from modules.pages import ppc_audit as audit_page
from tests.test_ppc_audit_frames import PROFILE_ID, AuditRest, job_row, read, synthetic_bulk_file

PAGE = "🛡️ PPC Audit"
TAB_LABELS = ["📊 KPIs Overview", "🛠️ Auditoría Estructura", "🎯 Performance Segmento", "🔍 Deep Checks",
              "📥 Export", "🎯 Target Graduation", "🤖 Análisis IA"]
ANSWER = {"synthesis": {"situation": "Una keyword no recibe tráfico.", "week_actions": ["Revisar U03"],
                        "mid_term": [], "risks": [], "executive_summary": "1 keyword sin impresiones."},
          "hallazgos": [{"row_id": "U03", "razon": "0 impresiones en una campaña con 1.500.", "veredicto": "INVESTIGAR",
                         "confianza": "media", "advertencia": None}]}


class _Upload:
    def __init__(self, name: str, data: bytes):
        self.name, self._data = name, data

    def getvalue(self) -> bytes:
        return self._data


@pytest.fixture(autouse=True)
def _single_select_button_groups(monkeypatch):
    # AppTest 1.43 serializes button groups as multi-select (st.feedback); a segmented control holds one value.
    def widget_state(group):
        state = WidgetState()
        state.id = group.id
        value = group.value
        selected = value if isinstance(value, list) else [] if value is None else [value]
        state.int_array_value.data[:] = [group.options.index(group.format_func(option)) for option in selected]
        return state
    monkeypatch.setattr(ButtonGroup, "_widget_state", property(widget_state))


@pytest.fixture(autouse=True)
def _empty_ai_registry():
    with ai_runtime._lock:
        ai_runtime._registry.clear()
    yield
    with ai_runtime._lock:
        ai_runtime._registry.clear()


_PAGE_SCRIPT = """
import streamlit as st
st.cache_data.clear()
st.session_state.setdefault("selected_page", "🛡️ PPC Audit")
from modules.pages.ppc_audit import render
render()
"""


def _run(monkeypatch, fake, *, uploads=None, ai_enabled=False) -> AppTest:
    import streamlit
    uploads = uploads or {}
    monkeypatch.setattr(streamlit, "file_uploader", lambda label, *args, key=None, **kwargs: uploads.get(key))
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: fake)
    monkeypatch.setattr("ai.config.AI_ENABLED", ai_enabled)
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.run()
    assert not app.exception, app.exception
    return app


def _text(app: AppTest) -> str:
    parts = [str(element.value) for element in app.markdown] + [str(element.value) for element in app.caption]
    parts += [str(element.value) for element in app.info] + [str(element.value) for element in app.success]
    parts += [str(element.value) for element in app.error]
    return " ".join(parts)


def test_one_account_choice_reads_the_search_terms_and_the_structure_of_the_same_period(monkeypatch):
    fake = AuditRest()

    app = _run(monkeypatch, fake)

    reads = dict(fake.reads)
    window = (reads["search_terms_between"]["p_from"], reads["search_terms_between"]["p_to"])
    for rpc in ("sp_structure_between", "product_campaigns_between", "sb_sd_targets_between",
                "sb_search_terms_between"):
        assert (reads[rpc]["p_from"], reads[rpc]["p_to"]) == window and reads[rpc]["p_profile_id"] == PROFILE_ID
    assert [tab.label for tab in app.tabs] == TAB_LABELS
    text = _text(app)
    assert "Estructura de las campañas" in text and "Listada" in text
    assert "2 campañas SP" in text and "2 keywords" in text and "SB 2" in text
    assert "✅ Datos de Amazon Ads cargados — SP: 2 campañas, 2 keywords, 1 PT" in text


def test_the_quiet_keyword_reaches_target_graduation_with_its_effective_bid(monkeypatch):
    app = _run(monkeypatch, AuditRest())

    graduation = app.tabs[5]
    table = graduation.dataframe[0].value
    assert list(table["Keyword Text"]) == ["kids sleep sack"] and list(table["Bid"]) == [0.6]
    assert audit_page.API_GRADUATION_NOTE in " ".join(str(caption.value) for caption in graduation.caption)


def test_sb_search_terms_and_sd_targets_reach_the_structure_checks(monkeypatch):
    app = _run(monkeypatch, AuditRest())

    captions = " ".join(str(caption.value) for caption in app.tabs[1].caption)
    assert "SB: $3.00 (1 terms)" in captions
    assert "SD: $8.00 (1 targets)" in captions


def test_a_structure_never_listed_says_so_and_offers_the_bulk_file(monkeypatch):
    app = _run(monkeypatch, AuditRest(structure=[], jobs=[job_row("sp_search_terms")]))

    assert not app.tabs
    assert NOT_LISTED_REASON in _text(app) and audit_source.BULK_FALLBACK_HINT in _text(app)
    assert app.button(key="audit_structure_src_upload_manual")
    assert PAGE not in app.session_state["app_chat_selections"]


def test_the_bulk_file_by_hand_audits_its_sheets(monkeypatch):
    uploads = {audit_source.BULK_UPLOADER_KEY: _Upload("bulk.xlsx", synthetic_bulk_file().read_bytes())}
    app = _run(monkeypatch, AuditRest(), uploads=uploads)

    app.button(key="audit_structure_src_upload_manual").click().run()

    assert not app.exception, app.exception
    assert audit_source.BULK_MODE_NOTE in _text(app)
    assert [tab.label for tab in app.tabs] == TAB_LABELS
    assert "✅ Bulk cargado — SP: 6 campañas, 5 keywords, 0 PT" in _text(app)
    selection = app.session_state["app_chat_selections"][PAGE]
    assert selection.source == FROM_HAND_UPLOAD and selection.calls == ()


def test_a_file_that_is_not_a_bulk_file_says_so(monkeypatch):
    uploads = {audit_source.BULK_UPLOADER_KEY: _Upload("notes.xlsx", b"not a workbook")}
    fake = AuditRest()
    fake.select = lambda table, params: [] if table == "ads_profile_sync" else AuditRest().select(table, params)

    app = _run(monkeypatch, fake, uploads=uploads)

    assert audit_source.UNREADABLE_FILE_MESSAGE in _text(app)
    assert not app.tabs


def test_without_connected_accounts_the_page_asks_for_the_bulk_file(monkeypatch):
    fake = AuditRest()
    fake.select = lambda table, params: [] if table == "ads_profile_sync" else AuditRest().select(table, params)

    app = _run(monkeypatch, fake)

    assert search_term_source.NO_CONNECTION_HINT in _text(app)
    assert "Subí el Bulk File" in _text(app)
    assert not app.tabs


def test_the_chat_learns_the_account_the_days_the_brand_terms_and_the_call(monkeypatch):
    app = _run(monkeypatch, AuditRest())

    app.text_input(key=audit_page.BRAND_TERMS_KEY).set_value("Luna, Kids ").run()

    selection = app.session_state["app_chat_selections"][PAGE]
    (call,) = selection.calls
    assert selection.source == FROM_AMAZON_ADS and selection.profile_id == PROFILE_ID
    assert call.name == "ppc_audit"
    assert dict(call.arguments) == {"profile_id": PROFILE_ID, "date_from": selection.window_start.isoformat(),
                                    "date_to": selection.window_end.isoformat(), "brand_terms": ["luna", "kids"]}
    assert dict(selection.values) == {"brand terms": "luna, kids"}


def test_a_bulk_file_selection_tells_the_chat_the_tools_cannot_see_it():
    frames = read(AuditRest()).frames
    source = audit_source.AuditSource(frames=frames, label="bulk.xlsx", signature="s")
    file_source = audit_source.AuditSource(frames=frames.__class__(**{**frames.__dict__, "source": "file"}),
                                           label="bulk.xlsx", signature="s")

    assert audit_page.screen_selection(file_source, ()).notes == (HAND_UPLOAD_NOTE,)
    older = audit_source.AuditSource(frames=frames, label="Luna Kids · US", signature="s", profile_id=PROFILE_ID,
                                     window_start=date(2026, 9, 9), window_end=date(2026, 9, 22), older_data=True)
    assert audit_page.screen_selection(older, ()).notes == (OLDER_DATA_NOTE,)
    assert source.from_amazon_ads and not file_source.from_amazon_ads


def test_the_excel_carries_every_sheet_of_the_audit():
    frames = read(AuditRest()).frames
    result = run_audit(frames, brand_terms=("luna",))
    source = audit_source.AuditSource(frames=frames, label="Luna Kids · US", signature="s", profile_id=PROFILE_ID,
                                      window_start=date(2026, 9, 9), window_end=date(2026, 9, 22))

    workbook = openpyxl.load_workbook(io.BytesIO(audit_page.build_audit_excel(
        result, audit_page.kpi_rows(result, source, "2026-09-23"))))

    assert workbook.sheetnames == ["Resumen KPIs", "Performance Segmento", "Auditoria", "Top Campanas",
                                   "Clasificacion Targets", "Target Graduation"]
    kpis = {row[0]: row[1] for row in workbook["Resumen KPIs"].iter_rows(min_row=2, values_only=True)}
    # SP 45, SB 12 (an archived campaign that spent in the period counts) and SD 8.
    assert kpis["Fuente"] == "Amazon Ads" and kpis["Total PPC Spend"] == "$65.00"
    assert kpis["Período"] == date_range_label(date(2026, 9, 9), date(2026, 9, 22))
    audit_rows = list(workbook["Auditoria"].iter_rows(min_row=2, values_only=True))
    assert audit_rows[2][2] == "SP: 5.00 (1 terms) | SB: 3.00 (1 terms)"


def test_the_ai_analysis_waits_for_the_click_and_reaches_the_chat_with_the_account(monkeypatch):
    calls = []

    def ask(**call):
        calls.append(call)
        return {"structured_output": ANSWER, "session_id": "audit-session"}

    monkeypatch.setattr(ai_client, "ask", ask)
    app = _run(monkeypatch, AuditRest(), ai_enabled=True)

    assert calls == []
    app.button(key="ppc_audit_ai_recalc").click().run()
    for _ in range(30):
        entry = app.session_state["app_chat_modules"].get("ppc_audit")
        if entry is not None and entry.state == "current":
            break
        time.sleep(0.2)
        app.run()

    assert not app.exception, app.exception
    assert len(calls) == 1
    shared = app.session_state["app_chat_modules"]["ppc_audit"].analysis
    titles = [document["title"] for document in shared.documents]
    assert titles[0] == "PPC Audit Pro · Luna Kids · US · Parámetros"
    assert titles[-1] == "PPC Audit Pro · Luna Kids · US · Lectura de la IA"
    assert "U03 · graduacion · kids sleep sack → INVESTIGAR · media" in shared.documents[-1]["content"]
    assert (shared.profile_id, shared.country_code) == (PROFILE_ID, "US")
    assert shared.annotate("Revisar U03") == "Revisar U03 (kids sleep sack)"
