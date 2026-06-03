"""Characterization tests para M30 F3.2 — parsers + lookups.

Lockean el comportamiento verbatim del port (HTML -> Python). Datos sintéticos
in-memory: CSV como bytes, XLSX construidos con openpyxl en el propio test.

Fuente del comportamiento esperado: pricing-dashboard.html
    buildFeeLookup (L941), buildCOGSLookup (L970), buildMaestroLookup (L986),
    parseFile / loadFile (L611), parseCSV (L704).
"""

import sys
from io import BytesIO
from pathlib import Path

# --- Path hardening (landmine pre-existente, ajeno a F3.2) ---------------------
# El runner de pytest recolecta `scripts/_scratch_M27/test_b5b_extract.py` (un
# script tracked, sin funciones de test) cuyo código top-level hace
# `sys.path.insert(0, <otro repo>)` + `from modules.pages.flat_file_migrator`.
# Eso deja `modules` / `modules.pages` cacheados en sys.modules apuntando a OTRO
# repo (que no tiene este módulo) y rompe el import de abajo bajo `pytest -q`.
# Forzamos que la raíz de ESTE worktree gane y purgamos las entradas cross-repo.
_ROOT = str(Path(__file__).resolve().parents[1])
sys.path.insert(0, _ROOT)
for _n in [n for n in list(sys.modules) if n == "modules" or n.startswith("modules.")]:
    _loc = str(getattr(sys.modules[_n], "__path__", "") or getattr(sys.modules[_n], "__file__", ""))
    if _ROOT not in _loc:
        del sys.modules[_n]
# ------------------------------------------------------------------------------

import pandas as pd
import pytest
from openpyxl import Workbook

from modules.pages.pricing_dashboard import (
    _parse_fba,
    _parse_fee,
    _parse_awd,
    _parse_pl,
    _parse_maestro,
    _parse_izzi,
    _build_cogs_lookup,
    _build_fee_lookup,
    _build_maestro_lookup,
)

# en-dash / em-dash para los casos "dash unicode"
EN_DASH = "–"
EM_DASH = "—"


# ---------------------------------------------------------------------
# Helpers de construcción de datos sintéticos
# ---------------------------------------------------------------------
def _csv_bytes(text: str, bom: bool = False) -> bytes:
    prefix = "﻿" if bom else ""
    return (prefix + text).encode("utf-8")


def _xlsx_bytes(columns, rows, sheet_name="Sheet1") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(list(columns))
    for row in rows:
        ws.append([row.get(c) for c in columns])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _xlsx_multi(sheets) -> bytes:
    """Construye un workbook multi-hoja. sheets = [(nombre, [fila, ...]), ...]
    donde cada fila es una lista posicional de celdas. La primera tupla es la
    hoja en índice 0."""
    wb = Workbook()
    first_ws = wb.active
    for idx, (name, rows) in enumerate(sheets):
        ws = first_ws if idx == 0 else wb.create_sheet()
        ws.title = name
        for row in rows:
            ws.append(list(row))
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


# =====================================================================
# B3 — Parsers CSV (_parse_fba / _parse_fee / _parse_awd)
# =====================================================================
class TestParseFba:
    def test_completo(self):
        data = _csv_bytes(
            "sku,available,asin,product-name\nABC,10,B00X,Widget\nDEF,0,B00Y,Gadget\n"
        )
        df = _parse_fba(data)
        assert list(df.columns) == ["sku", "available", "asin", "product-name"]
        assert len(df) == 2
        assert df.iloc[0]["sku"] == "ABC"

    def test_minimo_solo_core(self):
        df = _parse_fba(_csv_bytes("sku,available\nABC,10\n"))
        assert list(df.columns) == ["sku", "available"]
        assert len(df) == 1

    def test_dash_unicode_verbatim(self):
        name = f"Body Cream {EN_DASH} 1.76oz"
        df = _parse_fba(_csv_bytes(f"sku,product-name\nABC,{name}\n"))
        assert df.iloc[0]["product-name"] == name  # en-dash sin normalizar

    def test_col_core_ausente_no_crashea(self):
        df = _parse_fba(_csv_bytes("sku,asin\nABC,B00X\n"))
        assert "available" not in df.columns
        assert "sku" in df.columns

    def test_bom_se_elimina_del_header(self):
        df = _parse_fba(_csv_bytes("sku,available\nABC,10\n", bom=True))
        assert list(df.columns) == ["sku", "available"]  # no '﻿sku'

    def test_separador_semicolon_autodetect(self):
        df = _parse_fba(_csv_bytes("sku;available;asin\nABC;10;B00X\n"))
        assert list(df.columns) == ["sku", "available", "asin"]
        assert df.iloc[0]["sku"] == "ABC"


class TestParseFee:
    def test_completo(self):
        data = _csv_bytes(
            "MSKU,FBA fulfillment fees per unit,Referral fee per unit,Units sold\n"
            "ABC,2.0,1.0,10\nDEF,3.0,1.5,5\n"
        )
        df = _parse_fee(data)
        assert "MSKU" in df.columns
        assert len(df) == 2

    def test_minimo_solo_core(self):
        df = _parse_fee(_csv_bytes("MSKU,Units sold\nABC,10\n"))
        assert list(df.columns) == ["MSKU", "Units sold"]

    def test_dash_unicode_verbatim(self):
        sku = f"ABC{EM_DASH}1"
        df = _parse_fee(_csv_bytes(f"MSKU,Units sold\n{sku},10\n"))
        assert df.iloc[0]["MSKU"] == sku

    def test_col_core_ausente_no_crashea(self):
        df = _parse_fee(_csv_bytes("MSKU,otra\nABC,x\n"))
        assert "Units sold" not in df.columns

    def test_separador_semicolon_autodetect(self):
        df = _parse_fee(_csv_bytes("MSKU;Units sold\nABC;10\n"))
        assert list(df.columns) == ["MSKU", "Units sold"]
        assert df.iloc[0]["MSKU"] == "ABC"


class TestParseAwd:
    def test_completo(self):
        data = _csv_bytes("SKU,Available in AWD (units)\nABC,50\nDEF,0\n")
        df = _parse_awd(data)
        assert list(df.columns) == ["SKU", "Available in AWD (units)"]
        assert len(df) == 2

    def test_minimo_solo_core(self):
        df = _parse_awd(_csv_bytes("SKU\nABC\n"))
        assert list(df.columns) == ["SKU"]

    def test_dash_unicode_verbatim(self):
        sku = f"ABC{EN_DASH}9"
        df = _parse_awd(_csv_bytes(f"SKU,Available in AWD (units)\n{sku},50\n"))
        assert df.iloc[0]["SKU"] == sku

    def test_col_core_ausente_no_crashea(self):
        df = _parse_awd(_csv_bytes("SKU,nota\nABC,hola\n"))
        assert "Available in AWD (units)" not in df.columns

    def test_separador_semicolon_autodetect(self):
        df = _parse_awd(_csv_bytes("SKU;Available in AWD (units)\nABC;50\n"))
        assert list(df.columns) == ["SKU", "Available in AWD (units)"]
        assert df.iloc[0]["SKU"] == "ABC"


# =====================================================================
# B3 — Parsers XLSX (_parse_pl / _parse_maestro / _parse_izzi)
# =====================================================================
class TestParsePl:
    def test_completo(self):
        cols = ["SKU", "Total Cost Per Unit \n2026-01 (USD)"]
        data = _xlsx_bytes(cols, [{"SKU": "ABC", cols[1]: 5.0}])
        df = _parse_pl(data)
        assert list(df.columns) == cols
        assert df.iloc[0]["SKU"] == "ABC"

    def test_minimo_solo_core(self):
        data = _xlsx_bytes(["SKU"], [{"SKU": "ABC"}])
        df = _parse_pl(data)
        assert list(df.columns) == ["SKU"]

    def test_dash_unicode_verbatim(self):
        sku = f"ABC{EN_DASH}1"
        data = _xlsx_bytes(["SKU"], [{"SKU": sku}])
        df = _parse_pl(data)
        assert df.iloc[0]["SKU"] == sku

    def test_col_core_ausente_no_crashea(self):
        data = _xlsx_bytes(["otra"], [{"otra": 1}])
        df = _parse_pl(data)
        assert "SKU" not in df.columns


class TestParseMaestro:
    def test_completo(self):
        cols = ["SKU", "Modelo", "Talla", "Temporada", "Categoria", "Subcategoria"]
        data = _xlsx_bytes(
            cols,
            [{"SKU": "ABC", "Modelo": "M1", "Talla": "L",
              "Temporada": "Invierno", "Categoria": "C", "Subcategoria": "S"}],
        )
        df = _parse_maestro(data)
        assert list(df.columns) == cols
        assert df.iloc[0]["Temporada"] == "Invierno"

    def test_minimo_solo_core(self):
        data = _xlsx_bytes(["SKU"], [{"SKU": "ABC"}])
        df = _parse_maestro(data)
        assert list(df.columns) == ["SKU"]

    def test_dash_unicode_verbatim(self):
        val = f"Verano{EM_DASH}Otono"
        data = _xlsx_bytes(["SKU", "Temporada"], [{"SKU": "ABC", "Temporada": val}])
        df = _parse_maestro(data)
        assert df.iloc[0]["Temporada"] == val

    def test_col_core_ausente_no_crashea(self):
        data = _xlsx_bytes(["otra"], [{"otra": 1}])
        df = _parse_maestro(data)
        assert "SKU" not in df.columns


class TestParseIzzi:
    # header=None -> columnas posicionales (enteros) y NINGUNA fila consumida como
    # header. La hoja se selecciona por nombre 'Inventario 2526' (fallback a la 1ra).

    def test_selecciona_hoja_inventario_2526_no_indice_0(self):
        # hoja basura en índice 0; la data real vive en 'Inventario 2526'
        data = _xlsx_multi([
            ("Basura", [["x", "y"], ["1", "2"], ["3", "4"]]),
            ("Inventario 2526",
             [["SKU", "c1", "c2", "c3", "c4", "stock"],
              ["meta", None, None, None, None, None],
              ["ABC", None, None, None, None, 12]]),
        ])
        df = _parse_izzi(data)
        # devuelve la data de 'Inventario 2526', NO de 'Basura' (índice 0)
        assert df.iloc[2, 0] == "ABC"   # offset 2 + col 0 = SKU (lo consume F3.3)
        assert df.iloc[2, 5] == 12      # col 5 = stock
        assert df.iloc[0, 0] == "SKU"   # fila 0 preservada (header NO consumido)

    def test_fallback_primera_hoja_si_no_existe_inventario(self):
        data = _xlsx_multi([
            ("HojaUnica", [["SKU", "stock"], ["meta", None], ["ABC", 7]]),
        ])
        df = _parse_izzi(data)
        assert df.iloc[2, 0] == "ABC"   # cae a la primera hoja

    def test_header_none_preserva_todas_las_filas(self):
        data = _xlsx_multi([
            ("Inventario 2526", [["SKU", "stock"], ["ABC", 5], ["DEF", 9]]),
        ])
        df = _parse_izzi(data)
        assert len(df) == 3            # 3 filas: ninguna se pierde como header
        assert df.iloc[0, 0] == "SKU"  # la 1ra fila sigue siendo data
        assert df.iloc[1, 0] == "ABC"

    def test_dash_unicode_verbatim(self):
        sku = f"ABC{EN_DASH}7"
        data = _xlsx_multi([("Inventario 2526", [["SKU"], ["meta"], [sku]])])
        df = _parse_izzi(data)
        assert df.iloc[2, 0] == sku


# =====================================================================
# B4 — _build_cogs_lookup
# =====================================================================
COGS_JAN = "Total Cost Per Unit \nJan (USD)"
COGS_FEB = "Total Cost Per Unit \nFeb (USD)"


class TestBuildCogsLookup:
    def test_mapping_ok_y_month(self):
        col = "Total Cost Per Unit \n2026-01 (USD)"
        df = pd.DataFrame([{"SKU": "ABC", col: 5.0}], columns=["SKU", col])
        lk = _build_cogs_lookup(df)
        assert lk["ABC"] == 5.0
        assert lk["__month__ABC"] == "2026-01"

    def test_itera_de_ultima_a_primera_col(self):
        df = pd.DataFrame(
            [{"SKU": "ABC", COGS_JAN: 3.0, COGS_FEB: 7.0}],
            columns=["SKU", COGS_JAN, COGS_FEB],
        )
        lk = _build_cogs_lookup(df)
        assert lk["ABC"] == 7.0  # gana la última con valor > 0
        assert lk["__month__ABC"] == "Feb"

    def test_ultima_col_vacia_usa_anterior(self):
        df = pd.DataFrame(
            [{"SKU": "ABC", COGS_JAN: 3.0, COGS_FEB: 0}],
            columns=["SKU", COGS_JAN, COGS_FEB],
        )
        lk = _build_cogs_lookup(df)
        assert lk["ABC"] == 3.0
        assert lk["__month__ABC"] == "Jan"

    def test_sku_duplicado_last_valid_wins(self):
        df = pd.DataFrame(
            [{"SKU": "ABC", COGS_JAN: 5.0}, {"SKU": "ABC", COGS_JAN: 8.0}],
            columns=["SKU", COGS_JAN],
        )
        assert _build_cogs_lookup(df)["ABC"] == 8.0

    def test_sku_duplicado_segunda_sin_cogs_preserva_primera(self):
        df = pd.DataFrame(
            [{"SKU": "ABC", COGS_JAN: 5.0}, {"SKU": "ABC", COGS_JAN: 0}],
            columns=["SKU", COGS_JAN],
        )
        assert _build_cogs_lookup(df)["ABC"] == 5.0

    def test_sku_ausente_se_descarta(self):
        df = pd.DataFrame(
            [{"SKU": "", COGS_JAN: 5.0}, {"SKU": None, COGS_JAN: 6.0}],
            columns=["SKU", COGS_JAN],
        )
        assert _build_cogs_lookup(df) == {}

    def test_cast_a_str_confirmado(self):
        df = pd.DataFrame([{"SKU": 12345, COGS_JAN: 5.0}], columns=["SKU", COGS_JAN])
        lk = _build_cogs_lookup(df)
        assert "12345" in lk
        assert lk["12345"] == 5.0


# =====================================================================
# B4 — _build_fee_lookup
# =====================================================================
FF = "FBA fulfillment fees per unit"
RF = "Referral fee per unit"
PPC = "Sponsored Products charge per unit"
US = "Units sold"


class TestBuildFeeLookup:
    def test_mapping_ok(self):
        df = pd.DataFrame([{"MSKU": "ABC", FF: 2.0, RF: 1.0, PPC: 0.5, US: 10}])
        lk = _build_fee_lookup(df)
        assert lk["ABC"] == {
            "fulfillment_fee": 2.0,
            "referral_fee": 1.0,
            "ppc_fee": 0.5,
            "units_sold_week": 10.0,
        }

    def test_sku_duplicado_agrega_promedia_fees_suma_units(self):
        df = pd.DataFrame(
            [
                {"MSKU": "ABC", FF: 2.0, RF: 1.0, PPC: 0.0, US: 10},
                {"MSKU": "ABC", FF: 4.0, RF: 3.0, PPC: 0.0, US: 5},
            ]
        )
        lk = _build_fee_lookup(df)
        assert lk["ABC"]["fulfillment_fee"] == 3.0  # (2+4)/2
        assert lk["ABC"]["referral_fee"] == 2.0
        assert lk["ABC"]["ppc_fee"] is None  # ningún ppc > 0
        assert lk["ABC"]["units_sold_week"] == 15.0

    def test_fees_cero_dan_none(self):
        df = pd.DataFrame([{"MSKU": "ABC", FF: 0, RF: 0, PPC: 0, US: 4}])
        lk = _build_fee_lookup(df)
        assert lk["ABC"]["fulfillment_fee"] is None
        assert lk["ABC"]["units_sold_week"] == 4.0

    def test_sku_ausente_se_descarta(self):
        df = pd.DataFrame([{"MSKU": "", FF: 2.0, RF: 1.0, PPC: 0.5, US: 1}])
        assert _build_fee_lookup(df) == {}

    def test_cast_a_str_confirmado(self):
        df = pd.DataFrame([{"MSKU": 999, FF: 2.0, RF: 1.0, PPC: 0.5, US: 3}])
        assert "999" in _build_fee_lookup(df)


# =====================================================================
# B4 — _build_maestro_lookup
# =====================================================================
class TestBuildMaestroLookup:
    def test_mapping_ok_key_lowercase_value_fila_completa(self):
        df = pd.DataFrame(
            [{"SKU": "ABC123", "Modelo": "M1", "Temporada": "Invierno"}]
        )
        lk = _build_maestro_lookup(df)
        assert "abc123" in lk
        assert lk["abc123"]["Modelo"] == "M1"
        assert lk["abc123"]["Temporada"] == "Invierno"  # verbatim

    def test_sku_duplicado_last_wins(self):
        df = pd.DataFrame([{"SKU": "ABC", "Modelo": "M1"}, {"SKU": "ABC", "Modelo": "M2"}])
        assert _build_maestro_lookup(df)["abc"]["Modelo"] == "M2"

    def test_sku_ausente_se_descarta(self):
        df = pd.DataFrame([{"SKU": None, "Modelo": "M1"}, {"SKU": "", "Modelo": "M2"}])
        assert _build_maestro_lookup(df) == {}

    def test_cast_a_str_y_dash_unicode_verbatim(self):
        sku = f"ABC{EN_DASH}123"
        df = pd.DataFrame([{"SKU": sku, "Modelo": "M1"}])
        lk = _build_maestro_lookup(df)
        assert sku.lower() in lk  # en-dash preservado, key en minúscula


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
