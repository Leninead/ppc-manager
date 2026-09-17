"""Tests B3b — export Excel del Dashboard Global (M39).

`_build_agency_excel` es PURA: recibe el dict de `_build_agency_dashboard` y
devuelve bytes. Los tests abren esos bytes con openpyxl, sin tocar disco.

Diferencia deliberada con el export HTML: acá ACOS y TACOS van SIN relleno
(sigue el mockup de Dirección). Hay un test explícito de eso, porque es lo único
que separa a los dos exports.

Determinismo: los bytes NO se comparan crudos. openpyxl estampa
`dcterms:created` / `dcterms:modified` en `docProps/core.xml`, así que dos
corridas del mismo input dan archivos distintos byte a byte (medido). Se compara
el contenido: valores de celda y colores.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from core import agency_dashboard_export as exp
from core import agency_dashboard_format as fmt


_REPO_ROOT = str(Path(__file__).resolve().parents[1])
_GEN = date(2026, 9, 16)


# ─────────────────────────────────────────────────────────────────────────────
# Datos de prueba
# ─────────────────────────────────────────────────────────────────────────────

def _cell(actual=None, acco=None, partial=None) -> dict:
    keys = ("revenue", "ventasPPC", "spend", "acos", "tacos")
    return {
        "actual": actual,
        "forecast": None,
        "accomplishment": {k: (acco or {}).get(k) for k in keys},
        "partial": partial,
    }


_ACTUAL = {"revenue": 9000.0, "ventasPPC": 3000.0, "spend": 900.0,
           "acos": 30.0, "tacos": 10.0}
_ACCO_OK = {"revenue": 95.0, "ventasPPC": 92.0, "spend": 91.0,
            "acos": 1.0, "tacos": 2.5}


def _account(name="Alpha Brand", client_id="alpha", has_baseline=True,
             months=None, currency="USD") -> dict:
    return {
        "client_id": client_id,
        "name": name,
        "currency": currency,
        "has_baseline": has_baseline,
        "baseline_name": "Plan 2026" if has_baseline else None,
        "baseline_created_at": "2026-09-01" if has_baseline else None,
        "months": months or {
            "2026-08": _cell(actual=_ACTUAL, acco=_ACCO_OK, partial=False),
        },
    }


def _dash(accounts, periods=("2026-08",)) -> dict:
    return {"periods": list(periods), "accounts": accounts}


def _ws(data: dict, generated_at=_GEN):
    wb = load_workbook(BytesIO(exp._build_agency_excel(data, generated_at=generated_at)))
    return wb, wb[wb.sheetnames[0]]


def _celdas(ws) -> list:
    """Todas las celdas con valor, como lista de (coordenada, valor)."""
    return [(c.coordinate, c.value) for row in ws.iter_rows()
            for c in row if c.value not in (None, "")]


def _textos(ws) -> str:
    return "\n".join(str(v) for _, v in _celdas(ws))


def _celda(ws, metrica: str, sufijo: str):
    """Celda de una metrica en la columna cuyo encabezado termina en `sufijo`
    (ej. "Acco" o "Actual"). Busca el header de la tabla y su fila."""
    fila_header = next(c.row for row in ws.iter_rows() for c in row
                       if str(c.value or "") == "Métrica")
    col = next(c.column for c in ws[fila_header]
               if str(c.value or "").endswith(sufijo))
    fila = next(c.row for row in ws.iter_rows() for c in row
                if c.column == 1 and str(c.value or "") == metrica)
    return ws.cell(row=fila, column=col)


def _buscar(ws, texto: str):
    """Primera celda cuyo valor contiene `texto`."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and texto in str(c.value):
                return c
    return None


def _relleno(celda) -> str:
    """Color de relleno en RRGGBB, o '' si la celda no tiene fondo."""
    f = celda.fill
    if f is None or f.fill_type != "solid":
        return ""
    rgb = getattr(f.start_color, "rgb", None) or ""
    return str(rgb)[-6:].upper()


# ─────────────────────────────────────────────────────────────────────────────
# Workbook
# ─────────────────────────────────────────────────────────────────────────────

def test_sin_cuentas_devuelve_workbook_valido():
    """openpyxl explota si no queda ninguna hoja visible."""
    wb, ws = _ws(_dash([]))
    assert len(wb.sheetnames) >= 1
    assert ws.sheet_state == "visible"
    assert "sin cuentas" in _textos(ws).lower()


def test_portada_naranja_y_negra_del_estandar():
    _wb, ws = _ws(_dash([_account()]))
    assert "Dashboard Global" in str(ws["A1"].value)
    assert _relleno(ws["A1"]) == "E84000"
    assert "Capybaras" in str(ws["A2"].value)
    assert "2026-09-16" in str(ws["A2"].value)
    assert _relleno(ws["A2"]) == "1F1F1F"


def test_dos_cuentas_aparecen_como_bandas():
    wb, ws = _ws(_dash([_account(name="Alpha Brand", client_id="alpha"),
                        _account(name="Beta Labs", client_id="beta")]))
    texto = _textos(ws)
    assert "Alpha Brand" in texto and "Beta Labs" in texto
    # Apiladas: la banda de Beta va debajo de la de Alpha.
    assert _buscar(ws, "Alpha Brand").row < _buscar(ws, "Beta Labs").row


def test_las_cinco_metricas_en_el_orden_del_mockup():
    _wb, ws = _ws(_dash([_account()]))
    filas = [str(v) for _, v in _celdas(ws)]
    orden = [f for f in filas if f in ("Revenue", "Ad Sales", "Ad Spend", "ACOS", "TACOS")]
    assert orden == ["Revenue", "Ad Sales", "Ad Spend", "ACOS", "TACOS"]


def test_columnas_actual_y_acco_por_mes():
    _wb, ws = _ws(_dash([_account()], periods=("2026-08", "2026-09")))
    texto = _textos(ws)
    assert texto.count("Actual") >= 2
    assert texto.count("Acco") >= 2


def test_valores_son_numeros_operables_no_texto():
    """Ajuste B: la planilla se usa para sumar y graficar, asi que cada celda
    lleva el NUMERO crudo y el formato lo pone Excel."""
    _wb, ws = _ws(_dash([_account()]))

    rev = _celda(ws, "Revenue", "Actual")
    assert rev.value == 9000                       # no "USD 9,000"
    assert isinstance(rev.value, (int, float))
    assert "#,##0" in rev.number_format and "USD" in rev.number_format


def test_acos_actual_es_numero_con_formato_de_porcentaje():
    _wb, ws = _ws(_dash([_account()]))
    acos = _celda(ws, "ACOS", "Actual")
    assert acos.value == 30                        # no "30.0%"
    assert "%" in acos.number_format


def test_acco_de_porcentaje_se_guarda_como_fraccion():
    """95.0 de cumplimiento se escribe 0.95 con formato de porcentaje: asi
    Excel lo trata como porcentaje real y promediarlo entre cuentas da bien."""
    _wb, ws = _ws(_dash([_account()]))
    celda = _celda(ws, "Revenue", "Acco")
    assert celda.value == 0.95
    assert celda.number_format == "0.0%"
    assert _relleno(celda) == fmt._VERDE.lstrip("#").upper()   # sigue el semaforo


def test_acco_de_acos_es_delta_en_puntos_con_signo():
    _wb, ws = _ws(_dash([_account()]))
    celda = _celda(ws, "ACOS", "Acco")
    assert celda.value == 1.0                      # +1.0 punto, NO 0.01
    assert "+" in celda.number_format
    assert _relleno(celda) == ""                   # sin color, como pidio Direccion


def test_moneda_de_la_cuenta_en_el_formato():
    acc = _account(currency="MXN")
    _wb, ws = _ws(_dash([acc]))
    assert "MXN" in _celda(ws, "Revenue", "Actual").number_format


# ─────────────────────────────────────────────────────────────────────────────
# Semáforo: SOLO en las métricas de %
# ─────────────────────────────────────────────────────────────────────────────

def test_semaforo_en_revenue_verde_y_rojo():
    dash = _dash([_account(months={
        "2026-08": _cell(actual=_ACTUAL, acco={"revenue": 95.0}, partial=False),
        "2026-09": _cell(actual=_ACTUAL, acco={"revenue": 78.0}, partial=False),
    })], periods=("2026-08", "2026-09"))
    _wb, ws = _ws(dash)
    filas = [c for row in ws.iter_rows() for c in row if c.value == 0.95]
    rojas = [c for row in ws.iter_rows() for c in row if c.value == 0.78]
    assert filas and _relleno(filas[0]) == fmt._VERDE.lstrip("#").upper()
    assert rojas and _relleno(rojas[0]) == fmt._ROJO.lstrip("#").upper()


def test_semaforo_amarillo_en_la_franja_del_medio():
    dash = _dash([_account(months={
        "2026-08": _cell(actual=_ACTUAL, acco={"ventasPPC": 87.0}, partial=False)})])
    _wb, ws = _ws(dash)
    celda = _celda(ws, "Ad Sales", "Acco")
    assert celda.value == 0.87
    assert _relleno(celda) == fmt._AMARILLO.lstrip("#").upper()


def test_acos_y_tacos_van_sin_relleno():
    """Lo ÚNICO que separa este export del HTML. Si Guille pide color, se
    revierte sacando acos/tacos del filtro por _PCT_METRICS."""
    dash = _dash([_account(months={
        "2026-08": _cell(actual=_ACTUAL, acco={"acos": 5.0, "tacos": -3.0},
                         partial=False)})])
    _wb, ws = _ws(dash)
    acos = _celda(ws, "ACOS", "Acco")
    tacos = _celda(ws, "TACOS", "Acco")
    assert (acos.value, tacos.value) == (5.0, -3.0)
    assert _relleno(acos) == "" and _relleno(tacos) == ""


def test_celda_acco_sin_dato_no_lleva_relleno():
    dash = _dash([_account(has_baseline=False, months={
        "2026-08": _cell(actual=_ACTUAL, partial=False)})])
    _wb, ws = _ws(dash)
    rellenos = {_relleno(c) for row in ws.iter_rows() for c in row}
    for color in (fmt._VERDE, fmt._AMARILLO, fmt._ROJO):
        assert color.lstrip("#").upper() not in rellenos


# ─────────────────────────────────────────────────────────────────────────────
# Sin plan y mes en curso
# ─────────────────────────────────────────────────────────────────────────────

def test_cuenta_sin_baseline_banda_sin_plan_y_acco_vacio():
    acc = _account(name="Zeta Corp", client_id="zeta", has_baseline=False,
                   months={"2026-08": _cell(actual=_ACTUAL, partial=False)})
    _wb, ws = _ws(_dash([acc]))
    texto = _textos(ws)
    assert "sin plan" in texto.lower()
    assert _celda(ws, "Revenue", "Actual").value == 9000   # el real sí se muestra
    # Las celdas de la columna Acco quedan vacías (no hay plan contra qué medir).
    col_acco = next(c.column for row in ws.iter_rows() for c in row
                    if str(c.value or "").endswith(" Acco"))
    fila_header = next(c.row for row in ws.iter_rows() for c in row
                       if str(c.value or "") == "Métrica")
    acco_vals = [ws.cell(row=fila_header + i, column=col_acco).value
                 for i in range(1, 6)]
    assert all(v in (None, "") for v in acco_vals), acco_vals


def test_cuenta_con_baseline_muestra_el_plan():
    _wb, ws = _ws(_dash([_account()]))
    texto = _textos(ws)
    assert "Plan 2026" in texto and "2026-09-01" in texto


def test_mes_en_curso_marcado_y_nota_al_pie():
    con = _dash([_account(months={"2026-09": _cell(actual=_ACTUAL, partial=True)})],
                periods=("2026-09",))
    sin = _dash([_account(months={"2026-09": _cell(actual=_ACTUAL, partial=False)})],
                periods=("2026-09",))
    _wb_c, ws_c = _ws(con)
    _wb_s, ws_s = _ws(sin)
    assert "Sep*" in _textos(ws_c) and "MtD" in _textos(ws_c)
    assert "Sep*" not in _textos(ws_s) and "MtD" not in _textos(ws_s)


# ─────────────────────────────────────────────────────────────────────────────
# Presentación
# ─────────────────────────────────────────────────────────────────────────────

def test_autofit_deja_anchos_en_el_rango_del_estandar():
    _wb, ws = _ws(_dash([_account(name="Una cuenta con nombre largo de verdad")],
                        periods=("2026-08", "2026-09")))
    anchos = [d.width for d in ws.column_dimensions.values() if d.width]
    assert anchos, "no se seteó ningún ancho"
    assert all(8 <= w <= 40 for w in anchos)


def test_rango_de_meses_en_la_portada():
    _wb, ws = _ws(_dash([_account()], periods=("2026-08", "2026-09")))
    texto = _textos(ws)
    assert "2026-08" in texto and "2026-09" in texto


def test_determinismo_por_contenido_no_por_bytes():
    """openpyxl estampa dcterms:created/modified, así que los bytes NO son
    comparables entre corridas (medido). Se compara el contenido."""
    dash = _dash([_account(name="Alpha"), _account(name="Beta", client_id="beta")])
    _wb_a, ws_a = _ws(dash)
    _wb_b, ws_b = _ws(dash)
    assert _celdas(ws_a) == _celdas(ws_b)
    rell_a = [(c.coordinate, _relleno(c)) for row in ws_a.iter_rows() for c in row]
    rell_b = [(c.coordinate, _relleno(c)) for row in ws_b.iter_rows() for c in row]
    assert rell_a == rell_b


def test_generated_at_por_default_no_explota():
    assert exp._build_agency_excel(_dash([_account()]))[:2] == b"PK"


# ─────────────────────────────────────────────────────────────────────────────
# Wiring: los DOS botones
# ─────────────────────────────────────────────────────────────────────────────

_APP = """
import sys
sys.path.insert(0, r"__REPO_ROOT__")
import streamlit as st
from modules.pages import agency_dashboard_page as page

page._build_agency_dashboard = lambda periods, clients=None: __DASH__
page.render(username="quien", role="admin")
"""


def _run(dash_repr: str) -> AppTest:
    at = AppTest.from_string(
        _APP.replace("__REPO_ROOT__", _REPO_ROOT).replace("__DASH__", dash_repr)
    )
    at.run()
    return at


def test_con_cuentas_estan_los_dos_botones():
    at = _run(repr(_dash([_account()])))
    assert not at.exception
    labels = [b.proto.label for b in at.get("download_button")]
    assert len(labels) == 2
    assert any("HTML" in l for l in labels)
    assert any("Excel" in l for l in labels)


def test_sin_cuentas_no_hay_ningun_boton():
    at = _run(repr(_dash([])))
    assert not at.exception
    assert len(at.get("download_button")) == 0


def test_leyenda_del_desvio_en_la_planilla():
    """Ajuste A: Dirección pidió que quede explícito que el Acco de ACOS/TACOS
    es el desvío en puntos contra el plan, no un %."""
    _wb, ws = _ws(_dash([_account()]))
    texto = _textos(ws)
    assert "desvío en puntos" in texto
    assert "real − plan" in texto
