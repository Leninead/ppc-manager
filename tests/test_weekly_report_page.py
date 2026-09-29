"""Weekly Client Report over an in-memory PostgREST: the Advertising sheet from the chosen account or from a Campaign
CSV uploaded by hand, and the AI tab.

No network: `_open_rest` is replaced before every script run, the uploads are fakes that serve CSVs built here, and the
AI tab runs disabled except where a test fakes the provider.
"""
import csv
import hashlib
import io
import time
from datetime import date, datetime, timedelta, timezone

import pytest
import requests
from openpyxl import load_workbook
from streamlit.proto.Common_pb2 import FileURLs
from streamlit.proto.WidgetStates_pb2 import WidgetState
from streamlit.runtime.uploaded_file_manager import UploadedFile, UploadedFileRec
from streamlit.testing.v1 import AppTest
from streamlit.testing.v1.element_tree import ButtonGroup

import ai.client as ai_client
import ai.runtime as ai_runtime
import modules.pages.search_term_source as search_term_source
from core import ai_tab
from modules.pages import ad_account_block
from modules.pages import weekly_client_report as weekly_page

ACCOUNT = "Luna Kids"
CLIENT = "Luna Kids MX"
# Monday 3 to Sunday 16 August 2026: the prior week sells 400 a day and this one 500 (6,300 in all).
DAYS = [date(2026, 8, 3) + timedelta(days=offset) for offset in range(14)]
THIS_WEEK_START = date(2026, 8, 10)
SYNCED_AT = datetime(2026, 9, 23, 9, 30, tzinfo=timezone.utc)
WINDOW_HEADER = ["ad_product", "campaign_id", "campaign_name", "portfolio_id", "portfolio_name", "impressions", "clicks",
                 "cost", "purchases_7d", "sales_7d", "purchases_14d", "sales_14d", "purchases", "sales",
                 "purchases_clicks", "sales_clicks", "currency_code", "new_to_brand_purchases", "new_to_brand_sales"]
ANSWER = {"synthesis": {"situation": "Las ventas subieron 25%.", "week_actions": ["Mantener el budget"],
                        "mid_term": [], "risks": [], "executive_summary": "3.500 esta semana."},
          "lecturas": [{"tema": "VENTAS", "razon": "3.500 contra 2.800.", "veredicto": "OK", "advertencia": None},
                       {"tema": "PUBLICIDAD", "razon": "ACoS de 20%.", "veredicto": "VIGILAR",
                        "advertencia": "Los últimos días todavía pueden crecer."}],
          "resumen_cliente": {"situacion": "Fue una buena semana: vendimos 3.500, un 25% más.",
                              "highlights": ["Las ventas subieron 25%"], "atencion": [],
                              "proximos_pasos": ["Vamos a escalar la campaña de marca"]}}


def _daily_csv() -> bytes:
    lines = ["Date,Ordered Product Sales,Units Ordered,Sessions - Total"]
    for day in DAYS:
        sales = 500 if day >= THIS_WEEK_START else 400
        lines.append(f'{day.month}/{day.day}/26,"${sales:,.2f}",{sales // 20},{sales // 4}')
    return ("\n".join(lines) + "\n").encode("utf-8")


def _by_child_csv() -> bytes:
    return ("(Parent) ASIN,(Child) ASIN,Title,Sessions - Total,Units Ordered,Ordered Product Sales\n"
            'B0PARENT001,B0TEST0001,Producto Uno,700,140,"$2,800.00"\n'
            'B0PARENT002,B0TEST0002,Producto Dos,875,175,"$3,500.00"\n').encode("utf-8")


def _campaign_csv() -> bytes:
    """The account's campaigns over the report's days, as Campaign Manager exports them: 320 spent, 1,280 sold."""
    return ("Campaign name,Campaign ID,State,Type,Portfolio name,Impressions,Clicks,Total cost,Purchases,Sales\n"
            'Luna SP,111,ENABLED,Sponsored Products,Kids,2800,140,$280.00,28,"$1,120.00"\n'
            "Luna SB,222,PAUSED,Sponsored Brands,,800,40,$40.00,8,$160.00\n").encode("utf-8")


def _campaign_csv_without_portfolios() -> bytes:
    return ("Campaign name,Type,Impressions,Clicks,Total cost,Purchases,Sales\n"
            'Luna SP,Sponsored Products,2800,140,$280.00,28,"$1,120.00"\n').encode("utf-8")


def _campaign_csv_without_counts() -> bytes:
    """Campaign Manager lets the AM leave columns out: no impressions, clicks or purchases."""
    return ("Campaign name,Type,Total cost,Sales\n"
            'Luna SP,Sponsored Products,$280.00,"$1,120.00"\n').encode("utf-8")


def _upload(name: str, data: bytes) -> UploadedFile:
    """What Streamlit hands the page, a fresh file on every run: its cached parsers hash it by its content."""
    return UploadedFile(UploadedFileRec(file_id=name, name=name, type="text/csv", data=data), FileURLs())


def _profile_row() -> dict:
    return {"profile_id": "111", "account_id": 1, "cliente": ACCOUNT, "account_name": "Luna Kids US",
            "country_code": "US", "currency_code": "USD", "account_type": "seller", "timezone": "America/New_York",
            "status": "active", "data_from": "2026-07-20", "data_through": "2026-09-22", "refreshed_on": "2026-09-23",
            "last_success_at": SYNCED_AT.isoformat(), "last_error": ""}


def _campaign_job(window_end: date) -> dict:
    return {"id": 7, "integration_slug": "amazon_ads", "job_kind": "sp_campaigns", "trigger": "scheduled_daily",
            "external_account_id": "111", "status": "completed", "window_start": (window_end - timedelta(days=6))
            .isoformat(), "window_end": window_end.isoformat(), "local_day": (window_end + timedelta(days=1))
            .isoformat(), "finished_at": SYNCED_AT.isoformat(), "created_at": SYNCED_AT.isoformat()}


def _campaign_days(sp_sales: float = 40.0) -> list:
    """SP spends 10 and sells `sp_sales` every day; SB spends 5 and sells 20 on weekends."""
    rows = []
    for offset in range(65):
        day = date(2026, 7, 20) + timedelta(days=offset)
        rows.append(_totals_row(day, "SP", cost=10.0, sales_7d=sp_sales))
        if day.weekday() >= 5:
            rows.append(_totals_row(day, "SB", cost=5.0, sales=20.0))
    return rows


def _totals_row(day: date, product: str, *, cost: float, sales_7d: float = 0.0, sales: float = 0.0) -> dict:
    return {"report_date": day.isoformat(), "ad_product": product, "impressions": 100, "clicks": 5, "cost": cost,
            "purchases_7d": 1, "sales_7d": sales_7d, "purchases_14d": 1, "sales_14d": sales_7d, "purchases": 1,
            "sales": sales, "purchases_clicks": 1, "sales_clicks": sales, "currency_code": "USD",
            "campaign_names": None, "new_to_brand_purchases": None, "new_to_brand_sales": None}


def _window_row(product: str, campaign_id: str, name: str, *, cost: float, sales: float, orders: int,
                ntb: tuple = ("", "")) -> dict:
    row = dict.fromkeys(WINDOW_HEADER, 0)
    row.update(ad_product=product, campaign_id=campaign_id, campaign_name=name, portfolio_id="", portfolio_name="",
               impressions=900, clicks=40, cost=cost, currency_code="USD", new_to_brand_purchases=ntb[0],
               new_to_brand_sales=ntb[1])
    if product == "SP":
        row.update(purchases_7d=orders, sales_7d=sales, purchases_14d=orders, sales_14d=sales)
    else:
        row.update(purchases=orders, sales=sales, purchases_clicks=orders, sales_clicks=sales)
    return row


WINDOW_ROWS = [_window_row("SP", "1", "Luna - B0TEST0001 - SP - KW - EXACT - Brand", cost=120.0, sales=600.0, orders=12),
               _window_row("SP", "2", "Old auto", cost=30.0, sales=0.0, orders=0),
               _window_row("SB", "3", "Video Luna", cost=10.0, sales=40.0, orders=2, ntb=(1, 20.0))]


def _csv(header, rows) -> bytes:
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=header)
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode("utf-8")


class _FakeRest:
    """The profile table, the campaign sync jobs, and the campaign totals by day and by campaign."""

    def __init__(self, *, profiles=True, campaign_days=None, fail_totals=False):
        self._profiles, self._fail_totals = profiles, fail_totals
        self._campaign_days = _campaign_days() if campaign_days is None else campaign_days
        self.reads: list[tuple[str, dict]] = []

    def select(self, table, params):
        if table == "ads_profile_sync":
            return [_profile_row()] if self._profiles else []
        if table == "integration_sync_jobs":
            assert params["job_kind"] == "eq.sp_campaigns"
            return [_campaign_job(date(2026, 9, 22))]
        raise AssertionError(f"unexpected select on {table}")

    def rpc(self, name, args, *, timeout_s=8):
        self.reads.append((name, args))
        if self._fail_totals:
            raise requests.ConnectionError("gateway down")
        assert name == "campaign_daily_totals" and args["p_profile_id"] == "111"
        return [row for row in self._campaign_days if args["p_from"] <= row["report_date"] <= args["p_to"]]

    def rpc_csv(self, name, args, *, timeout_s=8):
        self.reads.append((name, args))
        assert name == "campaign_window_totals" and args["p_profile_id"] == "111"
        return _csv(WINDOW_HEADER, WINDOW_ROWS)


@pytest.fixture(autouse=True)
def _single_select_button_groups(monkeypatch):
    # AppTest 1.43 serializes button groups as multi-select (st.feedback); a segmented control holds one value.
    def widget_state(group):
        state = WidgetState()
        state.id = group.id
        value = group.value
        selected = value if isinstance(value, list) else [] if value is None else [value]
        state.int_array_value.data[:] = [group.options.index(group.format_func(option)) for option in selected]
        return state
    monkeypatch.setattr(ButtonGroup, "_widget_state", property(widget_state))


@pytest.fixture(autouse=True)
def _empty_ai_registry():
    with ai_runtime._lock:
        ai_runtime._registry.clear()
    yield
    with ai_runtime._lock:
        ai_runtime._registry.clear()


_PAGE_SCRIPT = """
import streamlit as st
from core.chat import app_chat
st.cache_data.clear()
if st.session_state.get("shared_before"):
    app_chat.report_failed("weekly_report")
from modules.pages.weekly_client_report import render
render()
"""


def _page(monkeypatch, rest=None, *, daily=True, by_child=False, ai_enabled=False,
          campaign_file: tuple[str, bytes] | None = None, downloads: dict | None = None) -> AppTest:
    """`downloads`, when given, collects each download button's bytes by its key."""
    import streamlit
    uploads = {}
    if daily:
        uploads["br_daily"] = ("BusinessReport-by-date.csv", _daily_csv())
    if by_child:
        uploads["br_child"] = ("BusinessReport-by-child.csv", _by_child_csv())
    if campaign_file is not None:
        uploads["wcr_src_file"] = campaign_file
    monkeypatch.setattr(streamlit, "file_uploader",
                        lambda label, *args, key=None, **kwargs: _upload(*uploads[key]) if key in uploads else None)
    if downloads is not None:
        def record_download(label, data=None, *, key=None, **kwargs):
            downloads[key] = data
            return False
        monkeypatch.setattr(streamlit, "download_button", record_download)
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: rest)
    monkeypatch.setattr("ai.config.AI_ENABLED", ai_enabled)
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.run()
    assert not app.exception, app.exception
    return app


def _choose_account(app: AppTest) -> AppTest:
    app.selectbox(key="wcr_src_account").set_value(ACCOUNT).run()
    assert not app.exception, app.exception
    return app


def _upload_offered(app: AppTest) -> bool:
    return any(button.key == "wcr_src_upload_manual" for button in app.button)


def _successes(app: AppTest) -> str:
    return " ".join(str(success.value) for success in app.success)


def _advertising_sheet(downloads: dict):
    return load_workbook(io.BytesIO(downloads["weekly_dl"]))["\U0001f4e3 Advertising"]


def _sheet_rows(sheet) -> dict:
    return {row[0].value: [cell.value for cell in row] for row in sheet.iter_rows(min_row=7) if row[0].value}


def _run_ai(app: AppTest) -> AppTest:
    app.button(key="weekly_report_ai_recalc").click().run()
    for _ in range(30):
        entry = app.session_state["app_chat_modules"].get("weekly_report")
        if entry is not None and entry.state == "current":
            break
        time.sleep(0.2)
        app.run()
    assert not app.exception, app.exception
    return app


def _text(app: AppTest) -> str:
    parts = [str(element.value) for element in app.markdown] + [str(element.value) for element in app.caption]
    parts += [str(element.value) for element in app.error] + [str(element.value) for element in app.info]
    parts += [str(element.value) for element in app.warning] + [str(element.value) for element in app.success]
    return " ".join(parts)


def test_without_reports_nothing_renders_and_the_analysis_is_withdrawn(monkeypatch):
    import streamlit
    monkeypatch.setattr(streamlit, "file_uploader", lambda *args, **kwargs: None)
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: _FakeRest())
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.session_state["shared_before"] = True

    app.run()

    assert not app.exception
    assert not app.tabs
    assert "weekly_report" not in app.session_state["app_chat_modules"]


def test_the_old_ai_button_is_gone_and_the_campaign_csv_comes_back_only_as_the_manual_upload(monkeypatch):
    app = _page(monkeypatch, _FakeRest())

    assert [tab.label for tab in app.tabs] == ["📊 Reporte", "🤖 Análisis IA"]
    assert "btn_weekly_ai" not in [button.key for button in app.button]
    assert _upload_offered(app)
    assert ad_account_block.FILE_HINT not in _text(app)


def test_before_choosing_an_account_nothing_is_read(monkeypatch):
    fake = _FakeRest()

    app = _page(monkeypatch, fake)

    assert fake.reads == []
    assert weekly_page._ADS_TEXTS.choose_account in _text(app)
    assert "campañas" not in " ".join(str(success.value) for success in app.success)


def test_the_chosen_account_is_read_over_the_daily_report_days(monkeypatch):
    fake = _FakeRest()

    app = _choose_account(_page(monkeypatch, fake))

    window = {"p_profile_id": "111", "p_from": "2026-08-03", "p_to": "2026-08-16"}
    assert fake.reads == [("campaign_daily_totals", {**window, "p_campaign": None}),
                          ("campaign_window_totals", window)]
    assert "3 campañas · 2,700 imps ✓" in " ".join(str(success.value) for success in app.success)
    assert "14 de 14 días del BR con datos de ads" in _text(app)


def test_without_the_daily_report_there_are_no_days_to_read(monkeypatch):
    fake = _FakeRest()

    app = _choose_account(_page(monkeypatch, fake, daily=False, by_child=True))

    assert fake.reads == []
    assert f"Sin hoja Advertising: {weekly_page.MISSING_NO_DAILY_REPORT}." in _text(app)


def test_ads_that_cannot_be_read_say_so_and_the_report_still_renders(monkeypatch):
    app = _choose_account(_page(monkeypatch, _FakeRest(fail_totals=True)))

    assert any(weekly_page._ADS_TEXTS.unreadable in str(error.value) for error in app.error)
    assert [tab.label for tab in app.tabs] == ["📊 Reporte", "🤖 Análisis IA"]
    assert not any("Traceback" in str(code.value) for code in app.code)


def test_more_ad_sales_than_report_sales_warns_that_the_account_may_not_be_the_reports(monkeypatch):
    app = _choose_account(_page(monkeypatch, _FakeRest(campaign_days=_campaign_days(sp_sales=500.0))))

    # 14 days x 500 of SP plus 4 weekend days x 20 of SB, against the report's 6,300.
    assert [str(warning.value) for warning in app.warning] == [
        "Las ventas de ads (\\$7,080.00) superan las del BR (\\$6,300.00) en los mismos días: revisá que la cuenta y "
        "el país sean los del BR."]


def test_the_ai_analysis_waits_for_the_click_and_hands_the_client_summary_and_the_chat_its_account(monkeypatch):
    calls = []

    def ask(**call):
        calls.append(call)
        return {"structured_output": ANSWER, "session_id": "weekly-session"}

    monkeypatch.setattr(ai_client, "ask", ask)
    app = _page(monkeypatch, _FakeRest(), ai_enabled=True)
    app.text_input(key="weekly_client").set_value(CLIENT).run()
    _choose_account(app)

    assert calls == []
    app.button(key="weekly_report_ai_recalc").click().run()
    for _ in range(30):
        entry = app.session_state["app_chat_modules"].get("weekly_report")
        if entry is not None and entry.state == "current":
            break
        time.sleep(0.2)
        app.run()

    assert not app.exception, app.exception
    assert len(calls) == 1
    shared = app.session_state["app_chat_modules"]["weekly_report"].analysis
    titles = [document["title"] for document in shared.documents]
    assert titles[0] == "Weekly Client Report · Luna Kids · US · Parámetros"
    assert titles[-1] == "Weekly Client Report · Luna Kids · US · Lectura de la IA"
    assert "📊 RESUMEN SEMANAL — Luna Kids MX" in shared.documents[-1]["content"]
    assert "Cliente: Luna Kids MX" in shared.documents[0]["content"]
    assert (shared.profile_id, shared.country_code) == ("111", "US")
    assert any(str(code.value).startswith("📊 RESUMEN SEMANAL — Luna Kids MX") for code in app.code)


def test_the_ai_rows_show_each_topic_with_its_verdict_and_skip_the_ones_the_module_did_not_compute():
    labels = ai_tab.ai_labels("es", weekly_page._AI_TEXTS["es"])
    records = [{"tema": "VENTAS", "item": "Ventas", "metrics": ["Sales TW MX$3,500.00", "PW MX$2,800.00", "+25.0%"]}]

    rows = weekly_page.weekly_ai_rows(ANSWER["lecturas"], records, labels)

    assert rows == [{"item": "Ventas", "metrics": ["Sales TW MX$3,500.00", "PW MX$2,800.00", "+25.0%"],
                     "badges": ["OK"], "warning": "", "reasoning": "3.500 contra 2.800."}]


def test_without_connected_accounts_the_block_asks_for_the_campaign_csv_and_the_report_still_renders(monkeypatch):
    downloads = {}

    app = _page(monkeypatch, _FakeRest(profiles=False), downloads=downloads)

    assert f"{ad_account_block.FILE_HINT} {weekly_page._ADS_TEXTS.no_accounts}" in _text(app)
    assert [box.key for box in app.selectbox] == ["weekly_stock_client"]
    assert [tab.label for tab in app.tabs] == ["📊 Reporte", "🤖 Análisis IA"]
    assert _advertising_sheet(downloads)["A1"].value == (
        f"⚠️ Sin datos de Amazon Ads: {ad_account_block.MISSING_NO_ACCOUNTS}.")


def test_without_connected_accounts_a_campaign_csv_fills_the_advertising_sheet_without_reading_amazon_ads(monkeypatch):
    fake, downloads = _FakeRest(profiles=False), {}

    app = _page(monkeypatch, fake, campaign_file=("campaigns.csv", _campaign_csv()), downloads=downloads)

    assert fake.reads == []
    assert "2 campañas · 3,600 imps ✓" in _successes(app)
    assert "Se compara con los 14 días del BR (3 – 16 ago 2026)" in _text(app)
    assert not app.warning
    sheet = _advertising_sheet(downloads)
    assert sheet["A2"].value == (
        "Fuente: Campaign CSV subido a mano (campaigns.csv) · SP · SB. El archivo no dice qué días cubre ni con qué "
        "atribución se exportó: tiene que estar exportado con los días del BR diario (03–16 ago 2026). Tampoco dice "
        "la moneda.")
    # The paused SB campaign still spent in the range: 280 + 40 spent and 1,120 + 160 sold.
    assert [cell.value for cell in sheet[4]][4:7] == ["$320.00", "$1,280.00", "25.0%"]
    assert sheet["A5"].value == ("New-to-brand (SB y SD): —  |  Vistas de la página de detalle: — (no se leen del "
                                 "Campaign CSV)")
    rows = _sheet_rows(sheet)
    assert rows["Luna SP"][:2] == ["Luna SP", "SP"]
    assert rows["Kids"][:4] == ["Kids", 280.0, 1120.0, 25.0]
    assert rows["(Sin Portfolio)"][:4] == ["(Sin Portfolio)", 40.0, 160.0, 25.0]


def test_a_campaign_csv_without_portfolios_leaves_the_portfolios_out_of_the_sheet(monkeypatch):
    downloads = {}

    _page(monkeypatch, _FakeRest(profiles=False), campaign_file=("campaigns.csv", _campaign_csv_without_portfolios()),
          downloads=downloads)

    rows = _sheet_rows(_advertising_sheet(downloads))
    assert "Luna SP" in rows
    assert "PORTFOLIOS" not in rows and "(Sin Portfolio)" not in rows


def test_a_campaign_csv_alone_fills_the_advertising_sheet_without_claiming_days(monkeypatch):
    downloads = {}

    app = _page(monkeypatch, _FakeRest(profiles=False), daily=False, ai_enabled=True,
                campaign_file=("campaigns.csv", _campaign_csv()), downloads=downloads)

    assert [tab.label for tab in app.tabs] == ["📊 Reporte", "🤖 Análisis IA"]
    assert "2 campañas · 3,600 imps ✓" in _successes(app)
    assert weekly_page.MISSING_NO_DAILY_REPORT not in _text(app)
    assert weekly_page._AI_TEXTS["es"]["needs_daily"] in _text(app)
    assert _advertising_sheet(downloads)["A2"].value == (
        "Fuente: Campaign CSV subido a mano (campaigns.csv) · SP · SB. El archivo no dice qué días cubre ni con qué "
        "atribución se exportó: sus cifras son las del rango con que se exportó. Tampoco dice la moneda.")


def test_the_advertising_section_names_both_of_its_sources(monkeypatch):
    app = _page(monkeypatch, _FakeRest(profiles=False))

    assert "#### 5️⃣ Publicidad — cuenta de Amazon Ads o Campaign CSV" in [md.value for md in app.markdown]
    app.radio(key="wlang").set_value("English").run()
    assert "#### 5️⃣ Advertising — Amazon Ads account or Campaign CSV" in [md.value for md in app.markdown]


def test_a_campaign_csv_without_the_counts_shows_them_unknown_instead_of_zero(monkeypatch):
    downloads = {}

    app = _page(monkeypatch, _FakeRest(profiles=False), campaign_file=("campaigns.csv", _campaign_csv_without_counts()),
                downloads=downloads)

    assert "1 campañas ✓" in _successes(app) and "imps" not in _successes(app)
    sheet = _advertising_sheet(downloads)
    assert [cell.value for cell in sheet[4]][:8] == ["—", "—", "—", "—", "$280.00", "$1,120.00", "25.0%", "—"]
    assert _sheet_rows(sheet)["Luna SP"] == ["Luna SP", "SP", "—", "—", "—", 280.0, 1120.0, 25.0, "—"]


def test_a_campaign_csv_alone_leaves_the_weekly_verdict_and_the_product_days_out_of_the_report(monkeypatch):
    downloads = {}

    _page(monkeypatch, _FakeRest(profiles=False), daily=False, campaign_file=("campaigns.csv", _campaign_csv()),
          downloads=downloads)

    workbook = load_workbook(io.BytesIO(downloads["weekly_dl"]))
    text = " ".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
                    if cell.value is not None)
    assert "Semana estable" not in text and "CONCLUSIÓN" not in text
    assert "?d" not in text and "Sin comparación semanal por producto. Los montos" in text


def test_a_campaign_csv_that_cannot_be_read_says_how_to_export_it_and_the_report_still_renders(monkeypatch):
    downloads = {}

    app = _page(monkeypatch, _FakeRest(profiles=False), campaign_file=("campaigns.xlsx", b"not a workbook"),
                downloads=downloads)

    assert any("No se pudo leer «campaigns.xlsx»" in str(error.value) and "Campaign Manager → Campaigns → Export"
               in str(error.value) for error in app.error)
    assert [tab.label for tab in app.tabs] == ["📊 Reporte", "🤖 Análisis IA"]
    assert not any("Traceback" in str(code.value) for code in app.code)
    assert _advertising_sheet(downloads)["A1"].value == (
        f"⚠️ Sin datos de Amazon Ads: {ad_account_block.MISSING_UNREADABLE_FILE}.")


@pytest.mark.parametrize("rest, choose, daily", [
    (_FakeRest(), False, True),
    (_FakeRest(), True, True),
    (_FakeRest(fail_totals=True), True, True),
    (_FakeRest(), True, False),
], ids=["not-chosen", "synced", "ads-unreadable", "no-daily-report"])
def test_every_state_of_the_account_card_offers_the_manual_upload(monkeypatch, rest, choose, daily):
    app = _page(monkeypatch, rest, daily=daily, by_child=True)
    if choose:
        _choose_account(app)

    assert _upload_offered(app)


def test_the_manual_upload_reads_the_campaign_csv_instead_of_the_account_and_going_back_restores_it(monkeypatch):
    fake, downloads = _FakeRest(), {}
    app = _choose_account(_page(monkeypatch, fake, campaign_file=("campaigns.csv", _campaign_csv()),
                                downloads=downloads))
    assert "14 de 14 días del BR con datos de ads" in _text(app)
    fake.reads.clear()

    app.button(key="wcr_src_upload_manual").click().run()

    assert not app.exception, app.exception
    assert fake.reads == []
    assert search_term_source.MANUAL_MODE_NOTE in _text(app)
    assert "2 campañas · 3,600 imps ✓" in _successes(app)
    assert _advertising_sheet(downloads)["A2"].value.startswith("Fuente: Campaign CSV subido a mano (campaigns.csv)")
    assert not _upload_offered(app) and not any(box.key == "wcr_src_account" for box in app.selectbox)

    app.button(key="wcr_src_back_to_api").click().run()

    assert not app.exception, app.exception
    assert app.selectbox(key="wcr_src_account").value == ACCOUNT
    assert "14 de 14 días del BR con datos de ads" in _text(app)
    assert _upload_offered(app)


def test_the_manual_mode_without_a_file_says_it_is_missing(monkeypatch):
    downloads = {}
    app = _page(monkeypatch, _FakeRest(), downloads=downloads)

    app.button(key="wcr_src_upload_manual").click().run()

    assert not app.exception, app.exception
    assert _advertising_sheet(downloads)["A1"].value == (
        f"⚠️ Sin datos de Amazon Ads: {ad_account_block.MISSING_NO_FILE}.")


def test_the_ai_analysis_of_a_campaign_csv_says_where_the_advertising_comes_from_and_signs_the_file(monkeypatch):
    monkeypatch.setattr(ai_client, "ask", lambda **call: {"structured_output": ANSWER, "session_id": "file-session"})
    content = _campaign_csv()
    app = _page(monkeypatch, _FakeRest(profiles=False), ai_enabled=True, campaign_file=("campaigns.csv", content))
    app.text_input(key="weekly_client").set_value(CLIENT).run()

    _run_ai(app)

    shared = app.session_state["app_chat_modules"]["weekly_report"].analysis
    assert shared.documents[0]["title"] == "Weekly Client Report · Luna Kids MX · Parámetros"
    params = shared.documents[0]["content"]
    assert "Cuenta de Amazon Ads del Business Report: ninguna: el AM subió el Campaign CSV a mano" in params
    assert "Publicidad de la cuenta: la del Campaign CSV «campaigns.csv»" in params
    assert "- Publicidad de la cuenta, origen: un Campaign CSV subido a mano" in params
    # 320 of spend over the 6,300 the whole daily report sold.
    assert "- TACoS de la cuenta (%): 5.1" in params
    assert "Atribución" not in params and "Publicidad de la cuenta, días" not in params
    assert (shared.profile_id, shared.country_code) == ("", "")
    assert app.session_state["weekly_report_ai_file_sig"].endswith(hashlib.sha256(content).hexdigest()[:16])
    # The file does not say its currency: its spend gets no assumed one, the report's sales keep the report's.
    assert "Spend &#36;320.00" in _text(app) and "Sales TW MX&#36;3,500.00" in _text(app)
