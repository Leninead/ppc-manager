"""PPC Forecast over an in-memory PostgREST: the Business Report's projection, and the ads of the chosen account or of
a Campaign CSV uploaded by hand.

No network: `_open_rest` is replaced before every script run, the uploads are fakes that serve CSVs built here, and
the AI tab runs disabled except where a test fakes the provider.
"""
import hashlib
import io
import time
from datetime import date, datetime, timedelta, timezone

import pandas as pd
import pytest
import requests
from openpyxl import load_workbook
from streamlit.proto.WidgetStates_pb2 import WidgetState
from streamlit.testing.v1 import AppTest
from streamlit.testing.v1.element_tree import ButtonGroup

import ai.client as ai_client
import ai.runtime as ai_runtime
import modules.pages.search_term_source as search_term_source
from core import ai_tab
from modules.pages import ad_account_block
from modules.pages import ppc_forecast as forecast_page

ACCOUNT = "Luna Kids"
# Four whole weeks, Monday 3 to Sunday 30 August 2026: weekdays sell 100 and Saturdays and Sundays 60, nothing else.
FIRST_DAY = date(2026, 8, 3)
LAST_DAY = date(2026, 8, 30)
SYNCED_AT = datetime(2026, 9, 23, 9, 30, tzinfo=timezone.utc)
ANSWER = {"synthesis": {"situation": "Las ventas se repiten semana a semana.", "week_actions": ["Mantener el budget"],
                        "mid_term": [], "risks": [], "executive_summary": "1.240 en 14 días."},
          "lecturas": [{"tema": "PROYECCION", "razon": "28 días idénticos.", "confianza": "alta", "advertencia": None},
                       {"tema": "PRESUPUESTO", "razon": "TACoS de 12,9%.", "confianza": "media",
                        "advertencia": "El TACoS puede subir con más spend."}]}


def _repeated_week(first: date = FIRST_DAY, last: date = LAST_DAY, weekday: float = 100.0,
                   weekend: float = 60.0) -> list[tuple[date, float]]:
    days = [first + timedelta(days=offset) for offset in range((last - first).days + 1)]
    return [(day, weekend if day.weekday() >= 5 else weekday) for day in days]


def _business_report_csv(days: list[tuple[date, float]]) -> bytes:
    frame = pd.DataFrame({
        "Date": [day.isoformat() for day, _ in days],
        "Ordered Product Sales": [f"${sales:,.2f}" for _, sales in days],
        "Units Ordered": [int(sales // 10) for _, sales in days],
        "Sessions - Total": [int(sales) for _, sales in days],
    })
    return frame.to_csv(index=False).encode("utf-8")


class _Upload:
    def __init__(self, name: str, data: bytes):
        self.name, self._data = name, data

    def getvalue(self) -> bytes:
        return self._data

    def read(self) -> bytes:
        return self._data


def _campaign_csv(sp_sales: str = "$1,120.00") -> bytes:
    """The account's campaigns over the report's days, as Campaign Manager exports them: 320 spent, 1,280 sold."""
    return ("Campaign name,Campaign ID,State,Type,Portfolio name,Impressions,Clicks,Total cost,Purchases,Sales\n"
            f'Luna SP,111,ENABLED,Sponsored Products,Kids,2800,140,$280.00,28,"{sp_sales}"\n'
            "Luna SB,222,PAUSED,Sponsored Brands,,800,40,$40.00,8,$160.00\n").encode("utf-8")


def _profile_row(status: str = "active") -> dict:
    return {"profile_id": "111", "account_id": 1, "cliente": ACCOUNT, "account_name": "Luna Kids US",
            "country_code": "US", "currency_code": "USD", "account_type": "seller", "timezone": "America/New_York",
            "status": status, "data_from": "2026-07-20", "data_through": "2026-09-22", "refreshed_on": "2026-09-23",
            "last_success_at": SYNCED_AT.isoformat(), "last_error": ""}


def _campaign_job(window_end: date, status: str = "completed") -> dict:
    return {"id": 7, "integration_slug": "amazon_ads", "job_kind": "sp_campaigns", "trigger": "scheduled_daily",
            "external_account_id": "111", "status": status, "window_start": (window_end - timedelta(days=6))
            .isoformat(), "window_end": window_end.isoformat(), "local_day": (window_end + timedelta(days=1))
            .isoformat(), "finished_at": SYNCED_AT.isoformat(), "created_at": SYNCED_AT.isoformat()}


def _campaign_days(first: date = date(2026, 7, 20), last: date = date(2026, 9, 22), *, sp_sales: float = 40.0) -> list:
    """SP spends 10 and sells `sp_sales` every day; SB spends 5 and sells 20 on Saturdays and Sundays."""
    rows = []
    for offset in range((last - first).days + 1):
        day = first + timedelta(days=offset)
        rows.append(_totals_row(day, "SP", cost=10.0, sales_7d=sp_sales))
        if day.weekday() >= 5:
            rows.append(_totals_row(day, "SB", cost=5.0, sales=20.0))
    return rows


def _totals_row(day: date, product: str, *, cost: float, sales_7d: float = 0.0, sales: float = 0.0) -> dict:
    return {"report_date": day.isoformat(), "ad_product": product, "impressions": 100, "clicks": 5, "cost": cost,
            "purchases_7d": 1, "sales_7d": sales_7d, "purchases_14d": 1, "sales_14d": sales_7d, "purchases": 1,
            "sales": sales, "purchases_clicks": 1, "sales_clicks": sales, "currency_code": "USD",
            "campaign_names": None, "new_to_brand_purchases": None, "new_to_brand_sales": None}


class _FakeRest:
    """The profile table, the campaign sync jobs and the daily campaign totals.

    `latest_status` is the status of a newer campaign request that has not completed (a first load, for one).
    """

    def __init__(self, *, profiles=True, synced_through: date | None = date(2026, 9, 22), campaign_days=None,
                 fail_totals=False, latest_status: str = "", profile_status: str = "active", fail_sync=False):
        self._profiles, self._synced_through, self._fail_totals = profiles, synced_through, fail_totals
        self._fail_sync = fail_sync
        self._latest_status, self._profile_status = latest_status, profile_status
        self._campaign_days = _campaign_days() if campaign_days is None else campaign_days
        self.reads: list[tuple[str, dict]] = []

    def select(self, table, params):
        if table == "ads_profile_sync":
            return [_profile_row(self._profile_status)] if self._profiles else []
        if table == "integration_sync_jobs":
            assert params["job_kind"] == "eq.sp_campaigns"
            if self._fail_sync:
                raise requests.ConnectionError("gateway down")
            if self._latest_status and params.get("status") != "eq.completed":
                return [_campaign_job(date(2026, 9, 22), self._latest_status)]
            return [_campaign_job(self._synced_through)] if self._synced_through else []
        raise AssertionError(f"unexpected select on {table}")

    def rpc(self, name, args, *, timeout_s=8):
        self.reads.append((name, args))
        if self._fail_totals:
            raise requests.ConnectionError("gateway down")
        assert name == "campaign_daily_totals" and args["p_profile_id"] == "111"
        return [row for row in self._campaign_days if args["p_from"] <= row["report_date"] <= args["p_to"]]


@pytest.fixture(autouse=True)
def _single_select_button_groups(monkeypatch):
    # AppTest 1.43 serializes button groups as multi-select (st.feedback); a segmented control holds one value.
    def widget_state(group):
        state = WidgetState()
        state.id = group.id
        value = group.value
        selected = value if isinstance(value, list) else [] if value is None else [value]
        state.int_array_value.data[:] = [group.options.index(group.format_func(option)) for option in selected]
        return state
    monkeypatch.setattr(ButtonGroup, "_widget_state", property(widget_state))


@pytest.fixture(autouse=True)
def _empty_ai_registry():
    with ai_runtime._lock:
        ai_runtime._registry.clear()
    yield
    with ai_runtime._lock:
        ai_runtime._registry.clear()


_PAGE_SCRIPT = """
import streamlit as st
from core.chat import app_chat
st.cache_data.clear()
if st.session_state.get("shared_before"):
    app_chat.report_failed("ppc_forecast")
from modules.pages.ppc_forecast import render
render()
"""


def _page(monkeypatch, rest=None, *, days=None, upload=True, ai_enabled=False,
          campaign_file: _Upload | None = None) -> AppTest:
    import streamlit
    report = _Upload("BusinessReport-by-date.csv", _business_report_csv(days or _repeated_week()))
    uploads = {"forecast_br": report} if upload else {}
    if campaign_file is not None:
        uploads["forecast_src_file"] = campaign_file
    monkeypatch.setattr(streamlit, "file_uploader", lambda label, *args, key=None, **kwargs: uploads.get(key))
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: rest)
    monkeypatch.setattr("ai.config.AI_ENABLED", ai_enabled)
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.run()
    assert not app.exception, app.exception
    return app


def _generate(app: AppTest) -> AppTest:
    app.button(key="forecast_run").click().run()
    assert not app.exception, app.exception
    return app


def _choose_account(app: AppTest) -> AppTest:
    app.selectbox(key="forecast_src_account").set_value(ACCOUNT).run()
    assert not app.exception, app.exception
    return app


def _upload_offered(app: AppTest) -> bool:
    return any(button.key == "forecast_src_upload_manual" for button in app.button)


def _run_ai(app: AppTest) -> AppTest:
    app.button(key="ppc_forecast_ai_recalc").click().run()
    for _ in range(30):
        entry = app.session_state["app_chat_modules"].get("ppc_forecast")
        if entry is not None and entry.state == "current":
            break
        time.sleep(0.2)
        app.run()
    assert not app.exception, app.exception
    return app


def _metric(app: AppTest, label: str) -> str:
    return next(str(metric.value) for metric in app.metric if metric.label == label)


def _text(app: AppTest) -> str:
    parts = [str(element.value) for element in app.markdown] + [str(element.value) for element in app.caption]
    parts += [str(element.value) for element in app.error] + [str(element.value) for element in app.info]
    parts += [str(element.value) for element in app.warning]
    return " ".join(parts)


def test_a_history_that_repeats_the_same_week_is_projected_as_that_same_week(monkeypatch):
    # The next 14 days are 10 weekdays and 4 weekend days: 10 x 100 + 4 x 60 = 1,240.
    app = _generate(_page(monkeypatch))

    assert _metric(app, "Ventas proyectadas (14d)") == "$1,240.00"


def test_without_the_business_report_the_page_shows_the_empty_state_and_withdraws_the_analysis(monkeypatch):
    import streamlit
    monkeypatch.setattr(streamlit, "file_uploader", lambda *args, **kwargs: None)
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: _FakeRest())
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.session_state["shared_before"] = True

    app.run()

    assert not app.exception
    assert not app.tabs
    assert "Sube el BR Diario para comenzar." in _text(app)
    assert "ppc_forecast" not in app.session_state["app_chat_modules"]


def test_the_forecast_waits_for_the_button_and_stays_while_other_choices_rerun_the_page(monkeypatch):
    fake = _FakeRest()
    app = _page(monkeypatch, fake)

    assert not app.tabs
    assert "Haz clic en 'Generar Forecast' para continuar." in _text(app)

    _choose_account(_generate(app))

    assert [tab.label for tab in app.tabs] == ["📈 Forecast", "🤖 Análisis IA"]
    assert _metric(app, "Spend estimado para objetivo") == "$176.00"


def test_a_new_growth_target_hides_the_forecast_until_it_is_generated_again(monkeypatch):
    app = _generate(_page(monkeypatch, _FakeRest()))

    app.slider[0].set_value(20).run()

    assert not app.tabs
    _generate(app)
    assert _metric(app, "Con crecimiento +20%") == "$1,488.00"


def test_before_choosing_an_account_nothing_is_read_and_the_spend_estimate_is_unknown(monkeypatch):
    fake = _FakeRest()

    app = _generate(_page(monkeypatch, fake))

    assert fake.reads == []
    assert app.selectbox(key="forecast_src_account").value is None
    assert forecast_page.CHOOSE_ACCOUNT_NOTE in _text(app)
    assert _metric(app, "Spend estimado para objetivo") == "—"
    assert f"Sin desglose: {ad_account_block.MISSING_NOT_CHOSEN}." in _text(app)


def test_the_chosen_account_splits_the_report_with_its_sp_sb_and_sd_campaigns(monkeypatch):
    fake = _FakeRest()

    app = _choose_account(_generate(_page(monkeypatch, fake)))

    assert fake.reads == [("campaign_daily_totals", {"p_profile_id": "111", "p_from": "2026-08-03",
                                                     "p_to": "2026-08-30", "p_campaign": None})]
    text = _text(app)
    # SP: 28 days x 10 spend and 40 sales; SB: 8 weekend days x 5 spend and 20 sales.
    assert "$320.00" in text and "$1,280.00" in text and "25.0%" in text
    # Organic = 2,480 of the report minus 1,280 of ads.
    assert "$1,200.00" in text
    assert "28 de 28 días del BR tienen datos de ads" in text and "SP · SB" in text
    # 1,240 x 1.10 = 1,364 at a TACoS of 320 / 2,480.
    assert _metric(app, "Spend estimado para objetivo") == "$176.00"


def test_an_account_synced_over_part_of_the_report_is_compared_over_those_days_only(monkeypatch):
    fake = _FakeRest(synced_through=date(2026, 8, 20))

    app = _choose_account(_generate(_page(monkeypatch, fake)))

    assert fake.reads[0][1]["p_to"] == "2026-08-20"
    assert "18 de 28 días del BR tienen datos de ads" in _text(app)


def test_an_account_without_days_in_common_with_the_report_reads_nothing_and_says_why(monkeypatch):
    fake = _FakeRest(synced_through=date(2026, 7, 31))

    app = _choose_account(_generate(_page(monkeypatch, fake)))

    assert fake.reads == []
    assert ("Sin desglose: el BR va del 3 ago al 30 ago y la cuenta tiene campañas sincronizadas del 28 may al 31 jul: "
            "no hay días en común.") in _text(app)
    assert _metric(app, "Spend estimado para objetivo") == "—"


def test_an_account_whose_campaigns_never_synced_says_so_and_the_forecast_still_renders(monkeypatch):
    app = _choose_account(_generate(_page(monkeypatch, _FakeRest(synced_through=None))))

    assert ad_account_block.NO_CAMPAIGN_DATA_NOTE in _text(app)
    assert _metric(app, "Ventas proyectadas (14d)") == "$1,240.00"
    assert _metric(app, "Spend estimado para objetivo") == "—"


def test_an_account_loading_its_campaigns_for_the_first_time_says_so(monkeypatch):
    app = _choose_account(_generate(_page(monkeypatch, _FakeRest(synced_through=None, latest_status="pending"))))

    assert forecast_page.FIRST_LOAD_NOTE in _text(app)
    assert _metric(app, "Spend estimado para objetivo") == "—"
    assert f"Sin desglose: {ad_account_block.MISSING_NOT_SYNCED}." in _text(app)


def test_a_first_campaign_load_that_failed_says_so(monkeypatch):
    app = _choose_account(_generate(_page(monkeypatch, _FakeRest(synced_through=None, latest_status="failed"))))

    assert any(ad_account_block.FIRST_LOAD_FAILED_NOTE in str(error.value) for error in app.error)
    assert _metric(app, "Spend estimado para objetivo") == "—"


def test_an_account_that_needs_reauthorization_warns_and_keeps_its_last_synced_ads(monkeypatch):
    app = _choose_account(_generate(_page(monkeypatch, _FakeRest(profile_status="needs_reauth"))))

    assert any(str(warning.value) == search_term_source.NEEDS_REAUTH_MESSAGE for warning in app.warning)
    assert _metric(app, "Spend estimado para objetivo") == "$176.00"


def test_ads_that_cannot_be_read_say_so_and_the_forecast_still_renders(monkeypatch):
    app = _choose_account(_generate(_page(monkeypatch, _FakeRest(fail_totals=True))))

    assert any(forecast_page.UNREADABLE_NOTE in str(error.value) for error in app.error)
    assert _metric(app, "Ventas proyectadas (14d)") == "$1,240.00"
    assert f"Sin desglose: {ad_account_block.MISSING_UNREADABLE}." in _text(app)


def test_more_ad_sales_than_report_sales_warns_that_the_account_may_not_be_the_reports(monkeypatch):
    fake = _FakeRest(campaign_days=_campaign_days(sp_sales=500.0))

    app = _choose_account(_generate(_page(monkeypatch, fake)))

    # 28 days x 500 of SP plus 8 weekend days x 20 of SB, against the report's 2,480; two bare $ would render as LaTeX.
    assert [str(warning.value) for warning in app.warning] == [
        "Las ventas de ads (\\$14,160.00) superan las del BR (\\$2,480.00) en los mismos días: revisá que la cuenta y "
        "el país sean los del BR."]


def test_without_connected_accounts_the_page_asks_for_the_campaign_csv_and_still_projects(monkeypatch):
    app = _generate(_page(monkeypatch, _FakeRest(profiles=False)))

    assert f"{ad_account_block.FILE_HINT} {forecast_page.NO_ACCOUNTS_NOTE}" in _text(app)
    assert [box.label for box in app.selectbox] == ["Horizonte de proyección (días)"]
    assert _metric(app, "Ventas proyectadas (14d)") == "$1,240.00"
    assert _metric(app, "Spend estimado para objetivo") == "—"
    assert f"Sin desglose: {ad_account_block.MISSING_NO_ACCOUNTS}." in _text(app)


def test_without_connected_accounts_a_campaign_csv_splits_the_report_over_all_its_days(monkeypatch):
    fake = _FakeRest(profiles=False)

    app = _generate(_page(monkeypatch, fake, campaign_file=_Upload("campaigns.csv", _campaign_csv())))

    assert fake.reads == []
    text = _text(app)
    # The paused SB campaign still spent in the range: 280 + 40 spent and 1,120 + 160 sold.
    assert "$320.00" in text and "$1,280.00" in text and "25.0%" in text
    assert "$1,200.00" in text
    assert forecast_page.FILE_SPLIT_CAPTION.format(file="campaigns.csv", products="SP · SB", total=28,
                                                   window="3 – 30 ago 2026") in text
    assert "días del BR tienen datos de ads" not in text and "Sponsored Products con atribución de" not in text
    assert "Se compara con los 28 días del BR (3 – 30 ago 2026)" in text and "el archivo no dice la moneda" in text
    assert _metric(app, "Spend estimado para objetivo") == "$176.00"
    help_text = next(metric.help for metric in app.metric if metric.label == "Spend estimado para objetivo")
    assert help_text == forecast_page.FILE_BUDGET_HELP


@pytest.mark.parametrize("csv, products", [
    ("Campaign name,Type,Total cost,Sales\nIdle,Sponsored Products,$0.00,$0.00\n", "sin campañas con actividad"),
    ("Campaign name,Total cost,Sales\nIdle,$0.00,$0.00\n", "sin campañas con actividad"),
    ("Campaign name,Type,Total cost,Sales\nTV,Sponsored TV,$50.00,$200.00\n", "sin decir si son SP, SB o SD"),
    ("Campaign name,Total cost,Sales\nAny,$50.00,$200.00\n", "sin decir si son SP, SB o SD"),
], ids=["idle-with-types", "idle-without-types", "active-unknown-type", "active-without-types"])
def test_a_campaign_csv_says_no_campaign_ran_only_when_none_did(monkeypatch, csv, products):
    app = _generate(_page(monkeypatch, _FakeRest(profiles=False),
                          campaign_file=_Upload("campaigns.csv", csv.encode("utf-8"))))

    assert forecast_page.FILE_SPLIT_CAPTION.format(file="campaigns.csv", products=products, total=28,
                                                   window="3 – 30 ago 2026") in _text(app)


def _built_excels(monkeypatch) -> list[bytes]:
    built = []
    build = forecast_page._build_forecast_excel

    def spy(*args, **kwargs):
        workbook = build(*args, **kwargs)
        built.append(workbook.getvalue())
        return workbook

    monkeypatch.setattr(forecast_page, "_build_forecast_excel", spy)
    return built


def _summary_rows(workbook: bytes) -> dict:
    sheet = load_workbook(io.BytesIO(workbook))["Resumen"]
    return {label: value for label, value in sheet.iter_rows(min_row=4, max_col=2, values_only=True)}


def test_the_excel_says_the_spend_estimate_of_a_campaign_csv_assumes_the_reports_range(monkeypatch):
    built = _built_excels(monkeypatch)

    _generate(_page(monkeypatch, _FakeRest(profiles=False), campaign_file=_Upload("campaigns.csv", _campaign_csv())))

    rows = _summary_rows(built[-1])
    assert rows["Spend estimado para objetivo"] == "$176.00"
    assert rows["Origen del spend estimado"] == ("Campaign CSV «campaigns.csv» comparado con los 28 días del BR "
                                                 "(3 – 30 ago 2026): supone que el archivo cubre ese mismo rango")
    assert list(rows)[-1] == "Origen del spend estimado"


def test_the_excel_of_an_account_has_no_file_origin_row(monkeypatch):
    built = _built_excels(monkeypatch)

    _choose_account(_generate(_page(monkeypatch, _FakeRest())))

    rows = _summary_rows(built[-1])
    assert rows["Spend estimado para objetivo"] == "$176.00" and "Origen del spend estimado" not in rows


def test_a_campaign_csv_that_says_its_currency_shows_the_split_in_it(monkeypatch):
    csv = _campaign_csv().decode("utf-8").replace("$", "MX$").encode("utf-8")

    app = _generate(_page(monkeypatch, _FakeRest(profiles=False), campaign_file=_Upload("campaigns.csv", csv)))

    assert _metric(app, "Spend estimado para objetivo") == "MX$176.00"
    assert "MXN" in _text(app)


# An encrypted .xlsx or a renamed .xls is an OLE2 compound file, which pandas cannot open without xlrd.
@pytest.mark.parametrize("content", [b"not a workbook", bytes.fromhex("D0CF11E0A1B11AE1") + bytes(504)],
                         ids=["text", "ole2"])
def test_a_campaign_csv_that_cannot_be_read_says_how_to_export_it_and_the_forecast_still_renders(monkeypatch,
                                                                                                  content):
    app = _generate(_page(monkeypatch, _FakeRest(profiles=False), campaign_file=_Upload("campaigns.xlsx", content)))

    assert any("No se pudo leer «campaigns.xlsx»" in str(error.value) and "Campaign Manager → Campaigns → Export"
               in str(error.value) for error in app.error)
    assert _metric(app, "Ventas proyectadas (14d)") == "$1,240.00"
    assert _metric(app, "Spend estimado para objetivo") == "—"
    assert f"Sin desglose: {ad_account_block.MISSING_UNREADABLE_FILE}." in _text(app)


def test_a_campaign_csv_that_sold_more_than_the_report_warns_that_it_may_be_another_account_or_range(monkeypatch):
    csv = _campaign_csv(sp_sales="$14,000.00")

    app = _generate(_page(monkeypatch, _FakeRest(profiles=False), campaign_file=_Upload("campaigns.csv", csv)))

    assert [str(warning.value) for warning in app.warning] == [
        "Las ventas de ads del Campaign CSV (\\$14,160.00) superan las del BR (\\$2,480.00): revisá que el archivo "
        "sea de la cuenta del BR y del mismo rango de fechas."]


@pytest.mark.parametrize("rest, choose", [
    (_FakeRest(), False),
    (_FakeRest(fail_sync=True), True),
    (_FakeRest(synced_through=None), True),
    (_FakeRest(synced_through=None, latest_status="pending"), True),
    (_FakeRest(synced_through=None, latest_status="failed"), True),
    (_FakeRest(), True),
    (_FakeRest(fail_totals=True), True),
], ids=["not-chosen", "sync-unreadable", "never-synced", "first-load", "first-load-failed", "synced",
        "ads-unreadable"])
def test_every_state_of_the_account_card_offers_the_manual_upload(monkeypatch, rest, choose):
    app = _generate(_page(monkeypatch, rest))
    if choose:
        _choose_account(app)

    assert _upload_offered(app)


def test_the_manual_upload_reads_the_campaign_csv_instead_of_the_account_and_going_back_restores_it(monkeypatch):
    fake = _FakeRest()
    app = _choose_account(_generate(_page(monkeypatch, fake,
                                          campaign_file=_Upload("campaigns.csv", _campaign_csv()))))
    assert "28 de 28 días del BR tienen datos de ads" in _text(app)
    fake.reads.clear()

    app.button(key="forecast_src_upload_manual").click().run()

    assert not app.exception, app.exception
    assert fake.reads == []
    text = _text(app)
    assert search_term_source.MANUAL_MODE_NOTE in text and ad_account_block.FILE_HINT in text
    assert "Ventas de ads del Campaign CSV «campaigns.csv»" in text
    assert not _upload_offered(app) and not any(box.key == "forecast_src_account" for box in app.selectbox)

    app.button(key="forecast_src_back_to_api").click().run()

    assert not app.exception, app.exception
    assert app.selectbox(key="forecast_src_account").value == ACCOUNT
    assert "28 de 28 días del BR tienen datos de ads" in _text(app)
    assert _upload_offered(app)


def test_the_manual_mode_without_a_file_says_it_is_missing(monkeypatch):
    app = _generate(_page(monkeypatch, _FakeRest()))

    app.button(key="forecast_src_upload_manual").click().run()

    assert _metric(app, "Spend estimado para objetivo") == "—"
    assert f"Sin desglose: {ad_account_block.MISSING_NO_FILE}." in _text(app)


def test_the_ai_analysis_of_a_campaign_csv_says_where_the_ads_come_from_and_signs_the_file(monkeypatch):
    monkeypatch.setattr(ai_client, "ask", lambda **call: {"structured_output": ANSWER, "session_id": "file-session"})
    content = _campaign_csv()
    app = _generate(_page(monkeypatch, _FakeRest(profiles=False), ai_enabled=True,
                          campaign_file=_Upload("campaigns.csv", content)))

    _run_ai(app)

    shared = app.session_state["app_chat_modules"]["ppc_forecast"].analysis
    assert shared.documents[0]["title"] == "PPC Forecast · BusinessReport-by-date.csv · Parámetros"
    params = shared.documents[0]["content"]
    assert "Cuenta de Amazon Ads del Business Report: ninguna: el AM subió el Campaign CSV a mano" in params
    assert "Datos de ads: los del Campaign CSV «campaigns.csv»" in params
    assert "- Origen de los datos de ads: un Campaign CSV subido a mano" in params
    assert "Días con datos de ads" not in params and "Atribución" not in params
    assert "- Spend estimado para el objetivo: 176" in params
    assert (shared.profile_id, shared.country_code) == ("", "")
    assert app.session_state["ppc_forecast_ai_file_sig"].endswith(hashlib.sha256(content).hexdigest()[:16])


def test_the_ai_analysis_of_a_campaign_csv_without_activity_says_it_had_no_products(monkeypatch):
    monkeypatch.setattr(ai_client, "ask", lambda **call: {"structured_output": ANSWER, "session_id": "idle-session"})
    csv = b"Campaign name,Type,Total cost,Sales\nIdle,Sponsored Products,$0.00,$0.00\n"
    app = _generate(_page(monkeypatch, _FakeRest(profiles=False), ai_enabled=True,
                          campaign_file=_Upload("campaigns.csv", csv)))

    _run_ai(app)

    params = app.session_state["app_chat_modules"]["ppc_forecast"].analysis.documents[0]["content"]
    assert "- Productos con actividad: ninguno" in params


def test_the_ai_analysis_waits_for_the_click_and_reaches_the_chat_with_the_account(monkeypatch):
    calls = []

    def ask(**call):
        calls.append(call)
        return {"structured_output": ANSWER, "session_id": "forecast-session"}

    monkeypatch.setattr(ai_client, "ask", ask)
    app = _choose_account(_generate(_page(monkeypatch, _FakeRest(), ai_enabled=True)))

    assert calls == []
    app.button(key="ppc_forecast_ai_recalc").click().run()
    for _ in range(30):
        entry = app.session_state["app_chat_modules"].get("ppc_forecast")
        if entry is not None and entry.state == "current":
            break
        time.sleep(0.2)
        app.run()

    assert not app.exception, app.exception
    assert len(calls) == 1
    shared = app.session_state["app_chat_modules"]["ppc_forecast"].analysis
    titles = [document["title"] for document in shared.documents]
    assert titles[0] == "PPC Forecast · Luna Kids · US · Parámetros"
    assert titles[-1] == "PPC Forecast · Luna Kids · US · Lectura de la IA"
    assert "Spend estimado para el objetivo → media: TACoS de 12,9%." in shared.documents[-1]["content"]
    assert "- Spend estimado para el objetivo: 176" in shared.documents[0]["content"]
    assert (shared.profile_id, shared.country_code) == ("111", "US")
    assert [tab.label for tab in app.tabs] == ["📈 Forecast", "🤖 Análisis IA"]


def test_the_ai_rows_show_each_figure_with_the_confidence_and_skip_the_ones_the_module_did_not_compute():
    labels = ai_tab.ai_labels("es", forecast_page._AI_TEXTS["es"])
    records = [{"tema": "PROYECCION", "item": "Ventas proyectadas (14 días)", "metrics": ["$1,240.00"]}]

    rows = forecast_page.forecast_ai_rows(ANSWER["lecturas"], records, labels)

    assert rows == [{"item": "Ventas proyectadas (14 días)", "metrics": ["$1,240.00"], "badges": ["Alta"],
                     "warning": "", "reasoning": "28 días idénticos."}]
