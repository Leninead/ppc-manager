"""Account Pulse over an in-memory PostgREST: the week against the prior one, and the ads of the chosen account.

No network: `_open_rest` is replaced before every script run, the uploads are fakes that serve CSVs built here, and the
AI tab runs disabled except where a test fakes the provider.
"""
import csv
import io
import time
from datetime import date, datetime, timedelta, timezone

import pandas as pd
import pytest
import requests
from openpyxl import load_workbook
from streamlit.proto.Common_pb2 import FileURLs
from streamlit.proto.WidgetStates_pb2 import WidgetState
from streamlit.runtime.uploaded_file_manager import UploadedFile, UploadedFileRec
from streamlit.testing.v1 import AppTest
from streamlit.testing.v1.element_tree import ButtonGroup

import ai.client as ai_client
import ai.runtime as ai_runtime
import modules.pages.search_term_source as search_term_source
from core import ai_tab
from core.account_pulse.ads_by_week import ads_by_week
from core.account_pulse.campaigns import campaign_rows
from core.amazon_ads import campaign_totals
from core.amazon_ads.campaign_totals import ProductDay, ProductSeries, Totals
from core.amazon_ads.report_provider import ProfileOption
from modules.pages import account_pulse as pulse_page
from modules.pages import ad_account_block

ACCOUNT = "Luna Kids"
# Monday 3 to Sunday 16 August 2026. The prior week sells 100 on weekdays and 60 on weekends (620); this one 120 and 80
# (760).
FIRST_DAY, LAST_DAY = date(2026, 8, 3), date(2026, 8, 16)
THIS_WEEK_START = date(2026, 8, 10)
SYNCED_AT = datetime(2026, 9, 23, 9, 30, tzinfo=timezone.utc)
ANSWER = {"synthesis": {"situation": "Las ventas subieron 22,6%.", "week_actions": ["Mantener el budget"],
                        "mid_term": [], "risks": [], "executive_summary": "760 esta semana."},
          "lecturas": [{"tema": "VENTAS", "razon": "760 contra 620.", "veredicto": "OK", "advertencia": None},
                       {"tema": "PUBLICIDAD", "razon": "TACoS de 10,5%.", "veredicto": "VIGILAR",
                        "advertencia": "Los últimos días todavía pueden crecer."}]}
WINDOW_HEADER = ["ad_product", "campaign_id", "campaign_name", "portfolio_id", "portfolio_name", "impressions", "clicks",
                 "cost", "purchases_7d", "sales_7d", "purchases_14d", "sales_14d", "purchases", "sales",
                 "purchases_clicks", "sales_clicks", "currency_code", "new_to_brand_purchases", "new_to_brand_sales"]


def _report_days(first: date = FIRST_DAY, last: date = LAST_DAY) -> list[tuple[date, float]]:
    days = [first + timedelta(days=offset) for offset in range((last - first).days + 1)]
    return [(day, (80.0 if day.weekday() >= 5 else 120.0) if day >= THIS_WEEK_START
             else (60.0 if day.weekday() >= 5 else 100.0)) for day in days]


def _business_report_csv(days: list[tuple[date, float]]) -> bytes:
    frame = pd.DataFrame({
        "Date": [day.isoformat() for day, _ in days],
        "Ordered Product Sales": [f"${sales:,.2f}" for _, sales in days],
        "Units Ordered": [int(sales // 10) for _, sales in days],
        "Sessions - Total": [int(sales) * 2 for _, sales in days],
    })
    return frame.to_csv(index=False).encode("utf-8")


def _by_child_csv() -> bytes:
    frame = pd.DataFrame({"(Child) ASIN": ["B0BIG00001", "B0HEALTHY1"], "Title": ["Big", "Healthy"],
                          "Ordered Product Sales": ["$1,000.00", "$500.00"], "Sessions - Total": [50, 30],
                          "Units Ordered": [10, 5], "Featured Offer (Buy Box) Percentage": ["90%", "99%"]})
    return frame.to_csv(index=False).encode("utf-8")


def _upload(name: str, data: bytes) -> UploadedFile:
    """What Streamlit hands the page: a fresh file on every run."""
    return UploadedFile(UploadedFileRec(file_id=name, name=name, type="text/csv", data=data), FileURLs())


def _profile_row(status: str = "active") -> dict:
    return {"profile_id": "111", "account_id": 1, "cliente": ACCOUNT, "account_name": "Luna Kids US",
            "country_code": "US", "currency_code": "USD", "account_type": "seller", "timezone": "America/New_York",
            "status": status, "data_from": "2026-07-20", "data_through": "2026-09-22", "refreshed_on": "2026-09-23",
            "last_success_at": SYNCED_AT.isoformat(), "last_error": ""}


def _campaign_job(window_end: date, status: str = "completed") -> dict:
    return {"id": 7, "integration_slug": "amazon_ads", "job_kind": "sp_campaigns", "trigger": "scheduled_daily",
            "external_account_id": "111", "status": status, "window_start": (window_end - timedelta(days=6))
            .isoformat(), "window_end": window_end.isoformat(), "local_day": (window_end + timedelta(days=1))
            .isoformat(), "finished_at": SYNCED_AT.isoformat(), "created_at": SYNCED_AT.isoformat()}


def _campaign_days(first: date = date(2026, 7, 20), last: date = date(2026, 9, 22), *, sp_sales: float | None = None):
    """SP spends 10 a day and sells 40 in the prior week and 50 in this one; SB spends 5 and sells 20 on weekends."""
    rows = []
    for offset in range((last - first).days + 1):
        day = first + timedelta(days=offset)
        sales = sp_sales if sp_sales is not None else 50.0 if day >= THIS_WEEK_START else 40.0
        rows.append(_totals_row(day, "SP", cost=10.0, sales_7d=sales))
        if day.weekday() >= 5:
            rows.append(_totals_row(day, "SB", cost=5.0, sales=20.0))
    return rows


def _totals_row(day: date, product: str, *, cost: float, sales_7d: float = 0.0, sales: float = 0.0) -> dict:
    return {"report_date": day.isoformat(), "ad_product": product, "impressions": 100, "clicks": 5, "cost": cost,
            "purchases_7d": 1, "sales_7d": sales_7d, "purchases_14d": 1, "sales_14d": sales_7d, "purchases": 1,
            "sales": sales, "purchases_clicks": 1, "sales_clicks": sales, "currency_code": "USD",
            "campaign_names": None, "new_to_brand_purchases": None, "new_to_brand_sales": None}


def _window_row(product: str, campaign_id: str, name: str, *, cost: float, sales: float, orders: int) -> dict:
    row = dict.fromkeys(WINDOW_HEADER, 0)
    row.update(ad_product=product, campaign_id=campaign_id, campaign_name=name, portfolio_id="", portfolio_name="",
               impressions=900, clicks=40, cost=cost, currency_code="USD", new_to_brand_purchases="",
               new_to_brand_sales="")
    if product == "SP":
        row.update(purchases_7d=orders, sales_7d=sales, purchases_14d=orders, sales_14d=sales)
    else:
        row.update(purchases=orders, sales=sales, purchases_clicks=orders, sales_clicks=sales)
    return row


WINDOW_ROWS = [_window_row("SP", "1", "Luna - B0TEST0001 - SP - KW - EXACT - Brand", cost=120.0, sales=600.0, orders=12),
               _window_row("SP", "2", "Old auto", cost=30.0, sales=0.0, orders=0),
               _window_row("SB", "3", "Video Luna", cost=10.0, sales=40.0, orders=2)]


def _csv(header, rows) -> bytes:
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=header)
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode("utf-8")


class _FakeRest:
    """The profile table, the campaign sync jobs, and the campaign totals by day and by campaign."""

    def __init__(self, *, profiles=True, synced_through: date | None = date(2026, 9, 22), campaign_days=None,
                 fail_totals=False, latest_status: str = ""):
        self._profiles, self._synced_through, self._fail_totals = profiles, synced_through, fail_totals
        self._latest_status = latest_status
        self._campaign_days = _campaign_days() if campaign_days is None else campaign_days
        self.reads: list[tuple[str, dict]] = []

    def select(self, table, params):
        if table == "ads_profile_sync":
            return [_profile_row()] if self._profiles else []
        if table == "integration_sync_jobs":
            assert params["job_kind"] == "eq.sp_campaigns"
            if self._latest_status and params.get("status") != "eq.completed":
                return [_campaign_job(date(2026, 9, 22), self._latest_status)]
            return [_campaign_job(self._synced_through)] if self._synced_through else []
        raise AssertionError(f"unexpected select on {table}")

    def rpc(self, name, args, *, timeout_s=8):
        self.reads.append((name, args))
        if self._fail_totals:
            raise requests.ConnectionError("gateway down")
        assert name == "campaign_daily_totals" and args["p_profile_id"] == "111"
        return [row for row in self._campaign_days if args["p_from"] <= row["report_date"] <= args["p_to"]]

    def rpc_csv(self, name, args, *, timeout_s=8):
        self.reads.append((name, args))
        assert name == "campaign_window_totals" and args["p_profile_id"] == "111"
        return _csv(WINDOW_HEADER, WINDOW_ROWS)


@pytest.fixture(autouse=True)
def _single_select_button_groups(monkeypatch):
    # AppTest 1.43 serializes button groups as multi-select (st.feedback); a segmented control holds one value.
    def widget_state(group):
        state = WidgetState()
        state.id = group.id
        value = group.value
        selected = value if isinstance(value, list) else [] if value is None else [value]
        state.int_array_value.data[:] = [group.options.index(group.format_func(option)) for option in selected]
        return state
    monkeypatch.setattr(ButtonGroup, "_widget_state", property(widget_state))


@pytest.fixture(autouse=True)
def _empty_ai_registry():
    with ai_runtime._lock:
        ai_runtime._registry.clear()
    yield
    with ai_runtime._lock:
        ai_runtime._registry.clear()


_PAGE_SCRIPT = """
import streamlit as st
from core.chat import app_chat
st.cache_data.clear()
if st.session_state.get("shared_before"):
    app_chat.report_failed("account_pulse")
from modules.pages.account_pulse import render
render()
"""


def _page(monkeypatch, rest=None, *, days=None, daily=True, by_child=False, ai_enabled=False) -> AppTest:
    import streamlit
    uploads = {}
    if daily:
        uploads["ap_br_daily"] = ("BusinessReport-by-date.csv", _business_report_csv(days or _report_days()))
    if by_child:
        uploads["ap_br_child"] = ("BusinessReport-by-child.csv", _by_child_csv())
    monkeypatch.setattr(streamlit, "file_uploader",
                        lambda label, *args, key=None, **kwargs: _upload(*uploads[key]) if key in uploads else None)
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: rest)
    monkeypatch.setattr("ai.config.AI_ENABLED", ai_enabled)
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.run()
    assert not app.exception, app.exception
    return app


def _choose_account(app: AppTest) -> AppTest:
    app.selectbox(key="ap_src_account").set_value(ACCOUNT).run()
    assert not app.exception, app.exception
    return app


def _text(app: AppTest) -> str:
    parts = [str(element.value) for element in app.markdown] + [str(element.value) for element in app.caption]
    parts += [str(element.value) for element in app.error] + [str(element.value) for element in app.info]
    parts += [str(element.value) for element in app.warning] + [str(element.value) for element in app.success]
    return " ".join(parts)


def _card(app: AppTest, label: str) -> str:
    return next(str(element.value) for element in app.markdown if f">{label}</div>" in str(element.value))


def test_without_reports_nothing_renders_and_the_analysis_is_withdrawn(monkeypatch):
    import streamlit
    monkeypatch.setattr(streamlit, "file_uploader", lambda *args, **kwargs: None)
    monkeypatch.setattr(search_term_source, "_open_rest", lambda: _FakeRest())
    app = AppTest.from_string(_PAGE_SCRIPT, default_timeout=30)
    app.session_state["shared_before"] = True

    app.run()

    assert not app.exception
    assert not app.tabs
    assert "account_pulse" not in app.session_state["app_chat_modules"]


def test_before_choosing_an_account_nothing_is_read_and_acos_and_tacos_are_unknown(monkeypatch):
    fake = _FakeRest()

    app = _page(monkeypatch, fake)

    assert fake.reads == []
    assert [tab.label for tab in app.tabs] == ["📊 Pulse", "🤖 Análisis IA"]
    assert app.selectbox(key="ap_src_account").value is None
    assert pulse_page._ADS_TEXTS.choose_account in _text(app)
    assert f"Sin ACoS, TACoS ni campañas: {ad_account_block.MISSING_NOT_CHOSEN}." in _text(app)
    assert ">—</div>" in _card(app, "ACoS TW") and ">—</div>" in _card(app, "TACoS TW")


def test_the_chosen_account_is_read_over_the_report_days_and_each_week_over_its_own(monkeypatch):
    fake = _FakeRest()

    app = _choose_account(_page(monkeypatch, fake))

    window = {"p_profile_id": "111", "p_from": "2026-08-03", "p_to": "2026-08-16"}
    assert fake.reads == [("campaign_daily_totals", {**window, "p_campaign": None}),
                          ("campaign_window_totals", window)]
    # This week: 80 of spend over 390 of ad sales and 760 of the report's sales; the prior one: 80 over 320 and 620.
    assert ">20.5%</div>" in _card(app, "ACoS TW") and "↓ 17.9%" in _card(app, "ACoS TW")
    assert ">10.5%</div>" in _card(app, "TACoS TW") and "↓ 18.4%" in _card(app, "TACoS TW")
    text = _text(app)
    assert "esta semana 7 de 7 días, la anterior 7 de 7 días" in text
    assert "14 de 14 días del BR con datos de ads" in text and "3 campañas con actividad" in text
    campaigns = app.dataframe[-1].value
    assert list(campaigns["Campaign"]) == ["Luna - B0TEST0001 - SP - KW - EXACT - Brand", "Old auto", "Video Luna"]
    assert list(campaigns["Tipo"]) == ["NUEVA", "HEREDADA", "HEREDADA"]
    assert list(campaigns["ACoS %"]) == ["20.0%", "—", "25.0%"]


def test_a_week_the_sync_covers_in_part_is_compared_over_those_days_only(monkeypatch):
    fake = _FakeRest(synced_through=date(2026, 8, 14))

    app = _choose_account(_page(monkeypatch, fake))

    assert fake.reads[0][1]["p_to"] == "2026-08-14"
    assert "esta semana 5 de 7 días, la anterior 7 de 7 días" in _text(app)


def test_an_account_whose_campaigns_never_synced_says_so_and_the_pulse_still_renders(monkeypatch):
    app = _choose_account(_page(monkeypatch, _FakeRest(synced_through=None)))

    assert ad_account_block.NO_CAMPAIGN_DATA_NOTE in _text(app)
    assert f"Sin ACoS, TACoS ni campañas: {ad_account_block.MISSING_NOT_SYNCED}." in _text(app)
    assert "Sales TW" in _card(app, "Sales TW")


def test_ads_that_cannot_be_read_say_so_and_the_pulse_still_renders(monkeypatch):
    app = _choose_account(_page(monkeypatch, _FakeRest(fail_totals=True)))

    assert any(pulse_page._ADS_TEXTS.unreadable in str(error.value) for error in app.error)
    assert f"Sin ACoS, TACoS ni campañas: {ad_account_block.MISSING_UNREADABLE}." in _text(app)
    assert [tab.label for tab in app.tabs] == ["📊 Pulse", "🤖 Análisis IA"]


def test_more_ad_sales_than_report_sales_warns_that_the_account_may_not_be_the_reports(monkeypatch):
    app = _choose_account(_page(monkeypatch, _FakeRest(campaign_days=_campaign_days(sp_sales=500.0))))

    # 14 days x 500 of SP plus 4 weekend days x 20 of SB, against the report's 1,380; two bare $ would render as LaTeX.
    assert [str(warning.value) for warning in app.warning] == [
        "Las ventas de ads (\\$7,080.00) superan las del BR (\\$1,380.00) en los mismos días: revisá que la cuenta y "
        "el país sean los del BR."]


def test_without_connected_accounts_the_page_says_so_and_still_renders(monkeypatch):
    app = _page(monkeypatch, _FakeRest(profiles=False))

    assert pulse_page._ADS_TEXTS.no_accounts in _text(app)
    assert not app.selectbox
    assert f"Sin ACoS, TACoS ni campañas: {ad_account_block.MISSING_NO_ACCOUNTS}." in _text(app)


def test_the_by_child_report_alone_lists_its_buybox_alerts_and_asks_for_the_daily_one(monkeypatch):
    app = _page(monkeypatch, _FakeRest(), daily=False, by_child=True)

    assert not app.tabs
    assert "BuyBox Alerts (1 ASINs < 95%)" in _text(app)
    assert "Cargá al menos el BR diario para generar el Excel." in _text(app)


def test_the_ai_analysis_waits_for_the_click_and_reaches_the_chat_with_the_account(monkeypatch):
    calls = []

    def ask(**call):
        calls.append(call)
        return {"structured_output": ANSWER, "session_id": "pulse-session"}

    monkeypatch.setattr(ai_client, "ask", ask)
    app = _choose_account(_page(monkeypatch, _FakeRest(), ai_enabled=True))

    assert calls == []
    app.button(key="account_pulse_ai_recalc").click().run()
    for _ in range(30):
        entry = app.session_state["app_chat_modules"].get("account_pulse")
        if entry is not None and entry.state == "current":
            break
        time.sleep(0.2)
        app.run()

    assert not app.exception, app.exception
    assert len(calls) == 1
    shared = app.session_state["app_chat_modules"]["account_pulse"].analysis
    titles = [document["title"] for document in shared.documents]
    assert titles[0] == "Account Pulse · Luna Kids · US · Parámetros"
    assert titles[-1] == "Account Pulse · Luna Kids · US · Lectura de la IA"
    assert "Publicidad → VIGILAR: TACoS de 10,5%." in shared.documents[-1]["content"]
    assert "- TACoS, semana actual (%): 10.5" in shared.documents[0]["content"]
    assert (shared.profile_id, shared.country_code) == ("111", "US")


def test_the_ai_rows_show_each_topic_with_its_verdict_and_skip_the_ones_the_module_did_not_compute():
    labels = ai_tab.ai_labels("es", pulse_page._AI_TEXTS["es"])
    records = [{"tema": "VENTAS", "item": "Ventas", "metrics": ["Sales TW $760.00", "PW $620.00", "+22.6%"]}]

    rows = pulse_page.pulse_ai_rows(ANSWER["lecturas"], records, labels)

    assert rows == [{"item": "Ventas", "metrics": ["Sales TW $760.00", "PW $620.00", "+22.6%"], "badges": ["OK"],
                     "warning": "", "reasoning": "760 contra 620."}]


def _history() -> pd.DataFrame:
    days = _report_days()
    return pd.DataFrame({"_date": pd.to_datetime([day for day, _ in days]), "_sales": [sales for _, sales in days]})


def _series(currency_code: str = "MXN") -> ProductSeries:
    days = [FIRST_DAY + timedelta(days=offset) for offset in range((LAST_DAY - FIRST_DAY).days + 1)]
    return ProductSeries(days=tuple(ProductDay(day, Totals(10.0, 50.0 if day >= THIS_WEEK_START else 40.0, 1, 5, 100,
                                                            0.0, 0)) for day in days),
                         campaigns=(), products=("SP",), currency_code=currency_code, attribution_days=7)


def _rows_by_label(worksheet) -> dict:
    return {row[0].value: [cell.value for cell in row] for row in worksheet.iter_rows() if row[0].value}


def test_the_excel_compares_acos_and_tacos_week_over_week_and_lists_the_campaigns_in_the_account_currency():
    daily_data = pulse_page._parse_br_daily(_upload("BusinessReport-by-date.csv", _business_report_csv(_report_days())))
    weeks = ads_by_week(_history(), _series(), THIS_WEEK_START)
    frame = campaign_totals.window_totals(_FakeRest(), ProfileOption.from_row(_profile_row()), FIRST_DAY, LAST_DAY)
    campaigns = campaign_rows(frame)

    workbook = load_workbook(pulse_page._build_account_pulse_excel(
        daily_data, {}, campaigns, "Luna", 30.0, weeks=weeks, currency_code="MXN", ads_period="3–16 ago 2026"))

    summary = _rows_by_label(workbook["Resumen Ejecutivo"])
    # 70 of spend over 350 of ad sales and 760 of the report's sales, against 70 over 280 and 620.
    assert summary["ACoS %"][1:5] == ["20.0%", "25.0%", "-20.0%", "OK"]
    assert summary["TACoS %"][1:5] == ["9.2%", "11.3%", "-18.4%", "OK"]
    assert summary["Sales"][1:3] == ["MX$760.00", "MX$620.00"]
    sheet = workbook["Campañas"]
    assert sheet["A1"].value == "Luna — Campañas · 3–16 ago 2026"
    assert [cell.value for cell in sheet[3]] == ["Luna - B0TEST0001 - SP - KW - EXACT - Brand", "SP", "NUEVA", 900, 40,
                                                 120.0, 600.0, "20.0%", 12]
    assert sheet["H4"].value == "—"
    assert sheet["F3"].number_format == '"MX$"#,##0.00'
