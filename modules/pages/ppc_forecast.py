"""PPC Forecast (M19): the Business Report's daily sales projected with their trend and their weekend difference.

The ad spend and ad sales behind the organic vs paid split come from the campaign reports of the Amazon Ads account
the AM picks, over the Business Report's own days. The rules live in core/ppc_forecast/.
"""
import hashlib
import html
import io
import logging
from dataclasses import dataclass
from datetime import date, datetime, timezone
from functools import partial
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.amazon_ads import campaign_totals
from core.amazon_ads.campaign_provider import campaign_sync_view
from core.amazon_ads.campaign_totals import ProductSeries
from core.amazon_ads.report_provider import ProfileOption, ReportReadError
from core.amazon_ads.sync_planner import PROFILE_NEEDS_REAUTH
from core.chat import app_chat
from core.currency_format import currency_symbol, money
from core.date_labels import date_range_label, short_date
from core.helpers import kpi_card
from core.ppc_forecast.analysis import ANALYSIS_MODULE, build_analysis_input
from core.ppc_forecast.paid_split import PaidSplit, covered_window, paid_split, spend_for_target
from core.ppc_forecast.projection import DATE, PROJECTED_SALES, SalesForecast, forecast_sales
from core.ui import palette
from modules.pages import campaign_source, search_term_source
from modules.pages.search_term_source import (
    NEEDS_REAUTH_MESSAGE,
    STATE_FIRST_LOAD_FAILED,
    country_labels,
    group_by_label,
    picker_key,
    source_state,
)

log = logging.getLogger(__name__)

MODULE_LABEL = "PPC Forecast"
KEY_PREFIX = "forecast"
AD_TOTALS_TTL_SECONDS = 60 * 60
_GENERATED_FOR_KEY = "forecast_generated_for"

ADS_BLOCK_TITLE = "Ventas de ads de la cuenta"
ADS_BLOCK_TAG = "Reportes de campañas de Amazon Ads"
ACCOUNT_PLACEHOLDER = "Elegí la cuenta del BR"
NO_ACCOUNTS_NOTE = ("No hay cuentas de Amazon Ads conectadas, así que no hay desglose orgánico vs paid ni spend "
                    "estimado. Se conectan en Sistema → Cuentas conectadas.")
CHOOSE_ACCOUNT_NOTE = ("Elegí la cuenta y el país del BR para separar las ventas de ads de las orgánicas y estimar el "
                       "spend para el objetivo. El BR no dice de qué cuenta es, así que no hay una por defecto.")
NO_CAMPAIGN_DATA_NOTE = "Todavía no hay métricas de campañas de esta cuenta: se sincronizan una vez por día."
FIRST_LOAD_NOTE = "Estamos trayendo las campañas de esta cuenta por primera vez; cuando termine aparece el desglose."
FIRST_LOAD_FAILED_NOTE = ("La primera carga de campañas de esta cuenta no se pudo completar. Si sigue así, avisale a "
                          "un admin.")
UNREADABLE_NOTE = "Sin desglose orgánico vs paid ni spend estimado."
NO_DATABASE_MESSAGE = "No hay base de datos configurada para leer las campañas de Amazon Ads."
ADS_EXCEED_BR_WARNING = ("Las ventas de ads ({ad_sales}) superan las del BR ({br_sales}) en los mismos días: revisá "
                         "que la cuenta y el país sean los del BR.")
SPLIT_CAPTION = ("{covered} de {total} días del BR tienen datos de ads ({window} · {products}). Sponsored Products "
                 "con atribución de {attribution} días; Sponsored Brands y Display como los cuenta Campaign Manager. "
                 "Las ventas de ads se atribuyen al día del click: las de los últimos días todavía pueden crecer.")
BUDGET_HELP = ("Spend de ads ÷ ventas del BR (TACoS) en los días con datos de ads, aplicado a las ventas con el "
               "crecimiento objetivo. Supone que el TACoS no cambia al subir el spend.")

# Why there are no ads figures: shown on the page and sent to the AI as is.
MISSING_NO_ACCOUNTS = "no hay cuentas de Amazon Ads conectadas"
MISSING_NOT_CHOSEN = "no se eligió la cuenta de Amazon Ads del BR"
MISSING_NOT_SYNCED = "la cuenta todavía no tiene campañas sincronizadas"
MISSING_UNREADABLE = "no se pudieron leer las campañas de la cuenta"
MISSING_NO_SHARED_DAYS = ("el BR va del {first} al {last} y la cuenta tiene campañas sincronizadas del {synced_from} "
                          "al {synced_through}: no hay días en común")

_CONFIDENCE_COLORS = ("background-color:#EAF3DE;color:#173404", "background-color:#FAEEDA;color:#412402",
                      "background-color:#FFEBEE;color:#9C0006")
_AI_TEXTS = {
    "es": {"title": "Análisis IA",
           "caption": "Lectura de la IA sobre las cifras que ya calculó el módulo: qué tanto confiar en la proyección "
                      "y en el spend estimado, y qué hacer esta semana",
           "disabled": "Análisis IA deshabilitado (AI_ENABLED=0).",
           "table_title": "Cifras del módulo — lectura IA",
           "col_item": "Cifra", "col_diag": "Confianza",
           "counts": "{n} cifras leídas",
           "stale_body": "Cambió lo que el análisis leyó, por ejemplo el crecimiento objetivo, el horizonte o los "
                         "datos de ads. Lo que se muestra abajo corresponde a los datos anteriores.",
           "projection_item": "Ventas proyectadas ({days} días)",
           "budget_item": "Spend estimado para el objetivo",
           "with_growth": "con +{growth}%: {amount}",
           "levels": {"alta": "Alta", "media": "Media", "baja": "Baja"}},
    "en": {"title": "AI analysis",
           "caption": "AI read on the figures the module already computed: how far to trust the projection and the "
                      "spend estimate, and what to do this week",
           "disabled": "AI analysis disabled (AI_ENABLED=0).",
           "table_title": "The module's figures — AI read",
           "col_item": "Figure", "col_diag": "Confidence",
           "counts": "{n} figures read",
           "stale_body": "What the analysis read changed, for example the growth target, the horizon or the ads "
                         "data. What is shown below belongs to the previous data.",
           "projection_item": "Projected sales ({days} days)",
           "budget_item": "Estimated spend for the target",
           "with_growth": "with +{growth}%: {amount}",
           "levels": {"alta": "High", "media": "Medium", "baja": "Low"}},
}


@dataclass(frozen=True)
class AdAccountChoice:
    """The account block's outcome: the chosen profile as the campaign sync sees it, or why there is none."""

    profile: ProfileOption | None = None
    info_line: Any = None  # the block's last line, filled in once the Business Report's days are known
    no_ads_reason: str = ""


@dataclass(frozen=True)
class AdsReading:
    """What the chosen account's campaign reports say about the Business Report's days."""

    account: str = ""
    profile_id: str = ""
    country_code: str = ""
    currency_code: str = ""
    series: ProductSeries | None = None
    split: PaidSplit | None = None
    no_ads_reason: str = ""


# ── Helpers numéricos ────────────────────────────────────────────────────────

def _clean_num(series):
    """Limpia $, %, comas y convierte a float."""
    return (
        pd.to_numeric(
            series.astype(str)
            .str.replace(r"[\$%,]", "", regex=True)
            .str.strip(),
            errors="coerce",
        ).fillna(0.0)
    )


# ── Parsers ──────────────────────────────────────────────────────────────────

@st.cache_data(max_entries=3, ttl=3600, show_spinner=False)
def _parse_br_daily(file_bytes, fname):
    """Parsea BR diario. Retorna DataFrame con columnas _date, _sales, _units, _sess (NaN si el BR no las trae)."""
    df = pd.read_excel(io.BytesIO(file_bytes)) if fname.endswith(".xlsx") else pd.read_csv(io.BytesIO(file_bytes))
    df.columns = df.columns.str.strip()

    date_col = next((c for c in df.columns if "date" in c.lower()), None)
    sales_col = next(
        (c for c in df.columns if "ordered product sales" in c.lower() and "b2b" not in c.lower()),
        None,
    )
    units_col = next(
        (c for c in df.columns if "units ordered" in c.lower() and "b2b" not in c.lower()),
        None,
    )
    sess_col = next(
        (c for c in df.columns if "sessions" in c.lower() and "total" in c.lower() and "b2b" not in c.lower()),
        None,
    )

    if not date_col:
        raise ValueError("No se encontró columna Date en el archivo.")
    if not sales_col:
        raise ValueError("No se encontró columna Ordered Product Sales en el archivo.")

    df["_date"] = pd.to_datetime(df[date_col], format="mixed", dayfirst=False, errors="coerce")
    df = df.dropna(subset=["_date"])
    df["_sales"] = _clean_num(df[sales_col])
    # A column the report does not carry is unknown, not zero: the AI reads these.
    df["_units"] = _clean_num(df[units_col]) if units_col else float("nan")
    df["_sess"] = _clean_num(df[sess_col]) if sess_col else float("nan")

    df = df.sort_values("_date").reset_index(drop=True)
    return df


# ── Excel export ─────────────────────────────────────────────────────────────

def _build_forecast_excel(forecast: SalesForecast, n_days: int, needed_spend: float | None, client_name: str,
                          currency_code: str):
    """Construye Excel con 2 hojas: Resumen + Proyección Diaria."""
    HDR_FILL = PatternFill("solid", fgColor="E84000")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
    ALT_FILL = PatternFill("solid", fgColor="FFF3E0")
    CENTER = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="DDDDDD")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    show = partial(money, currency_code=currency_code)

    wb = Workbook()

    # ── Hoja 1: Resumen ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumen"

    # Título branding
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    titulo = f"PPC Forecast — {client_name}" if client_name else "PPC Forecast"
    title_cell.value = titulo
    title_cell.fill = HDR_FILL
    title_cell.font = Font(bold=True, color="FFFFFF", size=13)
    title_cell.alignment = CENTER
    ws.row_dimensions[1].height = 24

    ws.append([])  # blank row

    # Métricas históricas
    ws.append(["Métrica", "Valor"])
    for cell in ws[ws.max_row]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    ratio = forecast.trend.weekend_ratio
    metric_rows = [
        ("Días de datos", n_days),
        ("Ventas promedio/día", show(forecast.average_daily_sales)),
        ("Tendencia por día", _signed_money(forecast.trend.slope, currency_code)),
        ("Ratio Finde/Laboral", f"{ratio:.2f}x" if ratio is not None else "sin dato"),
        ("Horizonte de proyección (días)", forecast.horizon),
        ("Total ventas proyectadas", show(forecast.projected_sales)),
        ("Total ventas con crecimiento objetivo", show(forecast.sales_with_growth)),
        ("Spend estimado para objetivo", show(needed_spend) if needed_spend is not None else "sin dato"),
    ]

    for i, (label, val) in enumerate(metric_rows):
        ws.append([label, val])
        row_idx = ws.max_row
        for cell in ws[row_idx]:
            cell.border = BORDER
            if i % 2 == 1:
                cell.fill = ALT_FILL

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 24

    # ── Hoja 2: Proyección Diaria ────────────────────────────────────────────
    ws2 = wb.create_sheet("Proyección Diaria")
    proj_df = forecast.projection

    headers = list(proj_df.columns)
    ws2.append(headers)
    for cell in ws2[1]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    for i, row in proj_df.iterrows():
        ws2.append(list(row))
        row_idx = ws2.max_row
        for cell in ws2[row_idx]:
            cell.border = BORDER
            if i % 2 == 1:
                cell.fill = ALT_FILL
        # Color tipo finde
        tipo_cell = ws2.cell(row=row_idx, column=4)
        if tipo_cell.value == "Fin de semana":
            tipo_cell.fill = PatternFill("solid", fgColor="FFF3E0")

    for col_idx, col in enumerate(["A", "B", "C", "D"], start=1):
        ws2.column_dimensions[col].width = 22

    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _signed_money(value: float, currency_code: str) -> str:
    return f"{'+' if value >= 0 else ''}{money(value, currency_code)}/día"


# ── Render ───────────────────────────────────────────────────────────────────

def render():
    st.markdown(
        "<div style='display:flex;align-items:center;gap:0.75rem;margin-bottom:0.25rem;'>"
        "<span style='font-size:2rem;'>📈</span>"
        "<div>"
        "<div style='font-size:1.3rem;font-weight:800;color:#1F1F1F;'>PPC Forecast</div>"
        "<div style='font-size:0.82rem;color:#888;'>Proyección de ventas y spend basada en tendencia histórica + estacionalidad.</div>"
        "</div>"
        "</div>",
        unsafe_allow_html=True,
    )
    st.divider()
    _how_to_use()

    # ── Inputs globales ──────────────────────────────────────────────────────
    col_a, col_b, col_c = st.columns(3)
    with col_a:
        target_growth = st.slider(
            "Crecimiento objetivo (%)",
            min_value=0, max_value=100, value=10, step=5,
            help="Cuánto quieres crecer sobre la proyección de tendencia base.",
        )
    with col_b:
        horizon = st.selectbox(
            "Horizonte de proyección (días)",
            options=[7, 14, 30],
            index=1,
        )
    with col_c:
        client_name = st.text_input("Nombre del cliente (para el Excel)", value="")

    st.markdown("---")

    file_br = st.file_uploader(
        "BR Diario (.csv o .xlsx) — requerido",
        type=["csv", "xlsx"],
        key="forecast_br",
        help="Business Report > By Date > Sales and Traffic. Mínimo 14 días.",
    )
    ad_account = _render_ad_account()

    if not file_br:
        _render_empty_state()
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return

    br_bytes = file_br.getvalue()
    try:
        history = _parse_br_daily(br_bytes, file_br.name)
    except Exception as e:  # an unreadable report must end in a message, not a traceback
        log.warning("PPC Forecast: business report %s unreadable: %s", file_br.name, e)
        st.error(f"Error al leer BR Diario: {e}")
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return

    n_days = len(history)
    if n_days < 7:
        st.error("El archivo tiene menos de 7 días de datos. Se necesitan al menos 7 días para proyectar.")
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return
    if n_days < 14:
        st.warning(f"Sólo hay {n_days} días de datos. La proyección será menos precisa. Se recomiendan al menos 14 días.")

    # ── Botón principal ──────────────────────────────────────────────────────
    signature = _inputs_signature(br_bytes, horizon, target_growth)
    if st.button("Generar Forecast", type="primary", key="forecast_run"):
        st.session_state[_GENERATED_FOR_KEY] = signature
    if st.session_state.get(_GENERATED_FOR_KEY) != signature:
        st.info(f"Archivo cargado: {file_br.name} ({n_days} días). Haz clic en 'Generar Forecast' para continuar.")
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return

    forecast = forecast_sales(history, horizon, target_growth)
    ads = _read_ads(ad_account, history)
    forecast_tab, analysis_tab = st.tabs(["📈 Forecast", "🤖 Análisis IA"])
    with forecast_tab:
        _render_forecast(forecast, n_days, ads, history, client_name)
    with analysis_tab:
        _render_ai_tab(history, forecast, ads, subject=ads.account or client_name.strip() or file_br.name,
                       data_signature=_data_signature(br_bytes, ads.profile_id))


def _how_to_use():
    with st.expander("❓ ¿Cómo usar este módulo?", expanded=False):
        col1, col2, col3 = st.columns(3)
        with col1:
            st.markdown("**🎯 Para qué sirve**")
            st.caption("Proyección de ventas con una tendencia lineal y la diferencia de fin de semana, ajustadas "
                       "juntas. Recomendación de budget con el TACoS de la cuenta.")
        with col2:
            st.markdown("**📂 De dónde salen los datos**")
            st.caption("BR Diario (By Date → Sales and Traffic) mínimo 14 días, ideal 30+. La cuenta de Amazon Ads "
                       "del BR (opcional) separa las ventas de ads de las orgánicas y estima el spend.")
        with col3:
            st.markdown("**➡️ Siguiente paso**")
            st.caption("Ajustar budgets en Campaign Manager según escenario elegido y revisar Account Pulse (M21) semanalmente.")
        st.markdown("**▶️ Pasos:**")
        st.markdown(
            "1. Configurá crecimiento objetivo + horizonte (7/14/30 días)\n"
            "2. Subí el BR diario (mínimo 14 días)\n"
            "3. Elegí la cuenta y el país del BR para el desglose orgánico vs paid (opcional)\n"
            "4. Tocá Generar Forecast: gráfico histórico + proyección, desglose y spend estimado\n"
            "5. Análisis IA: qué tanto confiar en la proyección y qué hacer esta semana\n"
            "6. Descargá el Excel con 2 hojas (resumen + proyección diaria)"
        )


def _render_empty_state():
    st.markdown(
        "<div style='text-align:center;padding:3rem 1rem;border:2px dashed #DDD;"
        "border-radius:12px;margin:1rem 0;'>"
        "<div style='font-size:2.5rem;margin-bottom:0.5rem;'>📂</div>"
        "<div style='font-size:0.95rem;color:#666;font-weight:600;'>Sube el BR Diario para comenzar.</div>"
        "<div style='font-size:0.78rem;color:#999;margin-top:0.3rem;'>"
        "Arrastrá o hacé click en el uploader de arriba</div>"
        "</div>",
        unsafe_allow_html=True,
    )


def _render_ad_account() -> AdAccountChoice:
    """Mounts the account block; returns the chosen profile as the campaign sync sees it, or why there is none."""
    search_term_source._keep_choices(KEY_PREFIX)
    profiles = search_term_source._available_profiles()
    if not profiles:
        st.caption(NO_ACCOUNTS_NOTE)
        return AdAccountChoice(no_ads_reason=MISSING_NO_ACCOUNTS)

    groups = group_by_label(profiles)
    card_key = picker_key(KEY_PREFIX, "card")
    st.markdown(search_term_source._card_css(card_key), unsafe_allow_html=True)
    with st.container(border=True, key=card_key):
        header = st.empty()
        account_col, country_col = st.columns([2.2, 1.3])
        account = _choose_account(account_col, list(groups))
        if account is None:
            header.markdown(_block_header("idle", "Sin cuenta elegida"), unsafe_allow_html=True)
            st.caption(CHOOSE_ACCOUNT_NOTE)
            return AdAccountChoice(no_ads_reason=MISSING_NOT_CHOSEN)

        countries = country_labels(groups[account])
        profile_id = search_term_source._resolve_choice(KEY_PREFIX, "profile", list(countries), fallback=None)
        country_col.segmented_control("País", options=list(countries), format_func=countries.get,
                                      key=picker_key(KEY_PREFIX, "profile"))
        option = next(profile for profile in groups[account] if profile.profile_id == profile_id)
        now = datetime.now(timezone.utc)
        try:
            latest_job, completed = campaign_source._load_campaign_sync(option.profile_id)
        except ReportReadError as exc:
            log.warning("PPC Forecast: campaign sync unreadable for profile %s: %s", option.profile_id, exc)
            header.markdown(_block_header("err", "No se pudo leer"), unsafe_allow_html=True)
            st.error(f"{exc} {UNREADABLE_NOTE}")
            return AdAccountChoice(no_ads_reason=MISSING_UNREADABLE)

        synced = campaign_sync_view(option, completed)
        header.markdown(_block_header(*campaign_source.campaign_pill(synced, latest_job, now)),
                        unsafe_allow_html=True)
        if option.status == PROFILE_NEEDS_REAUTH:
            st.warning(NEEDS_REAUTH_MESSAGE)
        if synced.data_through is None:
            if latest_job is None:
                st.info(NO_CAMPAIGN_DATA_NOTE)
            elif source_state(synced, latest_job, now) == STATE_FIRST_LOAD_FAILED:
                st.error(FIRST_LOAD_FAILED_NOTE)
            else:
                st.info(FIRST_LOAD_NOTE)
            return AdAccountChoice(no_ads_reason=MISSING_NOT_SYNCED)
        info_line = st.empty()
        info_line.markdown(_synced_line(synced), unsafe_allow_html=True)
    return AdAccountChoice(profile=synced, info_line=info_line)


def _choose_account(column, accounts: list[str]) -> str | None:
    """The chosen account, never a default one: the Business Report does not say whose it is."""
    key = picker_key(KEY_PREFIX, "account")
    if st.session_state.get(key) not in accounts:
        st.session_state[key] = None
    return column.selectbox("Cuenta", accounts, index=None, placeholder=ACCOUNT_PLACEHOLDER, key=key)


def _block_header(kind: str, label: str) -> str:
    return palette.band_header_html(title=ADS_BLOCK_TITLE, tag=ADS_BLOCK_TAG,
                                    right=palette.status_pill_html(kind, html.escape(label)))


def _account_label(option: ProfileOption) -> str:
    return f"{option.label} · {option.country_code}" if option.country_code else option.label


def _synced_line(profile: ProfileOption) -> str:
    parts = [html.escape(f"Campañas sincronizadas: {date_range_label(profile.data_from, profile.data_through)}")]
    if profile.currency_code:
        parts.append(palette.marketplace_chip_html(html.escape(profile.currency_code)))
    return search_term_source._muted_line_html(parts)


def _covered_line(split: PaidSplit) -> str:
    parts = [html.escape(date_range_label(split.start, split.end)),
             html.escape(f"{split.covered_days} de {split.history_days} días del BR con datos de ads")]
    if split.currency_code:
        parts.append(palette.marketplace_chip_html(html.escape(split.currency_code)))
    parts.append(html.escape(" · ".join(split.products) if split.products else "Sin campañas con actividad"))
    return search_term_source._muted_line_html(parts)


def _read_ads(choice: AdAccountChoice, history: pd.DataFrame) -> AdsReading:
    """The chosen account's ads over the Business Report's days, or why there are none."""
    profile = choice.profile
    if profile is None:
        return AdsReading(no_ads_reason=choice.no_ads_reason)
    identity = dict(account=_account_label(profile), profile_id=profile.profile_id, country_code=profile.country_code)
    first, last = history["_date"].min().date(), history["_date"].max().date()
    window = covered_window(first, last, profile.data_from, profile.data_through)
    if window is None:
        reason = MISSING_NO_SHARED_DAYS.format(first=short_date(first), last=short_date(last),
                                                synced_from=short_date(profile.data_from),
                                                synced_through=short_date(profile.data_through))
        choice.info_line.caption(f"Sin desglose: {reason}.")
        return AdsReading(**identity, currency_code=profile.currency_code, no_ads_reason=reason)
    try:
        series = _load_ad_series(profile, *window)
    except ReportReadError as exc:
        log.warning("PPC Forecast: ad totals unreadable for profile %s: %s", profile.profile_id, exc)
        choice.info_line.error(f"{exc} {UNREADABLE_NOTE}")
        return AdsReading(**identity, currency_code=profile.currency_code, no_ads_reason=MISSING_UNREADABLE)
    split = paid_split(history, series)
    choice.info_line.markdown(_covered_line(split), unsafe_allow_html=True)
    return AdsReading(**identity, currency_code=split.currency_code, series=series, split=split)


# The profile carries its last campaign sync: a new sync is a new read, never an old one served again.
@st.cache_data(ttl=AD_TOTALS_TTL_SECONDS, max_entries=16, show_spinner="Leyendo las ventas de ads de la cuenta…")
def _load_ad_series(profile: ProfileOption, start: date, end: date) -> ProductSeries:
    rest = search_term_source._open_rest()
    if rest is None:
        raise ReportReadError(NO_DATABASE_MESSAGE)
    return campaign_totals.daily_totals(rest, profile, start, end)


def _render_forecast(forecast: SalesForecast, n_days: int, ads: AdsReading, history: pd.DataFrame,
                     client_name: str):
    show = partial(money, currency_code=ads.currency_code)
    trend = forecast.trend
    needed_spend = spend_for_target(ads.split, forecast.sales_with_growth)

    # ── SECCIÓN 1: Métricas históricas ───────────────────────────────────────
    st.subheader("Métricas Históricas")
    ratio = trend.weekend_ratio
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.markdown(kpi_card("Días de datos", str(n_days)), unsafe_allow_html=True)
    with c2:
        st.markdown(kpi_card("Ventas promedio/día", show(forecast.average_daily_sales)), unsafe_allow_html=True)
    with c3:
        st.markdown(kpi_card("Tendencia", _signed_money(trend.slope, ads.currency_code)), unsafe_allow_html=True)
    with c4:
        st.markdown(kpi_card("Ratio Finde/Laboral", f"{ratio:.2f}x" if ratio is not None else "—"),
                    unsafe_allow_html=True)

    # ── SECCIÓN 2: Gráfico ───────────────────────────────────────────────────
    st.subheader("Tendencia Histórica + Proyección")

    hist_dates = [d.strftime("%Y-%m-%d") for d in history["_date"]]
    hist_sales = history["_sales"].tolist()
    proj_dates = forecast.projection[DATE].tolist()
    proj_sales = forecast.projection[PROJECTED_SALES].tolist()

    all_dates = hist_dates + proj_dates
    chart_df = pd.DataFrame(
        {
            "Ventas Reales ($)": hist_sales + [None] * forecast.horizon,
            "Ventas Proyectadas ($)": [None] * len(hist_sales) + proj_sales,
        },
        index=all_dates,
    )
    # Añadir el último punto histórico como primer punto proyectado para continuidad visual
    if hist_sales:
        chart_df.loc[hist_dates[-1], "Ventas Proyectadas ($)"] = hist_sales[-1]

    st.line_chart(chart_df)

    # ── SECCIÓN 3: Resumen de proyección ─────────────────────────────────────
    st.subheader(f"Resumen — Próximos {forecast.horizon} días")
    c1, c2, c3 = st.columns(3)
    c1.metric(
        f"Ventas proyectadas ({forecast.horizon}d)",
        show(forecast.projected_sales),
    )
    c2.metric(
        f"Con crecimiento +{forecast.target_growth}%",
        show(forecast.sales_with_growth),
        delta=f"+{show(forecast.sales_with_growth - forecast.projected_sales)}",
    )
    c3.metric(
        "Spend estimado para objetivo",
        show(needed_spend) if needed_spend is not None else "—",
        help=BUDGET_HELP if needed_spend is not None else f"Sin dato: {ads.no_ads_reason}.",
    )

    # ── SECCIÓN 4: Tabla de proyección diaria ────────────────────────────────
    with st.expander("Ver tabla de proyección diaria"):
        st.dataframe(
            forecast.projection,
            use_container_width=True,
            column_config={
                PROJECTED_SALES: st.column_config.NumberColumn(
                    PROJECTED_SALES, format=f"{currency_symbol(ads.currency_code)}%.2f"
                ),
            },
        )

    # ── SECCIÓN 5: Desglose orgánico vs paid ─────────────────────────────────
    _render_split(ads, show)

    # ── Export Excel ─────────────────────────────────────────────────────────
    st.markdown("---")
    st.subheader("Exportar")
    try:
        excel_buf = _build_forecast_excel(forecast, n_days, needed_spend, client_name, ads.currency_code)
        fname_out = f"PPC_Forecast_{client_name.replace(' ', '_') + '_' if client_name else ''}{forecast.horizon}d.xlsx"
        st.download_button(
            label="Descargar Forecast Excel",
            data=excel_buf,
            file_name=fname_out,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="forecast_dl",
        )
    except Exception as e:
        st.error(f"Error al generar Excel: {e}")


def _render_split(ads: AdsReading, show):
    st.subheader("Desglose Orgánico vs Paid")
    split = ads.split
    if split is None:
        st.caption(f"Sin desglose: {ads.no_ads_reason}.")
        return
    if split.ads_exceed_br:
        # Streamlit renders the text between two bare $ as LaTeX.
        st.warning(ADS_EXCEED_BR_WARNING.format(ad_sales=show(split.ad_sales).replace("$", "\\$"),
                                                br_sales=show(split.br_sales).replace("$", "\\$")))
    k1, k2, k3, k4 = st.columns(4)
    with k1:
        st.markdown(kpi_card("Spend de ads", show(split.ad_spend)), unsafe_allow_html=True)
    with k2:
        st.markdown(kpi_card("Ventas de ads", show(split.ad_sales)), unsafe_allow_html=True)
    with k3:
        st.markdown(kpi_card("ACoS", f"{split.acos:.1f}%" if split.acos is not None else "—"), unsafe_allow_html=True)
    with k4:
        st.markdown(kpi_card("Ventas orgánicas estimadas", show(split.organic_sales)), unsafe_allow_html=True)
    st.caption(SPLIT_CAPTION.format(covered=split.covered_days, total=split.history_days,
                                    window=date_range_label(split.start, split.end),
                                    products=" · ".join(split.products) or "sin campañas con actividad",
                                    attribution=split.attribution_days))
    if split.paid_share is not None:
        split_df = pd.DataFrame(
            {
                "Canal": ["Paid (Ads)", "Orgánico"],
                "Ventas": [show(split.ad_sales), show(split.organic_sales)],
                "% del Total": [round(split.paid_share, 1), round(100.0 - split.paid_share, 1)],
            }
        )
        st.dataframe(split_df, use_container_width=True, hide_index=True)


def _render_ai_tab(history: pd.DataFrame, forecast: SalesForecast, ads: AdsReading, *, subject: str,
                   data_signature: str):
    from ai.agents.ppc_forecast import chat_document
    from ai.config import AI_ENABLED
    from core import ai_tab

    lang = ai_tab.app_language()
    texts = _AI_TEXTS.get(lang, _AI_TEXTS["es"])
    st.subheader(texts["title"])
    st.caption(texts["caption"])
    if not AI_ENABLED:
        st.caption(texts["disabled"])
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return
    payload = build_analysis_input(history, forecast, split=ads.split, ads=ads.series, account=ads.account,
                                   ads_note=ads.no_ads_reason, currency_code=ads.currency_code, lang=lang)
    labels = ai_tab.ai_labels(lang, texts)
    st.markdown(ai_tab.AI_CSS, unsafe_allow_html=True)
    analysis = ai_tab.resolve_analysis(slug=ANALYSIS_MODULE, payload=payload, file_signature=data_signature,
                                       labels=labels, auto_fire=False)
    if analysis is not None:
        # A stale analysis read other figures: its rows show the ones it read.
        records = ai_tab.records_for_render(ANALYSIS_MODULE, analysis, payload, _ai_records(forecast, ads, labels))
        ai_tab.render_analysis(analysis, slug=ANALYSIS_MODULE, labels=labels,
                               render_result=partial(_render_ai_result, records=records, labels=labels))
    ai_tab.publish_analysis_to_chat(
        ANALYSIS_MODULE, analysis, payload, module_label=MODULE_LABEL, subject=subject,
        reading=lambda finished: chat_document.reading_text(finished.result),
        country_code=ads.country_code, profile_id=ads.profile_id)


def _ai_records(forecast: SalesForecast, ads: AdsReading, labels: dict) -> list[dict]:
    """The module's figures the AI reads, as the opinion table names them."""
    show = partial(money, currency_code=ads.currency_code)
    records = [{"tema": "PROYECCION", "item": labels["projection_item"].format(days=forecast.horizon),
                "metrics": [show(forecast.projected_sales),
                            labels["with_growth"].format(growth=forecast.target_growth,
                                                         amount=show(forecast.sales_with_growth))]}]
    needed_spend = spend_for_target(ads.split, forecast.sales_with_growth)
    if needed_spend is not None:
        records.append({"tema": "PRESUPUESTO", "item": labels["budget_item"], "metrics": [show(needed_spend)]})
    return records


def _render_ai_result(result, analysis, *, records, labels):
    from core import ai_tab

    rows = forecast_ai_rows(result.get("lecturas") or [], records, labels)
    warnings = sum(1 for row in rows if row["warning"])
    st.markdown(ai_tab.ai_chips_html(warnings, labels["counts"].format(n=len(rows)), analysis.elapsed, labels),
                unsafe_allow_html=True)
    st.markdown(ai_tab.synthesis_html(result.get("synthesis") or {}, labels), unsafe_allow_html=True)
    if rows:
        st.markdown(ai_tab.opinion_table_html(rows, labels["table_title"], labels, _confidence_colors(labels)),
                    unsafe_allow_html=True)


def forecast_ai_rows(readings: list, records: list, labels: dict) -> list[dict]:
    """Display rows for the opinion table: the module's figure, its value and the AI's confidence in it."""
    by_topic = {record["tema"]: record for record in records}
    rows = []
    for reading in readings:
        record = by_topic.get(reading.get("tema"))
        if record is None:
            continue
        level = str(reading.get("confianza", "")).lower()
        rows.append({
            "item": record["item"],
            "metrics": record["metrics"],
            "badges": [labels["levels"].get(level, level)],
            "warning": reading.get("advertencia") or "",
            "reasoning": reading.get("razon", ""),
        })
    return rows


def _confidence_colors(labels: dict) -> dict:
    return dict(zip((labels["levels"][level] for level in ("alta", "media", "baja")), _CONFIDENCE_COLORS))


def _digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()[:16]


def _inputs_signature(br_bytes: bytes, horizon: int, target_growth: int) -> str:
    """What "Generar Forecast" computed from: the forecast stays on screen until it changes."""
    return f"{_digest(br_bytes)}|{horizon}|{target_growth}"


def _data_signature(br_bytes: bytes, profile_id: str) -> str:
    """What changes when the Business Report or the account under the analysis do, not when a parameter does."""
    return f"{_digest(br_bytes)}|{profile_id}"
