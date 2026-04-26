"""
Módulo: Variation Builder (M26)
Sección: Account Manager
Input: Template flat file de Amazon (.xlsm) + datos parent + N children
Output: .xlsm con 1 parent + N children agrupados por variation_theme

Genera la estructura parent/child para Amazon (Pet Food u otra categoría
basada en fptcustom). Lee el template descargado de Seller Central, llena
las filas en la hoja Plantilla preservando macros, data validations y resto
de hojas, y devuelve el archivo listo para subir vía Add Products via Upload.
"""
import io
import re
from typing import Optional

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from core.helpers import kpi_card

# ── Constants ───────────────────────────────────────────────────────────
_TEMPLATE_SHEET = "Plantilla"
_VALUES_SHEET = "Valores válidos"
_HEADER_GROUP_ROW = 1
_HEADER_DISPLAY_ROW = 2
_HEADER_FIELD_ROW = 3
_DATA_START_ROW = 4

# Field names que necesitamos referenciar (no son hardcoded por columna,
# se mapean en runtime contra fila 3 de Plantilla)
_F_FEED_PRODUCT_TYPE = "feed_product_type"
_F_ITEM_SKU = "item_sku"
_F_BRAND_NAME = "brand_name"
_F_UPDATE_DELETE = "update_delete"
_F_EXTERNAL_PRODUCT_ID = "external_product_id"
_F_EXTERNAL_PRODUCT_ID_TYPE = "external_product_id_type"
_F_ITEM_NAME = "item_name"
_F_MANUFACTURER = "manufacturer"
_F_PRODUCT_DESCRIPTION = "product_description"
_F_GTIN_EXEMPTION = "gtin_exemption_reason"
_F_RECOMMENDED_BROWSE = "recommended_browse_nodes"
_F_AGE_RANGE = "age_range_description"
_F_MAIN_IMAGE = "main_image_url"
_F_OTHER_IMAGES = [f"other_image_url{i}" for i in range(1, 9)]
_F_PARENT_CHILD = "parent_child"
_F_PARENT_SKU = "parent_sku"
_F_RELATIONSHIP_TYPE = "relationship_type"
_F_VARIATION_THEME = "variation_theme"
_F_BULLET_POINTS = [f"bullet_point{i}" for i in range(1, 6)]
_F_BREED = "breed_recommendation"
_F_TARGET_AUDIENCE = "target_audience_keywords"
_F_FLAVOR_NAME = "flavor_name"
_F_SIZE_NAME = "size_name"
_F_PATTERN_NAME = "pattern_name"
_F_SCENT_NAME = "scent_name"
_F_UNIT_COUNT = "unit_count"
_F_UNIT_COUNT_TYPE = "unit_count_type"
_F_ITEM_FORM = "item_form"
_F_PACKAGE_LENGTH = "package_length"
_F_PACKAGE_WIDTH = "package_width"
_F_PACKAGE_HEIGHT = "package_height"
_F_PACKAGE_WEIGHT = "package_weight"
_F_PACKAGE_LENGTH_UOM = "package_length_unit_of_measure"
_F_PACKAGE_WIDTH_UOM = "package_width_unit_of_measure"
_F_PACKAGE_HEIGHT_UOM = "package_height_unit_of_measure"
_F_PACKAGE_WEIGHT_UOM = "package_weight_unit_of_measure"
_F_ITEM_WEIGHT = "item_weight"
_F_ITEM_WEIGHT_UOM = "item_weight_unit_of_measure"
_F_COUNTRY_OF_ORIGIN = "country_of_origin"
_F_CONDITION_TYPE = "condition_type"
_F_PRODUCT_TAX_CODE = "product_tax_code"

# MX marketplace (A1AM78C64UM0Y8)
_MX_MARKETPLACE_ID = "A1AM78C64UM0Y8"
_F_FULFILLMENT_QTY = "fulfillment_availability#1.quantity"
_F_FULFILLMENT_CHANNEL = "fulfillment_availability#1.fulfillment_channel_code"
_F_LEAD_TIME = "fulfillment_availability#1.lead_time_to_ship_max_days"

# Variation theme → atributo(s) que cada child debe llenar
_THEME_TO_ATTRS = {
    "Sabor": [_F_FLAVOR_NAME],
    "Nombre del Tamano": [_F_SIZE_NAME],
    "Nombre del Patron": [_F_PATTERN_NAME],
    "Scent": [_F_SCENT_NAME],
    "Nombre del Patron y del Tamano": [_F_PATTERN_NAME, _F_SIZE_NAME],
    "Tamano del Sabor": [_F_SIZE_NAME, _F_FLAVOR_NAME],
    "FlavorName-SizeName": [_F_FLAVOR_NAME, _F_SIZE_NAME],
}

# Label legible por atributo (para columnas de la tabla de children)
_ATTR_LABELS = {
    _F_FLAVOR_NAME: "Sabor",
    _F_SIZE_NAME: "Tamaño",
    _F_PATTERN_NAME: "Patrón",
    _F_SCENT_NAME: "Aroma",
}


# ── Parsers (cached) ────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def _parse_template(file_bytes: bytes) -> dict:
    """
    Parsea el .xlsm: mapea field_name → col_idx, extrae metadata.
    Devuelve dict con field_to_col, col_to_field, signature, version, template_type, sheets.
    """
    bio = io.BytesIO(file_bytes)
    wb = load_workbook(bio, keep_vba=True, data_only=True, read_only=False)

    if _TEMPLATE_SHEET not in wb.sheetnames:
        raise ValueError(
            f"El archivo no tiene la hoja '{_TEMPLATE_SHEET}'. "
            f"Verificá que sea un template flat file de Amazon."
        )

    ws = wb[_TEMPLATE_SHEET]

    # Metadata fila 1: TemplateType=fptcustom, Version=2026.0426, TemplateSignature=...
    # Cada token suele estar en su propia celda al inicio de la fila 1.
    meta = {}
    for c in range(1, min(ws.max_column + 1, 10)):
        v = ws.cell(row=1, column=c).value
        if not v:
            continue
        for tok in re.split(r"[\s,]+", str(v)):
            if "=" in tok:
                k, val = tok.split("=", 1)
                meta[k.strip()] = val.strip()

    # Mapeo field_name (fila 3) → col_idx
    field_to_col: dict[str, int] = {}
    col_to_field: dict[int, str] = {}
    display_to_col: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        fn = ws.cell(row=_HEADER_FIELD_ROW, column=c).value
        dn = ws.cell(row=_HEADER_DISPLAY_ROW, column=c).value
        if fn:
            fn_s = str(fn).strip()
            field_to_col[fn_s] = c
            col_to_field[c] = fn_s
            if dn:
                display_to_col[str(dn).strip()] = c

    return {
        "field_to_col": field_to_col,
        "col_to_field": col_to_field,
        "display_to_col": display_to_col,
        "template_type": meta.get("TemplateType", ""),
        "version": meta.get("Version", ""),
        "signature": meta.get("TemplateSignature", ""),
        "sheets": list(wb.sheetnames),
        "max_col": ws.max_column,
    }


@st.cache_data(show_spinner=False)
def _extract_valid_values(file_bytes: bytes) -> dict:
    """
    Lee 'Valores válidos': cada fila es 'Nombre del campo - [ producto ]' en col 2,
    seguido de los valores válidos en cols 3+.
    Devuelve dict { field_label: [valores] }.
    """
    bio = io.BytesIO(file_bytes)
    wb = load_workbook(bio, keep_vba=True, data_only=True, read_only=True)
    if _VALUES_SHEET not in wb.sheetnames:
        return {}

    ws = wb[_VALUES_SHEET]
    out: dict[str, list[str]] = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if len(row) < 3:
            continue
        label = row[1]
        if not label:
            continue
        # quitar el sufijo " - [ petfood ]"
        label_clean = re.sub(r"\s*-\s*\[[^\]]*\]\s*$", "", str(label)).strip()
        values = [str(v).strip() for v in row[2:] if v is not None and str(v).strip()]
        if values:
            out[label_clean] = values
    return out


# ── Render helpers ──────────────────────────────────────────────────────
def _header() -> None:
    st.markdown(
        "<div style='display:flex;align-items:center;gap:0.75rem;margin-bottom:0.25rem;'>"
        "<span style='font-size:2rem;'>🧬</span>"
        "<div><div style='font-size:1.3rem;font-weight:700;'>Variation Builder</div>"
        "<div style='font-size:0.82rem;color:#888;'>"
        "Generá la estructura parent + children para Amazon flat file</div>"
        "</div></div>",
        unsafe_allow_html=True,
    )
    st.divider()


def _empty_state() -> None:
    st.markdown(
        "<div style='border:2px dashed #FFD9B3;border-radius:12px;padding:2rem;"
        "text-align:center;background:#FFF3E0;margin-top:1rem;'>"
        "<div style='font-size:1.5rem;'>📂</div>"
        "<div style='font-weight:600;margin-top:0.5rem;'>Subí el template flat file (.xlsm)</div>"
        "<div style='font-size:0.82rem;color:#888;margin-top:0.25rem;'>"
        "Seller Central → Catálogo → Añadir productos a través de subida → "
        "Inventory Files → Pet Food (o categoría aplicable). Descargá el template vacío.</div>"
        "</div>",
        unsafe_allow_html=True,
    )


def _theme_attr_fields(theme: str) -> list[str]:
    """Devuelve los field_names de los atributos requeridos por el variation_theme."""
    return _THEME_TO_ATTRS.get(theme, [])


def _validate_inputs(
    parent: dict,
    children: pd.DataFrame,
    theme: str,
    field_to_col: dict,
) -> list[str]:
    """Devuelve lista de errores. Vacía = OK para generar."""
    errors: list[str] = []

    # Parent
    if not parent.get(_F_ITEM_SKU, "").strip():
        errors.append("Parent: SKU obligatorio.")
    if not parent.get(_F_ITEM_NAME, "").strip():
        errors.append("Parent: Nombre del producto obligatorio.")
    if not parent.get(_F_BRAND_NAME, "").strip():
        errors.append("Parent: Marca obligatoria.")
    if not parent.get(_F_MAIN_IMAGE, "").strip():
        errors.append("Parent: URL de imagen principal obligatoria.")

    # Theme
    if not theme:
        errors.append("Variation Theme obligatorio.")
        return errors  # sin theme no puedo validar children
    attrs = _theme_attr_fields(theme)
    for a in attrs:
        if a not in field_to_col:
            errors.append(
                f"El template no tiene la columna '{a}' requerida para el theme '{theme}'. "
                f"Es posible que el theme no aplique a esta categoría."
            )

    # Children
    if children is None or len(children) < 2:
        errors.append("Mínimo 2 children obligatorios para crear variations.")
        return errors

    # SKUs únicos (parent ≠ children, todos los children distintos)
    parent_sku = parent.get(_F_ITEM_SKU, "").strip()
    skus_ch = [str(s).strip() for s in children["item_sku"].fillna("").tolist()]
    skus_ch_filled = [s for s in skus_ch if s]
    if len(skus_ch_filled) != len(children):
        errors.append("Children: todos los SKUs son obligatorios.")
    if len(set(skus_ch_filled)) != len(skus_ch_filled):
        errors.append("Children: hay SKUs duplicados.")
    if parent_sku in skus_ch_filled:
        errors.append("Children: el SKU del parent no puede repetirse en un child.")

    # Valor del theme único por child (al menos uno de los atributos debe diferir entre children)
    if attrs:
        # construir tupla del theme por child
        tuples = []
        for _, r in children.iterrows():
            t = tuple(str(r.get(a, "")).strip() for a in attrs)
            tuples.append(t)
        if any(any(not v for v in t) for t in tuples):
            errors.append(
                f"Children: cada child debe tener {', '.join(_ATTR_LABELS.get(a, a) for a in attrs)} cargado."
            )
        elif len(set(tuples)) != len(tuples):
            errors.append(
                f"Children: la combinación de {', '.join(_ATTR_LABELS.get(a, a) for a in attrs)} "
                "tiene que ser única en cada child."
            )

    # Precio y cantidad
    for i, r in children.iterrows():
        try:
            p = float(r.get("price_mx", 0) or 0)
            if p <= 0:
                errors.append(f"Child fila {i+1}: precio MXN debe ser > 0.")
        except (TypeError, ValueError):
            errors.append(f"Child fila {i+1}: precio MXN inválido.")
        try:
            q = int(r.get("quantity", 0) or 0)
            if q < 0:
                errors.append(f"Child fila {i+1}: cantidad no puede ser negativa.")
        except (TypeError, ValueError):
            errors.append(f"Child fila {i+1}: cantidad inválida.")

    return errors


def _build_parent_row(parent: dict, theme: str, field_to_col: dict) -> dict:
    """Devuelve {col_idx: value} para escribir la fila parent."""
    row: dict[int, object] = {}

    def _set(field: str, value):
        c = field_to_col.get(field)
        if c is not None and value not in (None, ""):
            row[c] = value

    _set(_F_FEED_PRODUCT_TYPE, parent.get(_F_FEED_PRODUCT_TYPE, "petfood"))
    _set(_F_ITEM_SKU, parent.get(_F_ITEM_SKU, "").strip())
    _set(_F_BRAND_NAME, parent.get(_F_BRAND_NAME, "").strip())
    _set(_F_UPDATE_DELETE, "Actualizar")
    # Parent NO requiere external_product_id; se exenta
    _set(_F_GTIN_EXEMPTION, parent.get(_F_GTIN_EXEMPTION, "Pieza"))
    _set(_F_ITEM_NAME, parent.get(_F_ITEM_NAME, "").strip())
    _set(_F_MANUFACTURER, parent.get(_F_MANUFACTURER, parent.get(_F_BRAND_NAME, "")).strip())
    _set(_F_PRODUCT_DESCRIPTION, parent.get(_F_PRODUCT_DESCRIPTION, "").strip())
    _set(_F_AGE_RANGE, parent.get(_F_AGE_RANGE, ""))
    _set(_F_MAIN_IMAGE, parent.get(_F_MAIN_IMAGE, "").strip())
    # Imágenes secundarias del parent
    for i, key in enumerate(_F_OTHER_IMAGES):
        _set(key, (parent.get("other_images", []) or [None] * 8)[i] if i < len(parent.get("other_images", [])) else None)

    _set(_F_PARENT_CHILD, "Parent")
    _set(_F_RELATIONSHIP_TYPE, "")  # parent NO lleva relationship_type
    _set(_F_VARIATION_THEME, theme)

    # Bullets
    bullets = parent.get("bullets", []) or []
    for i, key in enumerate(_F_BULLET_POINTS):
        if i < len(bullets) and bullets[i]:
            _set(key, str(bullets[i]).strip())

    _set(_F_BREED, parent.get(_F_BREED, ""))
    _set(_F_TARGET_AUDIENCE, parent.get(_F_TARGET_AUDIENCE, ""))
    _set(_F_ITEM_FORM, parent.get(_F_ITEM_FORM, ""))
    _set(_F_COUNTRY_OF_ORIGIN, parent.get(_F_COUNTRY_OF_ORIGIN, ""))
    _set(_F_CONDITION_TYPE, parent.get(_F_CONDITION_TYPE, "Nuevo"))

    return row


def _build_child_row(
    child: dict,
    parent: dict,
    theme: str,
    field_to_col: dict,
) -> dict:
    """Devuelve {col_idx: value} para escribir una fila child."""
    row: dict[int, object] = {}

    def _set(field: str, value):
        c = field_to_col.get(field)
        if c is not None and value not in (None, ""):
            row[c] = value

    parent_sku = parent.get(_F_ITEM_SKU, "").strip()

    _set(_F_FEED_PRODUCT_TYPE, parent.get(_F_FEED_PRODUCT_TYPE, "petfood"))
    _set(_F_ITEM_SKU, str(child.get(_F_ITEM_SKU, "")).strip())
    _set(_F_BRAND_NAME, parent.get(_F_BRAND_NAME, "").strip())
    _set(_F_UPDATE_DELETE, "Actualizar")

    # External product id (UPC/EAN/GTIN/ASIN). Si no hay, exenta.
    upc = str(child.get("external_product_id", "")).strip()
    upc_type = str(child.get("external_product_id_type", "")).strip()
    if upc and upc_type:
        _set(_F_EXTERNAL_PRODUCT_ID, upc)
        _set(_F_EXTERNAL_PRODUCT_ID_TYPE, upc_type)
    else:
        _set(_F_GTIN_EXEMPTION, parent.get(_F_GTIN_EXEMPTION, "Pieza"))

    # Item name del child: usa el específico si lo dieron, sino concatena parent + atributos
    child_name = str(child.get(_F_ITEM_NAME, "")).strip()
    if not child_name:
        attrs = _theme_attr_fields(theme)
        suffix = " - ".join(str(child.get(a, "")).strip() for a in attrs if child.get(a))
        child_name = f"{parent.get(_F_ITEM_NAME, '').strip()} - {suffix}".strip(" -")
    _set(_F_ITEM_NAME, child_name)

    _set(_F_MANUFACTURER, parent.get(_F_MANUFACTURER, parent.get(_F_BRAND_NAME, "")).strip())
    _set(_F_PRODUCT_DESCRIPTION, parent.get(_F_PRODUCT_DESCRIPTION, "").strip())
    _set(_F_AGE_RANGE, parent.get(_F_AGE_RANGE, ""))

    # Imagen principal del child: específica si la dio, sino la del parent
    child_img = str(child.get(_F_MAIN_IMAGE, "")).strip()
    _set(_F_MAIN_IMAGE, child_img or parent.get(_F_MAIN_IMAGE, "").strip())

    _set(_F_PARENT_CHILD, "Child")
    _set(_F_PARENT_SKU, parent_sku)
    _set(_F_RELATIONSHIP_TYPE, "variation")
    _set(_F_VARIATION_THEME, theme)

    # Atributos del theme
    for a in _theme_attr_fields(theme):
        _set(a, str(child.get(a, "")).strip())

    # Bullets — heredan del parent
    bullets = parent.get("bullets", []) or []
    for i, key in enumerate(_F_BULLET_POINTS):
        if i < len(bullets) and bullets[i]:
            _set(key, str(bullets[i]).strip())

    _set(_F_BREED, parent.get(_F_BREED, ""))
    _set(_F_TARGET_AUDIENCE, parent.get(_F_TARGET_AUDIENCE, ""))
    _set(_F_ITEM_FORM, parent.get(_F_ITEM_FORM, ""))
    _set(_F_COUNTRY_OF_ORIGIN, parent.get(_F_COUNTRY_OF_ORIGIN, ""))
    _set(_F_CONDITION_TYPE, parent.get(_F_CONDITION_TYPE, "Nuevo"))

    # Unit count específico del child
    if child.get(_F_UNIT_COUNT):
        _set(_F_UNIT_COUNT, child.get(_F_UNIT_COUNT))
    if child.get(_F_UNIT_COUNT_TYPE):
        _set(_F_UNIT_COUNT_TYPE, child.get(_F_UNIT_COUNT_TYPE))

    # Peso/dimensiones del child
    if child.get(_F_PACKAGE_WEIGHT):
        _set(_F_PACKAGE_WEIGHT, child.get(_F_PACKAGE_WEIGHT))
        _set(_F_PACKAGE_WEIGHT_UOM, child.get(_F_PACKAGE_WEIGHT_UOM, "GR"))
    if child.get(_F_ITEM_WEIGHT):
        _set(_F_ITEM_WEIGHT, child.get(_F_ITEM_WEIGHT))
        _set(_F_ITEM_WEIGHT_UOM, child.get(_F_ITEM_WEIGHT_UOM, "GR"))

    # ── Oferta MX ──
    _set(_F_FULFILLMENT_QTY, int(child.get("quantity", 0) or 0))
    _set(_F_FULFILLMENT_CHANNEL, "DEFAULT")  # vendedor (Merchant fulfilled)
    _set(_F_LEAD_TIME, int(child.get("lead_time_days", 2) or 2))

    # Precio MX (col exacta varía si Amazon cambia la versión, así que buscamos por substring)
    price_mx = float(child.get("price_mx", 0) or 0)
    if price_mx > 0:
        for fn, c in field_to_col.items():
            if "marketplace_id=" + _MX_MARKETPLACE_ID in fn and "our_price" in fn and "value_with_tax" in fn:
                row[c] = price_mx
                break

    return row


# ── Writer (FUERA de render) ────────────────────────────────────────────
def _write_template_with_rows(
    template_bytes: bytes,
    parent_row: dict,
    child_rows: list,
) -> bytes:
    """
    Carga el .xlsm original con keep_vba=True, escribe parent en fila 4
    y children en filas 5+, y devuelve los bytes del archivo resultante.
    Preserva macros, data validations y resto de hojas.
    """
    bio = io.BytesIO(template_bytes)
    wb = load_workbook(bio, keep_vba=True, data_only=False)
    ws = wb[_TEMPLATE_SHEET]

    # Limpiar fila 4 en adelante por si viniera con datos
    for r in range(_DATA_START_ROW, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None

    # Parent en fila 4
    for col, val in parent_row.items():
        ws.cell(row=_DATA_START_ROW, column=col).value = val

    # Children desde fila 5
    for i, child_row in enumerate(child_rows, start=1):
        for col, val in child_row.items():
            ws.cell(row=_DATA_START_ROW + i, column=col).value = val

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _build_preview_df(
    parent_row: dict,
    child_rows: list,
    col_to_field: dict,
) -> pd.DataFrame:
    """
    Convierte dicts {col_idx: value} a un DataFrame legible para preview,
    mostrando solo columnas que tienen datos.
    """
    all_rows = [parent_row] + child_rows
    used_cols = sorted({c for r in all_rows for c in r.keys()})
    headers = [col_to_field.get(c, f"col_{c}") for c in used_cols]

    data = []
    for r in all_rows:
        data.append([r.get(c, "") for c in used_cols])

    df = pd.DataFrame(data, columns=headers)
    df.insert(0, "fila", [f"PARENT (fila {_DATA_START_ROW})"] +
              [f"CHILD {i} (fila {_DATA_START_ROW + i})" for i in range(1, len(child_rows) + 1)])
    return df


# ── Render principal ────────────────────────────────────────────────────
def render() -> None:
    _header()

    # Uploader
    up = st.file_uploader(
        "Template flat file de Amazon (.xlsm)",
        type=["xlsm", "xlsx"],
        key="vb_uploader",
        help="Descargado de Seller Central → Add Products via Upload → Inventory Files",
    )

    if up is None:
        _empty_state()
        return

    template_bytes = up.read()

    try:
        meta = _parse_template(template_bytes)
    except Exception as e:
        st.error(f"❌ No se pudo parsear el template: {e}")
        return

    field_to_col = meta["field_to_col"]
    col_to_field = meta["col_to_field"]

    # Verificar que es un template con soporte de variations
    missing_core = [f for f in (_F_PARENT_CHILD, _F_PARENT_SKU, _F_VARIATION_THEME) if f not in field_to_col]
    if missing_core:
        st.error(
            f"❌ El template no tiene los campos de variation requeridos: {', '.join(missing_core)}. "
            "Verificá que sea un template que soporte parent/child."
        )
        return

    # KPIs del template cargado
    cols = st.columns(4)
    with cols[0]:
        kpi_card("Versión", meta.get("version", "—"))
    with cols[1]:
        kpi_card("Tipo", meta.get("template_type", "—"))
    with cols[2]:
        kpi_card("Hojas", str(len(meta["sheets"])))
    with cols[3]:
        kpi_card("Columnas", f"{meta['max_col']:,}")

    # Valores válidos (themes, brands, etc.)
    valid_values = _extract_valid_values(template_bytes)

    themes_available = list(_THEME_TO_ATTRS.keys())
    # Filtrar solo los que el template realmente soporta (atributos existen como cols)
    themes_available = [
        t for t in themes_available
        if all(a in field_to_col for a in _theme_attr_fields(t))
    ]

    if not themes_available:
        st.error("❌ El template no soporta ninguno de los variation themes conocidos.")
        return

    # Valores válidos pre-extraídos del archivo (para dropdowns)
    valid_brands = valid_values.get("Brand Name", [])
    valid_age_ranges = valid_values.get("Age Range Description", [])
    valid_breeds = valid_values.get("breed-recommendation", [])
    valid_audiences = valid_values.get("Público objetivo", [])
    valid_item_forms = valid_values.get("item-form", [])
    valid_countries = valid_values.get("País de origen", [])
    valid_unit_count_types = valid_values.get("Tipo de recuento de unidades", [])
    valid_weight_uoms = valid_values.get("Package Weight Unit Of Measure", [])
    valid_flavors = valid_values.get("Flavor", [])
    valid_size_maps = valid_values.get("Size Map", [])
    valid_gtin_exemptions = valid_values.get("Motivo de exención del producto", [])

    # ── Tabs ──
    tab_parent, tab_children, tab_preview, tab_download = st.tabs(
        ["👤 Parent", "🧬 Children + Theme", "👀 Preview", "💾 Descargar"]
    )

    # ─── Tab: Parent ───
    with tab_parent:
        st.markdown("##### Datos del producto Parent")
        st.caption(
            "El parent es el SKU contenedor — no se vende, solo agrupa los children. "
            "No lleva precio ni stock."
        )

        c1, c2 = st.columns(2)
        with c1:
            st.session_state.setdefault("vb_p_sku", "")
            st.text_input("SKU del parent *", key="vb_p_sku", placeholder="MARCA-PROD-001")

            brand_default = valid_brands[0] if valid_brands else ""
            if valid_brands:
                st.selectbox("Marca *", options=valid_brands, key="vb_p_brand")
            else:
                st.text_input("Marca *", key="vb_p_brand")

            st.text_input("Manufacturer", key="vb_p_manufacturer", placeholder="(opcional, default = marca)")
            st.text_input("Nombre del producto *", key="vb_p_name", placeholder="Alimento balanceado para perro adulto")

        with c2:
            if valid_age_ranges:
                st.selectbox("Age range", options=[""] + valid_age_ranges, key="vb_p_age")
            else:
                st.text_input("Age range", key="vb_p_age")

            if valid_breeds:
                st.selectbox("Breed recommendation", options=[""] + valid_breeds, key="vb_p_breed")
            else:
                st.text_input("Breed recommendation", key="vb_p_breed")

            if valid_audiences:
                st.selectbox("Target audience", options=[""] + valid_audiences, key="vb_p_audience")
            else:
                st.text_input("Target audience", key="vb_p_audience")

            if valid_item_forms:
                st.selectbox("Item form", options=[""] + valid_item_forms, key="vb_p_form")
            else:
                st.text_input("Item form", key="vb_p_form")

        st.text_area(
            "Descripción del producto",
            key="vb_p_desc",
            height=100,
            placeholder="Descripción completa del producto (max ~2000 caracteres).",
        )

        st.markdown("**Bullets (Key Product Features)**")
        b1 = st.text_input("Bullet 1", key="vb_p_b1")
        b2 = st.text_input("Bullet 2", key="vb_p_b2")
        b3 = st.text_input("Bullet 3", key="vb_p_b3")
        b4 = st.text_input("Bullet 4", key="vb_p_b4")
        b5 = st.text_input("Bullet 5", key="vb_p_b5")

        st.markdown("**Imágenes**")
        st.text_input("URL imagen principal *", key="vb_p_main_img",
                      placeholder="https://...")
        with st.expander("Imágenes secundarias (hasta 8)"):
            for i in range(1, 9):
                st.text_input(f"Imagen secundaria {i}", key=f"vb_p_img_{i}")

        st.markdown("**Otros campos**")
        c3, c4 = st.columns(2)
        with c3:
            if valid_countries:
                st.selectbox("País de origen", options=[""] + valid_countries, key="vb_p_country")
            else:
                st.text_input("País de origen", key="vb_p_country")
        with c4:
            if valid_gtin_exemptions:
                st.selectbox(
                    "Motivo exención GTIN",
                    options=valid_gtin_exemptions,
                    key="vb_p_gtin_ex",
                    help="Para parents que no tienen UPC propio. Default: Pieza",
                )
            else:
                st.text_input("Motivo exención GTIN", key="vb_p_gtin_ex", value="Pieza")

    # ─── Tab: Children + Theme ───
    with tab_children:
        st.markdown("##### Variation Theme + Children")

        theme = st.selectbox(
            "Variation Theme *",
            options=themes_available,
            key="vb_theme",
            help="Define qué atributo diferencia a los children.",
        )

        attrs = _theme_attr_fields(theme)
        attr_labels = [_ATTR_LABELS.get(a, a) for a in attrs]

        st.info(
            f"📋 Theme **{theme}** — cada child debe tener: "
            + ", ".join(f"**{lbl}**" for lbl in attr_labels)
        )

        # Sugerencia de valores válidos
        if _F_FLAVOR_NAME in attrs and valid_flavors:
            with st.expander("💡 Valores válidos sugeridos para Sabor"):
                st.write(", ".join(valid_flavors[:30]))
        if _F_SIZE_NAME in attrs and valid_size_maps:
            with st.expander("💡 Valores válidos sugeridos para Tamaño (Size Map)"):
                st.write(", ".join(valid_size_maps))

        n_children = st.number_input(
            "Cantidad de children",
            min_value=2, max_value=20, value=3, step=1,
            key="vb_n_children",
        )

        # Esquema de columnas según theme actual
        base_cols_schema = ["item_sku"] + list(attrs) + [
            "price_mx", "quantity", "lead_time_days",
            "external_product_id", "external_product_id_type",
            _F_ITEM_NAME, _F_MAIN_IMAGE,
            _F_UNIT_COUNT, _F_UNIT_COUNT_TYPE,
            _F_PACKAGE_WEIGHT, _F_PACKAGE_WEIGHT_UOM,
        ]

        def _empty_children_df(n: int) -> pd.DataFrame:
            return pd.DataFrame({
                "item_sku": [""] * n,
                **{a: [""] * n for a in attrs},
                "price_mx": [0.0] * n,
                "quantity": [0] * n,
                "lead_time_days": [2] * n,
                "external_product_id": [""] * n,
                "external_product_id_type": [""] * n,
                _F_ITEM_NAME: [""] * n,
                _F_MAIN_IMAGE: [""] * n,
                _F_UNIT_COUNT: [None] * n,
                _F_UNIT_COUNT_TYPE: [""] * n,
                _F_PACKAGE_WEIGHT: [None] * n,
                _F_PACKAGE_WEIGHT_UOM: [""] * n,
            })

        # 3 keys separadas:
        # _vb_input_df    → input estable al editor (solo cambia en rebuild)
        # vb_data_editor  → key del widget (state interno persistente)
        # vb_children_df  → output merged que leen las otras tabs
        INPUT_KEY = "_vb_input_df"
        EDITOR_KEY = "vb_data_editor"
        OUTPUT_KEY = "vb_children_df"

        # Init INPUT_KEY una vez
        if INPUT_KEY not in st.session_state:
            st.session_state[INPUT_KEY] = _empty_children_df(n_children)

        # Rebuild solo si cambia n_children o cambia el theme (cambian columnas)
        prev_input = st.session_state[INPUT_KEY]
        prev_output = st.session_state.get(OUTPUT_KEY, prev_input)
        needs_rebuild = (
            len(prev_input) != n_children
            or list(prev_input.columns) != base_cols_schema
        )
        if needs_rebuild:
            new_df = _empty_children_df(n_children)
            # preservar valores del output anterior (que tiene los edits aplicados)
            if isinstance(prev_output, pd.DataFrame) and len(prev_output) > 0:
                common_cols = [c for c in prev_output.columns if c in new_df.columns]
                min_rows = min(len(prev_output), len(new_df))
                for c in common_cols:
                    new_df.loc[:min_rows - 1, c] = prev_output[c].iloc[:min_rows].values
            st.session_state[INPUT_KEY] = new_df
            # IMPORTANTE: borrar la key del editor para que se reinicialice con el nuevo input
            if EDITOR_KEY in st.session_state:
                del st.session_state[EDITOR_KEY]

        df_default = st.session_state[INPUT_KEY]

        # Configurar column types
        col_config: dict = {
            "item_sku": st.column_config.TextColumn("SKU child *", required=True),
            "price_mx": st.column_config.NumberColumn("Precio MXN *", min_value=0.0, step=1.0, format="%.2f"),
            "quantity": st.column_config.NumberColumn("Cantidad *", min_value=0, step=1),
            "lead_time_days": st.column_config.NumberColumn("Lead time (días)", min_value=1, step=1),
            "external_product_id": st.column_config.TextColumn("UPC/EAN/GTIN", help="Opcional. Si lo dejás vacío, usa exención GTIN del parent."),
            "external_product_id_type": st.column_config.SelectboxColumn(
                "Tipo ID", options=["", "GTIN", "EAN", "UPC", "ASIN"]
            ),
            _F_ITEM_NAME: st.column_config.TextColumn(
                "Título child", help="Opcional. Si vacío, se autogenera: parent_name - atributo(s)."
            ),
            _F_MAIN_IMAGE: st.column_config.TextColumn(
                "URL imagen child", help="Opcional. Si vacío, hereda imagen del parent."
            ),
            _F_UNIT_COUNT: st.column_config.NumberColumn("Unit count", min_value=0.0, step=1.0),
            _F_UNIT_COUNT_TYPE: st.column_config.SelectboxColumn(
                "Unit count type", options=[""] + (valid_unit_count_types or ["unidad", "gramo", "mililitro"])
            ),
            _F_PACKAGE_WEIGHT: st.column_config.NumberColumn("Peso paquete", min_value=0.0, step=0.1),
            _F_PACKAGE_WEIGHT_UOM: st.column_config.SelectboxColumn(
                "Unidad peso", options=[""] + (valid_weight_uoms or ["GR", "KG", "LB", "OZ"])
            ),
        }

        # Atributos del theme (ej: flavor_name, size_name)
        for a in attrs:
            label = _ATTR_LABELS.get(a, a)
            col_config[a] = st.column_config.TextColumn(
                f"{label} *",
                required=True,
                help="Valores sugeridos arriba en el expander. Podés usar valores custom."
            )

        edited = st.data_editor(
            df_default,
            column_config=col_config,
            num_rows="fixed",
            use_container_width=True,
            hide_index=True,
            key=EDITOR_KEY,
        )
        # Guardar merged para que preview/download lo lean
        st.session_state[OUTPUT_KEY] = edited

    # ─── Tab: Preview ───
    with tab_preview:
        st.markdown("##### Preview de las filas a escribir")

        # Reconstruir parent dict y children df desde session state
        parent_data = {
            _F_FEED_PRODUCT_TYPE: "petfood",
            _F_ITEM_SKU: st.session_state.get("vb_p_sku", ""),
            _F_BRAND_NAME: st.session_state.get("vb_p_brand", ""),
            _F_MANUFACTURER: st.session_state.get("vb_p_manufacturer", ""),
            _F_ITEM_NAME: st.session_state.get("vb_p_name", ""),
            _F_PRODUCT_DESCRIPTION: st.session_state.get("vb_p_desc", ""),
            _F_AGE_RANGE: st.session_state.get("vb_p_age", ""),
            _F_BREED: st.session_state.get("vb_p_breed", ""),
            _F_TARGET_AUDIENCE: st.session_state.get("vb_p_audience", ""),
            _F_ITEM_FORM: st.session_state.get("vb_p_form", ""),
            _F_MAIN_IMAGE: st.session_state.get("vb_p_main_img", ""),
            "other_images": [st.session_state.get(f"vb_p_img_{i}", "") for i in range(1, 9)],
            "bullets": [
                st.session_state.get("vb_p_b1", ""),
                st.session_state.get("vb_p_b2", ""),
                st.session_state.get("vb_p_b3", ""),
                st.session_state.get("vb_p_b4", ""),
                st.session_state.get("vb_p_b5", ""),
            ],
            _F_COUNTRY_OF_ORIGIN: st.session_state.get("vb_p_country", ""),
            _F_GTIN_EXEMPTION: st.session_state.get("vb_p_gtin_ex", "Pieza"),
            _F_CONDITION_TYPE: "Nuevo",
        }

        children_df = st.session_state.get("vb_children_df")
        theme = st.session_state.get("vb_theme", themes_available[0])

        if children_df is None or not isinstance(children_df, pd.DataFrame):
            st.warning("⚠️ Cargá los children en la tab anterior antes de ver el preview.")
        else:
            errors = _validate_inputs(parent_data, children_df, theme, field_to_col)

            if errors:
                st.error("**No se puede generar — corregí estos puntos:**\n\n" + "\n".join(f"- {e}" for e in errors))
            else:
                st.success(f"✅ Validación OK — {len(children_df)} children listos para generar.")

            # Preview siempre, aunque haya errores (útil para depurar)
            parent_row = _build_parent_row(parent_data, theme, field_to_col)
            child_rows = [
                _build_child_row(r.to_dict(), parent_data, theme, field_to_col)
                for _, r in children_df.iterrows()
            ]
            preview_df = _build_preview_df(parent_row, child_rows, col_to_field)

            st.markdown(f"**Filas que se van a escribir en `Plantilla` desde fila {_DATA_START_ROW}:**")
            st.dataframe(preview_df, use_container_width=True, hide_index=True)

            # Mostrar columnas clave en KPIs
            st.markdown("##### Resumen rápido")
            c1, c2, c3, c4 = st.columns(4)
            with c1:
                kpi_card("Parent SKU", parent_data[_F_ITEM_SKU] or "—")
            with c2:
                kpi_card("Children", str(len(children_df)))
            with c3:
                kpi_card("Theme", theme)
            with c4:
                stock_total = int(children_df["quantity"].fillna(0).sum())
                kpi_card("Stock total", f"{stock_total:,}")

    # ─── Tab: Descargar ───
    with tab_download:
        st.markdown("##### Generar y descargar el template completo")

        children_df = st.session_state.get("vb_children_df")
        theme = st.session_state.get("vb_theme", themes_available[0])

        if children_df is None:
            st.warning("⚠️ Cargá los children en la tab '🧬 Children + Theme' primero.")
            return

        # Reconstruir parent_data igual que en preview
        parent_data = {
            _F_FEED_PRODUCT_TYPE: "petfood",
            _F_ITEM_SKU: st.session_state.get("vb_p_sku", ""),
            _F_BRAND_NAME: st.session_state.get("vb_p_brand", ""),
            _F_MANUFACTURER: st.session_state.get("vb_p_manufacturer", ""),
            _F_ITEM_NAME: st.session_state.get("vb_p_name", ""),
            _F_PRODUCT_DESCRIPTION: st.session_state.get("vb_p_desc", ""),
            _F_AGE_RANGE: st.session_state.get("vb_p_age", ""),
            _F_BREED: st.session_state.get("vb_p_breed", ""),
            _F_TARGET_AUDIENCE: st.session_state.get("vb_p_audience", ""),
            _F_ITEM_FORM: st.session_state.get("vb_p_form", ""),
            _F_MAIN_IMAGE: st.session_state.get("vb_p_main_img", ""),
            "other_images": [st.session_state.get(f"vb_p_img_{i}", "") for i in range(1, 9)],
            "bullets": [
                st.session_state.get("vb_p_b1", ""),
                st.session_state.get("vb_p_b2", ""),
                st.session_state.get("vb_p_b3", ""),
                st.session_state.get("vb_p_b4", ""),
                st.session_state.get("vb_p_b5", ""),
            ],
            _F_COUNTRY_OF_ORIGIN: st.session_state.get("vb_p_country", ""),
            _F_GTIN_EXEMPTION: st.session_state.get("vb_p_gtin_ex", "Pieza"),
            _F_CONDITION_TYPE: "Nuevo",
        }

        errors = _validate_inputs(parent_data, children_df, theme, field_to_col)
        if errors:
            st.error("**No se puede generar — corregí estos puntos:**\n\n" + "\n".join(f"- {e}" for e in errors))
            return

        parent_row = _build_parent_row(parent_data, theme, field_to_col)
        child_rows = [
            _build_child_row(r.to_dict(), parent_data, theme, field_to_col)
            for _, r in children_df.iterrows()
        ]

        try:
            output_bytes = _write_template_with_rows(template_bytes, parent_row, child_rows)
        except Exception as e:
            st.error(f"❌ Error al escribir el template: {e}")
            return

        parent_sku_safe = re.sub(r"[^A-Za-z0-9_-]", "_", parent_data[_F_ITEM_SKU])
        filename = f"FlatFile_Variations_{parent_sku_safe}_{len(child_rows)}children.xlsm"

        st.success(f"✅ Template generado — {len(child_rows)} children agrupados al parent **{parent_data[_F_ITEM_SKU]}**.")

        st.download_button(
            label="💾 Descargar template completo (.xlsm)",
            data=output_bytes,
            file_name=filename,
            mime="application/vnd.ms-excel.sheet.macroEnabled.12",
            key="vb_download_xlsm",
            type="primary",
        )

        st.caption(
            f"**Cómo subirlo:** Seller Central → Catálogo → Añadir productos a través de subida → "
            f"Cargar archivo de inventario → seleccioná este `.xlsm`. "
            f"El feed va a procesar 1 parent + {len(child_rows)} children agrupados por **{theme}**."
        )

        # Resumen final
        st.markdown("##### Resumen del archivo generado")
        c1, c2, c3 = st.columns(3)
        with c1:
            kpi_card("Tamaño", f"{len(output_bytes)/1024:.0f} KB")
        with c2:
            kpi_card("Filas", f"{1 + len(child_rows)}")
        with c3:
            kpi_card("Hojas preservadas", str(len(meta["sheets"])))
