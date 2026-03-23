import io
import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Constants ──────────────────────────────────────────────────────────────

_DEFAULT_ASINS = [
    {"ASIN": "B0CYLMJJJC", "Producto": "Moisturizing Cream",  "Precio": 9.99},
    {"ASIN": "B0CYK4G2Y8", "Producto": "Facial Cleanser",     "Precio": 9.89},
    {"ASIN": "B0CYLM4L23", "Producto": "Body Lotion",          "Precio": 18.89},
    {"ASIN": "B0CYLDSQ5L", "Producto": "Body Cream",           "Precio": 13.49},
    {"ASIN": "B0F4KXZVNM", "Producto": "2-Pack Cream",         "Precio": 16.99},
    {"ASIN": "B0CYKDSDJX", "Producto": "Hyaluronic Serum",     "Precio": 18.89},
    {"ASIN": "B0CYL1RLNQ", "Producto": "Night Cream",          "Precio": 19.79},
    {"ASIN": "B0F548KTXD", "Producto": "2-Pack Lotion",        "Precio": 32.11},
    {"ASIN": "B0F6VZMF2V", "Producto": "Skincare Set",         "Precio": 32.90},
]

_TIER_PARAMS = {
    "LOW":  {"clicks_neg": 18, "spend_stop": 15, "cvr_ref": 0.11},
    "MID":  {"clicks_neg": 22, "spend_stop": 22, "cvr_ref": 0.10},
    "HIGH": {"clicks_neg": 28, "spend_stop": 30, "cvr_ref": 0.08},
}

_OBJECTIVES = ["DISCOVERY", "RANKING", "CONQUEST", "DEFENSIVE", "PROFIT", "REMARKETING"]

_OBJ_COLORS = {
    "DISCOVERY":   "FFD54F",
    "RANKING":     "66BB6A",
    "CONQUEST":    "EF5350",
    "DEFENSIVE":   "42A5F5",
    "PROFIT":      "AB47BC",
    "REMARKETING": "FF7043",
    "SCAVENGER":   "BDBDBD",
    "SP | DISCOVERY":   "FFD54F",
    "SP | RANKING":     "66BB6A",
    "SP | CONQUEST":    "EF5350",
    "SP | DEFENSIVE":   "42A5F5",
    "SP | PROFIT":      "AB47BC",
    "SB | DEFENSIVE":   "42A5F5",
    "SB | RANKING":     "66BB6A",
    "SD | REMARKETING": "FF7043",
    "SD | CONQUEST":    "EF5350",
    "SD | DEFENSIVE":   "42A5F5",
    "SCAVENGER":        "BDBDBD",
}

_OBJ_MULTIPLIERS = {
    "DISCOVERY":   1.20,
    "RANKING":     1.00,
    "CONQUEST":    0.86,
    "DEFENSIVE":   0.71,
    "PROFIT":      0.50,
    "REMARKETING": 0.71,
}


# ── Helpers ────────────────────────────────────────────────────────────────

def _assign_tier(precio: float) -> str:
    if precio < 12:
        return "LOW"
    elif precio <= 22:
        return "MID"
    return "HIGH"


def _calc_objective_targets(target_cuenta: int) -> dict:
    return {obj: max(5, int(target_cuenta * mult))
            for obj, mult in _OBJ_MULTIPLIERS.items()}


def _classify_campaign(name: str, ad_type: str, prefix: str, brand_terms: list) -> str:
    nl = name.lower()
    pfx = prefix.upper()

    if "scavenger" in nl:
        return "SCAVENGER"

    if ad_type in ("SB2", "SB", "Sponsored Brands"):
        if any(x in nl for x in ["defensive", f"{pfx.lower()} defensive"]):
            return "SB | DEFENSIVE"
        return "SB | RANKING"

    if ad_type in ("SD", "Sponsored Display"):
        if "remarketing" in nl:
            return "SD | REMARKETING"
        if "prospecting" in nl:
            return "SD | CONQUEST"
        if any(x in nl for x in ["defensive", "difensive", "listing defense",
                                  f"{pfx.lower()} products"]):
            return "SD | DEFENSIVE"
        return "SD | REMARKETING"

    # SP
    if name.startswith(f"{pfx} | DISCOVERY"):
        return "SP | DISCOVERY"
    if "auto" in nl and any(x in nl for x in ["discovery", "close match"]):
        return "SP | DISCOVERY"

    if any(x in nl for x in ["brand defensive", "defensive", "difensive",
                              f"{pfx.lower()} defensive"]):
        return "SP | DEFENSIVE"

    brand_kw_patterns = []
    for bt in brand_terms:
        brand_kw_patterns.extend([
            f"phrase - {bt}", f"broad - {bt}", f"exact - {bt}",
            f"exact - {bt} cream", f"exact - {bt} crema",
            f"exact - {bt} face", "exact - brand",
        ])
    if any(x in nl for x in brand_kw_patterns):
        return "SP | DEFENSIVE"

    if "harvest" in nl:
        return "SP | PROFIT"

    if any(x in nl for x in ["top comp", "top competitors", "pt - asin",
                              "pt - category", "sp pr - categ",
                              "products-categories", "category -",
                              "related", "retinol competitors"]):
        return "SP | CONQUEST"
    if "pat" in nl and "competitor" in nl:
        return "SP | CONQUEST"

    if "spanish" in nl:
        return "SP | RANKING"
    if any(x in nl for x in ["vitamin a", "retinol", "vitamin e"]):
        return "SP | RANKING"
    if "broad kws" in nl:
        return "SP | DISCOVERY"
    if any(x in nl for x in [
        "core hero", "dry skin", "oily skin", "hydrating", "non greasy",
        "thick intensive", "winter", "women use", "body lotion core",
        "dermatological cleanser", "best 1", "best 2", "phrase kws",
        "exact - 2", "acid f1", "lotion f3", "kws validation",
    ]):
        return "SP | RANKING"
    if "phrase" in nl and not any(bt in nl for bt in brand_terms):
        return "SP | RANKING"

    return "SP | DISCOVERY"


def _get_base_objective(group: str) -> str:
    """Extrae el objetivo base de un grupo (ej. 'SP | DISCOVERY' → 'DISCOVERY')."""
    parts = group.split(" | ")
    return parts[-1] if len(parts) > 1 else group


def _build_sheet_name(group: str, prefix: str) -> str:
    _SHEET_MAP = {
        "SP | DISCOVERY":   f"{prefix} | DISCOVERY | SP",
        "SP | RANKING":     f"{prefix} | RANKING | SP",
        "SP | CONQUEST":    f"{prefix} | CONQUEST | SP",
        "SP | DEFENSIVE":   f"{prefix} | DEFENSIVE | SP",
        "SP | PROFIT":      f"{prefix} | PROFIT | SP",
        "SB | DEFENSIVE":   f"{prefix} | DEFENSIVE | SB",
        "SB | RANKING":     f"{prefix} | RANKING | SB",
        "SD | REMARKETING": f"{prefix} | REMARKETING | SD",
        "SD | CONQUEST":    f"{prefix} | CONQUEST | SD",
        "SD | DEFENSIVE":   f"{prefix} | DEFENSIVE | SD",
        "SCAVENGER":        f"{prefix} | SCAVENGER",
    }
    return _SHEET_MAP.get(group, group)


# ── Rules generation ───────────────────────────────────────────────────────

def _generate_rules(target_cuenta: int, prefix: str) -> list[dict]:
    """Genera las 274 rules dinámicamente."""
    targets = _calc_objective_targets(target_cuenta)
    rules = []

    for obj in _OBJECTIVES:
        t = targets[obj]

        for tier_name, tier in _TIER_PARAMS.items():
            min_clicks = max(3, tier["clicks_neg"] // 3)

            # ── Bid Rules (7 per objective per tier)
            bid_rules = [
                {
                    "Rule Name": f"{prefix} | {obj} | BID INC AGG | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS < {int(t*0.5)}% AND Orders ≥ 2 AND Clicks ≥ {min_clicks}",
                    "Action": "+15% bid", "Color": "INC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | BID INC SOFT | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS {int(t*0.5)}%-{int(t*0.85)}% AND Orders ≥ 1",
                    "Action": "+8% bid", "Color": "INC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | BID FLAT | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS {int(t*0.85)}%-{t}% (zona segura)",
                    "Action": "Sin cambio", "Color": "FLAT", "Window": "—",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | BID DEC SOFT | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS > {int(t*1.3)}%",
                    "Action": "-10% bid", "Color": "DEC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | BID DEC RISK | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS > {int(t*1.5)}%",
                    "Action": "-15% bid", "Color": "DEC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | BID DEC CTRL | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS > {int(t*1.8)}%",
                    "Action": "-25% bid", "Color": "DEC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | BID DEC HARD | {tier_name}",
                    "Objective": obj, "Type": "Bid", "Tier": tier_name,
                    "Condition": f"ACoS > {int(t*2.2)}%",
                    "Action": "-40% bid", "Color": "DEC", "Window": "7 días",
                },
            ]

            # ── Placement Rules (6 per objective per tier)
            placement_rules = [
                {
                    "Rule Name": f"{prefix} | {obj} | TOS INC | {tier_name}",
                    "Objective": obj, "Type": "Placement", "Tier": tier_name,
                    "Condition": f"TOS ACoS < {int(t*0.65)}%",
                    "Action": "TOS +10%", "Color": "INC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | TOS DEC SOFT | {tier_name}",
                    "Objective": obj, "Type": "Placement", "Tier": tier_name,
                    "Condition": f"TOS ACoS > {t}%",
                    "Action": "TOS -15%", "Color": "DEC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | TOS DEC HARD | {tier_name}",
                    "Objective": obj, "Type": "Placement", "Tier": tier_name,
                    "Condition": f"TOS ACoS > {int(t*1.4)}%",
                    "Action": "TOS -30%", "Color": "DEC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | PP INC | {tier_name}",
                    "Objective": obj, "Type": "Placement", "Tier": tier_name,
                    "Condition": f"PP ACoS < {int(t*0.6)}%",
                    "Action": "PP +10%", "Color": "INC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | PP DEC SOFT | {tier_name}",
                    "Objective": obj, "Type": "Placement", "Tier": tier_name,
                    "Condition": f"PP ACoS > {int(t*0.9)}%",
                    "Action": "PP -15%", "Color": "DEC", "Window": "14 días",
                },
                {
                    "Rule Name": f"{prefix} | {obj} | PP DEC HARD | {tier_name}",
                    "Objective": obj, "Type": "Placement", "Tier": tier_name,
                    "Condition": f"PP ACoS > {int(t*1.3)}%",
                    "Action": "PP -30%", "Color": "DEC", "Window": "14 días",
                },
            ]

            # ── Negate Rule (1 per tier)
            negate_rule = {
                "Rule Name": f"{prefix} | {obj} | NEGATE | {tier_name}",
                "Objective": obj, "Type": "Negate", "Tier": tier_name,
                "Condition": f"Clicks ≥ {tier['clicks_neg']} AND Orders = 0",
                "Action": "Add Negative Exact", "Color": "NEG", "Window": "30 días",
            }

            # ── Hard Stop (1 per tier)
            hard_stop = {
                "Rule Name": f"{prefix} | {obj} | HARD STOP | {tier_name}",
                "Objective": obj, "Type": "Hard Stop", "Tier": tier_name,
                "Condition": (f"Spend ≥ ${tier['spend_stop']} AND Orders = 0 "
                              f"AND Clicks ≥ {tier['clicks_neg']}"),
                "Action": "Bid -50%", "Color": "DEC", "Window": "14 días",
            }

            rules.extend(bid_rules)
            rules.extend(placement_rules)
            rules.append(negate_rule)
            rules.append(hard_stop)

    # ── Harvest rules (only DISCOVERY and RANKING)
    for obj in ["DISCOVERY", "RANKING"]:
        t = targets[obj]
        rules.append({
            "Rule Name": f"{prefix} | {obj} | HARVEST AUTO→PHRASE",
            "Objective": obj, "Type": "Harvest", "Tier": "ALL",
            "Condition": f"Orders ≥ 2 AND ACoS < {int(t*1.2)}% AND Clicks ≥ 5",
            "Action": "Add as Phrase Match", "Color": "INC", "Window": "30 días",
        })
        rules.append({
            "Rule Name": f"{prefix} | {obj} | HARVEST PHRASE→EXACT",
            "Objective": obj, "Type": "Harvest", "Tier": "ALL",
            "Condition": f"Orders ≥ 3 AND ACoS < {t}% AND Clicks ≥ 8",
            "Action": "Add as Exact Match", "Color": "INC", "Window": "30 días",
        })

    return rules


# ── Excel builders ─────────────────────────────────────────────────────────

def _hex_fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _build_groups_excel(df_campaigns: pd.DataFrame, prefix: str) -> bytes:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "RESUMEN"

    # Summary header
    ws_summary["A1"] = f"{prefix} | Campaign Groups — Atom11"
    ws_summary["A1"].font = Font(bold=True, size=13, color="E84000")
    ws_summary["A2"] = f"Total campañas: {len(df_campaigns)}"

    obj_groups = df_campaigns.groupby("Objetivo")
    ws_summary["A4"] = "Grupo"
    ws_summary["B4"] = "Sheet Atom11"
    ws_summary["C4"] = "Campañas"
    for col in ["A", "B", "C"]:
        ws_summary[f"{col}4"].font = Font(bold=True)
        ws_summary[f"{col}4"].fill = _hex_fill("1F1F1F")
        ws_summary[f"{col}4"].font = Font(bold=True, color="FFFFFF")

    row = 5
    for grupo, grp_df in sorted(obj_groups):
        sheet_name = _build_sheet_name(grupo, prefix)
        color = _OBJ_COLORS.get(grupo, "EEEEEE")
        ws_summary.cell(row, 1, grupo).fill = _hex_fill(color)
        ws_summary.cell(row, 2, sheet_name)
        ws_summary.cell(row, 3, len(grp_df))
        row += 1

    ws_summary.column_dimensions["A"].width = 22
    ws_summary.column_dimensions["B"].width = 36
    ws_summary.column_dimensions["C"].width = 12

    # One sheet per group
    for grupo, grp_df in sorted(obj_groups):
        sheet_name = _build_sheet_name(grupo, prefix)[:31]  # Excel 31-char limit
        ws = wb.create_sheet(title=sheet_name)
        color = _OBJ_COLORS.get(grupo, "EEEEEE")

        ws["A1"] = "Campaign Name"
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = _hex_fill(color)
        ws["A1"].alignment = Alignment(horizontal="center")

        for i, camp in enumerate(grp_df["Campaign Name"].tolist(), start=2):
            ws.cell(i, 1, camp)

        ws.column_dimensions["A"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _build_rules_excel(rules: list[dict], prefix: str, target_cuenta: int) -> bytes:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "RESUMEN"

    ws_summary["A1"] = f"{prefix} | Atom11 Rules — {len(rules)} rules totales"
    ws_summary["A1"].font = Font(bold=True, size=13, color="E84000")
    ws_summary["A2"] = f"Target ACoS cuenta: {target_cuenta}%"

    headers = ["Tipo", "Rules"]
    for ci, h in enumerate(headers, 1):
        c = ws_summary.cell(4, ci, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _hex_fill("1F1F1F")

    type_counts = {}
    for r in rules:
        type_counts[r["Type"]] = type_counts.get(r["Type"], 0) + 1
    for row_i, (rtype, cnt) in enumerate(sorted(type_counts.items()), start=5):
        ws_summary.cell(row_i, 1, rtype)
        ws_summary.cell(row_i, 2, cnt)

    ws_summary.column_dimensions["A"].width = 16
    ws_summary.column_dimensions["B"].width = 10

    cols = ["Rule Name", "Condition", "Action", "Tier", "Window"]
    col_widths = [55, 55, 22, 8, 12]

    for obj in _OBJECTIVES:
        obj_rules = [r for r in rules if r["Objective"] == obj]
        if not obj_rules:
            continue
        ws = wb.create_sheet(title=f"RULES {obj}"[:31])
        color = _OBJ_COLORS.get(obj, "EEEEEE")

        for ci, (col, width) in enumerate(zip(cols, col_widths), 1):
            c = ws.cell(1, ci, col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _hex_fill(color)
            ws.column_dimensions[get_column_letter(ci)].width = width

        _COLOR_MAP = {
            "INC":  "E8F5E9",
            "DEC":  "FFEBEE",
            "FLAT": "F5F5F5",
            "NEG":  "FFF3E0",
        }

        for row_i, r in enumerate(obj_rules, start=2):
            row_color = _COLOR_MAP.get(r.get("Color", "FLAT"), "FFFFFF")
            for ci, col in enumerate(cols, 1):
                c = ws.cell(row_i, ci, r.get(col, ""))
                c.fill = _hex_fill(row_color)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Main render ────────────────────────────────────────────────────────────

def render():
    st.header("🤖 Atom11 Rules Builder")
    st.caption("Clasifica campañas por objetivo y genera rules dinámicas listas para importar en Atom11.")
    st.divider()

    tab1, tab2, tab3 = st.tabs(["⚙️ Configuración", "📁 Campaign Groups", "📋 Rules Generator"])

    # ═══════════════════════════════════════════════════════════════
    # TAB 1 — Configuración
    # ═══════════════════════════════════════════════════════════════
    with tab1:
        st.subheader("Configuración de marca")

        col_a, col_b = st.columns([1, 1])
        with col_a:
            prefix = st.text_input(
                "Prefijo de marca",
                value=st.session_state.get("atom11_rb_prefix", "DG"),
                help="Ej: DG, MB, STX, LTD — se usa en nombres de rules y sheets",
            )
            st.session_state["atom11_rb_prefix"] = prefix.strip().upper()

        with col_b:
            target_acos = st.slider(
                "Target ACoS cuenta (%)",
                min_value=5, max_value=100, value=st.session_state.get("atom11_rb_target", 70),
                step=1,
            )
            st.session_state["atom11_rb_target"] = target_acos

        st.divider()
        st.subheader("Catálogo de ASINs")

        init_df = pd.DataFrame(st.session_state.get("atom11_rb_asins", _DEFAULT_ASINS))
        edited_df = st.data_editor(
            init_df[["ASIN", "Producto", "Precio"]],
            num_rows="dynamic",
            use_container_width=True,
            column_config={
                "ASIN":     st.column_config.TextColumn("ASIN", width="medium"),
                "Producto": st.column_config.TextColumn("Producto", width="large"),
                "Precio":   st.column_config.NumberColumn("Precio ($)", format="$%.2f", width="small"),
            },
            key="atom11_rb_editor",
        )
        edited_df["Tier"] = edited_df["Precio"].apply(_assign_tier)
        st.session_state["atom11_rb_asins"] = edited_df.to_dict("records")

        st.markdown("**Tier asignado:**")
        st.dataframe(
            edited_df[["ASIN", "Producto", "Precio", "Tier"]],
            use_container_width=True, hide_index=True,
        )

        st.divider()
        st.subheader("Target ACoS por objetivo")
        targets = _calc_objective_targets(target_acos)
        tgt_df = pd.DataFrame([
            {"Objetivo": obj, "Target ACoS": f"{tgt}%",
             "Multiplicador": f"{_OBJ_MULTIPLIERS[obj]:.0%} del target cuenta",
             "Lógica": {
                 "DISCOVERY": "Comprando data — permite más gasto",
                 "RANKING": "Posicionando en búsquedas clave",
                 "CONQUEST": "PAT/ASIN — menor CVR esperado",
                 "DEFENSIVE": "Brand terms deben convertir más barato",
                 "PROFIT": "Harvested winners — máxima eficiencia",
                 "REMARKETING": "Retargeting — buen CVR esperado",
             }[obj]}
            for obj, tgt in targets.items()
        ])
        st.dataframe(tgt_df, use_container_width=True, hide_index=True)

        st.divider()
        st.subheader("Parámetros de Tier")
        tier_df = pd.DataFrame([
            {"Tier": tier, "Rango precio": "<$12" if tier == "LOW" else "$12-$22" if tier == "MID" else ">$22",
             "Clicks Negate": p["clicks_neg"], "Spend Hard Stop": f"${p['spend_stop']}", "CVR ref": f"{p['cvr_ref']:.0%}"}
            for tier, p in _TIER_PARAMS.items()
        ])
        st.dataframe(tier_df, use_container_width=True, hide_index=True)

    # ═══════════════════════════════════════════════════════════════
    # TAB 2 — Campaign Groups
    # ═══════════════════════════════════════════════════════════════
    with tab2:
        st.subheader("Clasificación de campañas por objetivo")
        st.caption("Sube el Campaign CSV de Amazon Campaign Manager. Detecta el tipo de campaña y asigna el grupo Atom11.")

        uploaded = st.file_uploader(
            "Campaign CSV (.csv)",
            type=["csv"],
            key="atom11_rb_csv",
        )

        brand_terms_raw = st.text_input(
            "Términos de marca (separados por coma)",
            value="dermaglos",
            help="Ej: dermaglos, vitamina a — usado para detectar campañas DEFENSIVE",
        )
        brand_terms = [t.strip().lower() for t in brand_terms_raw.split(",") if t.strip()]
        st.session_state["atom11_rb_brand_terms"] = brand_terms

        if uploaded:
            try:
                raw = pd.read_csv(uploaded, encoding="utf-8-sig", low_memory=False)
            except Exception:
                raw = pd.read_csv(uploaded, encoding="latin-1", low_memory=False)

            # Detect name + type columns
            name_col = next((c for c in raw.columns if "campaign name" in c.lower()), None)
            type_col = next((c for c in raw.columns if "type" in c.lower() and "ad" in c.lower()), None)
            state_col = next((c for c in raw.columns if c.strip().lower() in ("state", "status", "estado")), None)

            if name_col is None:
                st.error("No se encontró columna 'Campaign Name' en el CSV.")
                return

            pfx = st.session_state.get("atom11_rb_prefix", "DG")

            # Get unique campaigns
            grp_cols = [name_col]
            if type_col:
                grp_cols.append(type_col)
            if state_col:
                grp_cols.append(state_col)

            df_camps = raw[grp_cols].drop_duplicates(subset=[name_col]).copy()
            df_camps = df_camps.rename(columns={name_col: "Campaign Name"})

            if type_col:
                df_camps = df_camps.rename(columns={type_col: "Ad Type"})
            else:
                df_camps["Ad Type"] = "SP"

            if state_col:
                df_camps = df_camps.rename(columns={state_col: "Estado"})
            else:
                df_camps["Estado"] = "ENABLED"

            df_camps["Objetivo"] = df_camps.apply(
                lambda r: _classify_campaign(
                    str(r["Campaign Name"]),
                    str(r["Ad Type"]),
                    pfx,
                    brand_terms,
                ),
                axis=1,
            )
            df_camps["Sheet Atom11"] = df_camps["Objetivo"].apply(
                lambda g: _build_sheet_name(g, pfx)
            )

            # Summary metrics
            st.markdown(f"**Total campañas encontradas: {len(df_camps)}**")
            obj_counts = df_camps["Objetivo"].value_counts()
            cols_m = st.columns(min(len(obj_counts), 4))
            for i, (obj, cnt) in enumerate(obj_counts.items()):
                cols_m[i % len(cols_m)].metric(obj, cnt)

            st.divider()

            # Manual reclassification
            st.markdown("**Reclasificación manual (opcional)**")
            all_groups = sorted(_OBJ_COLORS.keys())
            manual_overrides = {}
            with st.expander("Reclasificar campañas individualmente"):
                for idx, row in df_camps.iterrows():
                    new_obj = st.selectbox(
                        row["Campaign Name"][:80],
                        options=all_groups,
                        index=all_groups.index(row["Objetivo"]) if row["Objetivo"] in all_groups else 0,
                        key=f"reclass_{idx}",
                    )
                    manual_overrides[idx] = new_obj

            for idx, new_obj in manual_overrides.items():
                df_camps.at[idx, "Objetivo"] = new_obj
                df_camps.at[idx, "Sheet Atom11"] = _build_sheet_name(new_obj, pfx)

            st.dataframe(
                df_camps[["Campaign Name", "Estado", "Objetivo", "Sheet Atom11"]],
                use_container_width=True, hide_index=True,
            )

            st.session_state["atom11_rb_campaigns_df"] = df_camps

            xlsx_bytes = _build_groups_excel(df_camps, pfx)
            st.download_button(
                "📥 Exportar Campaign Groups para Atom11",
                data=xlsx_bytes,
                file_name=f"{pfx}_Atom11_Campaign_Groups.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        else:
            st.info("Sube el Campaign CSV para clasificar las campañas en grupos Atom11.")

    # ═══════════════════════════════════════════════════════════════
    # TAB 3 — Rules Generator
    # ═══════════════════════════════════════════════════════════════
    with tab3:
        st.subheader("Generador de rules Atom11")
        pfx = st.session_state.get("atom11_rb_prefix", "DG")
        t_cuenta = st.session_state.get("atom11_rb_target", 70)
        targets = _calc_objective_targets(t_cuenta)

        rules = _generate_rules(t_cuenta, pfx)

        # Summary
        type_counts = {}
        for r in rules:
            type_counts[r["Type"]] = type_counts.get(r["Type"], 0) + 1

        st.markdown(f"**Total rules generadas: {len(rules)}** | Target cuenta: {t_cuenta}%")
        cols_s = st.columns(len(type_counts))
        for i, (rtype, cnt) in enumerate(sorted(type_counts.items())):
            cols_s[i].metric(rtype, cnt)

        st.divider()

        _COLOR_LABEL = {
            "INC":  "🟢 INCREASE",
            "DEC":  "🔴 DECREASE",
            "FLAT": "⚪ FLAT",
            "NEG":  "🟠 NEGATE",
        }

        for obj in _OBJECTIVES:
            obj_rules = [r for r in rules if r["Objective"] == obj]
            t_obj = targets[obj]
            color_hex = _OBJ_COLORS.get(obj, "EEEEEE")

            with st.expander(
                f"**{obj}** — Target {t_obj}% — {len(obj_rules)} rules",
                expanded=False,
            ):
                df_obj = pd.DataFrame(obj_rules)[["Rule Name", "Type", "Tier", "Condition", "Action", "Window"]]
                st.dataframe(df_obj, use_container_width=True, hide_index=True)

        st.divider()

        xlsx_rules = _build_rules_excel(rules, pfx, t_cuenta)
        st.download_button(
            f"📥 Exportar Rules Atom11 — {len(rules)} rules ({pfx})",
            data=xlsx_rules,
            file_name=f"{pfx}_Atom11_Rules_Complete_v2026.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        st.info(
            "💡 **Cómo usar:** Importar el Excel en Atom11 → Campaign Groups tab → "
            "cargar cada sheet en el grupo correspondiente. Luego importar las rules "
            "en el Rules tab de Atom11."
        )
