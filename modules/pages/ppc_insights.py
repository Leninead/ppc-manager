"""PPC Insights: health score por ASIN desde el Search Term Report, con SQP, BR y Campaign CSV opcionales.

El STR llega del picker de Amazon Ads o de un archivo subido a mano, que se lee con el parser propio del
módulo. Las reglas por ASIN viven en core/ppc_insights: el worker de análisis arma con ellas el mismo payload.
"""
import hashlib
import io
import logging
from dataclasses import dataclass, field
from functools import partial

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.amazon_ads.advertised_asins import FROM_AD_GROUP, FROM_CAMPAIGN_NAME, SEVERAL_ASINS, WITHOUT_ASIN
from core.amazon_ads.report_provider import ReportProvider, ReportReadError
from core.currency_format import ZERO_DECIMAL, currency_symbol, money
from core.helpers import kpi_card
from core.integrations.store import StoreError
from core.ppc_insights.analysis import (
    ANALYSIS_MODULE,
    CANONICAL_LANG,
    InsightsAnalysisParams,
    build_analysis_input,
    insights_row_labels,
)
from core.ppc_insights.asin_health import FROM_FILE, NO_ASINS, analyze_asins, resolve_asins
from core.search_term.analysis import uses_dollar_price
from core.search_term.file import SearchTermFileError
from core.search_term.frame import SOURCE_FILE
from modules.pages import search_term_source
from modules.pages.search_term_source import date_range_label, render_source_picker

log = logging.getLogger(__name__)

# ── Paleta Capybaras ──────────────────────────────────────────────────────────
_ORG   = "E84000"; _ORG_P = "FFF3E0"
_BLK   = "1F1F1F"; _WHT   = "FAFAFA"
_GRN   = "1B6B2F"; _GRN_L = "E8F5E9"
_RED   = "B71C1C"; _RED_L = "FFEBEE"
_YEL   = "9C5700"; _YEL_L = "FFEB9C"
_DGRAY = "2D3748"; _MGRAY = "CBD5E0"
_WHITE = "FFFFFF"

_TARGET_KEY = "insights_target_acos"
_PRICE_KEY = "insights_precio"
_GENERATED_FOR_KEY = "insights_generated_for"
_RESULT_KEY = "insights_result"
_SEEDED_ACCOUNT_KEY = "insights_seeded_account"

_ORIGIN_LABELS = {FROM_FILE: "columna de ASIN del archivo", FROM_AD_GROUP: "producto anunciado del ad group",
                  FROM_CAMPAIGN_NAME: "nombre de la campaña",
                  SEVERAL_ASINS: "ad groups con varios ASINs sin ASIN en el nombre", WITHOUT_ASIN: "sin ASIN"}
_UNATTRIBUTED_ORIGINS = (SEVERAL_ASINS, WITHOUT_ASIN)
_NO_ASIN_CAUSES = {SEVERAL_ASINS: "ad groups que anuncian varios ASINs",
                   WITHOUT_ASIN: "ad groups que el listado de productos anunciados no vio"}
NO_ASIN_FROM_FILE = ("El archivo no trae la columna de ASIN y ningún nombre de campaña lleva uno. Se muestra la "
                     "cuenta entera como una sola fila (ALL).")
ADVERTISED_ASINS_UNREADABLE = ("No se pudieron leer los productos anunciados de la cuenta: el ASIN sale sólo del "
                               "nombre de la campaña.")


@dataclass(frozen=True)
class InsightsResult:
    """What "Generar Insights" computed for one set of inputs, kept until they change."""

    asin_data: dict
    asin_source: str
    spend_share: dict
    sqp_df: pd.DataFrame | None
    br_df: pd.DataFrame | None
    camp_df: pd.DataFrame | None
    file_warnings: tuple = ()
    report_spend: float = 0.0
    grouped_asins: dict = field(default_factory=dict)


# ── OpenPyXL helpers ─────────────────────────────────────────────────────────
def _fill(c):
    return PatternFill("solid", fgColor=c)


def _font(bold=False, color="000000", size=9, name="Arial"):
    return Font(bold=bold, color=color, size=size, name=name)


def _bd():
    s = Side(style="thin", color=_MGRAY)
    return Border(left=s, right=s, top=s, bottom=s)


def _al(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _cell(ws, r, c, val, bg=None, fg="000000", bold=False, fmt=None, size=9, left=False):
    cell = ws.cell(row=r, column=c, value=val)
    if bg:
        cell.fill = _fill(bg)
    cell.font = _font(bold=bold, color=fg, size=size)
    cell.alignment = _al("left" if left else "center")
    cell.border = _bd()
    if fmt:
        cell.number_format = fmt
    return cell


def _hdr(ws, rn, cols, h=14):
    for i, col in enumerate(cols, 1):
        c = ws.cell(row=rn, column=i, value=col)
        c.fill = _fill(_ORG)
        c.font = _font(True, _WHITE, 9)
        c.alignment = _al()
        c.border = _bd()
    ws.row_dimensions[rn].height = h


def _sec(ws, rn, text, ncols, bg=None, fg=None, h=16):
    bg = bg or _DGRAY
    fg = fg or _WHITE
    ws.merge_cells(start_row=rn, start_column=1, end_row=rn, end_column=ncols)
    c = ws.cell(row=rn, column=1, value=text)
    c.fill = _fill(bg)
    c.font = _font(True, fg, 10)
    c.alignment = _al("left")
    c.border = _bd()
    ws.row_dimensions[rn].height = h
    return rn + 1


def _autofit(ws, min_w=8, max_w=40):
    for col_cells in ws.columns:
        best = min_w
        for cell in col_cells:
            try:
                if cell.value:
                    best = max(best, min(max_w, len(str(cell.value)) + 2))
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = best


def _excel_money_format(currency_code):
    code = (currency_code or "").strip().upper()
    decimals = "" if code in ZERO_DECIMAL else ".00"
    return f'"{currency_symbol(code)}"#,##0{decimals}'


# ── Parsers ───────────────────────────────────────────────────────────────────
@st.cache_data(max_entries=3, ttl=3600, show_spinner=False)
def _parse_str(file_bytes, fname):
    try:
        buf = io.BytesIO(file_bytes)
        df = pd.read_excel(buf) if fname.endswith(".xlsx") else pd.read_csv(buf)
        df.columns = df.columns.str.strip()
        for col in df.columns:
            if df[col].dtype == object:
                try:
                    cleaned = df[col].astype(str).str.replace(r"[$%,]", "", regex=True).str.strip()
                    numeric = pd.to_numeric(cleaned, errors="coerce")
                    if numeric.notna().sum() / max(len(numeric), 1) > 0.5:
                        df[col] = numeric.fillna(0)
                except Exception:
                    pass
        return df, None
    except Exception as e:
        return None, str(e)


def _read_manual_str(file_bytes, file_name):
    """The picker's manual upload, read with this module's own STR parser."""
    df, err = _parse_str(file_bytes, file_name)
    if err or df is None:
        raise SearchTermFileError(f"Error al leer el STR: {err}")
    return df


@st.cache_data(max_entries=3, ttl=3600, show_spinner=False)
def _parse_sqp_cached(file_bytes, fname):
    try:
        buf = io.BytesIO(file_bytes)
        buf.name = fname
        df = pd.read_excel(buf, skiprows=1) if fname.endswith(".xlsx") else pd.read_csv(buf, skiprows=1)
        df.columns = df.columns.str.strip()
        return df, None
    except Exception as e:
        return None, str(e)


@st.cache_data(max_entries=3, ttl=3600, show_spinner=False)
def _parse_br(file_bytes, fname):
    try:
        buf = io.BytesIO(file_bytes)
        df = pd.read_excel(buf) if fname.endswith(".xlsx") else pd.read_csv(buf)
        df.columns = df.columns.str.strip()
        return df, None
    except Exception as e:
        return None, str(e)


@st.cache_data(max_entries=3, ttl=3600, show_spinner=False)
def _parse_campaigns(file_bytes):
    try:
        buf = io.BytesIO(file_bytes)
        df = pd.read_csv(buf)
        df.columns = df.columns.str.strip()
        return df, None
    except Exception as e:
        return None, str(e)


# ── Excel export ──────────────────────────────────────────────────────────────
def _build_insights_excel(asin_data, client_name, target_acos, currency_code=""):
    wb = Workbook()
    money_fmt = _excel_money_format(currency_code)
    show_money = partial(money, currency_code=currency_code)

    # ── Sheet 1: Resumen ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumen"
    ws.freeze_panes = "A3"

    # Title row
    ws.merge_cells("A1:L1")
    title_cell = ws.cell(row=1, column=1, value=f"PPC Insights Engine — {client_name}")
    title_cell.fill = _fill(_ORG)
    title_cell.font = _font(True, _WHITE, 13)
    title_cell.alignment = _al("left")
    ws.row_dimensions[1].height = 22

    hdr_cols = [
        "ASIN", "Health Score", "Spend", "Sales", "ACoS %",
        "CVR %", "BuyBox %", "Imp Share %", "Wasted Spend",
        "Campañas", "Funnel", "SQP Gaps",
    ]
    _hdr(ws, 2, hdr_cols, h=15)

    sorted_asins = sorted(asin_data.keys(), key=lambda a: asin_data[a]["spend"], reverse=True)

    for i, asin in enumerate(sorted_asins, 3):
        d = asin_data[asin]
        score = d["health_score"]
        acos  = d["acos"]
        score_bg = _GRN_L if score >= 70 else (_YEL_L if score >= 40 else _RED_L)
        acos_bg  = None
        if acos is not None:
            acos_bg = _GRN_L if acos <= target_acos else (_YEL_L if acos <= target_acos * 1.5 else _RED_L)

        row_bg = "F7F7F7" if i % 2 == 0 else _WHITE

        _cell(ws, i, 1,  asin,                                      bg=row_bg, left=True)
        _cell(ws, i, 2,  score,                                      bg=score_bg, bold=True)
        _cell(ws, i, 3,  d["spend"],                                 bg=row_bg, fmt=money_fmt)
        _cell(ws, i, 4,  d["sales"],                                 bg=row_bg, fmt=money_fmt)
        _cell(ws, i, 5,  round(acos, 1) if acos is not None else None, bg=acos_bg, fmt="0.0")
        _cell(ws, i, 6,  round(d["cvr"], 1) if d["cvr"] is not None else None, bg=row_bg, fmt="0.0")
        _cell(ws, i, 7,  round(d["buybox"], 1) if d["buybox"] is not None else None, bg=row_bg, fmt="0.0")
        _cell(ws, i, 8,  round(d["imp_share"], 1) if d["imp_share"] is not None else None, bg=row_bg, fmt="0.0")
        _cell(ws, i, 9,  d["wasted_spend"],                          bg=row_bg, fmt=money_fmt)
        _cell(ws, i, 10, d["n_campaigns"],                           bg=row_bg)
        _cell(ws, i, 11, "Completo" if d["funnel_complete"] else ("Parcial" if d["funnel_complete"] is not None else "—"),
              bg=row_bg)
        _cell(ws, i, 12, d["sqp_gaps"],                              bg=row_bg)

    _autofit(ws)

    # ── Sheets per ASIN (max 10) ──────────────────────────────────────────────
    for asin in sorted_asins[:10]:
        d = asin_data[asin]
        ws2 = wb.create_sheet(title=f"ASIN_{asin[:8]}")
        r = 1

        # Title
        ws2.merge_cells(f"A{r}:F{r}")
        tc = ws2.cell(row=r, column=1, value=f"Insights — {asin} — {client_name}")
        tc.fill = _fill(_ORG)
        tc.font = _font(True, _WHITE, 12)
        tc.alignment = _al("left")
        ws2.row_dimensions[r].height = 20
        r += 1

        # KPI Summary
        r = _sec(ws2, r, "KPI Summary", 6, bg=_DGRAY)
        _hdr(ws2, r, ["Métrica", "Valor"], h=13)
        r += 1
        kpi_rows = [
            ("Spend",         show_money(d["spend"])),
            ("Sales",         show_money(d["sales"])),
            ("ACoS",          f"{d['acos']:.1f}%" if d["acos"] is not None else "—"),
            ("CVR",           f"{d['cvr']:.1f}%" if d["cvr"] is not None else "—"),
            ("Orders",        f"{d['orders']:.0f}"),
            ("Clicks",        f"{d['clicks']:.0f}"),
            ("Wasted Spend",  show_money(d["wasted_spend"])),
            ("BuyBox %",      f"{d['buybox']:.1f}%" if d["buybox"] is not None else "—"),
            ("Sessions",      f"{d['sessions']:,.0f}" if d["sessions"] is not None else "—"),
            ("Imp Share %",   f"{d['imp_share']:.1f}%" if d["imp_share"] is not None else "—"),
            ("Purchase Share %", f"{d['purchase_share']:.1f}%" if d["purchase_share"] is not None else "—"),
            ("SQP Gaps",      str(d["sqp_gaps"])),
            ("Campañas",      str(d["n_campaigns"]) if d["n_campaigns"] is not None else "—"),
            ("Tipos Campañas",str(d["campaign_types"]) if d["campaign_types"] else "—"),
            ("Funnel",        "Completo" if d["funnel_complete"] else ("Parcial" if d["funnel_complete"] is not None else "—")),
            ("Health Score",  str(d["health_score"])),
        ]
        for label, val in kpi_rows:
            bg = "F7F7F7" if r % 2 == 0 else _WHITE
            _cell(ws2, r, 1, label, bg=bg, left=True, bold=True)
            _cell(ws2, r, 2, val,   bg=bg, left=True)
            r += 1

        r += 1

        # Top Keywords
        if not d["top_kws"].empty:
            r = _sec(ws2, r, "Top Keywords (por Ventas)", 6, bg=_GRN, fg=_WHITE)
            kw_cols = list(d["top_kws"].columns)
            _hdr(ws2, r, kw_cols)
            r += 1
            for _, krow in d["top_kws"].iterrows():
                bg = "F0FFF4" if r % 2 == 0 else _WHITE
                for ci, val in enumerate(krow.values, 1):
                    _cell(ws2, r, ci, val, bg=bg, left=(ci == 1))
                r += 1
            r += 1

        # Bleeders
        if not d["bleeders"].empty:
            r = _sec(ws2, r, "Bleeders — Gasto sin Conversión", 6, bg=_RED, fg=_WHITE)
            bl_cols = list(d["bleeders"].columns)
            _hdr(ws2, r, bl_cols)
            r += 1
            for _, brow in d["bleeders"].iterrows():
                bg = "FFF5F5" if r % 2 == 0 else _WHITE
                for ci, val in enumerate(brow.values, 1):
                    _cell(ws2, r, ci, val, bg=bg, left=(ci == 1))
                r += 1

        _autofit(ws2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── Inputs ────────────────────────────────────────────────────────────────────
def _price_key(currency_code):
    """One price input per currency: a price typed for dollars must never carry over to pesos or yen."""
    code = (currency_code or "").strip().upper()
    return _PRICE_KEY if uses_dollar_price(code) else f"{_PRICE_KEY}_{code}"


def _seed_input(key, default):
    if key not in st.session_state:
        st.session_state[key] = default


@st.cache_data(ttl=60, show_spinner=False)
def _load_account_settings(profile_id):
    from core.ai_analysis.store import AiAnalysisStore

    rest = search_term_source._open_rest()
    return AiAnalysisStore(rest).settings(ANALYSIS_MODULE, profile_id) if rest is not None else None


def _seed_account_parameters(source):
    """Loads the account's saved parameters, or its defaults, into the inputs when the AM opens an account."""
    if st.session_state.get(_SEEDED_ACCOUNT_KEY) == source.profile_id:
        return
    try:
        settings = _load_account_settings(source.profile_id)
    except (requests.RequestException, StoreError) as exc:
        log.warning("ppc insights parameters for %s could not be read: %s", source.profile_id, exc)
        return
    st.session_state[_SEEDED_ACCOUNT_KEY] = source.profile_id
    params = (InsightsAnalysisParams.from_dict(settings.params, source.currency_code) if settings
              else InsightsAnalysisParams.defaults(source.currency_code))
    st.session_state[_TARGET_KEY] = params.target_acos
    key = _price_key(source.currency_code)
    if params.price is None:
        st.session_state.pop(key, None)
    else:
        st.session_state[key] = params.price


def _price_input(currency_code):
    key = _price_key(currency_code)
    label = "Precio promedio del producto"
    if uses_dollar_price(currency_code):
        _seed_input(key, InsightsAnalysisParams.defaults(currency_code).price)
        return st.number_input(label, min_value=1.0, step=0.5, key=key)
    return st.number_input(label, min_value=1.0, value=None, step=0.5, key=key,
                           placeholder=f"Precio en {currency_code}")


@st.cache_data(ttl=120, show_spinner=False)
def _advertised_asins(profile_id):
    rest = search_term_source._open_rest()
    return ReportProvider(rest).advertised_asins(profile_id) if rest is not None else {}


def _ad_group_asins(source):
    """Amazon's ASIN for each ad group of an API source; a manual file has no ad groups to look up."""
    if source.source == SOURCE_FILE or not source.profile_id:
        return {}
    try:
        return _advertised_asins(source.profile_id)
    except ReportReadError as exc:
        log.warning("advertised asins for %s could not be read: %s", source.profile_id, exc)
        st.warning(ADVERTISED_ASINS_UNREADABLE)
        return {}


def _file_digest(uploaded):
    return hashlib.sha256(uploaded.getvalue()).hexdigest()[:16] if uploaded is not None else ""


def _inputs_signature(source, uploads, target_acos):
    """What the insights on screen were computed from; "Generar Insights" keeps them until it changes."""
    parts = [source.signature, str(int(target_acos))] + [f"{name}:{_file_digest(upload)}"
                                                         for name, upload in sorted(uploads.items())]
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def _compute_insights(source, ad_group_asins, f_sqp, f_br, f_camp, target_acos):
    warnings = []
    sqp_df = None
    if f_sqp:
        sqp_df, sqp_err = _parse_sqp_cached(f_sqp.getvalue(), f_sqp.name)
        if sqp_err:
            warnings.append(f"No se pudo leer el SQP: {sqp_err}")
            sqp_df = None

    br_df = None
    if f_br:
        br_df, br_err = _parse_br(f_br.getvalue(), f_br.name)
        if br_err:
            warnings.append(f"No se pudo leer el BR: {br_err}")
            br_df = None

    camp_df = None
    if f_camp:
        camp_df, camp_err = _parse_campaigns(f_camp.getvalue())
        if camp_err:
            warnings.append(f"No se pudo leer el Campaign CSV: {camp_err}")
            camp_df = None

    resolved = resolve_asins(source.frame.copy(), ad_group_asins)
    asin_data = analyze_asins(resolved.frame.copy(), None if sqp_df is None else sqp_df.copy(),
                              None if br_df is None else br_df.copy(), None if camp_df is None else camp_df.copy(),
                              target_acos, resolved.column)
    return InsightsResult(asin_data, resolved.source, resolved.spend_share, sqp_df, br_df, camp_df, tuple(warnings),
                          resolved.report_spend, resolved.grouped_asins)


def _insights_for(signature, source, f_sqp, f_br, f_camp, target_acos):
    """The insights of these inputs, computed once: the AI tab's reruns must not recompute every ASIN."""
    kept = st.session_state.get(_RESULT_KEY)
    if kept is not None and kept[0] == signature:
        return kept[1]
    with st.spinner("Procesando archivos..."):
        result = _compute_insights(source, _ad_group_asins(source), f_sqp, f_br, f_camp, target_acos)
    st.session_state[_RESULT_KEY] = (signature, result)
    return result


def asin_coverage_caption(asin_source, spend_share):
    """Where the ASINs came from, as shares of the report's spend; "" when there is nothing to say."""
    if asin_source == NO_ASINS or not spend_share:
        return ""
    outside = sum(share for origin, share in spend_share.items() if origin in _UNATTRIBUTED_ORIGINS)
    if asin_source == FROM_FILE and outside == 0:
        return ""
    ordered = sorted(spend_share.items(), key=lambda item: -item[1])
    parts = [f"{_ORIGIN_LABELS.get(origin, origin)} {share:.1f}%" for origin, share in ordered]
    text = "Gasto por origen del ASIN: " + " · ".join(parts)
    if outside > 0:
        text += f". El {outside:.1f}% sin ASIN no entra en las cards."
    return text


def no_asin_notice(spend_share):
    """Why no term of an Amazon Ads report got an ASIN, with each cause's share of the spend."""
    ordered = sorted(spend_share.items(), key=lambda item: -item[1])
    causes = [f"{_NO_ASIN_CAUSES[origin]} ({share:.1f}%)" for origin, share in ordered if origin in _NO_ASIN_CAUSES]
    where = f" El gasto está en {' y '.join(causes)}." if causes else ""
    return ("Ningún término del período tiene ASIN: ningún nombre de campaña lleva uno." + where
            + " Se muestra la cuenta entera como una sola fila (ALL).")


def spend_kpi(cards_spend, report_spend, currency_code):
    """Label, value and caption of the spend card; the caption says of how much when the cards are not all of it."""
    value = money(cards_spend, currency_code)
    if report_spend <= 0 or cards_spend >= report_spend - 0.005:
        return "Spend total", value, None
    share = cards_spend / report_spend * 100
    return "Spend en cards", value, f"de {money(report_spend, currency_code)} ({share:.1f}%)"


# ── UI helpers ────────────────────────────────────────────────────────────────
def _score_emoji(score):
    if score >= 70:
        return "🟢"
    if score >= 40:
        return "🟡"
    return "🔴"


def _fmt_opt(val, fmt=".1f", suffix="", none_str="—"):
    if val is None:
        return none_str
    try:
        return f"{val:{fmt}}{suffix}"
    except Exception:
        return str(val)


def _period_label(source):
    if source.window_start and source.window_end:
        return date_range_label(source.window_start, source.window_end)
    return ""


def _empty_state():
    st.markdown(
        "<div style='text-align:center;padding:3rem 1rem;border:2px dashed #DDD;"
        "border-radius:12px;margin:1rem 0;'>"
        "<div style='font-size:2.5rem;margin-bottom:0.5rem;'>📂</div>"
        "<div style='font-size:0.95rem;color:#666;font-weight:600;'>Elegí una cuenta de Amazon Ads o subí el "
        "Search Term Report para comenzar el análisis.</div>"
        "</div>",
        unsafe_allow_html=True,
    )


# ── AI analysis (stored with Amazon Ads data; in memory with a file) ──────────
_FOCUS_BADGE_COLORS = {
    "ESCALAR": "background-color:#EAF3DE;color:#173404",
    "MONITOREAR": "background-color:#F1EFE8;color:#2C2C2A",
    "ACOS": "background-color:#FAEEDA;color:#412402",
    "CONVERSION": "background-color:#FAEEDA;color:#412402",
    "BUYBOX": "background-color:#FAEEDA;color:#412402",
    "FUNNEL": "background-color:#FAEEDA;color:#412402",
    "DESPERDICIO": "background-color:#FCEBEB;color:#501313",
}

INSIGHTS_FIELD_NAMES = {
    "es": {"health_score": "salud del producto", "pts_cvr": "puntos de conversión", "pts_buybox": "puntos de Buy Box",
           "pts_acos": "puntos de ACoS", "pts_funnel": "puntos de estructura de campañas",
           "pts_imp_share": "puntos de visibilidad", "spend": "gasto", "sales": "ventas", "orders": "órdenes",
           "clicks": "clicks", "acos": "ACoS", "cvr": "conversión",
           "gasto_sin_venta": "gasto sin venta de sus peores términos",
           "terminos_sin_venta": "términos que gastaron sin vender", "top_termino": "término que más vende",
           "sessions": "sesiones", "buybox": "Buy Box", "cvr_br": "conversión del Business Report",
           "campanas": "campañas habilitadas", "tipos_campana": "tipos de campaña", "funnel": "estructura de campañas",
           "asins_agrupados": "ASINs que agrupa"},
    "en": {"health_score": "product health", "pts_cvr": "conversion points", "pts_buybox": "Buy Box points",
           "pts_acos": "ACoS points", "pts_funnel": "campaign structure points", "pts_imp_share": "visibility points",
           "spend": "spend", "sales": "sales", "orders": "orders", "clicks": "clicks", "acos": "ACoS",
           "cvr": "conversion", "gasto_sin_venta": "spend without sales of its worst terms",
           "terminos_sin_venta": "terms that spent without selling", "top_termino": "best-selling term",
           "sessions": "sessions", "buybox": "Buy Box", "cvr_br": "Business Report conversion",
           "campanas": "enabled campaigns", "tipos_campana": "campaign types", "funnel": "campaign structure",
           "asins_agrupados": "ASINs it groups"},
}

_AI_LABELS = {
    "es": {"title": "Análisis IA",
           "caption": "Lectura de la IA sobre la salud de cada ASIN que ya calculó el módulo",
           "disabled": "Análisis IA deshabilitado (AI_ENABLED=0).",
           "table_title": "ASINs priorizados — lectura IA",
           "col_item": "ASIN",
           "counts": "{n} ASINs priorizados",
           "no_rows": "No hay ASINs con datos suficientes para analizar.",
           "api_scope": ("El análisis IA guardado usa sólo los datos de Amazon Ads: no incluye el SQP, el Business "
                         "Report ni el Campaign CSV que subiste.")},
    "en": {"title": "AI analysis",
           "caption": "AI read on the health of each ASIN the module already computed",
           "disabled": "AI analysis disabled (AI_ENABLED=0).",
           "table_title": "Prioritized ASINs — AI read",
           "col_item": "ASIN",
           "counts": "{n} ASINs prioritized",
           "no_rows": "No ASINs with enough data to analyze.",
           "api_scope": ("The stored AI analysis uses Amazon Ads data only: it does not include the SQP, the Business "
                         "Report or the Campaign CSV you uploaded.")},
}


def insights_ai_rows(opinions, records, glossary, currency_code):
    """Display rows for the opinion table: the ASIN, the figures the reason cites and the AI's read."""
    by_id = dict(zip(insights_row_labels(records), records))
    rows = []
    for opinion in opinions:
        row_id = str(opinion.get("row_id", ""))
        record = by_id.get(row_id)
        if record is None:
            continue
        metrics = [f"salud {record.get('health_score')}/100", f"gasto {money(record.get('spend'), currency_code)}",
                   f"ACoS {_fmt_opt(record.get('acos'), suffix='%')}", f"CVR {_fmt_opt(record.get('cvr'), suffix='%')}"]
        if record.get("gasto_sin_venta"):
            metrics.append(f"sin venta {money(record.get('gasto_sin_venta'), currency_code)}")
        if record.get("asins_agrupados"):
            metrics.append(f"agrupa {record['asins_agrupados']} ASINs")
        rows.append({
            "row_id": row_id,
            "item": record.get("asin", ""),
            "metrics": metrics,
            "badges": [opinion.get("foco", "")],
            "confidence": str(opinion.get("confianza", "")).upper(),
            "warning": _humanized(opinion.get("advertencia") or "", glossary),
            "reasoning": _humanized(opinion.get("razon", ""), glossary),
        })
    return rows


def _humanized(text, glossary):
    from core import ai_tab

    return ai_tab.humanize_fields(text, glossary) if text else ""


def _render_insights_ai_result(result, elapsed_s, records, labels, currency_code):
    from core import ai_tab

    opinions = result.get("asins") or []
    warnings = sum(1 for opinion in opinions if opinion.get("advertencia"))
    st.markdown(ai_tab.ai_chips_html(warnings, labels["counts"].format(n=len(opinions)), elapsed_s, labels),
                unsafe_allow_html=True)
    glossary = labels["field_names"]
    row_labels = insights_row_labels(records)
    synthesis = ai_tab.map_synthesis_text(
        result.get("synthesis") or {},
        lambda text: ai_tab.annotate_row_ids(ai_tab.humanize_fields(text, glossary), row_labels))
    st.markdown(ai_tab.synthesis_html(synthesis, labels), unsafe_allow_html=True)
    rows = insights_ai_rows(opinions, records, glossary, currency_code)
    if rows:
        st.markdown(ai_tab.opinion_table_html(rows, labels["table_title"], labels, _FOCUS_BADGE_COLORS),
                    unsafe_allow_html=True)


def stored_difference_text(stored, params, window_start, window_end):
    """How a stored analysis differs from what is on screen, in one or two short sentences."""
    differences = []
    if (stored.window_start, stored.window_end) != (window_start, window_end) and stored.window_start:
        differences.append(f"es del {date_range_label(stored.window_start, stored.window_end)}")
    stored_params = InsightsAnalysisParams.from_dict(stored.params, "")
    if stored_params.target_acos != params.target_acos:
        differences.append(f"usa un target de {stored_params.target_acos}%")
    if stored_params.price != params.price:
        differences.append(f"usa un precio de {stored_params.price:g}" if stored_params.price else "no tenía precio")
    if not differences:
        return ("Se generó antes de que llegaran los datos nuevos de Amazon Ads de este período. "
                "Recalculalo para leer lo que estás viendo.")
    return f"Este análisis {' y '.join(differences)}. Recalculalo para leer lo que estás viendo."


def _stored_chat_documents(stored, subject):
    """What the chat reads of the analysis on screen: its own rows and the AI's read, whatever data it covers."""
    from ai.agents.ppc_insights import chat_document

    rows = pd.DataFrame(stored.records)
    if not rows.empty:
        rows.insert(0, "row_id", list(insights_row_labels(stored.records)))
    period = date_range_label(stored.window_start, stored.window_end) if stored.window_start else "sin período"
    target = stored.params.get("target_acos")
    return (
        {"title": f"PPC Insights · {subject} · Salud por ASIN ({period}, target {target}%)",
         "content": rows.to_csv(index=False, lineterminator="\n")},
        {"title": f"PPC Insights · {subject} · Lectura de la IA",
         "content": chat_document.reading_text(stored.result or {}, stored.records)},
    )


def _profile_country(profile_id):
    for option in search_term_source._available_profiles():
        if option.profile_id == profile_id:
            return option.country_code
    return ""


def _render_stored_analysis(source, params, labels, currency_code, uploaded_files):
    """Amazon Ads data: the analysis the worker stored, or the account's latest, with Recalcular."""
    from ai.agent_call import build_agent_call
    from core import ai_tab
    from core.ai_analysis import stored_tab
    from core.chat import app_chat

    analysis_input = build_analysis_input(
        source.frame, params=params, account_label=source.label, period_label=_period_label(source),
        currency_code=currency_code, lang=CANONICAL_LANG, ad_group_asins=_ad_group_asins(source))
    if analysis_input.data is None:
        st.info(labels["no_rows"])
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return
    if uploaded_files:
        st.caption(labels["api_scope"])

    call = build_agent_call(ANALYSIS_MODULE, analysis_input.data)
    settings = stored_tab._settings(ANALYSIS_MODULE, source.profile_id, search_term_source._open_rest)
    account_params = (InsightsAnalysisParams.from_dict(settings.params, currency_code) if settings
                      else InsightsAnalysisParams.defaults(currency_code))

    def _render(stored):
        # The rows that analysis read, not this run's: the row_ids it cites are its own.
        _render_insights_ai_result(stored.result or {}, int((stored.duration_ms or 0) / 1000), stored.records,
                                   labels, currency_code)

    result = stored_tab.render_recalculable_analysis(
        module=ANALYSIS_MODULE, key_prefix="insights", source=source, input_digest=call.input_digest,
        agent_version=call.agent_version, params=params, account_params=account_params,
        open_rest=search_term_source._open_rest, current_username=search_term_source._current_username,
        render_result=_render,
        describe_difference=partial(stored_difference_text, params=params, window_start=source.window_start,
                                    window_end=source.window_end),
        timezone=search_term_source.DISPLAY_TIMEZONE)

    if result.state == stored_tab.STATE_RUNNING:
        app_chat.report_running(ANALYSIS_MODULE)
    elif result.state == stored_tab.STATE_FAILED:
        app_chat.report_failed(ANALYSIS_MODULE)
    elif result.analysis is None:
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
    else:
        stored = result.analysis
        state = (app_chat.AnalysisState.CURRENT if result.state == stored_tab.STATE_CURRENT
                 else app_chat.AnalysisState.OUTDATED)
        app_chat.share_analysis(app_chat.ChatAnalysis(
            module=ANALYSIS_MODULE, key=f"{ANALYSIS_MODULE}:{stored.id}", subject=source.label,
            documents=_stored_chat_documents(stored, source.label),
            annotate=partial(ai_tab.annotate_row_ids, labels_by_id=insights_row_labels(stored.records)),
            country_code=_profile_country(source.profile_id), profile_id=source.profile_id), state)


def _render_file_analysis(source, params, labels, currency_code, insights, signature, lang):
    """A hand-uploaded report: nothing to store, so the analysis runs when the AM asks for it."""
    from ai.agents.ppc_insights import chat_document
    from core import ai_tab

    analysis_input = build_analysis_input(
        source.frame, params=params, account_label=source.label, period_label="", currency_code=currency_code,
        lang=lang, sqp_df=insights.sqp_df, br_df=insights.br_df, camp_df=insights.camp_df)
    if analysis_input.data is None:
        from core.chat import app_chat

        st.info(labels["no_rows"])
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return
    payload = analysis_input.data
    analysis = ai_tab.resolve_analysis(slug=ANALYSIS_MODULE, payload=payload, file_signature=signature,
                                       labels=labels, show_previous=True, auto_fire=False)
    records = analysis_input.records
    if analysis is not None:
        records = ai_tab.records_for_render(ANALYSIS_MODULE, analysis, payload, analysis_input.records)

        def _render_result(result, a, _records=records):
            _render_insights_ai_result(result, a.elapsed, _records, labels, currency_code)

        ai_tab.render_analysis(analysis, slug=ANALYSIS_MODULE, labels=labels, render_result=_render_result)
    ai_tab.publish_analysis_to_chat(
        ANALYSIS_MODULE, analysis, payload, module_label="PPC Insights", subject=source.label,
        reading=lambda a, _records=records: chat_document.reading_text(a.result, _records),
        annotate=partial(ai_tab.annotate_row_ids, labels_by_id=insights_row_labels(records)))


def _render_ai_tab(source, params, currency_code, insights, signature, uploaded_files):
    from ai.config import AI_ENABLED
    from core import ai_tab
    from core.chat import app_chat

    lang = ai_tab.app_language()
    texts = _AI_LABELS.get(lang, _AI_LABELS["es"])
    st.subheader(texts["title"])
    st.caption(texts["caption"])
    if not AI_ENABLED:
        st.caption(texts["disabled"])
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return
    labels = ai_tab.ai_labels(lang, {**texts, "field_names": INSIGHTS_FIELD_NAMES.get(lang, INSIGHTS_FIELD_NAMES["es"])})
    st.markdown(ai_tab.AI_CSS, unsafe_allow_html=True)
    if source.source == SOURCE_FILE:
        _render_file_analysis(source, params, labels, currency_code, insights, signature, lang)
    else:
        _render_stored_analysis(source, params, labels, currency_code, uploaded_files)


# ── render ────────────────────────────────────────────────────────────────────
def render():
    from core.chat import app_chat

    st.markdown(
        "<div style='display:flex;align-items:center;gap:0.75rem;margin-bottom:0.25rem;'>"
        "<span style='font-size:2rem;'>🔎</span>"
        "<div>"
        "<div style='font-size:1.3rem;font-weight:800;color:#1F1F1F;'>PPC Insights</div>"
        "<div style='font-size:0.82rem;color:#888;'>Análisis integral por ASIN — cruza STR, SQP, BR y Campañas para un health score completo.</div>"
        "</div>"
        "</div>",
        unsafe_allow_html=True,
    )
    st.divider()

    with st.expander("❓ ¿Cómo usar este módulo?", expanded=False):
        col1, col2, col3 = st.columns(3)
        with col1:
            st.markdown("**🎯 Para qué sirve**")
            st.caption("Health score 0-100 por ASIN cruzando STR + SQP + BR + Campañas. Detecta wasted spend y bleeders.")
        with col2:
            st.markdown("**📂 De dónde salen los datos**")
            st.caption("STR de la cuenta conectada de Amazon Ads o subido a mano (requerido). SQP + BR by ASIN + "
                       "Campaign CSV (opcionales — mínimo 2 fuentes para score confiable).")
        with col3:
            st.markdown("**➡️ Siguiente paso**")
            st.caption("PPC Audit (M20) para auditoría estructural o Bid Optimizer (M9) para ajustar bids.")
        st.markdown("**▶️ Pasos:**")
        st.markdown(
            "1. Elegí la cuenta y el período (o subí el STR a mano)\n"
            "2. Configurá Target ACoS, precio promedio y nombre del cliente\n"
            "3. Sumá los archivos opcionales que tengas y tocá Generar Insights\n"
            "4. Revisá los cards por ASIN — score 🟢 ≥70 / 🟡 40-69 / 🔴 <40\n"
            "5. Pestaña Análisis IA: la lectura guardada de la cuenta, o Recalcular\n"
            "6. Descargá el Excel multi-sheet (hasta 10 hojas por ASIN)"
        )

    source = render_source_picker(key_prefix="insights", module_label="PPC Insights", manual_reader=_read_manual_str)
    currency_code = source.currency_code if source is not None else ""
    if source is not None and source.source != SOURCE_FILE and source.profile_id:
        _seed_account_parameters(source)

    # ── Global inputs ─────────────────────────────────────────────────────────
    _seed_input(_TARGET_KEY, InsightsAnalysisParams.defaults(currency_code).target_acos)
    col_a, col_b = st.columns(2)
    with col_a:
        target_acos = st.slider("Target ACoS (%)", min_value=5, max_value=80, step=1, key=_TARGET_KEY)
    with col_b:
        precio_promedio = _price_input(currency_code)

    client_name = st.text_input(
        "Nombre del cliente (para el Excel)",
        value="Cliente",
        key="insights_client",
    )

    st.markdown("---")

    # ── File uploaders ────────────────────────────────────────────────────────
    st.markdown("**Archivos opcionales**")
    c1, c2, c3 = st.columns(3)
    with c1:
        f_sqp  = st.file_uploader("Search Query Performance (opcional)", type=["xlsx", "csv"], key="insights_sqp")
    with c2:
        f_br   = st.file_uploader("Business Report by ASIN (opcional)", type=["csv", "xlsx"], key="insights_br")
    with c3:
        f_camp = st.file_uploader("Campaign CSV (opcional)", type=["csv"], key="insights_camp")

    if source is None:
        _empty_state()
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return

    uploads = {"sqp": f_sqp, "br": f_br, "camp": f_camp}
    signature = _inputs_signature(source, uploads, target_acos)
    if st.button("Generar Insights", type="primary", key="insights_run"):
        st.session_state[_GENERATED_FOR_KEY] = signature
    if st.session_state.get(_GENERATED_FOR_KEY) != signature:
        app_chat.withdraw_analysis(ANALYSIS_MODULE)
        return

    try:
        insights = _insights_for(signature, source, f_sqp, f_br, f_camp, target_acos)
    except Exception as e:  # a malformed optional file must end in a message, not a traceback
        log.exception("ppc insights could not be computed")
        st.error(f"Error durante el análisis: {e}")
        return
    asin_data = insights.asin_data
    for warning in insights.file_warnings:
        st.warning(warning)

    if not asin_data:
        st.warning("No se encontraron ASINs en el STR.")
        return

    if insights.asin_source == NO_ASINS:
        st.info(NO_ASIN_FROM_FILE if source.source == SOURCE_FILE else no_asin_notice(insights.spend_share))
    else:
        coverage = asin_coverage_caption(insights.asin_source, insights.spend_share)
        if coverage:
            st.caption(coverage)

    sorted_asins = sorted(asin_data.keys(), key=lambda a: asin_data[a]["spend"], reverse=True)
    show_money = partial(money, currency_code=currency_code)
    money_column = st.column_config.NumberColumn(format=f"{currency_symbol(currency_code)}%.2f")
    params = InsightsAnalysisParams(int(target_acos), precio_promedio)

    tab_asins, tab_ai = st.tabs(["Insights por ASIN", "Análisis IA"])

    with tab_asins:
        # ── Summary metrics ───────────────────────────────────────────────────
        total_spend     = sum(d["spend"] for d in asin_data.values())
        total_wasted    = sum(d["wasted_spend"] for d in asin_data.values())
        avg_score       = sum(d["health_score"] for d in asin_data.values()) / len(asin_data)
        n_asins         = len(asin_data)

        st.markdown("### Resumen de cuenta")
        mc1, mc2, mc3, mc4 = st.columns(4)
        with mc1:
            st.markdown(kpi_card("ASINs analizados", str(n_asins)), unsafe_allow_html=True)
        with mc2:
            st.markdown(kpi_card("Health Score prom.", f"{avg_score:.0f}/100"), unsafe_allow_html=True)
        with mc3:
            spend_label, spend_value, spend_caption = spend_kpi(total_spend, insights.report_spend, currency_code)
            st.markdown(kpi_card(spend_label, spend_value, caption=spend_caption), unsafe_allow_html=True)
        with mc4:
            pct = float(total_wasted / total_spend * 100) if total_spend > 0 else 0
            st.markdown(kpi_card("Wasted spend", show_money(total_wasted), delta=-pct, delta_good=False),
                        unsafe_allow_html=True)

        st.divider()

        # ── Per-ASIN cards ────────────────────────────────────────────────────
        st.markdown("### Análisis por ASIN")

        for asin in sorted_asins:
            d = asin_data[asin]
            score = d["health_score"]
            emoji = _score_emoji(score)

            c1, c2, c3, c4 = st.columns([3, 1, 1, 1])
            with c1:
                st.markdown(f"**{emoji} {asin}**")
                if insights.grouped_asins.get(asin):
                    st.caption(f"Nombre de campaña · agrupa {insights.grouped_asins[asin]} ASINs")
            with c2:
                st.metric("Health Score", f"{score}/100", label_visibility="visible")
            with c3:
                st.metric("ACoS", _fmt_opt(d["acos"], suffix="%"))
            with c4:
                st.metric("Spend", show_money(d["spend"]))

            with st.expander(f"Detalle {asin}"):
                tab_str, tab_sqp, tab_br, tab_camp = st.tabs(
                    ["STR", "SQP", "Business Report", "Campañas"]
                )

                with tab_str:
                    dc1, dc2, dc3, dc4 = st.columns(4)
                    dc1.metric("Spend",    show_money(d["spend"]))
                    dc2.metric("Sales",    show_money(d["sales"]))
                    dc3.metric("ACoS",     _fmt_opt(d["acos"],  suffix="%"))
                    dc4.metric("CVR",      _fmt_opt(d["cvr"],   suffix="%"))
                    dc1b, dc2b, dc3b, dc4b = st.columns(4)
                    dc1b.metric("Orders",  f"{d['orders']:.0f}")
                    dc2b.metric("Clicks",  f"{d['clicks']:.0f}")
                    dc3b.metric("Wasted",  show_money(d["wasted_spend"]))
                    dc4b.metric("", "")

                    if not d["top_kws"].empty:
                        st.markdown("**Top Keywords por ventas**")
                        st.dataframe(d["top_kws"], use_container_width=True, hide_index=True,
                                     column_config={"Sales": money_column, "Spend": money_column})

                    if not d["bleeders"].empty:
                        st.markdown("**Bleeders — gasto sin conversion**")
                        st.dataframe(d["bleeders"], use_container_width=True, hide_index=True,
                                     column_config={"Spend": money_column})

                with tab_sqp:
                    if insights.sqp_df is None:
                        st.info("No se subio el archivo SQP.")
                    else:
                        dc1, dc2, dc3 = st.columns(3)
                        dc1.metric("Imp Share",      _fmt_opt(d["imp_share"],      suffix="%"))
                        dc2.metric("Purchase Share", _fmt_opt(d["purchase_share"], suffix="%"))
                        dc3.metric("Gaps detectados", str(d["sqp_gaps"]))

                with tab_br:
                    if insights.br_df is None:
                        st.info("No se subio el Business Report.")
                    else:
                        dc1, dc2, dc3 = st.columns(3)
                        dc1.metric("Sessions",  _fmt_opt(d["sessions"], ".0f"))
                        dc2.metric("BuyBox %",  _fmt_opt(d["buybox"],   suffix="%"))
                        dc3.metric("CVR (BR)",  _fmt_opt(
                            (d["br_units"] / d["sessions"] * 100)
                            if d["br_units"] is not None and d["sessions"] and d["sessions"] > 0
                            else None,
                            suffix="%",
                        ))

                with tab_camp:
                    if insights.camp_df is None:
                        st.info("No se subio el Campaign CSV.")
                    else:
                        dc1, dc2 = st.columns(2)
                        dc1.metric("Campañas (ENABLED)", str(d["n_campaigns"]) if d["n_campaigns"] is not None else "—")
                        dc2.metric("Tipos detectados", d["campaign_types"] or "—")
                        if d["funnel_complete"] is not None:
                            if d["funnel_complete"]:
                                st.success("Funnel completo detectado (Auto + Exact presentes).")
                            else:
                                st.warning("Funnel incompleto — revisar cobertura de match types.")

            st.divider()

        # ── Excel export ──────────────────────────────────────────────────────
        st.markdown("### Exportar")

        try:
            excel_bytes = _build_insights_excel(asin_data, client_name, target_acos, currency_code)
            st.download_button(
                label="Descargar Insights Excel",
                data=excel_bytes,
                file_name=f"PPC_Insights_{client_name}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="insights_dl",
            )
        except Exception as e:
            st.error(f"Error al generar el Excel: {e}")

    with tab_ai:
        uploaded_files = any(upload is not None for upload in uploads.values())
        _render_ai_tab(source, params, currency_code, insights, signature, uploaded_files)
