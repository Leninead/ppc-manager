import io
from datetime import datetime
import re

import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.helpers import kpi_card


# ── Keyword dictionaries ──────────────────────────────────────────────────

_CRITICAL_EN = [
    "weighted", "weighted blanket", "weighted swaddle", "weighted sleep",
    "weighted sleep sack", "weighted sack", "weight therapy",
]
_CRITICAL_ES = [
    "con peso", "ponderado", "ponderada", "ponderados", "ponderadas",
    "manta con peso", "saco con peso", "cobija con peso",
    "saquito con peso", "manta ponderada", "saco ponderado",
]
_HIGH_EN = [
    "gentle pressure", "deep pressure", "pressure therapy",
    "calming weight", "soothing weight", "therapeutic weight",
    "added weight", "extra weight", "heavier feel",
]
_HIGH_ES = [
    "presion suave", "presión suave", "presion profunda", "presión profunda",
    "presion calmante", "presión calmante", "terapia de presion",
    "terapia de presión", "peso calmante", "peso terapeutico",
    "peso terapéutico", "sensacion de peso", "sensación de peso",
]
_MEDIUM_EN = [
    "heavy", "heavier", "mimics touch", "like being held",
    "feels like a hug", "hug-like", "snug pressure",
]
_MEDIUM_ES = [
    "pesado", "pesada", "mas pesado", "más pesado", "mas pesada",
    "más pesada", "como estar abrazado", "como un abrazo",
    "sensacion de abrazo", "sensación de abrazo", "como abrazo",
]

_SEVERITY_MAP = [
    ("CRITICAL", _CRITICAL_EN + _CRITICAL_ES),
    ("HIGH", _HIGH_EN + _HIGH_ES),
    ("MEDIUM", _MEDIUM_EN + _MEDIUM_ES),
]


# ── Fields to scan ────────────────────────────────────────────────────────

_LISTING_FIELDS = [
    "item-name", "item_name", "title", "product-title",
    "bullet-point1", "bullet_point1", "bullet-point-1", "bullet_point_1",
    "bullet-point2", "bullet_point2", "bullet-point-2", "bullet_point_2",
    "bullet-point3", "bullet_point3", "bullet-point-3", "bullet_point_3",
    "bullet-point4", "bullet_point4", "bullet-point-4", "bullet_point_4",
    "bullet-point5", "bullet_point5", "bullet-point-5", "bullet_point_5",
    "product-description", "product_description", "description",
    "generic-keywords", "generic_keywords", "search-terms", "search_terms",
    "backend-keywords", "backend_keywords",
    "subject-keywords", "subject_keywords",
    "special-features", "special_features",
    "intended-use", "intended_use",
]
_ASIN_COLS = ["asin", "ASIN", "asin1", "product-id", "product_id", "child-asin"]
_SKU_COLS = ["sku", "SKU", "seller-sku", "seller_sku", "item-sku", "item_sku"]


# ── Helpers ───────────────────────────────────────────────────────────────

_ACCENT_MAP = str.maketrans("áéíóúñü", "aeiounu")


def _normalize(text):
    """Lowercase, strip accents, collapse whitespace."""
    if not isinstance(text, str) or not text.strip():
        return ""
    return re.sub(r"\s+", " ", text.lower().translate(_ACCENT_MAP).strip())


def _find_matches(text, keywords):
    """Return list of (keyword, position, context_snippet) for matches in text."""
    norm = _normalize(text)
    if not norm:
        return []
    results = []
    for kw in keywords:
        norm_kw = _normalize(kw)
        if not norm_kw:
            continue
        # Single word → word boundary; multi-word → substring
        if " " not in norm_kw:
            pattern = r"\b" + re.escape(norm_kw) + r"\b"
        else:
            pattern = re.escape(norm_kw)
        for m in re.finditer(pattern, norm):
            pos = m.start()
            start = max(0, pos - 30)
            end = min(len(norm), m.end() + 30)
            snippet = norm[start:pos] + ">>>" + norm[pos:m.end()] + "<<<" + norm[m.end():end]
            results.append((kw, pos, snippet))
    return results


def _detect_asin_col(df):
    """Detect ASIN column in df."""
    for c in _ASIN_COLS:
        if c in df.columns:
            return c
    for c in df.columns:
        if "asin" in c.lower():
            return c
    return None


def _detect_sku_col(df):
    """Detect SKU column in df."""
    for c in _SKU_COLS:
        if c in df.columns:
            return c
    for c in df.columns:
        if "sku" in c.lower():
            return c
    return None


def _detect_listing_fields(df):
    """Detect listing content columns to scan."""
    cols = []
    lower_map = {c.lower(): c for c in df.columns}
    # Exact matches first
    for f in _LISTING_FIELDS:
        if f.lower() in lower_map:
            real = lower_map[f.lower()]
            if real not in cols:
                cols.append(real)
    # Fuzzy catch
    fuzzy_terms = ["title", "bullet", "descr", "keyword", "search-term", "search_term", "feature"]
    for cl, real in lower_map.items():
        if real in cols:
            continue
        if any(t in cl for t in fuzzy_terms):
            cols.append(real)
    return cols


# ── Parser ────────────────────────────────────────────────────────────────

@st.cache_data(max_entries=3, ttl=3600, show_spinner=False)
def _parse_listings(data_bytes, filename):
    """Parse listing export file to DataFrame."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("xlsx", "xls"):
        return pd.read_excel(io.BytesIO(data_bytes))
    if ext == "txt":
        return pd.read_csv(io.BytesIO(data_bytes), sep="\t", encoding="utf-8-sig")
    # CSV — try multiple encodings
    for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            return pd.read_csv(io.BytesIO(data_bytes), encoding=enc, sep=None, engine="python")
        except Exception:
            continue
    return pd.read_csv(io.BytesIO(data_bytes), encoding="latin-1")


# ── Scanner ───────────────────────────────────────────────────────────────

def _scan_dataframe(df, asin_filter=None):
    """Scan listing fields for compliance keywords. Returns list of finding dicts."""
    asin_col = _detect_asin_col(df)
    sku_col = _detect_sku_col(df)
    fields = _detect_listing_fields(df)
    if not fields:
        return []

    findings = []
    for _, row in df.iterrows():
        asin_val = str(row.get(asin_col, "")) if asin_col else ""
        sku_val = str(row.get(sku_col, "")) if sku_col else ""
        if asin_val == "nan":
            asin_val = ""
        if sku_val == "nan":
            sku_val = ""

        # Apply ASIN/SKU filter
        if asin_filter:
            af = asin_filter.upper()
            if af not in asin_val.upper() and af not in sku_val.upper():
                continue

        for field in fields:
            text = str(row.get(field, ""))
            if text == "nan" or not text.strip():
                continue

            # Collect all matches across severities, dedup by position
            all_matches = {}  # pos -> (severity, keyword, context, full_text)
            for severity, kw_list in _SEVERITY_MAP:
                for kw, pos, ctx in _find_matches(text, kw_list):
                    existing = all_matches.get(pos)
                    if existing is None or len(kw) > len(existing[1]):
                        all_matches[pos] = (severity, kw, ctx, text)

            for _pos, (severity, kw, ctx, full_text) in all_matches.items():
                findings.append({
                    "ASIN": asin_val,
                    "SKU": sku_val,
                    "Field": field,
                    "Severity": severity,
                    "Keyword": kw,
                    "Position": _pos,
                    "Context": ctx,
                    "Full Text": full_text,
                })

    return findings


# ── Excel export ──────────────────────────────────────────────────────────

def _build_compliance_excel(findings):
    """Build compliance report XLSX. Returns bytes."""
    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    wrap_align = Alignment(wrap_text=True, vertical="top")

    sev_fills = {
        "CRITICAL": PatternFill(start_color="FBEAE7", end_color="FBEAE7", fill_type="solid"),
        "HIGH": PatternFill(start_color="FDF3E3", end_color="FDF3E3", fill_type="solid"),
        "MEDIUM": PatternFill(start_color="FEF9E3", end_color="FEF9E3", fill_type="solid"),
    }

    def _write_headers(ws, headers, widths):
        for i, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = hdr_align
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

    # ── Sheet 1: Findings ──
    ws1 = wb.active
    ws1.title = "🚨 Findings"
    hdrs1 = ["Severity", "ASIN", "SKU", "Field", "Keyword", "Context", "Full Text"]
    widths1 = [12, 14, 18, 22, 20, 55, 60]
    _write_headers(ws1, hdrs1, widths1)

    for r, f in enumerate(findings, 2):
        vals = [f["Severity"], f["ASIN"], f["SKU"], f["Field"], f["Keyword"], f["Context"], f["Full Text"]]
        fill = sev_fills.get(f["Severity"])
        for c, v in enumerate(vals, 1):
            cell = ws1.cell(row=r, column=c, value=v)
            cell.border = thin_border
            cell.alignment = wrap_align
            if fill:
                cell.fill = fill

    # ── Sheet 2: Summary by ASIN ──
    ws2 = wb.create_sheet("📊 Summary by ASIN")
    hdrs2 = ["ASIN/SKU", "CRITICAL", "HIGH", "MEDIUM", "Total", "Fields Affected", "Action"]
    widths2 = [18, 12, 10, 12, 10, 40, 25]
    _write_headers(ws2, hdrs2, widths2)

    summary = _build_summary(findings)
    for r, row in enumerate(summary, 2):
        for c, v in enumerate(row, 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = thin_border
            cell.alignment = wrap_align

    # ── Sheet 3: Keyword Dictionary ──
    ws3 = wb.create_sheet("📖 Keyword Dictionary")
    hdrs3 = ["Severity", "Language", "Keyword", "Why it triggers"]
    widths3 = [12, 10, 30, 60]
    _write_headers(ws3, hdrs3, widths3)

    explanations = {
        "CRITICAL": "Direct match for 'weighted product' category — instant flag",
        "HIGH": "Strongly associated with weighted-product marketing language",
        "MEDIUM": "Borderline — OK only if referring to swaddle snug fit, not added weight",
    }
    kw_rows = []
    for sev, en_list, es_list in [
        ("CRITICAL", _CRITICAL_EN, _CRITICAL_ES),
        ("HIGH", _HIGH_EN, _HIGH_ES),
        ("MEDIUM", _MEDIUM_EN, _MEDIUM_ES),
    ]:
        for kw in en_list:
            kw_rows.append((sev, "EN", kw, explanations[sev]))
        for kw in es_list:
            kw_rows.append((sev, "ES", kw, explanations[sev]))

    for r, row in enumerate(kw_rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws3.cell(row=r, column=c, value=v)
            cell.border = thin_border
            cell.alignment = wrap_align

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_summary(findings):
    """Build summary rows sorted by severity counts desc."""
    by_id = {}
    for f in findings:
        key = f["ASIN"] or f["SKU"] or "UNKNOWN"
        if key not in by_id:
            by_id[key] = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "fields": set()}
        by_id[key][f["Severity"]] += 1
        by_id[key]["fields"].add(f["Field"])

    rows = []
    for asin_sku, d in by_id.items():
        total = d["CRITICAL"] + d["HIGH"] + d["MEDIUM"]
        fields_str = ", ".join(sorted(d["fields"]))
        if d["CRITICAL"] > 0:
            action = "🔴 FIX IMMEDIATELY"
        elif d["HIGH"] > 0:
            action = "🟠 Review this week"
        else:
            action = "🟡 Review in context"
        rows.append((asin_sku, d["CRITICAL"], d["HIGH"], d["MEDIUM"], total, fields_str, action))

    rows.sort(key=lambda x: (-x[1], -x[2], -x[3]))
    return rows


# ── Render ────────────────────────────────────────────────────────────────

def render():
    st.markdown("### 🛡️ Listing Compliance Scanner")
    st.caption(
        "Detecta keywords que el bot de compliance de Amazon usa para flaggear "
        "\"weighted infant and toddler sleep products\" (categoría prohibida desde 2024). "
        "Scanea título, bullets, descripción, backend keywords y atributos en EN+ES."
    )

    uploaded = st.file_uploader(
        "Subir export de listings (Seller Central)",
        type=["csv", "xlsx", "xls", "txt"],
        key="compliance_upload",
    )

    asin_filter = st.text_input(
        "Filtrar por ASIN o SKU (opcional)",
        placeholder="Ej: B00MJXHM48 o L10 01 001 GR M",
        key="compliance_asin_filter",
    ).strip() or None

    if not uploaded:
        return

    df = _parse_listings(uploaded.getvalue(), uploaded.name)
    asin_col = _detect_asin_col(df)
    sku_col = _detect_sku_col(df)
    fields = _detect_listing_fields(df)

    st.caption(f"ASIN col: **{asin_col or 'no detectada'}** · SKU col: **{sku_col or 'no detectada'}** · {len(df)} filas · {len(fields)} campos a escanear")

    findings = _scan_dataframe(df, asin_filter=asin_filter)

    if not findings:
        st.success("✅ NO COMPLIANCE ISSUES FOUND. Listings are clean.")
        return

    # KPI cards
    n_crit = len({f["ASIN"] or f["SKU"] for f in findings if f["Severity"] == "CRITICAL"})
    n_high = len({f["ASIN"] or f["SKU"] for f in findings if f["Severity"] == "HIGH"})
    n_med = len({f["ASIN"] or f["SKU"] for f in findings if f["Severity"] == "MEDIUM"})

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.markdown(kpi_card("Total Findings", len(findings)), unsafe_allow_html=True)
    with c2:
        st.markdown(kpi_card("ASINs CRITICAL", n_crit), unsafe_allow_html=True)
    with c3:
        st.markdown(kpi_card("ASINs HIGH", n_high), unsafe_allow_html=True)
    with c4:
        st.markdown(kpi_card("ASINs MEDIUM", n_med), unsafe_allow_html=True)

    # Tabs
    tab_crit, tab_high, tab_med, tab_summary = st.tabs(
        ["🔴 Critical", "🟠 High", "🟡 Medium", "📊 Summary"]
    )

    display_cols = ["ASIN", "SKU", "Field", "Keyword", "Context"]

    with tab_crit:
        df_c = pd.DataFrame([f for f in findings if f["Severity"] == "CRITICAL"])
        if df_c.empty:
            st.caption("Sin hallazgos en este nivel.")
        else:
            st.dataframe(df_c[display_cols], use_container_width=True, hide_index=True)

    with tab_high:
        df_h = pd.DataFrame([f for f in findings if f["Severity"] == "HIGH"])
        if df_h.empty:
            st.caption("Sin hallazgos en este nivel.")
        else:
            st.dataframe(df_h[display_cols], use_container_width=True, hide_index=True)

    with tab_med:
        df_m = pd.DataFrame([f for f in findings if f["Severity"] == "MEDIUM"])
        if df_m.empty:
            st.caption("Sin hallazgos en este nivel.")
        else:
            st.dataframe(df_m[display_cols], use_container_width=True, hide_index=True)

    with tab_summary:
        summary_rows = _build_summary(findings)
        df_sum = pd.DataFrame(
            summary_rows,
            columns=["ASIN/SKU", "CRITICAL", "HIGH", "MEDIUM", "Total", "Fields Affected", "Action"],
        )
        st.dataframe(df_sum, use_container_width=True, hide_index=True)

    # Download
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_bytes = _build_compliance_excel(findings)
    st.download_button(
        "⬇️ Descargar Reporte de Compliance (Excel)",
        data=excel_bytes,
        file_name=f"compliance_scan_{ts}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="compliance_dl",
    )
