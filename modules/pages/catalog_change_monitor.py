"""M33 — Catalog Change Monitor (port del HTML standalone a módulo Streamlit).

FASES INGESTA + DIFF (núcleo).
Sin severidades (D3), sin comparación de unidades (R10), sin UI, sin router, sin
persistencia (esas fases vienen después). Este archivo todavía NO expone render()
— nada routea acá.

Fuente de verdad del comportamiento: `.claude/porting-sources/m33/monitor_cambios.html`
(gitignored). Validado contra archivos reales de Setex MX (2026-08-18), ver
`notes/modules/m33-catalog-change-monitor.md`.

Contrato hasta acá — 4 parsers puros + 1 ensamblador + 2 funciones de diff:

    _parse_all_listings(data)                  -> {seller-sku: {...}}
    _parse_fee_preview(data, marketplace)      -> {sku: {...}}
    _parse_category_listings(data)             -> {sku: {...}}
    _parse_business_report(data)               -> {(Child) ASIN: {...}}
    _build_snapshot(inv, fee, cat, br)         -> {sku: {...}}
    _diff_snapshots(viejo, nuevo)              -> [{cambio}, ...]
    _agrupar_alertas_por_asin(cambios)         -> [{alerta}, ...]

Riesgos de la nota que este archivo implementa:

    R2  — hoja `Plantilla` buscada con match ANCLADO `^plantilla$` (case-insensitive).
          El mismo .xlsm trae `Cambios en la plantilla`: un match por "contiene"
          agarraría la hoja equivocada.
    R3  — encodings verificados sobre los archivos reales: All Listings utf-8-sig
          (BOM real) TSV · Fee Preview ISO-8859-1 con un byte ASCII 0x3F ('?')
          literal pegado antes del primer header · Business Report utf-8-sig.
    R4  — todo valor sale como str trimmed. Única excepción documentada:
          `cat_es_padre`, que es un flag DERIVADO (no se diffea) y sale bool.
    R6  — el dedupe por SKU del Fee Preview cuenta y loguea las filas descartadas
          (el `dedupeBySku` del HTML las tiraba en silencio).
    R7  — los (Child) ASIN del Business Report que no cruzan contra ningún SKU se
          cuentan y loguean; no rompen el ensamblado.
    R11 — `parentage_level` en MX viene como 'Principal.' (CON punto final) y
          'Niños'. Se normaliza (lower + sin acentos + sin punto) contra un
          diccionario multi-idioma; el prototipo comparaba contra
          parent/eltern/padre y no matcheaba NADA en MX.
    R13 — el Fee Preview de Norteamérica trae US/CA/MX en el mismo archivo, con
          3 monedas y 2 sistemas de unidades. El filtro por `amazon-store` es
          OBLIGATORIO y corre ANTES del dedupe.
    R14 — `product-size-tier` NO existe en el Fee Preview de NA. `fee_size_tier`
          queda vacío EXPLÍCITO, no ausente: así el diff no lo lee como campo nuevo.
    R15 — `expected-fulfillment-fee-per-unit` no existe en NA. `fee_fulfillment`
          se compone sumando pick-pack + weight-handling.
    R17 — PROHIBIDO el fallback posicional del prototipo (asumía imagen principal
          en 28 y bullets en 39-43; en MX son 33 y 44-48). Si el mapeo por nombre
          técnico falla, se levanta excepción RUIDOSA.
    R20 — buy_box es propiedad del ASIN, no del SKU: el ensamblado lo REPLICA a
          todos los SKUs que comparten el ASIN.
    R22 — consecuencia directa de R20: sin agrupar, un solo cambio de Buy Box
          genera N alertas idénticas (una por SKU del ASIN). Es el modo de falla
          de R8 por otra vía. `_agrupar_alertas_por_asin` las colapsa en una.

Decisiones de forma:

    - Los parsers reciben `bytes`, no rutas — convención del repo (M30
      `_parse_fba(data: bytes)`, M28 `_parse_csv_bytes(raw, filename)`): es lo que
      entrega el uploader de Streamlit y lo que `@st.cache_data` puede hashear.
    - `@st.cache_data` NO se aplica todavía: el out-param `diag` es un dict mutable
      (no hasheable) y rompería el cache. Se envuelve en un shim cacheado cuando se
      escriba la UI.
    - El mapeo del Category Listings es por NOMBRE BASE, no exacto: la fila 5 trae
      los atributos con sufijos (`contribution_sku#1.value`,
      `parentage_level[marketplace_id=A1AM78C64UM0Y8]#1.value`,
      `bullet_point[...][language_tag=es_MX]#3.value`). Se corta en el primer
      '[' o '#'. Los repetidos (other_product_image_locator_N x8, bullet_point x5)
      preservan el orden de columna para poder diffear slot por slot (D5).
"""

from __future__ import annotations

import csv
import io
import logging
import re
import unicodedata
from io import BytesIO

from openpyxl import load_workbook

_LOG = logging.getLogger(__name__)


# =====================================================================
# Constantes
# =====================================================================

# Fila (1-indexed) de la hoja `Plantilla` que trae los nombres técnicos, y
# primera fila de datos. La fila 6 es el ejemplo que Amazon precarga (SKU
# 'ABC123') y se saltea.
_CAT_TECH_ROW = 5
_CAT_FIRST_DATA_ROW = 7

# Atributos de una sola columna. nombre base -> campo del snapshot.
_CAT_SINGLE_FIELDS: dict[str, str] = {
    "::listing_status": "cat_status",
    "::title": "cat_title",
    "contribution_sku": "cat_sku",
    "product_type": "cat_product_type",
    "parentage_level": "cat_parentage_raw",
    "child_parent_sku_relationship": "cat_parent_sku",
    "title_differentiation": "cat_title_differentiation",
    "main_product_image_locator": "cat_main_image",
    "swatch_product_image_locator": "cat_swatch_image",
}

# Atributos repetidos: prefijo del nombre base -> prefijo del campo.
_CAT_REPEATED_FIELDS: dict[str, str] = {
    "other_product_image_locator_": "cat_other_image_",
    "bullet_point": "cat_bullet_",
}

# R11 — vocabulario de parentesco por marketplace. Se compara normalizado
# (lower, sin acentos, sin punto final).
_PARENTAGE_PADRE = {"principal", "parent", "padre", "eltern", "principale", "pai", "genitore"}
_PARENTAGE_HIJO = {"ninos", "nino", "child", "hijo", "kind", "enfant", "figlio", "filho"}

# Columnas del Fee Preview (NA). fee_size_tier NO está acá: no existe (R14).
_FEE_FIELDS: dict[str, str] = {
    "estimated-referral-fee-per-unit": "fee_referral",
    "estimated-fee-total": "fee_total",
    "longest-side": "fee_longest_side",
    "median-side": "fee_median_side",
    "shortest-side": "fee_shortest_side",
    "unit-of-dimension": "fee_dim_unit",
    "item-package-weight": "fee_weight",
    "unit-of-weight": "fee_weight_unit",
    "currency": "fee_currency",
}

# R15 — fee_fulfillment se compone sumando estas dos.
_FEE_FULFILLMENT_PARTS = (
    "estimated-pick-pack-fee-per-unit",
    "estimated-weight-handling-fee-per-unit",
)

# Campos del snapshot, en orden. Toda fila del snapshot tiene TODAS estas keys
# (vacías donde la fuente no aporta), para que el diff no confunda "campo ausente"
# con "campo que cambió".
_SNAPSHOT_FIELDS: list[str] = (
    ["sku", "inv_status", "inv_price", "inv_asin"]
    + ["fee_referral", "fee_fulfillment", "fee_total", "fee_size_tier",
       "fee_longest_side", "fee_median_side", "fee_shortest_side",
       "fee_dim_unit", "fee_weight", "fee_weight_unit", "fee_currency"]
    + ["cat_status", "cat_title", "cat_product_type", "cat_parentage_raw",
       "cat_parent_sku", "cat_title_differentiation", "cat_main_image",
       "cat_swatch_image"]
    + [f"cat_other_image_{i}" for i in range(1, 9)]
    + [f"cat_bullet_{i}" for i in range(1, 6)]
    + ["br_asin", "br_buy_box", "br_sessions", "br_page_views",
       "br_units_ordered", "br_sales"]
)

# R21 — campos ADITIVOS del Business Report: cuando un child ASIN aparece en
# varias filas (colgado de distintos parents) estos se SUMAN. buy_box no está
# acá: se pondera por sesiones, no se suma.
_BR_ADITIVOS = ("br_sessions", "br_page_views", "br_units_ordered", "br_sales")


# =====================================================================
# Helpers
# =====================================================================

def _s(v) -> str:
    """Todo valor del snapshot sale como str trimmed (R4).

    None y el string 'None' que devuelve openpyxl para celdas vacías -> ''.
    Nunca devuelve None: el diff compara strings y un None suelto rompería la
    comparación.
    """
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() == "none" else s


def _num(v) -> str:
    """Limpia numéricos de Amazon dejándolos como str comparable.

    Saca separador de miles ('1,414' -> '1414'), símbolo de porcentaje
    ('99.79%' -> '99.79') y prefijo de moneda ('MX$35,991.75' -> '35991.75').
    Si no matchea nada numérico devuelve el valor trimmed tal cual: preferimos
    conservar el original a inventar un 0 (R4, comparar verbatim).
    """
    s = _s(v)
    if not s:
        return ""
    s = s.replace(",", "").replace("%", "")
    s = re.sub(r"^[A-Za-z]{0,3}\s*[$€£]\s*", "", s)
    return s.strip()


def _to_float(v) -> float | None:
    """float o None. None significa 'no había dato', distinto de 0.0."""
    s = _num(v)
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _fmt_num(v: float) -> str:
    """Formatea un numérico agregado como str estable para el diff.

    Entero exacto -> sin decimales ('165'); si no, hasta 2 decimales sin ceros
    de relleno ('99.79'). El punto es que el MISMO valor produzca SIEMPRE la
    misma string: la agregación corre para todos los ASIN (incluso los de una
    sola fila), así que el formato queda uniforme y el diff no reporta cambios
    fantasma por representación.
    """
    r = round(v + 0.0, 2)
    return str(int(r)) if r == int(r) else f"{r:.2f}".rstrip("0").rstrip(".")


def _is_status_estructural(inv_status, cat_es_padre) -> bool:
    """H1 — ¿este status es normal por estructura, y no una anomalía?

    Un SKU padre SIEMPRE figura 'Incomplete' en el All Listings: no es un listing
    vendible, es el contenedor de la variación. Validado en Setex MX: de los 9
    'Incomplete', 8 son padres. Marcarlos como alerta sería ruido permanente —
    exactamente el modo de falla de R8.

    Helper para el motor de diff (fase siguiente); la ingesta no lo usa todavía.
    """
    return _s(inv_status).lower() == "incomplete" and cat_es_padre is True


def _norm_key(s) -> str:
    """Normaliza un nombre de columna para matcheo tolerante (lower + colapso
    de espacios). Amazon varía el espaciado alrededor de los guiones entre
    exports del Business Report.
    """
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def _tech_base(name) -> str:
    """Nombre base de un atributo del Category Listings.

    Corta en el primer '[' o '#', que es donde arrancan los sufijos de
    marketplace / índice / language_tag:

        'contribution_sku#1.value'                                -> 'contribution_sku'
        'parentage_level[marketplace_id=A1AM78C64UM0Y8]#1.value'  -> 'parentage_level'
        'bullet_point[...][language_tag=es_MX]#3.value'           -> 'bullet_point'
        '::listing_status'                                        -> '::listing_status'
    """
    return re.split(r"[\[#]", _s(name))[0].strip()


def _es_padre(parentage_raw) -> bool | None:
    """R11 — resuelve el parentesco contra el diccionario multi-idioma.

    True (padre) / False (hijo) / None (sin dato o vocabulario desconocido).
    El None es deliberado: un valor que no matchea NINGUNA de las dos listas es
    un idioma que no cubrimos, y hay que verlo, no asumir hijo.
    """
    s = _strip_accents(_s(parentage_raw).lower()).rstrip(".").strip()
    if not s:
        return None
    if s in _PARENTAGE_PADRE:
        return True
    if s in _PARENTAGE_HIJO:
        return False
    _LOG.warning("parentage_level desconocido (no matchea padre ni hijo): %r", parentage_raw)
    return None


def _bump(diag: dict | None, key: str, n: int = 1) -> None:
    """Acumula un contador de diagnóstico si el caller pasó un dict."""
    if diag is not None:
        diag[key] = diag.get(key, 0) + n


# =====================================================================
# Parser 1 — All Listings (.txt TSV)
# =====================================================================

def _parse_all_listings(data: bytes, diag: dict | None = None) -> dict[str, dict]:
    """All Listings Report de Seller Central. Clave: `seller-sku`.

    Encoding utf-8-sig (BOM real, verificado) y delimitador tab.

    Aporta el campo que R12 identificó como el ÚNICO con señal real de status:
    `status` (Active / Inactive / Incomplete). El `::listing_status` del Category
    Listings es la constante 'Activa' y no sirve para diffear.

    También aporta `asin1`, que es el puente para cruzar el Business Report (que
    viene por ASIN, sin SKU).

    NO captura `quantity`: validado vacío en 59 de 63 SKUs porque los items FBA no
    reportan cantidad en este reporte.
    """
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text), delimiter="\t")

    out: dict[str, dict] = {}
    for row in reader:
        sku = _s(row.get("seller-sku"))
        if not sku:
            _bump(diag, "inv_filas_sin_sku")
            continue
        if sku in out:
            _bump(diag, "inv_skus_duplicados")
            _LOG.warning("All Listings: SKU duplicado %r, se conserva la primera fila", sku)
            continue
        out[sku] = {
            "inv_status": _s(row.get("status")),
            "inv_price": _num(row.get("price")),
            "inv_asin": _s(row.get("asin1")),
        }

    _bump(diag, "inv_skus", len(out))
    _LOG.info("All Listings: %d SKUs", len(out))
    return out


# =====================================================================
# Parser 2 — FBA Fee Preview (.csv)
# =====================================================================

def _parse_fee_preview(data: bytes, marketplace: str,
                       diag: dict | None = None) -> dict[str, dict]:
    """FBA Fee Preview. Clave: `sku`. Filtrado por `amazon-store` (R13).

    El archivo de Norteamérica trae US + CA + MX juntos: 3 monedas (USD/CAD/MXN) y
    2 sistemas de unidades (US en inches/pounds, MX y CA en centimeters/grams).
    Sin el filtro, el dedupe se queda con una fila arbitraria y el diff termina
    comparando pulgadas contra centímetros. Por eso el filtro corre PRIMERO, antes
    del dedupe.

    Encoding ISO-8859-1: el archivo no decodifica como utf-8. Además el primer
    header viene como '?"sku"' — un byte ASCII 0x3F literal pegado adelante, no un
    BOM; se limpia junto con las comillas.

    `fee_size_tier` se emite VACÍO EXPLÍCITO: `product-size-tier` no existe en el
    Fee Preview de NA (R14), y omitir la key haría que el diff la leyera como campo
    nuevo el día que aparezca.
    """
    text = data.decode("ISO-8859-1")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise ValueError("Fee Preview vacío: no se pudo leer ninguna fila.")

    header = [_s(h).lstrip("?").strip().strip('"') for h in rows[0]]
    records = [dict(zip(header, r)) for r in rows[1:]]
    _bump(diag, "fee_filas_totales", len(records))

    if "amazon-store" not in header:
        raise ValueError(
            "Fee Preview sin columna 'amazon-store': no se puede filtrar por "
            f"marketplace (R13). Headers leídos: {header}"
        )

    # R13 — filtro por marketplace ANTES de cualquier otra cosa.
    wanted = _s(marketplace).upper()
    filtered = [r for r in records if _s(r.get("amazon-store")).upper() == wanted]
    _bump(diag, "fee_filas_marketplace", len(filtered))
    _bump(diag, "fee_filas_otros_marketplaces", len(records) - len(filtered))
    if not filtered:
        vistos = sorted({_s(r.get("amazon-store")) for r in records})
        _LOG.warning("Fee Preview: 0 filas para marketplace %r. Presentes: %s", wanted, vistos)

    out: dict[str, dict] = {}
    for r in filtered:
        sku = _s(r.get("sku"))
        if not sku:
            _bump(diag, "fee_filas_sin_sku")
            continue
        if sku in out:
            # R6 — el dedupeBySku del HTML descartaba en silencio.
            _bump(diag, "fee_filas_dedupeadas")
            _LOG.warning("Fee Preview: SKU duplicado %r dentro de %s, se conserva la primera",
                         sku, wanted)
            continue

        rec = {campo: _num(r.get(col)) for col, campo in _FEE_FIELDS.items()}

        # R15 — fulfillment no viene entero: se suma pick-pack + weight-handling.
        partes = [_to_float(r.get(c)) for c in _FEE_FULFILLMENT_PARTS]
        presentes = [p for p in partes if p is not None]
        rec["fee_fulfillment"] = f"{sum(presentes):g}" if presentes else ""

        rec["fee_size_tier"] = ""  # R14 — vacío explícito, no ausente.
        out[sku] = rec

    _bump(diag, "fee_skus", len(out))
    _LOG.info("Fee Preview [%s]: %d filas totales, %d del marketplace, %d SKUs",
              wanted, len(records), len(filtered), len(out))
    return out


# =====================================================================
# Parser 3 — Category Listings (.xlsm)
# =====================================================================

def _find_plantilla_sheet(sheetnames: list[str]) -> str:
    """R2 — match ANCLADO `^plantilla$`, case-insensitive.

    El mismo archivo trae una hoja `Cambios en la plantilla`: un match por
    "contiene" la agarraría a ella. Si no hay hoja exacta, excepción explícita.
    """
    for name in sheetnames:
        if _s(name).lower() == "plantilla":
            return name
    raise ValueError(
        "Category Listings sin hoja 'Plantilla' (match exacto, case-insensitive). "
        f"Hojas encontradas: {sheetnames}. "
        "En otros marketplaces la hoja se llama Template (US) / Modele (CA) / "
        "Vorlage (DE): agregar el nombre acá, NO relajar el match a 'contiene'."
    )


def _map_tech_columns(tech_row: tuple) -> dict[str, int | list[int]]:
    """Mapea nombre técnico -> índice de columna. Match por NOMBRE BASE.

    R17 — sin fallback posicional. Si falta un atributo obligatorio se levanta
    excepción: el fallback del prototipo asumía posiciones de US (imagen principal
    en 28, bullets en 39-43) que en MX son 33 y 44-48, y leía columnas equivocadas
    EN SILENCIO.

    Los repetidos preservan el orden de columna para poder diffear slot por slot
    (D5): 'Imagen secundaria 3' en vez de 'alguna imagen cambió'.
    """
    bases = [(i, _tech_base(c)) for i, c in enumerate(tech_row) if _s(c)]

    mapping: dict[str, int | list[int]] = {}
    faltantes: list[str] = []
    for base, campo in _CAT_SINGLE_FIELDS.items():
        idxs = [i for i, b in bases if b == base]
        if not idxs:
            faltantes.append(base)
            continue
        if len(idxs) > 1:
            _LOG.warning("Category Listings: %r aparece %d veces, se usa la primera (col %d)",
                         base, len(idxs), idxs[0])
        mapping[campo] = idxs[0]

    if faltantes:
        raise ValueError(
            "Category Listings: no se encontraron los atributos técnicos "
            f"{faltantes} en la fila {_CAT_TECH_ROW}. "
            "NO hay fallback posicional a propósito (R17): las posiciones varían "
            "por marketplace y leer la columna equivocada es un error silencioso. "
            "Revisar que el export sea el Category Listings Report completo."
        )

    for prefijo, campo_prefijo in _CAT_REPEATED_FIELDS.items():
        if prefijo.endswith("_"):
            # other_product_image_locator_1..8 -> el número va en el nombre base.
            def _slot(i: int) -> int:
                m = re.search(r"(\d+)$", _tech_base(tech_row[i]))
                return int(m.group(1)) if m else 0
            idxs = sorted((i for i, b in bases if b.startswith(prefijo)), key=_slot)
        else:
            # bullet_point x5 -> mismo nombre base, se ordena por columna.
            idxs = sorted(i for i, b in bases if b == prefijo)
        if not idxs:
            _LOG.warning("Category Listings: 0 columnas para %r", prefijo)
        mapping[campo_prefijo] = idxs

    return mapping


def _parse_category_listings(data: bytes, diag: dict | None = None) -> dict[str, dict]:
    """Category Listings Report (.xlsm). Clave: `contribution_sku`.

    Hoja `Plantilla` (R2), mapeo por nombre base de la fila 5, datos desde la fila
    7 — la 6 es el ejemplo precargado de Amazon (SKU 'ABC123').

    Emite `cat_es_padre` como bool: es un flag DERIVADO para filtrar la vista
    ("Ocultar padres", D6), no un campo que se diffee. Es la única excepción a la
    regla de str-trimmed (R4), y está acotada a propósito.
    """
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = _find_plantilla_sheet(list(wb.sheetnames))
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()

    if len(rows) < _CAT_FIRST_DATA_ROW:
        raise ValueError(
            f"Category Listings: la hoja {sheet!r} tiene {len(rows)} filas, se "
            f"esperaban al menos {_CAT_FIRST_DATA_ROW} (nombres técnicos en la fila "
            f"{_CAT_TECH_ROW}, datos desde la {_CAT_FIRST_DATA_ROW})."
        )

    mapping = _map_tech_columns(rows[_CAT_TECH_ROW - 1])
    imagenes = mapping["cat_other_image_"]
    bullets = mapping["cat_bullet_"]

    def _cell(row: tuple, idx: int) -> str:
        return _s(row[idx]) if idx < len(row) else ""

    out: dict[str, dict] = {}
    for row in rows[_CAT_FIRST_DATA_ROW - 1:]:
        sku = _cell(row, mapping["cat_sku"])
        if not sku:
            _bump(diag, "cat_filas_sin_sku")
            continue
        if sku in out:
            _bump(diag, "cat_skus_duplicados")
            _LOG.warning("Category Listings: SKU duplicado %r, se conserva la primera fila", sku)
            continue

        rec = {campo: _cell(row, idx) for campo, idx in mapping.items()
               if isinstance(idx, int)}
        rec.pop("cat_sku", None)

        for n, idx in enumerate(imagenes[:8], start=1):
            rec[f"cat_other_image_{n}"] = _cell(row, idx)
        for n, idx in enumerate(bullets[:5], start=1):
            rec[f"cat_bullet_{n}"] = _cell(row, idx)

        rec["cat_es_padre"] = _es_padre(rec.get("cat_parentage_raw", ""))
        out[sku] = rec

    padres = sum(1 for r in out.values() if r["cat_es_padre"] is True)
    hijos = sum(1 for r in out.values() if r["cat_es_padre"] is False)
    sin_dato = sum(1 for r in out.values() if r["cat_es_padre"] is None)
    _bump(diag, "cat_skus", len(out))
    _bump(diag, "cat_padres", padres)
    _bump(diag, "cat_hijos", hijos)
    _bump(diag, "cat_parentage_sin_dato", sin_dato)
    _LOG.info("Category Listings [%s]: %d SKUs (%d padres, %d hijos, %d sin parentage)",
              sheet, len(out), padres, hijos, sin_dato)
    return out


# =====================================================================
# Parser 4 — Business Report by Child Item (.csv)
# =====================================================================

def _parse_business_report(data: bytes, diag: dict | None = None) -> dict[str, dict]:
    """Business Report 'Detail Page Sales and Traffic By Child Item'.

    Clave: `(Child) ASIN`. Encoding utf-8-sig.

    TODAS las filas entran y cruzan por child ASIN. Las filas donde
    `(Parent) ASIN == (Child) ASIN` son productos STANDALONE legítimos (sin
    variaciones), no filas de padre puro: este reporte no trae padres. Lo que sí
    queda vacío legítimamente es el buy_box de un SKU padre del catálogo
    (`parentage_level = 'Principal.'`), que no aparece acá — eso NO es dato faltante.

    R21 — un child ASIN puede colgar de varios parents y aparecer en más de una
    fila; los aditivos se suman, buy_box se pondera por sessions. Descartar filas
    pierde sesiones/ventas reales: en Setex MX, B09HVXDH7M viene en 2 filas y
    quedarse con la primera tiraba 53 sesiones y MX$1,040.

    La agregación corre para TODOS los ASIN, incluso los de una sola fila, para
    que el formato numérico de salida sea uniforme.

    El match de columnas es tolerante al espaciado (Amazon varía los espacios
    alrededor de los guiones entre exports) y excluye siempre las variantes B2B.
    """
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    campos = list(reader.fieldnames or [])
    if not campos:
        raise ValueError("Business Report vacío: no se pudo leer el header.")

    normalizados = [(_norm_key(c), c) for c in campos]

    def _col(needle: str) -> str | None:
        """Primera columna cuyo nombre normalizado contiene `needle`, sin B2B."""
        n = _norm_key(needle)
        for norm, real in normalizados:
            if n in norm and "b2b" not in norm:
                return real
        return None

    col_child = _col("(child) asin")
    col_buybox = _col("featured offer (buy box) percentage")
    col_sessions = _col("sessions - total")
    col_page_views = _col("page views - total")
    col_units = _col("units ordered")
    col_sales = _col("ordered product sales")

    if not col_child:
        raise ValueError(
            "Business Report sin columna '(Child) ASIN'. Probablemente se bajó "
            "'Sales and Traffic by Date' en vez de 'By ASIN -> Detail Page Sales "
            f"and Traffic By Child Item'. Headers leídos: {campos}"
        )
    if not col_buybox:
        _LOG.warning("Business Report sin columna de Buy Box: el chequeo 3 queda sin dato.")

    # Paso 1 — agrupar las filas crudas por child ASIN.
    crudas: dict[str, list[dict]] = {}
    for row in reader:
        asin = _s(row.get(col_child))
        if not asin:
            _bump(diag, "br_filas_sin_asin")
            continue
        crudas.setdefault(asin, []).append({
            "buy_box": _to_float(row.get(col_buybox)) if col_buybox else None,
            "br_sessions": _to_float(row.get(col_sessions)) if col_sessions else None,
            "br_page_views": _to_float(row.get(col_page_views)) if col_page_views else None,
            "br_units_ordered": _to_float(row.get(col_units)) if col_units else None,
            "br_sales": _to_float(row.get(col_sales)) if col_sales else None,
        })

    # Paso 2 — R21: agregar. Aditivos suman; buy_box pondera por sesiones.
    out: dict[str, dict] = {}
    for asin, filas in crudas.items():
        if len(filas) > 1:
            _bump(diag, "br_asins_multi_parent")
            _LOG.info("Business Report: %r viene en %d filas (varios parents), se agregan",
                      asin, len(filas))

        rec: dict[str, str] = {}
        for campo in _BR_ADITIVOS:
            valores = [f[campo] for f in filas if f[campo] is not None]
            rec[campo] = _fmt_num(sum(valores)) if valores else ""

        con_bb = [f for f in filas if f["buy_box"] is not None]
        if not con_bb:
            rec["br_buy_box"] = ""
        else:
            peso_total = sum(f["br_sessions"] or 0.0 for f in con_bb)
            if peso_total > 0:
                # Promedio PONDERADO por sesiones: un parent con 112 sesiones
                # pesa más que uno con 53.
                rec["br_buy_box"] = _fmt_num(
                    sum(f["buy_box"] * (f["br_sessions"] or 0.0) for f in con_bb) / peso_total
                )
            else:
                # Sin sesiones en ninguna fila no hay con qué ponderar: simple.
                rec["br_buy_box"] = _fmt_num(
                    sum(f["buy_box"] for f in con_bb) / len(con_bb)
                )
        out[asin] = rec

    _bump(diag, "br_filas", sum(len(f) for f in crudas.values()))
    _bump(diag, "br_asins", len(out))
    _LOG.info("Business Report: %d filas -> %d ASINs child (%d con varios parents)",
              sum(len(f) for f in crudas.values()), len(out),
              (diag or {}).get("br_asins_multi_parent", 0))
    return out


# =====================================================================
# Ensamblado del snapshot
# =====================================================================

def _build_snapshot(all_listings: dict[str, dict],
                    fee_preview: dict[str, dict],
                    category: dict[str, dict],
                    business_report: dict[str, dict],
                    diag: dict | None = None) -> dict[str, dict]:
    """Une las 4 fuentes en una foto del catálogo. Clave primaria: SKU.

    El Business Report viene por ASIN y se cruza vía el mapa asin -> SKUs que aporta
    el All Listings (`asin1`). Los ASIN que no cruzan contra ningún SKU se CUENTAN y
    se loguean (R7); no rompen el ensamblado, pero tampoco desaparecen en silencio
    como en el prototipo.

    R20 — buy_box es propiedad del ASIN/oferta, no del SKU: se replica a TODOS los
    SKUs que comparten el ASIN. No depender del orden del archivo. En Setex MX hay
    10 ASIN repartidos en 21 SKUs (6 de ellos con datos de BR): quedarse con "el
    primer SKU del All Listings" hacía que el buy_box saltara de SKU el día que
    Amazon reordenara el export, y el diff lo reportaba como cambio real.

    Un SKU presente en una fuente y ausente en otra NO es un error: se completa con
    strings vacíos. Toda fila termina con el mismo juego de keys
    (`_SNAPSHOT_FIELDS`), para que el motor de diff no confunda "la fuente no trajo
    el campo" con "el campo cambió".
    """
    universo = sorted(set(all_listings) | set(fee_preview) | set(category))

    # R20 — mapa asin -> TODOS sus SKUs. Sin "primero gana": un ASIN compartido
    # por varios SKUs no es ambigüedad a resolver, es un hecho del catálogo.
    asin_a_skus: dict[str, list[str]] = {}
    for sku, rec in all_listings.items():
        asin = _s(rec.get("inv_asin"))
        if asin:
            asin_a_skus.setdefault(asin, []).append(sku)

    compartidos = {a: s for a, s in asin_a_skus.items() if len(s) > 1}
    if compartidos:
        _bump(diag, "asins_compartidos", len(compartidos))
        if diag is not None:
            diag["asins_compartidos_mapa"] = compartidos
        _LOG.info("%d ASIN compartidos por más de un SKU: los datos de BR se "
                  "replican a todos (R20)", len(compartidos))

    br_por_sku: dict[str, dict] = {}
    asins_sin_cruce: list[str] = []
    for asin, rec in business_report.items():
        skus = asin_a_skus.get(asin)
        if not skus:
            asins_sin_cruce.append(asin)
            continue
        for sku in skus:
            br_por_sku[sku] = {**rec, "br_asin": asin}

    snapshot: dict[str, dict] = {}
    for sku in universo:
        fila = {campo: "" for campo in _SNAPSHOT_FIELDS}
        fila["sku"] = sku
        fila.update(all_listings.get(sku, {}))
        fila.update(fee_preview.get(sku, {}))
        fila.update(category.get(sku, {}))
        fila.update(br_por_sku.get(sku, {}))
        fila.setdefault("cat_es_padre", None)
        snapshot[sku] = fila

    _bump(diag, "snapshot_skus", len(snapshot))
    _bump(diag, "br_asins_sin_cruce", len(asins_sin_cruce))
    _bump(diag, "skus_con_datos_br", len(br_por_sku))
    if diag is not None:
        diag["br_asins_sin_cruce_lista"] = asins_sin_cruce
        diag["skus_solo_en_fee"] = sorted(set(fee_preview) - set(all_listings) - set(category))
        diag["skus_sin_fee"] = sorted(set(all_listings) - set(fee_preview))

    if asins_sin_cruce:
        _LOG.warning("Business Report: %d ASIN sin SKU en el All Listings: %s",
                     len(asins_sin_cruce), asins_sin_cruce)
    _LOG.info("Snapshot: %d SKUs (inv=%d, fee=%d, cat=%d, con datos de BR=%d)",
              len(snapshot), len(all_listings), len(fee_preview), len(category),
              len(br_por_sku))
    return snapshot


# =====================================================================
# Motor de diff — núcleo
# =====================================================================

# Campos que el diff NO compara.
#   sku          — es la CLAVE del snapshot. Para un SKU presente en los dos lados
#                  siempre es igual a sí mismo: compararlo es una tautología, no un
#                  chequeo.
#   cat_es_padre — flag DERIVADO (R4), no un dato del catálogo; tampoco está en
#                  _SNAPSHOT_FIELDS. Se lista acá para dejar dicho que la omisión
#                  es deliberada y no un olvido.
_DIFF_EXCLUIDOS = ("sku", "cat_es_padre")

# Orden de comparación = orden de _SNAPSHOT_FIELDS. Es lo que hace determinista la
# salida del diff, junto con el recorrido de SKUs ordenado.
_DIFF_FIELDS: list[str] = [c for c in _SNAPSHOT_FIELDS if c not in _DIFF_EXCLUIDOS]

# R22 — prefijo de los campos que el ensamblado REPLICA a todos los SKUs que
# comparten un ASIN (R20). `br_asin` entra a propósito: también se replica, y viaja
# pegado a los valores (`_build_snapshot` los escribe juntos).
_BR_PREFIJO = "br_"


def _asin_del_cambio(viejo: dict | None, nuevo: dict | None) -> str:
    """ASIN de referencia de un SKU: `br_asin` del snapshot nuevo, con fallback al
    viejo.

    El fallback NO es cosmético, es lo que hace funcionar R22 en el caso feo: cuando
    un ASIN desaparece del export del Business Report, `br_asin` va 'B09...' -> ''
    y los cinco br_* se vacían de golpe. Tomando solo el lado nuevo, esas N alertas
    tendrían ASIN vacío y no agruparían — volveríamos a las N alertas idénticas que
    R22 existe para evitar. Con el fallback colapsan bajo el ASIN viejo.
    """
    for rec in (nuevo, viejo):
        if rec:
            asin = _s(rec.get("br_asin"))
            if asin:
                return asin
    return ""


def _es_padre_del_cambio(viejo: dict | None, nuevo: dict | None) -> bool:
    """Marca si el SKU del cambio es un padre de variación.

    Se lee del snapshot nuevo, y del viejo para las bajas (donde el nuevo no existe)
    — misma regla que el `padreDe` del HTML. `cat_es_padre` es True/False/None (R11);
    acá colapsa a bool porque el consumidor es el filtro binario "Ocultar padres"
    (D6): un parentesco desconocido no se esconde.
    """
    rec = nuevo if nuevo is not None else viejo
    return (rec or {}).get("cat_es_padre") is True


def _diff_snapshots(snap_viejo: dict[str, dict],
                    snap_nuevo: dict[str, dict]) -> list[dict]:
    """Compara dos snapshots `{sku: {campo: str}}` y devuelve los cambios CRUDOS.

    Cada cambio:

        {"sku", "asin", "es_padre", "tipo", "campo", "valor_viejo", "valor_nuevo"}

    donde `tipo` es 'alta_sku' | 'baja_sku' | 'cambio_campo'.

    Detecta, no juzga: sin severidad (D3) y sin comparar unidades antes que valores
    (R10). Ambas cosas van en la fase siguiente.

    Comparación STRING-TRIMMED, sin reconvertir tipos (R4): los valores ya salen
    normalizados de la ingesta (`_num` limpia miles/moneda/porcentaje, `_fmt_num`
    canoniza los agregados del Business Report). DESVÍO CONSCIENTE vs el HTML de
    Marcos, que además tenía una guarda numérica (19.99 contra 19.990 no disparaba).
    Los campos que pasan por `_num` sin canonizar decimales — `inv_price`, `fee_*` —
    quedan expuestos a un falso positivo si Amazon cambia el formato del decimal sin
    cambiar el valor. Anotado para la tanda de severidades, que va a comparar
    magnitudes numéricamente igual.

    Otro desvío consciente: el HTML descartaba el cambio cuando un lado venía vacío
    en Dimensiones/Fee ('dato faltante'). Acá vacío -> valor y valor -> vacío SÍ son
    cambios; vacío en AMBOS lados no lo es, por igualdad.

    Salida determinista: SKUs en orden alfabético, campos en el orden de
    `_SNAPSHOT_FIELDS`.
    """
    cambios: list[dict] = []

    for sku in sorted(set(snap_viejo) | set(snap_nuevo)):
        viejo = snap_viejo.get(sku)
        nuevo = snap_nuevo.get(sku)
        base = {
            "sku": sku,
            "asin": _asin_del_cambio(viejo, nuevo),
            "es_padre": _es_padre_del_cambio(viejo, nuevo),
        }

        if viejo is None:
            # Alta: una sola entrada por SKU. Los valores van vacíos a propósito —
            # el '(no existía)' / 'Aparece en catálogo' del HTML es presentación y
            # vive en la UI, no en el dato.
            cambios.append({**base, "tipo": "alta_sku", "campo": "",
                            "valor_viejo": "", "valor_nuevo": ""})
            continue

        if nuevo is None:
            cambios.append({**base, "tipo": "baja_sku", "campo": "",
                            "valor_viejo": "", "valor_nuevo": ""})
            continue

        for campo in _DIFF_FIELDS:
            va = _s(viejo.get(campo))
            vn = _s(nuevo.get(campo))
            if va == vn:
                continue
            cambios.append({**base, "tipo": "cambio_campo", "campo": campo,
                            "valor_viejo": va, "valor_nuevo": vn})

    altas = sum(1 for c in cambios if c["tipo"] == "alta_sku")
    bajas = sum(1 for c in cambios if c["tipo"] == "baja_sku")
    _LOG.info("Diff: %d cambios (%d altas, %d bajas, %d de campo) sobre %d SKUs "
              "viejos -> %d nuevos", len(cambios), altas, bajas,
              len(cambios) - altas - bajas, len(snap_viejo), len(snap_nuevo))
    return cambios


def _agrupar_alertas_por_asin(cambios: list[dict]) -> list[dict]:
    """R22 — colapsa las alertas replicadas por ASIN y uniforma la salida.

    Los campos `br_*` (Buy Box, sessions, page views, units, sales, y el propio
    `br_asin`) son propiedad del ASIN, no del SKU: `_build_snapshot` los REPLICA a
    todos los SKUs que comparten el ASIN (R20). Sin agrupar, un único cambio de Buy
    Box de un ASIN con 5 SKUs produce 5 alertas idénticas — el mismo modo de falla
    de R8 (el equipo aprende a ignorar el ruido) por otra vía.

    Agrupa por `(asin, campo, valor_viejo, valor_nuevo)` los `cambio_campo` de campos
    `br_*`. El resto — campos no br_*, altas y bajas — queda una alerta por SKU.

    TODA alerta sale con `"skus": list[str]`, así el consumidor tiene una interfaz
    uniforme y no necesita saber si hubo agrupación. En una alerta agrupada `"sku"`
    conserva el primer SKU como representante estable y `"skus"` trae la lista
    completa.

    Guarda: si el ASIN resuelto viene vacío, NO agrupa. Agrupar bajo la clave vacía
    fusionaría SKUs que no tienen nada que ver. En la práctica no debería dispararse
    (`br_asin` viaja pegado a los valores br_*), pero fusionar en silencio es peor
    que una alerta de más.

    `es_padre` de un grupo se resuelve con `all(...)`, no `any(...)`: el consumidor
    es el filtro "Ocultar padres" (D6), y si un solo SKU hijo real está afectado la
    alerta NO se puede esconder.
    """
    grupos: dict[tuple, dict] = {}
    salida: list[dict] = []

    for c in cambios:
        campo = _s(c.get("campo"))
        asin = _s(c.get("asin"))
        agrupable = (c.get("tipo") == "cambio_campo"
                     and campo.startswith(_BR_PREFIJO)
                     and bool(asin))

        if not agrupable:
            salida.append({**c, "skus": [c["sku"]]})
            continue

        clave = (asin, campo, c.get("valor_viejo", ""), c.get("valor_nuevo", ""))
        grupo = grupos.get(clave)
        if grupo is None:
            grupo = {**c, "skus": [c["sku"]]}
            grupos[clave] = grupo
            # Se agrega ACÁ para preservar el orden de primera aparición; las
            # ocurrencias siguientes mutan este mismo objeto por referencia.
            salida.append(grupo)
            continue

        if c["sku"] not in grupo["skus"]:
            grupo["skus"].append(c["sku"])
        grupo["es_padre"] = bool(grupo["es_padre"]) and bool(c.get("es_padre"))

    colapsadas = len(cambios) - len(salida)
    if colapsadas:
        _LOG.info("R22: %d cambios -> %d alertas (%d replicadas por ASIN colapsadas)",
                  len(cambios), len(salida), colapsadas)
    return salida
