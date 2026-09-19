"""Parser one-shot de las 5 propuestas semilla.

Procesa los archivos originales (PDF/DOCX) y los convierte en JSONs estructurados
en `data/sales/_seed/`. Los JSONs son referencias internas para entender qué
bloques del catálogo aparecen en propuestas reales históricas.

Heurística (deliberadamente simple — NO buscar 100% de precisión):
- Para cada bloque del catálogo, define una lista de matchers (sustrings, regex
  laxos) que sugieren su presencia.
- Para cada source file, extrae texto plano y busca matchers.
- Bloques detectados → array `detected_blocks`.
- Bloques no mapeados → array `unmapped_sections` con el texto crudo.

Si un archivo fuente no existe, genera un stub con TODO comment para revisión.

Uso:
    python scripts/parse_seed_proposals.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

# Asegurarse de que el script funcione tanto si se corre desde la raíz como desde scripts/
REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from core.proposals.paths import SEED_DIR, ensure_dirs  # noqa: E402

# ─────────────────────────────────────────────────────────────────────────────
# Lookup paths para los archivos fuente
# ─────────────────────────────────────────────────────────────────────────────

# Posibles ubicaciones donde podrían vivir los seed files (Claude.ai project mount,
# carpeta local, downloads). El primero que matchee gana.
SEARCH_DIRS = [
    REPO_ROOT,
    REPO_ROOT / "data" / "sales" / "_seed_sources",
    Path("/mnt/project"),  # Claude.ai project directory
    Path.home() / "Downloads",
]


SEED_SOURCES = [
    {
        "seed_id": "sunny_zebra",
        "client_name": "Sunny Zebra",
        "archetype": "cvr",
        "language": "en",
        "filenames": [
            "Copia_de_The_Sunny_Zebra_Análisis_Amazon.pdf",
            "The_Sunny_Zebra_Análisis_Amazon.pdf",
            "sunny_zebra.pdf",
        ],
    },
    {
        "seed_id": "happy_mammoth",
        "client_name": "Happy Mammoth",
        "archetype": "scale_seo",
        "language": "en",
        "filenames": [
            "Copia_de_Happy_Mammoth___Commercial_Proposal.pdf",
            "Happy_Mammoth_Commercial_Proposal.pdf",
            "happy_mammoth.pdf",
        ],
    },
    {
        "seed_id": "garland_rug",
        "client_name": "Garland Rug",
        "archetype": "launch",
        "language": "en",
        "filenames": [
            "Garland_Rug_New_Brand_Launching_Proposal.pdf",
            "garland_rug.pdf",
        ],
    },
    {
        "seed_id": "nandog",
        "client_name": "Nandog",
        "archetype": "scale_seo",
        "language": "es",
        "filenames": [
            "Nandog_Análisis_Amazon.docx",
            "nandog.docx",
        ],
    },
    {
        "seed_id": "nobl_travel",
        "client_name": "Nobl Travel",
        "archetype": "defense",
        "language": "en",
        "filenames": [
            "Copia_de_Nobl_Travel_Amazon_Brand_Opportunity.docx",
            "Nobl_Travel_Amazon_Brand_Opportunity.docx",
            "nobl_travel.docx",
        ],
    },
]


# ─────────────────────────────────────────────────────────────────────────────
# Catálogo de matchers — sustrings que sugieren la presencia de un module_id
# ─────────────────────────────────────────────────────────────────────────────

BLOCK_MATCHERS: dict[str, list[str]] = {
    "F1_cover": ["commercial proposal", "análisis amazon", "brand opportunity", "amazon proposal"],
    "F2_about_stats": ["years of experience", "monthly revenue", "monthly ad spend", "average roas", "$100m"],
    "F3_brand_stages": ["create", "launch", "scale", "brands we work with", "marcas con las que trabajamos"],
    "F4_operation_pillars": [
        "strategic management",
        "advertising",
        "creative",
        "supply",
        "how we operate",
        "cómo operamos",
        "4 pillars",
        "4 pilares",
    ],
    "F5_case_studies": ["case study", "case studies", "casos de éxito", "tattoo care", "wamery", "shapermint", "aimzone"],
    "F6_why_capybaras": ["why capybaras", "por qué capybaras", "amateur hour", "operating system", "capital intelectual"],
    "F7_team": ["assigned team", "equipo asignado", "account manager", "ppc expert"],
    "F8_lets_scale": ["let's scale", "hagámoslo crecer", "let us scale", "next steps"],
    "V1_brand_overview": ["brand overview", "resumen de marca", "founded", "mission statement", "value proposition"],
    "V2_category_overview": ["category overview", "resumen de categoría", "market summary", "competitor analysis", "median price"],
    "V3_seo_opportunity": ["seo opportunity", "missing keywords", "launch score", "page 1 domination"],
    "V4_listing_improvements_current_state": ["current state", "estado actual", "listing checklist", "main image", "infographics", "a+ content"],
    "V5_listing_comparison_competitor": ["side-by-side", "side by side", "vs competitor", "vs competidor"],
    "V6_growth_plan_phases": ["growth plan", "plan de crecimiento", "phase 1", "phase 2", "phase 3", "foundations", "expansion", "dsp"],
    "V7_ppc_audit_overview": ["ppc audit", "auditoría ppc"],
    "V8_ppc_audit_target_types": ["exact match", "broad match", "target types", "match type"],
    "V9_ppc_audit_campaign_type": ["sponsored brand", "sponsored display", "campaign type", "tipo de campaña"],
    "V10_ppc_audit_harvesting": ["harvesting", "harvest", "promote to exact"],
    "V11_ppc_audit_placement": ["placement", "top of search", "product pages", "rest of search"],
    "V12_ppc_audit_negative_keywords": ["negative keyword", "negativo", "wasted spend"],
    "V13_sales_forecast_table": ["sales forecast", "forecast de ventas", "projected", "delta %"],
    "V14_market_trend_growth": ["market trend", "quarterly", "year-over-year", "yoy growth"],
    "V15_competitive_position_table": ["competitive position", "competidores", "bsr"],
    "V16_listing_main_image_case_dermaglos": ["dermaglos", "ctr +", "cvr +", "before/after"],
    "V17_made_in_country_advantage": ["made in", "tariff", "restock", "arancel"],
    "V18_modular_launch_strategy": ["modular launch", "channels", "amazon shopify meta"],
    "V19_amazon_launch_grid": ["amazon launch", "listing pcc", "storefront brand story"],
    "V20_shopify_d2c_channel": ["shopify", "d2c", "customer journey"],
    "V21_meta_ads_growth": ["meta ads", "instagram ads", "facebook ads"],
    "V22_walmart_marketplaces": ["walmart", "ebay", "wayfair", "target+", "target plus"],
    "V23_branded_terms_overview": ["branded terms", "brand attack", "atacando la marca"],
    "V24_branded_terms_chart": ["attackers chart", "branded terms chart"],
    "V25_defensive_campaigns_strategy": ["defensive campaigns", "campañas defensivas"],
    "V26_branded_terms_summary": ["branded terms summary", "branded conclusions"],
    "V27_listing_audit_quality_score": ["quality score", "helium10 score", "helium 10 score"],
    "V28_listing_audit_infographics_detail": ["infographics detail", "slide-by-slide infographics"],
    "V29_listing_audit_aplus_comparison": ["a+ generic", "a+ premium", "premium a+"],
}


# ─────────────────────────────────────────────────────────────────────────────
# Extracción de texto
# ─────────────────────────────────────────────────────────────────────────────


def _find_source_file(filenames: list[str]) -> Path | None:
    for fname in filenames:
        for d in SEARCH_DIRS:
            candidate = d / fname
            if candidate.is_file():
                return candidate
    return None


def _extract_text_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError:
        return f"[ERROR] pdfplumber no instalado — no se pudo procesar {path.name}"
    try:
        with pdfplumber.open(path) as pdf:
            parts = []
            for page in pdf.pages:
                t = page.extract_text() or ""
                parts.append(t)
            return "\n".join(parts)
    except Exception as e:
        return f"[ERROR] al leer {path.name}: {e}"


def _extract_text_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return f"[ERROR] python-docx no instalado — no se pudo procesar {path.name}"
    try:
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:
        return f"[ERROR] al leer {path.name}: {e}"


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_text_pdf(path)
    if suffix == ".docx":
        return _extract_text_docx(path)
    return f"[ERROR] formato no soportado: {suffix}"


# ─────────────────────────────────────────────────────────────────────────────
# Detección de bloques
# ─────────────────────────────────────────────────────────────────────────────


def _detect_blocks(text: str) -> tuple[list[str], list[str]]:
    """Recorre BLOCK_MATCHERS y detecta cuáles aparecen en el texto.

    Returns:
        (detected_module_ids, unmapped_keywords)
    """
    lower = text.lower()
    detected: list[str] = []
    for module_id, matchers in BLOCK_MATCHERS.items():
        for needle in matchers:
            if needle in lower:
                detected.append(module_id)
                break

    # Buscar headers candidatos (líneas en mayúsculas o terminadas en ':') que no matchearon
    unmapped: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped) > 80 or len(stripped) < 5:
            continue
        # Heurística: linea corta tipo título
        is_titlelike = stripped.isupper() or stripped.endswith(":") or stripped.istitle()
        if not is_titlelike:
            continue
        line_lower = stripped.lower()
        any_matched = any(
            any(needle in line_lower for needle in matchers)
            for matchers in BLOCK_MATCHERS.values()
        )
        if not any_matched:
            unmapped.append(stripped)

    # Dedupe manteniendo orden
    unmapped = list(dict.fromkeys(unmapped))[:30]  # cap a 30 para no inflar el JSON
    return detected, unmapped


# ─────────────────────────────────────────────────────────────────────────────
# Construcción del JSON seed
# ─────────────────────────────────────────────────────────────────────────────


def _build_seed_proposal(
    seed_id: str,
    client_name: str,
    archetype: str,
    language: str,
    source_path: Path | None,
    text: str | None,
    detected: list[str],
    unmapped: list[str],
) -> dict:
    blocks = []
    for module_id in detected:
        blocks.append(
            {
                "module_id": module_id,
                "is_fixed": module_id.startswith("F"),
                "data": {},
                "copy_overrides": {},
                "_source": "auto-detected from seed file",
            }
        )

    return {
        "seed_id": seed_id,
        "client_name": client_name,
        "archetype": archetype,
        "language": language,
        "source_file": source_path.name if source_path else None,
        "source_found": source_path is not None,
        "_TODO": (
            None
            if source_path
            else "Source file no encontrado en repo local. Cuando esté disponible, "
            "correr de nuevo `python scripts/parse_seed_proposals.py` para poblar "
            "detected_blocks + unmapped_sections."
        ),
        "detected_blocks": blocks,
        "unmapped_sections": unmapped,
        "_text_sample": (text[:500] if text else None),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────


def main() -> int:
    ensure_dirs()
    SEED_DIR.mkdir(parents=True, exist_ok=True)

    generated = []
    stubs = []
    for spec in SEED_SOURCES:
        seed_id = spec["seed_id"]
        source = _find_source_file(spec["filenames"])

        if source is None:
            seed = _build_seed_proposal(
                seed_id=seed_id,
                client_name=spec["client_name"],
                archetype=spec["archetype"],
                language=spec["language"],
                source_path=None,
                text=None,
                detected=[],
                unmapped=[],
            )
            stubs.append(seed_id)
        else:
            text = _extract_text(source)
            detected, unmapped = _detect_blocks(text)
            seed = _build_seed_proposal(
                seed_id=seed_id,
                client_name=spec["client_name"],
                archetype=spec["archetype"],
                language=spec["language"],
                source_path=source,
                text=text,
                detected=detected,
                unmapped=unmapped,
            )

        out_path = SEED_DIR / f"{seed_id}.json"
        out_path.write_text(
            json.dumps(seed, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        generated.append(out_path.name)

    # Summary
    print("=" * 72)
    print("parse_seed_proposals.py — Summary")
    print("=" * 72)
    print(f"Archivos generados ({len(generated)}):")
    for fname in generated:
        marker = " [STUB]" if fname.replace(".json", "") in stubs else ""
        print(f"  - {fname}{marker}")
    if stubs:
        print()
        print(f"{len(stubs)} stubs generados — falta el archivo fuente para:")
        for sid in stubs:
            print(f"  - {sid}")
        print("Cuando los archivos estén disponibles, correr de nuevo este script.")
    print("=" * 72)
    return 0


if __name__ == "__main__":
    sys.exit(main())
