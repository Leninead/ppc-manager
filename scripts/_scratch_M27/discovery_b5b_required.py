# discovery_b5b_required.py — throwaway, no commitear
from openpyxl import load_workbook
import sys
sys.path.insert(0, r"C:\proyectos\ppc-manager")
from modules.pages.flat_file_migrator import (
    _parse_data_definitions,
    _detect_schema,
    _locate_template_headers,
)

OLD_PATH = r"C:\proyectos\ppc-manager\data\account-health\flat-file-migrator\ALRBB093_p_USA_2026__1_.xlsm"
NEW_PATH = r"C:\proyectos\ppc-manager\data\account-health\flat-file-migrator\COAT__5_.xlsm"

for label, path in [("OLD", OLD_PATH), ("NEW", NEW_PATH)]:
    print(f"\n{'='*60}\n{label}: {path}\n{'='*60}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        schema = _detect_schema(wb)
        dd = _parse_data_definitions(wb)
        headers = _locate_template_headers(wb)
        print(f"schema={schema} | dd_fields={len(dd)} | headers={headers}")

        # Q1: valores únicos en columna Required
        required_values = {}
        for fid, info in dd.items():
            r = info.get("required", "")
            required_values[r] = required_values.get(r, 0) + 1
        print(f"\nQ1 - Valores unicos columna 'Required?' (con count):")
        for v, c in sorted(required_values.items(), key=lambda x: -x[1]):
            print(f"  {v!r:30s} -> {c}")

        # Q2: primeros 5 fields marcados como Required-equivalente
        req_keywords = {"yes", "required", "y", "true", "1"}
        req_fields = [
            (fid, info) for fid, info in dd.items()
            if info.get("required", "").strip().lower() in req_keywords
        ]
        print(f"\nQ2 - Fields con required-equivalente (total {len(req_fields)}, muestra 5):")
        for fid, info in req_fields[:5]:
            print(f"  {fid:35s} label={info.get('label', '')!r}")

        # Q3: rows entre field_id_row+1 y data_start_row+5
        ws = wb["Template"]
        fid_row = headers["field_id_row"]
        data_start = headers["data_start_row"]
        print(f"\nQ3 - Rows entre field_id_row={fid_row} y data_start_row+5={data_start+5}:")
        for r_idx in range(fid_row + 1, min(data_start + 6, ws.max_row + 1)):
            cells = []
            for c_idx in range(1, 6):
                v = ws.cell(row=r_idx, column=c_idx).value
                cells.append(f"{str(v)[:25]!r}" if v else "''")
            print(f"  row {r_idx}: [{', '.join(cells)}]")
    finally:
        wb.close()
