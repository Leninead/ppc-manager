"""
Smoke E2E M27 v1.1 — Pipeline orquestador B1 → B5-c sin UI.

Objetivo: validar que el pipeline cross-schema completo (B1 detect + B2 dd +
B3 field map + B4a valid values + B4b value map + B5-a headers + B5-b
extract + B5-c migrate) funciona end-to-end contra par real.

Output:
- Stats agregadas (rows, new_fields, diagnostics by code/level)
- Resumen ejecutivo estilo lo que B6 va a mostrar en UI
- Excel migrado para inspección manual
- Exit code 0 si pipeline E2E PASS, 1 si fail.

Usage:
    python scripts/smoke_b6a_e2e_pipeline.py

Configuración: editá las constantes OLD_PATH y NEW_PATH abajo.
"""
import sys
import time
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, r'C:\proyectos\ppc-manager')

from openpyxl import load_workbook, Workbook
from modules.pages.flat_file_migrator import (
    _detect_schema,
    _parse_data_definitions,
    _parse_valid_values,
    _build_field_map,
    _build_value_map,
    _locate_template_headers,
    _extract_old_rows,
    _migrate_row,
)

# Configuración — paths del par discovery (ALRBB093 fptcustom OLD + COAT PTD NEW)
OLD_PATH = r'C:\proyectos\ppc-manager\data\account-health\flat-file-migrator\ALRBB093_p_USA_2026__1_.xlsm'
NEW_PATH = r'C:\proyectos\ppc-manager\data\account-health\flat-file-migrator\COAT__5_.xlsm'

OUTPUT_XLSX = Path(r'C:\proyectos\ppc-manager\scripts\_smoke_output_migrated.xlsx')


def banner(title: str, char: str = "=") -> None:
    print(f"\n{char * 78}")
    print(f"  {title}")
    print(char * 78)


def main() -> int:
    t_start = time.perf_counter()
    banner("M27 v1.1 - Smoke E2E B1 -> B5-c")

    # B1 schema detect
    t_b1 = time.perf_counter()
    old_wb = load_workbook(OLD_PATH, read_only=True, data_only=True)
    new_wb = load_workbook(NEW_PATH, read_only=True, data_only=True)
    old_schema = _detect_schema(old_wb)
    new_schema = _detect_schema(new_wb)
    print(f"\n[B1] Schema detect: OLD={old_schema}, NEW={new_schema}")
    print(f"     Tiempo: {(time.perf_counter() - t_b1) * 1000:.1f}ms")

    if old_schema == "unknown" or new_schema == "unknown":
        print("FAIL: schema desconocido en algun workbook")
        return 1

    # B2 Data Definitions
    t_b2 = time.perf_counter()
    old_dd = _parse_data_definitions(old_wb)
    new_dd = _parse_data_definitions(new_wb)
    old_required_count = sum(
        1 for v in old_dd.values()
        if v.get("required", "").strip().lower() in {"yes", "required"}
    )
    new_required_count = sum(
        1 for v in new_dd.values()
        if v.get("required", "").strip().lower() in {"yes", "required"}
    )
    print(f"\n[B2] Data Definitions:")
    print(f"     OLD fields: {len(old_dd)} ({old_required_count} Required)")
    print(f"     NEW fields: {len(new_dd)} ({new_required_count} Required)")
    print(f"     Tiempo: {(time.perf_counter() - t_b2) * 1000:.1f}ms")

    # B3 Field map cross-schema
    t_b3 = time.perf_counter()
    field_map, fm_warnings = _build_field_map(old_dd, new_dd)
    coverage_pct = (len(field_map) / max(len(old_dd), 1)) * 100
    print(f"\n[B3] Field map (cross-schema):")
    print(f"     Matches: {len(field_map)}/{len(old_dd)} ({coverage_pct:.1f}% cobertura)")
    print(f"     Warnings: {len(fm_warnings)}")
    print(f"     Tiempo: {(time.perf_counter() - t_b3) * 1000:.1f}ms")

    # B4a Valid Values
    t_b4a = time.perf_counter()
    old_vv = _parse_valid_values(old_wb)
    new_vv = _parse_valid_values(new_wb)
    print(f"\n[B4a] Valid Values:")
    print(f"      OLD enums: {len(old_vv)}")
    print(f"      NEW enums: {len(new_vv)}")
    print(f"      Tiempo: {(time.perf_counter() - t_b4a) * 1000:.1f}ms")

    # B4b Value map
    t_b4b = time.perf_counter()
    value_map, vm_warnings = _build_value_map(old_vv, new_vv)
    print(f"\n[B4b] Value map:")
    print(f"      Entries: {len(value_map)}")
    print(f"      Warnings: {len(vm_warnings)}")
    print(f"      Tiempo: {(time.perf_counter() - t_b4b) * 1000:.1f}ms")

    # B5-a Headers
    t_b5a = time.perf_counter()
    old_headers = _locate_template_headers(old_wb)
    new_headers = _locate_template_headers(new_wb)
    print(f"\n[B5-a] Headers locator:")
    print(f"       OLD: data_start_row={old_headers['data_start_row']}, "
          f"method={old_headers['detection_method']}")
    print(f"       NEW: data_start_row={new_headers['data_start_row']}, "
          f"method={new_headers['detection_method']}")
    print(f"       Tiempo: {(time.perf_counter() - t_b5a) * 1000:.1f}ms")

    # B5-b Extract old rows
    t_b5b = time.perf_counter()
    old_rows, b5b_diags = _extract_old_rows(old_wb, old_headers, old_dd)
    b5b_codes: dict[str, int] = defaultdict(int)
    b5b_levels: dict[str, int] = defaultdict(int)
    for d in b5b_diags:
        b5b_codes[d.get("code", "<no-code>")] += 1
        b5b_levels[d.get("level", "<no-level>")] += 1
    print(f"\n[B5-b] Extract old rows:")
    print(f"       Rows extraidas: {len(old_rows)}")
    print(f"       Diagnostics: {len(b5b_diags)}")
    print(f"       By code: {dict(b5b_codes)}")
    print(f"       By level: {dict(b5b_levels)}")
    print(f"       Tiempo: {(time.perf_counter() - t_b5b) * 1000:.1f}ms")

    # B5-c Migrate rows (loop)
    t_b5c = time.perf_counter()
    new_rows: list[dict[str, str]] = []
    all_b5c_diags: list[dict] = []
    b5c_codes: dict[str, int] = defaultdict(int)
    b5c_levels: dict[str, int] = defaultdict(int)

    for row_offset, old_row in enumerate(old_rows):
        new_row, diags = _migrate_row(
            old_row, field_map, value_map, old_dd, new_dd,
        )
        new_rows.append(new_row)
        for d in diags:
            d_with_idx = {**d, "row_offset": row_offset}
            all_b5c_diags.append(d_with_idx)
            b5c_codes[d.get("code", "<no-code>")] += 1
            b5c_levels[d.get("level", "<no-level>")] += 1

    total_new_fields = sum(len(r) for r in new_rows)
    avg_new_fields = total_new_fields / max(len(new_rows), 1)

    print(f"\n[B5-c] Migrate rows:")
    print(f"       New rows: {len(new_rows)}")
    print(f"       Total new_fields escritos: {total_new_fields}")
    print(f"       Avg new_fields/row: {avg_new_fields:.1f}")
    print(f"       Diagnostics: {len(all_b5c_diags)}")
    print(f"       By code: {dict(b5c_codes)}")
    print(f"       By level: {dict(b5c_levels)}")
    print(f"       Tiempo: {(time.perf_counter() - t_b5c) * 1000:.1f}ms")

    # Build output Excel
    t_xlsx = time.perf_counter()
    all_new_fields_set: set[str] = set()
    for r in new_rows:
        all_new_fields_set.update(r.keys())
    ordered_known = [f for f in new_dd.keys() if f in all_new_fields_set]
    extras = sorted(all_new_fields_set - set(ordered_known))
    ordered_cols = ordered_known + extras

    out_wb = Workbook()
    ws = out_wb.active
    ws.title = "Migrated"
    ws.append(ordered_cols)
    for r in new_rows:
        ws.append([r.get(f, "") for f in ordered_cols])
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(OUTPUT_XLSX)
    print(f"\n[XLSX] Output guardado: {OUTPUT_XLSX}")
    print(f"       Columnas: {len(ordered_cols)}")
    print(f"       Rows: {len(new_rows)}")
    print(f"       Tiempo: {(time.perf_counter() - t_xlsx) * 1000:.1f}ms")

    # Resumen ejecutivo estilo UI B6
    banner("RESUMEN EJECUTIVO (preview de lo que B6 mostraria en UI)", "-")
    print(f"""
Migracion cross-schema completada

  Schemas:
    * OLD: {old_schema}
    * NEW: {new_schema}

  Cobertura schema:
    * {len(field_map)}/{len(old_dd)} fields OLD mapeados ({coverage_pct:.1f}%)
    * {len(value_map)} enum translators activos

  Migracion:
    * {len(old_rows)} rows OLD procesadas
    * {len(new_rows)} rows NEW generadas
    * {total_new_fields} valores migrados (avg {avg_new_fields:.1f}/row)

  Diagnostics:
    B5-b (extraccion): {len(b5b_diags)} eventos""")
    for code, count in sorted(b5b_codes.items()):
        print(f"      * {code}: {count}")
    print(f"\n    B5-c (migracion): {len(all_b5c_diags)} eventos")
    for code, count in sorted(b5c_codes.items()):
        print(f"      * {code}: {count}")

    print(f"""
  Required gaps:
    * OLD Required missing: {b5b_codes.get('missing_required_in_old', 0)}
    * NEW Required missing post-migracion: {b5c_codes.get('missing_required_in_new', 0)}

  Data perdida (warnings criticos):
    * unmapped_field (sin destino): {b5c_codes.get('unmapped_field', 0)}
    * deprecated_value (mapea a None): {b5c_codes.get('deprecated_value', 0)}
    * deprecated_enum_no_target: {b5c_codes.get('deprecated_enum_no_target', 0)}
    * unknown_enum_value (pasthrough crudo): {b5c_codes.get('unknown_enum_value', 0)}
""")

    # Sanity checks finales
    banner("SANITY CHECKS", "-")
    checks = [
        ("OLD schema detected", old_schema != "unknown"),
        ("NEW schema detected", new_schema != "unknown"),
        ("OLD dd non-empty", len(old_dd) > 0),
        ("NEW dd non-empty", len(new_dd) > 0),
        ("field_map non-empty", len(field_map) > 0),
        ("value_map non-empty", len(value_map) > 0),
        ("OLD headers detected", old_headers.get("data_start_row") is not None),
        ("NEW headers detected", new_headers.get("data_start_row") is not None),
        ("Old rows extracted", len(old_rows) > 0),
        ("New rows generated", len(new_rows) > 0),
        ("Total new_fields > 0", total_new_fields > 0),
        ("Output XLSX exists", OUTPUT_XLSX.exists()),
    ]
    all_pass = True
    for name, ok in checks:
        mark = "OK" if ok else "FAIL"
        print(f"  [{mark}] {name}")
        if not ok:
            all_pass = False

    t_total = (time.perf_counter() - t_start) * 1000
    status = "PASS" if all_pass else "FAIL"
    banner(f"PIPELINE E2E {status} - total {t_total:.0f}ms")

    # Cerrar workbooks
    old_wb.close()
    new_wb.close()

    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
